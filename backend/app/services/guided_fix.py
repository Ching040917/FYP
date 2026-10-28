"""Guided Fixes: deterministic direct-format edits to a verified DOCX copy."""
import hashlib
import io
import math
import re
from decimal import Decimal, InvalidOperation
import unicodedata
import urllib.parse
from typing import Any, Dict, List, Optional, Tuple

from pydantic import BaseModel, ConfigDict, Field

from app.models.audit import AuditRecord, Violation
from app.services.document_parser import extract_paragraphs, parse_document
from app.services.profile_schema import (
    _FONT_SIZE_PT_RANGE,
    _LINE_SPACING_RANGE,
    _SPACING_PT_RANGE,
)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SAFE_RUN_RULES = frozenset({"FONT_CONSISTENCY", "FONT_SIZE", "REFERENCE_HEADING_BOLD"})
SAFE_PARA_RULES = frozenset({"ALIGNMENT", "LINE_SPACING", "SPACE_BEFORE", "SPACE_AFTER", "REFERENCE_HEADING_ALIGNMENT"})

MISMATCH_DETAIL = (
    "The selected document does not match the document used for this audit. "
    "Upload the original DOCX file and try again."
)
LEGACY_DETAIL = "Guided Fixes is unavailable for audits created before document matching was introduced."

# WordprocessingML namespaces (explicit — no qn() needed for the tag set).
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

# Structural safety gate: paragraphs/runs containing any of these elements
# (as self, descendant, or ancestor) are never modified. Tracked changes,
# fields, hyperlinks, comments, structured content, text boxes, subdocs,
# math, and AlternateContent are all conservative skips.
_UNSAFE_TAGS = frozenset({
    _W + "ins", _W + "del", _W + "fldChar", _W + "instrText",
    _W + "hyperlink", _W + "commentRangeStart", _W + "commentRangeEnd",
    _W + "commentReference", _W + "sdt", _W + "txbxContent", _W + "subDoc",
    _M + "oMath", _MC + "AlternateContent",
})

_ALIGNMENTS = frozenset({"left", "center", "right", "justify"})

_PT_RE = re.compile(r"^([0-9]+(?:\.[0-9]+)?)\s*pt$", re.IGNORECASE)
_RATIO_RE = re.compile(r"^([0-9]+(?:\.[0-9]+)?)$")
_TRAILING_SIZE_RE = re.compile(r"\s+\d+(?:\.\d+)?pt$", re.IGNORECASE)


class GuidedFixFailed(Exception):
    """A fix could not be applied or the corrected output could not be
    verified. No partial output is ever returned."""
    pass


class ReferenceHeadingBoldUnavailable(GuidedFixFailed):
    """Only these candidate-specific reasons may cross the public Review boundary."""

    DETAILS = {
        "ambiguous_heading": (
            "ACA cannot identify one unambiguous Reference List heading. "
            "Correct this finding manually; your original document is unchanged."
        ),
        "unverifiable_bold": (
            "ACA cannot safely confirm bold formatting for every text segment in this heading. "
            "Correct its bold formatting manually; your original document is unchanged."
        ),
        "unsupported_package": (
            "This correction is unavailable for signed, protected or macro-enabled documents. "
            "Correct the heading manually; your original document is unchanged."
        ),
    }

    def __init__(self, reason: str):
        self.reason = reason
        super().__init__(self.DETAILS[reason])


# ---------------------------------------------------------------------------
# Request schema
# ---------------------------------------------------------------------------

class GuidedFixSelectionRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    violation_ids: List[str] = Field(min_length=1, max_length=20)


# ---------------------------------------------------------------------------
# Parse / compare helpers
# ---------------------------------------------------------------------------

def _parse_pt(s) -> Optional[float]:
    """Parse a 'Npt' string (case-insensitive) into finite points."""
    if not isinstance(s, str):
        return None
    m = _PT_RE.match(s)
    if m is None:
        return None
    value = float(m.group(1))
    return value if math.isfinite(value) else None


def _point_twips(value: str, *, requirement: bool) -> Optional[int]:
    """Return exact twentieths of a point without rounding the saved value."""
    if not isinstance(value, str):
        return None
    match = _PT_RE.match(value)
    if match is None:
        return None
    try:
        points = Decimal(match.group(1))
        if (not points.is_finite() or points < 0 or
                (requirement and points > Decimal(str(_SPACING_PT_RANGE[1])))):
            return None
        numerator, denominator = points.as_integer_ratio()
        twips, remainder = divmod(numerator * 20, denominator)
        return twips if remainder == 0 else None
    except (InvalidOperation, ValueError, OverflowError):
        return None


def _line_spacing_encoded(value: str) -> Optional[int]:
    """Return exact 240ths for a saved decimal multiplier; never quantize it."""
    if not isinstance(value, str) or _RATIO_RE.fullmatch(value) is None:
        return None
    try:
        decimal = Decimal(value)
        if not decimal.is_finite() or not Decimal(1) <= decimal <= Decimal(4):
            return None
        numerator, denominator = decimal.as_integer_ratio()
        encoded, remainder = divmod(numerator * 240, denominator)
        return encoded if remainder == 0 else None
    except (InvalidOperation, ValueError, OverflowError):
        return None


def _in_bounds(v: Optional[float], lo: float, hi: float) -> bool:
    return v is not None and math.isfinite(v) and lo <= v <= hi


def _family_from_actual(s) -> Optional[str]:
    """Lowercase family token from a persisted FONT_CONSISTENCY actual_value.

    Exact-family profiles persist a bare family name; allowed-pair profiles
    persist a combo string like "Calibri 11pt". Strip the trailing
    ' <number>pt' and lowercase; anything else yields None (unverifiable).
    """
    if not isinstance(s, str):
        return None
    token = _TRAILING_SIZE_RE.sub("", s).strip().lower()
    return token or None




# ---------------------------------------------------------------------------
# Structural safety gate
# ---------------------------------------------------------------------------

def _unsafe_element(el) -> bool:
    """True when the element, any descendant, or any ancestor carries an
    unsafe tag (tracked changes, fields, hyperlinks, comments, sdt, text
    boxes, subdocs, math, AlternateContent)."""
    for node in el.iter():
        if node.tag in _UNSAFE_TAGS:
            return True
    for node in el.iterancestors():
        if node.tag in _UNSAFE_TAGS:
            return True
    return False


def _unsafe_paragraph(para) -> bool:
    return _unsafe_element(para._p)


def _unsafe_run(run) -> bool:
    return _unsafe_element(run._r)


def _require_unique_reference_heading(doc, paragraphs, index: int) -> None:
    """The same main-story identity gate used by Reference-heading alignment."""
    from app.services.reference_section import reference_section_bounds, normalise_reference_heading

    bounds = reference_section_bounds(paragraphs)
    recognized = [p["index"] for p in paragraphs if normalise_reference_heading(p["text"]) is not None]
    table_heading = any(
        normalise_reference_heading("".join(node.text or "" for node in cell_para.iter(_W + "t"))) is not None
        for table in doc.element.body.iter(_W + "tbl")
        for cell_para in table.iter(_W + "p")
    )
    if bounds is None or recognized != [index] or bounds.heading_index != index or table_heading:
        raise GuidedFixFailed("ambiguous Reference heading")


def _reference_heading_bold_run_indexes(doc, paragraphs, index: int) -> List[int]:
    """Enumerate complete physical text-run addresses, also on corrected output.

    No persisted Actual comparison or noncompliance decision belongs here.
    Direct booleans make local overrides provable; used style/default chains
    are inspected only to exclude numbering and unsupported script formatting.
    """
    from app.services.reference_section import NUMERIC_SECTION_PREFIX

    try:
        _require_unique_reference_heading(doc, paragraphs, index)
    except GuidedFixFailed as exc:
        raise ReferenceHeadingBoldUnavailable("ambiguous_heading") from exc
    para = doc.paragraphs[index]
    if para._p.getparent() is not doc.element.body:
        raise ReferenceHeadingBoldUnavailable("ambiguous_heading")
    pprs = para._p.findall(_W + "pPr")
    if (len(pprs) != 1 or para._p[0] is not pprs[0] or
            any(text.strip() for text in para._p.xpath("text()")) or NUMERIC_SECTION_PREFIX.match(para.text.strip())):
        raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
    paragraph_properties = {
        "pStyle", "keepNext", "keepLines", "pageBreakBefore", "widowControl",
        "suppressLineNumbers", "shd", "suppressAutoHyphens",
        "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
        "autoSpaceDN", "snapToGrid", "spacing", "ind", "contextualSpacing",
        "mirrorIndents", "suppressOverlap", "jc", "textDirection",
        "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    }
    run_properties = {
        "rStyle", "rFonts", "b", "i", "iCs", "caps", "smallCaps", "strike",
        "dstrike", "outline", "shadow", "emboss", "imprint", "noProof",
        "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
        "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
        "fitText", "vertAlign", "lang", "eastAsianLayout", "specVanish",
    }

    def check_properties(container, allowed):
        if container is None:
            return
        tags = [node.tag for node in container]
        if (container.attrib or len(tags) != len(set(tags)) or
                any(tag not in {_W + name for name in allowed} for tag in tags) or
                any(node.tag in {_W + "bCs", _W + "cs", _W + "rtl", _W + "bidi",
                                 _W + "numPr", _W + "sectPr", _W + "rPrChange",
                                 _W + "pPrChange"} for node in container.iter())):
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        # Normal lang/@bidi and rFonts/@csTheme are metadata, not script flags.
        for node in container:
            if node.tag == _W + "rPr":
                check_properties(node, run_properties - {"rStyle"})
            elif len(node) or (node.text or "").strip():
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            if node.tag == _W + "b" and (
                set(node.attrib) - {_W + "val"} or node.text or
                node.get(_W + "val") not in (None, "1", "0", "true", "false", "on", "off")
            ):
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")

    check_properties(pprs[0], paragraph_properties)
    styles = doc.styles.element
    style_nodes = {}
    for node in styles.findall(_W + "style"):
        style_nodes.setdefault(node.get(_W + "styleId"), []).append(node)

    def check_style(style_id, kind):
        if not style_id:
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        seen = set()
        while style_id is not None:
            matches = style_nodes.get(style_id, [])
            if style_id in seen or len(matches) != 1 or matches[0].get(_W + "type") != kind:
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            seen.add(style_id)
            style = matches[0]
            for tag, allowed in (("pPr", paragraph_properties), ("rPr", run_properties - {"rStyle"})):
                properties = style.findall(_W + tag)
                if len(properties) > 1:
                    raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
                if properties:
                    check_properties(properties[0], allowed)
            parents = style.findall(_W + "basedOn")
            if len(parents) > 1 or (parents and (
                not parents[0].get(_W + "val") or len(parents[0]) or
                set(parents[0].attrib) != {_W + "val"} or (parents[0].text or "").strip()
            )):
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            style_id = parents[0].get(_W + "val") if parents else None

    def used_style(container, tag, kind):
        refs = container.findall(_W + tag) if container is not None else []
        if len(refs) > 1:
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        if refs:
            style_id = refs[0].get(_W + "val")
            if not style_id or len(refs[0]) or set(refs[0].attrib) != {_W + "val"}:
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            check_style(style_id, kind)
        else:
            if any(node.get(_W + "type") == kind and
                   node.get(_W + "default") not in (None, "1", "0", "true", "false", "on", "off")
                   for node in styles.findall(_W + "style")):
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            defaults = [node for node in styles.findall(_W + "style")
                        if node.get(_W + "type") == kind and node.get(_W + "default") in ("1", "true", "on")]
            if len(defaults) > 1 or (kind == "paragraph" and not defaults):
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            if defaults:
                check_style(defaults[0].get(_W + "styleId"), kind)

    defaults = styles.findall(_W + "docDefaults")
    if len(defaults) > 1:
        raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
    if defaults:
        if (defaults[0].attrib or
                any(node.tag not in {_W + "pPrDefault", _W + "rPrDefault"} for node in defaults[0])):
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        for outer, inner, allowed in (
            ("pPrDefault", "pPr", paragraph_properties),
            ("rPrDefault", "rPr", run_properties - {"rStyle"}),
        ):
            wrappers = defaults[0].findall(_W + outer)
            if len(wrappers) > 1:
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            if wrappers:
                properties = wrappers[0].findall(_W + inner)
                if (len(properties) != 1 or len(wrappers[0]) != 1 or wrappers[0].attrib or
                        (wrappers[0].text or "").strip()):
                    raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
                if properties:
                    check_properties(properties[0], allowed)
    used_style(pprs[0], "pStyle", "paragraph")
    bookmarks = []
    seen_bookmarks = set()
    indexes = []
    run_index = 0
    for child in para._p:
        if child.tag == _W + "pPr":
            continue
        if child.tag in (_W + "bookmarkStart", _W + "bookmarkEnd"):
            identity = child.get(_W + "id")
            allowed = {_W + "id", _W + "name"} if child.tag == _W + "bookmarkStart" else {_W + "id"}
            if (not identity or not identity.isascii() or not identity.isdecimal() or
                    len(child) or child.text or set(child.attrib) != allowed):
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            if child.tag == _W + "bookmarkStart":
                if identity in seen_bookmarks or not child.get(_W + "name"):
                    raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
                seen_bookmarks.add(identity)
                bookmarks.append(identity)
                for tag in ("bookmarkStart", "bookmarkEnd"):
                    if sum(node.get(_W + "id") == identity for node in doc.element.body.iter(_W + tag)) != 1:
                        raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            elif not bookmarks or bookmarks.pop() != identity:
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            continue
        if child.tag != _W + "r" or _unsafe_element(child):
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        if (any(text.strip() for text in child.xpath("text()")) or
                any(node.tag not in (_W + "rPr", _W + "t") for node in child)):
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        rprs = child.findall(_W + "rPr")
        if (len(rprs) > 1 or (rprs and child[0] is not rprs[0]) or
                any(len(node) or set(node.attrib) - {"{http://www.w3.org/XML/1998/namespace}space"}
                    for node in child.findall(_W + "t"))):
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        rpr = rprs[0] if rprs else None
        check_properties(rpr, run_properties)
        used_style(rpr, "rStyle", "character")
        run = para.runs[run_index]
        if any(unicodedata.bidirectional(ch) in ("R", "AL", "AN") for ch in run.text):
            raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
        if run.text != "":
            bold = rpr.findall(_W + "b") if rpr is not None else []
            if (len(bold) != 1 or len(bold[0]) or bold[0].text or
                    set(bold[0].attrib) - {_W + "val"} or
                    bold[0].get(_W + "val") not in (None, "1", "0", "true", "false", "on", "off") or
                    type(run.font.bold) is not bool):
                raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
            indexes.append(run_index)
        run_index += 1
    if bookmarks or not indexes:
        raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
    return indexes


def _reference_heading_bold_targets(audit, doc, paragraphs, violation) -> Tuple[List[int], bool]:
    location = violation.location
    if (not isinstance(location, dict) or set(location) != {"paragraph_index"} or
            type(location["paragraph_index"]) is not int or
            not 0 <= location["paragraph_index"] < len(doc.paragraphs)):
        raise ReferenceHeadingBoldUnavailable("ambiguous_heading")
    index = location["paragraph_index"]
    indexes = _reference_heading_bold_run_indexes(doc, paragraphs, index)
    expected = _resolve_expected(violation, audit, doc, None, paragraphs[index])
    if type(expected) is not bool:
        raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
    states = [doc.paragraphs[index].runs[i].font.bold for i in indexes]
    if (any(value is not states[0] for value in states) or
            violation.actual_value != ("Bold" if states[0] else "Not bold") or
            states[0] is expected):
        raise ReferenceHeadingBoldUnavailable("unverifiable_bold")
    return indexes, expected


def _fix_reference_heading_bold(run, required: bool) -> bool:
    run.font.bold = required
    return type(run.font.bold) is bool and run.font.bold is required


# ---------------------------------------------------------------------------
# Expected-value resolution — never guess a font
# ---------------------------------------------------------------------------

def _font_half_points(value, *, requirement=True) -> Optional[int]:
    """Parse font points exactly; neither rounding nor detector tolerance is proof."""
    if not isinstance(value, str):
        return None
    match = _PT_RE.fullmatch(value.strip())
    if match is None:
        return None
    try:
        points = Decimal(match.group(1))
        numerator, denominator = points.as_integer_ratio()
        if (points <= 0 or (numerator * 2) % denominator or
                (requirement and not Decimal(str(_FONT_SIZE_PT_RANGE[0])) <= points <=
                 Decimal(str(_FONT_SIZE_PT_RANGE[1])))):
            return None
        return numerator * 2 // denominator
    except (InvalidOperation, ValueError, OverflowError):
        return None


def _resolve_font_expected(expected_value, audit, role, code) -> Optional[Any]:
    """Independently restore one saved, role-specific font requirement."""
    from app.services.profile_snapshot import snapshot_from_dict
    from app.services.profile_preset_adapter import EffectiveProfileConfig
    from app.services.role_eligibility import is_body_role, is_list_item_role

    levels = {"HEADING_1": 1, "HEADING_2": 2, "HEADING_3": 3, "APPENDIX_HEADING": 1}
    level = levels.get(role)
    if level is None and not (is_body_role(role) or is_list_item_role(role)):
        return None
    try:
        cfg = EffectiveProfileConfig(snapshot_from_dict(audit.profile_snapshot))
        combos = cfg.HEADING_ALLOWED_FONT_COMBOS if level else cfg.BODY_ALLOWED_FONT_COMBOS
        if code == "FONT_SIZE":
            saved = cfg.heading_expected_size(level) if level else cfg.FONT_SIZE_BODY
            saved_half_points = _font_half_points(f"{saved}pt")
            if saved_half_points is None:
                return None
            if combos:
                sizes = {_font_half_points(f"{size}pt") for _, size in combos}
                if sizes != {saved_half_points}:
                    return None
                persisted = ", ".join(f"{family} {size:g}pt" for family, size in combos)
                if expected_value != persisted:
                    return None
            elif _font_half_points(expected_value) != saved_half_points:
                return None
            return saved_half_points / 2
        family = cfg.HEADING_FONT_FAMILY if level else cfg.FONT_FAMILY
        if combos:
            families = {name for name, _ in combos}
            if len(families) != 1:
                return None
            family = next(iter(families))
            if expected_value != ", ".join(f"{name} {size:g}pt" for name, size in combos):
                return None
        elif not isinstance(expected_value, str) or expected_value.strip() != family:
            return None
        if (not isinstance(family, str) or not family or family != family.strip() or
                len(family) > 64 or family.isdigit() or
                any(ord(ch) < 32 or ord(ch) == 127 for ch in family)):
            return None
        return family
    except Exception:
        return None


def _resolve_line_spacing_expected(violation: Violation, audit, doc, role) -> Optional[float]:
    """Bind multiplier units to the immutable, role-specific saved requirement."""
    from app.services.document_parser import get_heading_level
    from app.services.profile_snapshot import snapshot_from_dict
    from app.services.reference_section import reference_section_bounds
    from app.services.role_eligibility import body_prose_paragraph

    try:
        snapshot = snapshot_from_dict(audit.profile_snapshot)
        location = violation.location
        index = location.get("paragraph_index") if isinstance(location, dict) else None
        if type(index) is not int or not 0 <= index < len(doc.paragraphs):
            return None
        para = doc.paragraphs[index]
        if role in ("BODY", "APPENDIX_BODY"):
            if (not body_prose_paragraph(role, para.text) or
                    get_heading_level(para.style.name) is not None):
                return None
            saved = snapshot.body_line_spacing
        elif role == "REFERENCE_ENTRY":
            if (snapshot.schema_version < 3 or
                    snapshot.references_check_line_spacing is not True):
                return None
            paragraphs = extract_paragraphs(doc)
            bounds = reference_section_bounds(paragraphs)
            if bounds is None or index not in bounds.entry_indexes:
                return None
            _require_unique_reference_heading(doc, paragraphs, bounds.heading_index)
            saved = snapshot.references_line_spacing
        else:
            return None
        if (type(saved) not in (int, float) or not _in_bounds(saved, *_LINE_SPACING_RANGE) or
                violation.expected_value != str(saved) or
                _line_spacing_encoded(violation.expected_value) is None):
            return None
        return float(saved)
    except Exception:
        return None


def _resolve_point_spacing_expected(violation, audit, doc, role, para_dict) -> Optional[int]:
    """Resolve the detector's point requirement from the immutable snapshot."""
    from app.services.layout_engine import _heading_style_level, _is_caption_paragraph, _is_list_item
    from app.services.profile_preset_adapter import EffectiveProfileConfig
    from app.services.profile_snapshot import snapshot_from_dict
    from app.services.role_eligibility import (
        body_prose_paragraph, is_caption_role, is_reference_role, typography_skipped,
    )

    try:
        location = violation.location
        index = location.get("paragraph_index") if isinstance(location, dict) else None
        if (type(index) is not int or not 0 <= index < len(doc.paragraphs) or
                not isinstance(para_dict, dict) or para_dict.get("index") != index):
            return None
        side = "BEFORE" if violation.rule_code == "SPACE_BEFORE" else "AFTER"
        para = doc.paragraphs[index]
        config = EffectiveProfileConfig(snapshot_from_dict(audit.profile_snapshot))
        if role is not None and (typography_skipped(role) or role == "REFERENCES_HEADING"):
            return None
        if is_caption_role(role):
            saved = getattr(config, "CAPTION_SPACE_" + side)
        else:
            if role is not None and (is_reference_role(role) or
                                     not body_prose_paragraph(role, para_dict.get("text") or "")):
                return None
            if _is_list_item(para_dict, para):
                saved = config.LIST_SPACE_AFTER if side == "AFTER" else None
            elif _is_caption_paragraph(para_dict, para, config):
                saved = getattr(config, "CAPTION_SPACE_" + side)
            else:
                level = _heading_style_level(para_dict.get("style_name", ""), role)
                saved = getattr(config, "SPACE_" + side + ("_HEADING" if level is not None else "_BODY"))
        if type(saved) not in (int, float):
            return None
        saved_decimal = Decimal(str(saved))
        expected = violation.expected_value
        twips = _point_twips(expected, requirement=True)
        if (not saved_decimal.is_finite() or
                not Decimal(str(_SPACING_PT_RANGE[0])) <= saved_decimal <= Decimal(str(_SPACING_PT_RANGE[1])) or
                twips is None or saved_decimal != Decimal(_PT_RE.match(expected).group(1))):
            return None
        return twips
    except Exception:
        return None


def _resolve_expected(
    violation: Violation, audit, doc, role, para_dict
) -> Optional[Any]:
    """Resolve the persisted expected_value for a whitelisted rule, or None
    (= skip `expected_unresolved`)."""
    rule_code = violation.rule_code
    expected_value = violation.expected_value

    if rule_code == "REFERENCE_HEADING_BOLD":
        from app.services.profile_snapshot import snapshot_from_dict
        from app.services.profile_preset_adapter import EffectiveProfileConfig
        try:
            config = EffectiveProfileConfig(snapshot_from_dict(audit.profile_snapshot))
            required = config.REFERENCES_HEADING_BOLD
            if (config.REFERENCES_CHECK_HEADING_BOLD is True and type(required) is bool and
                    expected_value == ("Bold" if required else "Not bold")):
                return required
        except Exception:
            return None
        return None
    if rule_code in ("FONT_SIZE", "FONT_CONSISTENCY"):
        return _resolve_font_expected(expected_value, audit, role, rule_code)
    if rule_code == "LINE_SPACING":
        return _resolve_line_spacing_expected(violation, audit, doc, role)
    if rule_code in ("SPACE_BEFORE", "SPACE_AFTER"):
        return _resolve_point_spacing_expected(violation, audit, doc, role, para_dict)
    if rule_code == "ALIGNMENT":
        if expected_value in _ALIGNMENTS:
            return expected_value
        return None
    if rule_code == "REFERENCE_HEADING_ALIGNMENT":
        from app.services.profile_snapshot import snapshot_from_dict
        from app.services.profile_preset_adapter import EffectiveProfileConfig
        try:
            config = EffectiveProfileConfig(snapshot_from_dict(audit.profile_snapshot))
            required = config.REFERENCES_HEADING_ALIGNMENT
            if (config.REFERENCES_CHECK_HEADING_ALIGNMENT and
                    isinstance(required, str) and required.lower() in _ALIGNMENTS and
                    isinstance(expected_value, str) and expected_value.strip().lower() == required.lower()):
                return required.lower()
        except Exception:
            return None
        return None
    return None


# ---------------------------------------------------------------------------
# Fixers — python-docx direct formatting only; verify after each
# ---------------------------------------------------------------------------

def _fix_font_size(run, pt: float) -> bool:
    from docx.shared import Pt
    run.font.size = Pt(pt)
    try:
        actual = run.font.size
        if actual is None:
            return False
        return actual.pt * 2 == _font_half_points(f"{pt}pt")
    except Exception:
        return False


def _fix_font_family(run, family: str) -> bool:
    from docx.oxml.ns import qn
    run.font.name = family  # writes w:ascii + w:hAnsi
    # Only these three literal slots are authorized; no alternate slot is cleared.
    run._element.rPr.rFonts.set(qn("w:eastAsia"), family)
    rfonts = run._element.rPr.rFonts
    return (
        rfonts.get(qn("w:ascii")) == family
        and rfonts.get(qn("w:hAnsi")) == family
        and rfonts.get(qn("w:eastAsia")) == family
    )


def _fix_alignment(para, name: str) -> bool:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    para.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[name]
    return para.alignment == {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[name]


def _fix_line_spacing(para, ratio: float) -> bool:
    encoded = _line_spacing_encoded(str(ratio))
    if encoded is None:
        return False
    _direct_line_spacing(para)
    spacing = para._p.pPr.find(_W + "spacing")
    spacing.set(_W + "line", str(encoded))
    spacing.set(_W + "lineRule", "auto")
    return _direct_line_spacing(para) == encoded


def _fix_space_before(para, twips: int) -> bool:
    if type(twips) is not int or twips < 0:
        return False
    spacing = para._p.pPr.find(_W + "spacing")
    value = str(twips)
    spacing.set(_W + "before", value)
    return spacing.get(_W + "before") == value


def _fix_space_after(para, twips: int) -> bool:
    if type(twips) is not int or twips < 0:
        return False
    spacing = para._p.pPr.find(_W + "spacing")
    value = str(twips)
    spacing.set(_W + "after", value)
    return spacing.get(_W + "after") == value


def _spacing_property_layers(para) -> List[Any]:
    """Read only the used paragraph-style chain; never create missing parts."""
    def fail():
        raise GuidedFixFailed("unverifiable paragraph spacing")

    def leaf(node, attributes):
        if (len(node) or (node.text or "").strip() or
                set(node.attrib) - attributes):
            fail()

    def properties(container):
        nodes = container.findall(_W + "pPr")
        if len(nodes) > 1:
            fail()
        if not nodes:
            return None
        node = nodes[0]
        if (node.attrib or any((child.tail or "").strip() for child in node) or
                (node.text or "").strip() or
                any(child.tag in {_W + "numPr", _W + "pPrChange", _MC + "AlternateContent",
                                  _W + "ins", _W + "del"} for child in node.iter())):
            fail()
        return node

    def reference(container, tag):
        nodes = container.findall(_W + tag) if container is not None else []
        if len(nodes) > 1:
            fail()
        if not nodes:
            return None
        leaf(nodes[0], {_W + "val"})
        value = nodes[0].get(_W + "val")
        if not value or value != value.strip():
            fail()
        return value

    try:
        part = para.part
        if (part.element.tag != _W + "document" or
                para._p.getparent() is not part.element.find(_W + "body")):
            fail()
        direct = properties(para._p)
        if direct is None or para._p[0] is not direct:
            fail()
        relationships = [rel for rel in part.rels.values() if rel.reltype.endswith("/styles")]
        if len(relationships) != 1 or relationships[0].is_external:
            fail()
        styles = relationships[0].target_part.element
        if styles.tag != _W + "styles":
            fail()
        style_nodes = {}
        for node in styles.findall(_W + "style"):
            style_nodes.setdefault(node.get(_W + "styleId"), []).append(node)
        style_id = reference(direct, "pStyle")
        if style_id is None:
            candidates = []
            for node in styles.findall(_W + "style"):
                if node.get(_W + "type") != "paragraph":
                    continue
                default = node.get(_W + "default")
                if default not in (None, "1", "0", "true", "false", "on", "off"):
                    fail()
                if default in ("1", "true", "on"):
                    candidates.append(node)
            if len(candidates) != 1:
                fail()
            style_id = candidates[0].get(_W + "styleId")
            if not style_id:
                fail()
        layers = [direct]
        seen = set()
        while style_id is not None:
            matches = style_nodes.get(style_id, [])
            if (not style_id or style_id != style_id.strip() or style_id in seen or
                    len(matches) != 1 or matches[0].get(_W + "type") != "paragraph"):
                fail()
            seen.add(style_id)
            style = matches[0]
            if any(node.tag in {_MC + "AlternateContent", _W + "pPrChange"}
                   for node in style.iter()):
                fail()
            props = properties(style)
            if props is not None:
                if reference(props, "pStyle") is not None:
                    fail()
                layers.append(props)
            style_id = reference(style, "basedOn")
        defaults = styles.findall(_W + "docDefaults")
        if len(defaults) > 1:
            fail()
        if defaults:
            default = defaults[0]
            tags = [node.tag for node in default]
            if (default.attrib or (default.text or "").strip() or
                    len(tags) != len(set(tags)) or
                    any(tag not in {_W + "pPrDefault", _W + "rPrDefault"} for tag in tags) or
                    any((node.tail or "").strip() for node in default)):
                fail()
            wrappers = default.findall(_W + "pPrDefault")
            if wrappers:
                wrapper = wrappers[0]
                if (wrapper.attrib or (wrapper.text or "").strip() or
                        any(node.tag != _W + "pPr" or (node.tail or "").strip() for node in wrapper)):
                    fail()
                props = properties(wrapper)
                if props is not None:
                    if reference(props, "pStyle") is not None:
                        fail()
                    layers.append(props)
        return layers
    except GuidedFixFailed:
        raise
    except Exception as exc:
        raise GuidedFixFailed("unverifiable paragraph spacing") from exc


_LINE_PPR_LEAVES = frozenset(_W + name for name in (
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "widowControl",
    "suppressLineNumbers", "shd", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "spacing",
    "ind", "mirrorIndents", "suppressOverlap", "jc", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
))


def _direct_line_spacing(para) -> int:
    """Certify a plain, explicit auto multiplier without interpreting inheritance."""
    def fail():
        raise GuidedFixFailed("unverifiable line spacing")

    try:
        layers = _spacing_property_layers(para)
        paragraph = para._p
        if (not para.text.strip() or _unsafe_paragraph(para) or
                any(text.strip() for text in paragraph.xpath("text()")) or
                any((child.tail or "").strip() for child in paragraph) or
                any(child.tag != _W + "r" for child in list(paragraph)[1:]) or
                any(node.tag == _W + "fldSimple" or
                    (isinstance(node.tag, str) and node.tag.startswith(_W) and
                     node.tag.endswith("PrChange")) for node in paragraph.iter()) or
                next(para.part.element.iter(_W + "docGrid"), None) is not None):
            fail()
        for run in para.runs:
            children = list(run._r)
            properties = run._r.findall(_W + "rPr")
            if (len(properties) > 1 or
                    (properties and children[0] is not properties[0]) or
                    any(text.strip() for text in run._r.xpath("text()")) or
                    any((child.tail or "").strip() for child in children)):
                fail()
            for child in children:
                if child.tag == _W + "rPr":
                    if any(text.strip() for text in child.xpath(".//text()")):
                        fail()
                elif (child.tag not in {_W + "t", _W + "tab", _W + "br", _W + "cr"} or
                      len(child) or
                      (child.tag != _W + "t" and (child.text or "").strip())):
                    fail()
        # The raw style graph has already been validated above. Linked styles
        # remain unsupported even when their spacing would be shadowed.
        style = para.style
        while style is not None:
            if next(style.element.iter(_W + "link"), None) is not None:
                fail()
            style = style.base_style
        for layer_index, properties in enumerate(layers):
            tags = [child.tag for child in properties]
            if (len(tags) != len(set(tags)) or
                    any(child.tag not in _LINE_PPR_LEAVES or len(child) or
                        (child.text or "").strip() or (child.tail or "").strip()
                        for child in properties)):
                fail()
            spacing = properties.find(_W + "spacing")
            if spacing is None:
                if layer_index == 0:
                    fail()
                continue
            allowed = {_W + "before", _W + "after"}
            if layer_index == 0:
                allowed |= {_W + "line", _W + "lineRule"}
            if set(spacing.attrib) - allowed:
                fail()
            for name in ("before", "after"):
                value = spacing.get(_W + name)
                if value is not None and re.fullmatch(r"[0-9]+", value) is None:
                    fail()
        direct = layers[0].find(_W + "spacing")
        raw = direct.get(_W + "line")
        if (raw is None or re.fullmatch(r"[1-9][0-9]*", raw) is None or
                len(raw) > 10 or int(raw) > 2147483647 or
                direct.get(_W + "lineRule") != "auto"):
            fail()
        encoded = int(raw)
        current = para.paragraph_format.line_spacing
        if type(current) is not float or not math.isfinite(current) or current != encoded / 240:
            fail()
        return encoded
    except GuidedFixFailed:
        raise
    except Exception as exc:
        raise GuidedFixFailed("unverifiable line spacing") from exc


def _preflight_line_spacing_targets(doc, violations) -> None:
    """Reject selected malformed raw values before the scalar extraction pass."""
    paragraphs = doc.paragraphs
    for violation in violations:
        if violation.rule_code != "LINE_SPACING":
            continue
        location = violation.location
        index = location.get("paragraph_index") if isinstance(location, dict) else None
        if type(index) is not int or not 0 <= index < len(paragraphs):
            raise GuidedFixFailed("invalid paragraph")
        _direct_line_spacing(paragraphs[index])


def _direct_spacing_points(para, rule_code: str) -> float:
    """Certify a direct point setting, not a rendered inter-paragraph gap."""
    if rule_code not in ("SPACE_BEFORE", "SPACE_AFTER"):
        raise GuidedFixFailed("unverifiable paragraph spacing")
    side = "before" if rule_code == "SPACE_BEFORE" else "after"
    layers = _spacing_property_layers(para)
    known = {_W + name for name in (
        "before", "after", "beforeLines", "afterLines",
        "beforeAutospacing", "afterAutospacing", "line", "lineRule",
    )}
    booleans = {"1": True, "true": True, "on": True,
                "0": False, "false": False, "off": False}
    automatic = contextual = None
    direct_twips = None
    for index, layer in enumerate(layers):
        spacing = layer.findall(_W + "spacing")
        context = layer.findall(_W + "contextualSpacing")
        if len(spacing) > 1 or len(context) > 1:
            raise GuidedFixFailed("unverifiable paragraph spacing")
        if spacing:
            node = spacing[0]
            if len(node) or (node.text or "").strip() or set(node.attrib) - known:
                raise GuidedFixFailed("unverifiable paragraph spacing")
            for name, value in node.attrib.items():
                local = name[len(_W):]
                if local.endswith("Autospacing"):
                    valid = value in booleans
                elif local == "lineRule":
                    valid = value in ("auto", "exact", "atLeast")
                else:
                    valid = re.fullmatch(r"[0-9]+", value) is not None
                if not valid:
                    raise GuidedFixFailed("unverifiable paragraph spacing")
            if node.get(_W + side + "Lines") is not None:
                raise GuidedFixFailed("unverifiable paragraph spacing")
            flag = node.get(_W + side + "Autospacing")
            if flag is not None and automatic is None:
                automatic = booleans[flag]
            if index == 0:
                value = node.get(_W + side)
                if value is not None:
                    try:
                        direct_twips = int(value)
                    except ValueError as exc:
                        raise GuidedFixFailed("unverifiable paragraph spacing") from exc
        if context:
            node = context[0]
            value = node.get(_W + "val", "true")
            if (len(node) or (node.text or "").strip() or
                    set(node.attrib) - {_W + "val"} or value not in booleans):
                raise GuidedFixFailed("unverifiable paragraph spacing")
            if contextual is None:
                contextual = booleans[value]
    if direct_twips is None or automatic is True or contextual is True:
        raise GuidedFixFailed("unverifiable paragraph spacing")
    try:
        points = direct_twips / 20
        if not math.isfinite(points):
            raise ValueError("nonfinite spacing")
        return points
    except (ValueError, OverflowError) as exc:
        raise GuidedFixFailed("unverifiable paragraph spacing") from exc


def _direct_spacing_twips(para, rule_code: str) -> int:
    """Prove the selected direct setting and complete point-target structure."""
    _direct_spacing_points(para, rule_code)
    paragraph = para._p
    allowed = {_W + "pPr", _W + "r", _W + "bookmarkStart", _W + "bookmarkEnd"}
    if (any(child.tag not in allowed for child in paragraph) or
            next(paragraph.iter(_W + "fldSimple"), None) is not None or
            paragraph.pPr.find(_W + "sectPr") is not None):
        raise GuidedFixFailed("unverifiable paragraph spacing")
    starts, ends = set(), set()
    for child in paragraph:
        if child.tag not in (_W + "bookmarkStart", _W + "bookmarkEnd"):
            continue
        raw = child.get(_W + "id")
        if raw is None or re.fullmatch(r"[0-9]+", raw) is None:
            raise GuidedFixFailed("unverifiable paragraph spacing")
        try:
            bookmark = int(raw)
        except ValueError as exc:
            raise GuidedFixFailed("unverifiable paragraph spacing") from exc
        if child.tag == _W + "bookmarkStart":
            if bookmark in starts:
                raise GuidedFixFailed("unverifiable paragraph spacing")
            starts.add(bookmark)
        else:
            if bookmark in ends or bookmark not in starts:
                raise GuidedFixFailed("unverifiable paragraph spacing")
            ends.add(bookmark)
    if starts != ends:
        raise GuidedFixFailed("unverifiable paragraph spacing")
    side = "before" if rule_code == "SPACE_BEFORE" else "after"
    return int(paragraph.pPr.find(_W + "spacing").get(_W + side))


# ---------------------------------------------------------------------------
# Current-value comparators (point and font units are exact)
# ---------------------------------------------------------------------------

def _align_to_name(align) -> str:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    return {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }.get(align, "unknown")


def _current_matches_expected(rule_code: str, expected: Any, para_dict: Dict, run_dict: Optional[Dict], para=None) -> bool:
    """True when the recomputed current value already equals the resolved
    expected value, using the rule's own tolerances. Paragraph rules read
    the paragraph object directly when provided — `extract_paragraphs`
    maps explicit zero lengths to None, so the dict view is unreliable for
    SPACE_* when the expected value is 0pt."""
    if rule_code == "REFERENCE_HEADING_BOLD":
        current = run_dict.get("bold") if run_dict else None
        return type(current) is bool and current is expected
    if rule_code in SAFE_RUN_RULES:
        if run_dict is None:
            return False
        if rule_code == "FONT_SIZE":
            current = run_dict.get("effective_font_size")
            return current is not None and _font_half_points(
                f"{current}pt", requirement=False,
            ) == _font_half_points(f"{expected}pt")
        # FONT_CONSISTENCY
        current = run_dict.get("effective_font_name")
        return current is not None and current.lower() == expected.lower()
    if rule_code == "REFERENCE_HEADING_ALIGNMENT":
        from app.services.formatting_evidence import resolve_alignment
        return para is not None and resolve_alignment(para).alignment == expected
    if rule_code in ("SPACE_BEFORE", "SPACE_AFTER"):
        return abs(_direct_spacing_twips(para, rule_code) - expected) <= 20
    if rule_code == "LINE_SPACING":
        return (para is not None and
                _direct_line_spacing(para) == _line_spacing_encoded(str(expected)))
    if para is not None:
        if rule_code == "ALIGNMENT":
            return _align_to_name(para.alignment) == expected
    if rule_code == "ALIGNMENT":
        return _align_to_name(para_dict.get("alignment")) == expected
    return False


def _current_changed_from_actual(rule_code: str, para_dict: Dict, run_dict: Optional[Dict], actual_value, para=None) -> bool:
    """True when the recomputed current value does NOT match the persisted
    actual_value (same extraction + tolerances). An actual_value that fails
    its own parser is treated as changed — an unverifiable current-vs-actual
    comparison is never applied."""
    if rule_code == "REFERENCE_HEADING_BOLD":
        current = run_dict.get("bold") if run_dict else None
        return type(current) is not bool or actual_value != ("Bold" if current else "Not bold")
    if rule_code == "FONT_SIZE":
        token = actual_value.rsplit(" ", 1)[-1] if isinstance(actual_value, str) else actual_value
        actual_half_points = _font_half_points(token, requirement=False)
        current = run_dict.get("effective_font_size") if run_dict else None
        return current is None or actual_half_points != _font_half_points(f"{current}pt", requirement=False)
    if rule_code == "FONT_CONSISTENCY":
        actual_family = _family_from_actual(actual_value)
        if actual_family is None:
            return True
        current = run_dict.get("effective_font_name") if run_dict else None
        return current is None or current.lower() != actual_family
    if rule_code == "LINE_SPACING":
        return para is None or actual_value != str(_direct_line_spacing(para) / 240)
    if rule_code in ("SPACE_BEFORE", "SPACE_AFTER"):
        actual_twips = _point_twips(actual_value, requirement=False)
        return actual_twips is None or _direct_spacing_twips(para, rule_code) != actual_twips
    if rule_code == "REFERENCE_HEADING_ALIGNMENT":
        from app.services.formatting_evidence import resolve_alignment
        return (para is None or not isinstance(actual_value, str) or
                actual_value.strip().lower() not in _ALIGNMENTS or
                resolve_alignment(para).alignment != actual_value.strip().lower())
    if rule_code == "ALIGNMENT":
        if actual_value not in _ALIGNMENTS:
            return True
        return _align_to_name(para_dict.get("alignment")) != actual_value
    return True


# ---------------------------------------------------------------------------
# Filename / header helpers
# ---------------------------------------------------------------------------

_FILENAME_UNSAFE_RE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def build_corrected_filename(original_name: str) -> str:
    """Sanitized download name, never a path."""
    stem = (original_name or "").rsplit(".", 1)[0] if "." in (original_name or "") else (original_name or "")
    stem = _FILENAME_UNSAFE_RE.sub("_", stem)
    stem = re.sub(r"_{3,}", "_", stem).strip(" .")[:150]
    return f"{stem or 'document'}_ACA-fixed.docx"


def _content_disposition(filename: str) -> str:
    """ASCII-safe plain form + RFC 5987 filename* only when they differ."""
    ascii_form = filename.encode("ascii", errors="replace").decode("ascii").replace("?", "_")
    header = f'attachment; filename="{ascii_form}"'
    encoded = urllib.parse.quote(filename, safe="")
    if encoded != ascii_form:
        header += f"; filename*=UTF-8''{encoded}"
    return header

# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def _font_styles(part):
    """Check the original relationship before library getters can create defaults."""
    relationships = [rel for rel in part.rels.values() if rel.reltype.endswith("/styles")]
    if len(relationships) != 1 or relationships[0].is_external:
        raise GuidedFixFailed("unresolved alternate font slot")
    try:
        styles = relationships[0].target_part.element
        if styles.tag == _W + "styles":
            return styles
    except (AttributeError, KeyError, ValueError) as exc:
        raise GuidedFixFailed("unresolved alternate font slot") from exc
    raise GuidedFixFailed("unresolved alternate font slot")


def _font_property_layers(para, run) -> List[Any]:
    """Validate only reachable styles/defaults, without resolving shadowed fonts."""
    def fail():
        raise GuidedFixFailed("unresolved alternate font slot")

    def leaf(node, attrs):
        if len(node) or (node.text or "").strip() or set(node.attrib) - attrs:
            fail()

    def properties(container, tag):
        nodes = container.findall(_W + tag)
        if len(nodes) > 1:
            fail()
        if not nodes:
            return None
        node = nodes[0]
        tags = [child.tag for child in node]
        if (node.attrib or (node.text or "").strip() or len(tags) != len(set(tags)) or
                any((child.tail or "").strip() for child in node) or
                any(child.tag in {_W + "rPrChange", _W + "pPrChange", _W + "ins",
                                  _W + "del", _MC + "AlternateContent"} for child in node.iter())):
            fail()
        if tag == "pPr" and node.find(_W + "bidi") is not None:
            fail()
        if tag == "rPr":
            allowed = {
                _W + name for name in (
                    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
                    "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
                    "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
                    "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
                    "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
                    "eastAsianLayout", "specVanish", "oMath", "bidi", "script",
                )
            }
            if any(child.tag not in allowed or len(child) or (child.text or "").strip()
                   for child in node):
                fail()
        elif any(descendant.tag in {_W + "rFonts", _W + "sz", _W + "szCs", _W + "cs",
                                    _W + "rtl", _W + "script", _W + "bidi", _W + "rPr"}
                 for child in node for descendant in child.iterdescendants()):
            fail()
        return node

    def reference(container, tag):
        nodes = container.findall(_W + tag) if container is not None else []
        if len(nodes) > 1:
            fail()
        if not nodes:
            return None
        leaf(nodes[0], {_W + "val"})
        value = nodes[0].get(_W + "val")
        if not value or value != value.strip():
            fail()
        return value

    try:
        part = para.part
        if (part.element.tag != _W + "document" or
                para._p.getparent() is not part.element.find(_W + "body") or
                run._r.getparent() is not para._p or
                any(child.tag not in {_W + "pPr", _W + "r", _W + "bookmarkStart",
                                      _W + "bookmarkEnd"} for child in para._p) or
                any(child.tag not in {_W + "rPr", _W + "t", _W + "tab", _W + "br",
                                      _W + "cr"} for child in run._r)):
            fail()
        if (any(text.strip() for text in para._p.xpath("text()")) or
                any(text.strip() for text in run._r.xpath("text()")) or
                any((child.tail or "").strip() for child in para._p) or
                any((child.tail or "").strip() or len(child)
                    for child in run._r if child.tag != _W + "rPr")):
            fail()
        direct = properties(run._r, "rPr")
        if direct is None or run._r[0] is not direct:
            fail()
        paragraph_props = properties(para._p, "pPr")
        if paragraph_props is not None and para._p[0] is not paragraph_props:
            fail()
        styles = _font_styles(part)
        style_nodes = {}
        for node in styles.findall(_W + "style"):
            style_nodes.setdefault(node.get(_W + "styleId"), []).append(node)

        def default_style(kind):
            candidates = []
            for node in styles.findall(_W + "style"):
                if node.get(_W + "type") != kind:
                    continue
                value = node.get(_W + "default")
                if value not in (None, "1", "0", "true", "false", "on", "off"):
                    fail()
                if value in ("1", "true", "on"):
                    candidates.append(node)
            if len(candidates) != 1 or not candidates[0].get(_W + "styleId"):
                fail()
            return candidates[0].get(_W + "styleId")

        layers = [direct]

        def chain(style_id, kind):
            seen = set()
            while style_id is not None:
                matches = style_nodes.get(style_id, [])
                if (style_id in seen or len(matches) != 1 or
                        matches[0].get(_W + "type") != kind):
                    fail()
                seen.add(style_id)
                style = matches[0]
                tags = [node.tag for node in style]
                if (len(tags) != len(set(tags)) or (style.text or "").strip() or
                        any((node.tail or "").strip() for node in style) or
                        any(node.tag not in {_W + "rPr", _W + "pPr"} and
                            (len(node) or (node.text or "").strip()) for node in style) or
                        any(node.tag not in {_W + name for name in (
                            "name", "aliases", "basedOn", "next", "link", "autoRedefine",
                            "hidden", "uiPriority", "semiHidden", "unhideWhenUsed", "qFormat",
                            "locked", "personal", "personalCompose", "personalReply", "rsid",
                            "rPr", "pPr",
                        )} for node in style)):
                    fail()
                if (style.findall(_W + "link") or
                        any(node.tag in {_MC + "AlternateContent", _W + "rPrChange",
                                         _W + "pPrChange"} for node in style.iter())):
                    fail()
                props = properties(style, "rPr")
                if props is not None:
                    if reference(props, "rStyle") is not None:
                        fail()
                    layers.append(props)
                ppr = properties(style, "pPr")
                if ppr is not None and (kind == "character" or reference(ppr, "pStyle") is not None or
                                        ppr.find(_W + "rPr") is not None):
                    fail()
                style_id = reference(style, "basedOn")

        character = reference(direct, "rStyle")
        chain(character if character is not None else default_style("character"), "character")
        # The default character style is a used fallback even with an explicit style.
        if character is not None:
            chain(default_style("character"), "character")
        paragraph = reference(paragraph_props, "pStyle")
        chain(paragraph if paragraph is not None else default_style("paragraph"), "paragraph")
        if paragraph_props is not None and paragraph_props.find(_W + "rPr") is not None:
            fail()
        defaults = styles.findall(_W + "docDefaults")
        if len(defaults) > 1:
            fail()
        if defaults:
            default = defaults[0]
            tags = [child.tag for child in default]
            if (default.attrib or (default.text or "").strip() or len(tags) != len(set(tags)) or
                    any(tag not in {_W + "rPrDefault", _W + "pPrDefault"} for tag in tags) or
                    any((child.tail or "").strip() for child in default)):
                fail()
            for tag, prop in (("rPrDefault", "rPr"), ("pPrDefault", "pPr")):
                wrappers = default.findall(_W + tag)
                if not wrappers:
                    continue
                wrapper = wrappers[0]
                if (wrapper.attrib or (wrapper.text or "").strip() or
                        any(child.tag != _W + prop or (child.tail or "").strip() for child in wrapper)):
                    fail()
                props = properties(wrapper, prop)
                if props is not None:
                    if prop == "rPr":
                        if reference(props, "rStyle") is not None:
                            fail()
                        layers.append(props)
                    elif reference(props, "pStyle") is not None or props.find(_W + "rPr") is not None:
                        fail()
        return layers
    except GuidedFixFailed:
        raise
    except Exception as exc:
        raise GuidedFixFailed("unresolved alternate font slot") from exc


def _unsupported_font_slots(run, para, code: str) -> bool:
    """One strict direct-slot predicate shared by source and reopened output."""
    if not run.text or any(not (32 <= ord(ch) <= 126 or ch in "\t\r\n") for ch in run.text):
        return True
    try:
        layers = _font_property_layers(para, run)
        for index, rpr in enumerate(layers):
            if any(rpr.find(_W + tag) is not None for tag in (
                "cs", "rtl", "bidi", "script", "szCs", "eastAsianLayout", "oMath",
            )):
                return True
            language = rpr.find(_W + "lang")
            if language is not None:
                value = language.get(_W + "val")
                if (set(language.attrib) != {_W + "val"} or not value or value != value.strip() or
                        any(ord(ch) < 32 or ord(ch) == 127 for ch in value)):
                    return True
            fonts = rpr.findall(_W + "rFonts")
            sizes = rpr.findall(_W + "sz")
            if index:
                if fonts or (code == "FONT_SIZE" and sizes):
                    return True
                continue
            if len(fonts) != 1:
                return True
            font = fonts[0]
            ascii_name = font.get(_W + "ascii")
            if (len(font) or (font.text or "").strip() or
                    set(font.attrib) - {_W + "ascii", _W + "hAnsi", _W + "eastAsia"} or
                    not ascii_name or ascii_name != ascii_name.strip() or
                    any(ord(ch) < 32 or ord(ch) == 127 for ch in ascii_name) or
                    font.get(_W + "hAnsi") != ascii_name or
                    font.get(_W + "eastAsia", ascii_name) != ascii_name):
                return True
            if sizes or code == "FONT_SIZE":
                if len(sizes) != 1:
                    return True
                size = sizes[0]
                raw = size.get(_W + "val", "")
                if (len(size) or (size.text or "").strip() or set(size.attrib) != {_W + "val"} or
                        re.fullmatch(r"[0-9]+", raw) is None or int(raw) <= 0):
                    return True
        return False
    except (GuidedFixFailed, ValueError):
        return True


def _preflight_font_targets(doc, violations):
    """Check selected raw font evidence before scalar getters can parse it."""
    font_violations = [v for v in violations if v.rule_code in ("FONT_SIZE", "FONT_CONSISTENCY")]
    if not font_violations:
        return
    _font_styles(doc.part)
    paragraphs = doc.paragraphs
    for violation in font_violations:
        location = violation.location
        if not isinstance(location, dict):
            raise GuidedFixFailed("missing location")
        index = location.get("paragraph_index")
        if type(index) is not int or not 0 <= index < len(paragraphs):
            raise GuidedFixFailed("invalid paragraph")
        paragraph = paragraphs[index]
        runs = paragraph.runs
        raw = location.get("run_indexes", location.get("run_index"))
        indexes = raw if isinstance(raw, list) else [raw]
        if (not indexes or any(type(i) is not int or not 0 <= i < len(runs) for i in indexes) or
                len(set(indexes)) != len(indexes)):
            raise GuidedFixFailed("invalid run targets")
        if any(_unsupported_font_slots(runs[i], paragraph, violation.rule_code) for i in indexes):
            raise GuidedFixFailed("unresolved alternate font slot")


def review_targets(audit: AuditRecord, source_bytes: bytes, violations: List[Violation]):
    """Resolve the complete persisted targets before any mutation.

    A review does not certify the edit: apply repeats these checks against
    the exact source and later verifies the saved package and rule findings.
    """
    from app.services.role_classifier import classify_paragraphs

    if any(v.rule_code == "REFERENCE_HEADING_BOLD" for v in violations):
        from app.services.docx_preservation import (
            validate_reference_heading_bold_package, ReferenceHeadingBoldPackageUnsupported,
        )
        try:
            validate_reference_heading_bold_package(source_bytes)
        except ReferenceHeadingBoldPackageUnsupported as exc:
            raise ReferenceHeadingBoldUnavailable("unsupported_package") from exc
        except ValueError as exc:
            raise GuidedFixFailed("unverifiable DOCX package") from exc
    doc = parse_document(source_bytes)
    _preflight_font_targets(doc, violations)
    _preflight_line_spacing_targets(doc, violations)
    paragraphs = extract_paragraphs(doc)
    roles = classify_paragraphs(doc, paragraphs)
    if len(paragraphs) != len(doc.paragraphs):
        raise GuidedFixFailed("paragraph traversal is ambiguous")
    items = []
    targets = set()
    property_targets = set()
    for violation in violations:
        code = violation.rule_code
        if code not in SAFE_RUN_RULES | SAFE_PARA_RULES:
            raise GuidedFixFailed("unsupported rule")
        if code == "REFERENCE_HEADING_BOLD":
            bold_indexes, bold_expected = _reference_heading_bold_targets(audit, doc, paragraphs, violation)
        loc = violation.location
        if not isinstance(loc, dict):
            raise GuidedFixFailed("missing location")
        if code == "REFERENCE_HEADING_ALIGNMENT" and set(loc) != {"paragraph_index"}:
            raise GuidedFixFailed("ambiguous heading location")
        index = loc.get("paragraph_index")
        if type(index) is not int or not 0 <= index < len(paragraphs):
            raise GuidedFixFailed("invalid paragraph")
        if "paragraph_indexes" in loc or "occurrences" in loc:
            indexes = loc.get("paragraph_indexes")
            if (not isinstance(indexes, list) or
                    any(type(i) is not int or not 0 <= i < len(paragraphs) for i in indexes) or
                    loc.get("occurrences") != len(set(indexes)) or index not in indexes or
                    len(indexes) != 1):
                raise GuidedFixFailed("incomplete aggregate")
        para = doc.paragraphs[index]
        if _unsafe_paragraph(para):
            raise GuidedFixFailed("unsafe paragraph")
        if len(para.runs) != len(paragraphs[index]["runs"]):
            raise GuidedFixFailed("run traversal is ambiguous")
        if any(run.text != parsed.get("text") for run, parsed in zip(para.runs, paragraphs[index]["runs"])):
            raise GuidedFixFailed("run traversal changed")
        if code == "REFERENCE_HEADING_ALIGNMENT":
            from app.services.formatting_evidence import resolve_alignment
            _require_unique_reference_heading(doc, paragraphs, index)
            if not resolve_alignment(para).is_resolved:
                raise GuidedFixFailed("ambiguous Reference heading")
        raw = loc.get("run_indexes", loc.get("run_index"))
        run_indexes = raw if isinstance(raw, list) else [raw]
        if code == "REFERENCE_HEADING_BOLD":
            run_indexes = bold_indexes
        elif code in SAFE_RUN_RULES:
            if (not run_indexes or any(type(i) is not int or not 0 <= i < len(para.runs)
                                       for i in run_indexes) or len(set(run_indexes)) != len(run_indexes)):
                raise GuidedFixFailed("invalid run targets")
            if any(_unsafe_run(para.runs[i]) for i in run_indexes):
                raise GuidedFixFailed("unsafe run")
            if any(_unsupported_font_slots(para.runs[i], para, code) for i in run_indexes):
                raise GuidedFixFailed("unresolved alternate font slot")
        else:
            if raw is not None:
                raise GuidedFixFailed("paragraph target contains run indexes")
            run_indexes = [None]
        expected = bold_expected if code == "REFERENCE_HEADING_BOLD" else _resolve_expected(
            violation, audit, doc, roles[index] if index < len(roles) else None, paragraphs[index],
        )
        if expected is None:
            raise GuidedFixFailed("unresolved requirement")
        for run_index in run_indexes:
            key = (index, run_index, code)
            if key in targets:
                raise GuidedFixFailed("duplicate target")
            targets.add(key)
            property_key = (index, run_index, "ALIGNMENT" if code == "REFERENCE_HEADING_ALIGNMENT" else code)
            if property_key in property_targets:
                raise GuidedFixFailed("conflicting property targets")
            property_targets.add(property_key)
            parsed = paragraphs[index]["runs"][run_index] if run_index is not None else None
            if code in ("FONT_SIZE", "FONT_CONSISTENCY"):
                raw_run = para.runs[run_index]
                if code == "FONT_SIZE":
                    raw_size = int(raw_run._r.rPr.find(_W + "sz").get(_W + "val"))
                    actual = violation.actual_value
                    token = actual.rsplit(" ", 1)[-1] if isinstance(actual, str) else actual
                    if (_font_half_points(token, requirement=False) != raw_size or
                            (token != actual and _family_from_actual(actual) !=
                             raw_run._r.rPr.rFonts.get(_W + "ascii").lower())):
                        raise GuidedFixFailed("source differs from finding")
                elif _family_from_actual(violation.actual_value) != raw_run._r.rPr.rFonts.get(_W + "ascii").lower():
                    raise GuidedFixFailed("source differs from finding")
                if code == "FONT_CONSISTENCY" and isinstance(violation.actual_value, str):
                    combo_size = _TRAILING_SIZE_RE.search(violation.actual_value)
                    if combo_size is not None:
                        size = raw_run._r.rPr.find(_W + "sz")
                        if (size is None or _font_half_points(combo_size.group().strip(), requirement=False) !=
                                int(size.get(_W + "val"))):
                            raise GuidedFixFailed("source differs from finding")
            if _current_changed_from_actual(code, paragraphs[index], parsed, violation.actual_value, para=para):
                raise GuidedFixFailed("source differs from finding")
            if _current_matches_expected(code, expected, paragraphs[index], parsed, para=para):
                raise GuidedFixFailed("already compliant")
        items.append({
            "violation_id": violation.id,
            "before": violation.actual_value,
            "after": violation.expected_value,
            "locations": ([f'Reference List heading "{para.text}" — Paragraph {index + 1}']
                          if code == "REFERENCE_HEADING_BOLD" else [f"Paragraph {index + 1}"]),
            "affected_count": 1 if code == "REFERENCE_HEADING_BOLD" else len(run_indexes),
            "property": ("Reference List heading bold (all heading text)"
                         if code == "REFERENCE_HEADING_BOLD" else code.lower().replace("_", " ")),
        })
    return items, targets


def apply_guided_fixes(
    audit: AuditRecord,
    source_bytes: bytes,
    violations: List[Violation],
) -> Tuple[bytes, List[Violation], List[Tuple[Violation, str]]]:
    """Apply whitelisted fixes to a parsed copy; return
    (corrected_bytes, applied, skipped). Never mutates source_bytes; never
    persists anything. Raises GuidedFixFailed on any unverifiable state."""
    _, reviewed_targets = review_targets(audit, source_bytes, violations)
    from app.services.role_classifier import classify_paragraphs

    doc = parse_document(source_bytes)
    _preflight_line_spacing_targets(doc, violations)
    paragraphs = extract_paragraphs(doc)
    roles = classify_paragraphs(doc, paragraphs)

    applied: List[Violation] = []
    skipped: List[Tuple[Violation, str]] = []
    failed: List[str] = []

    for violation in violations:
        rule_code = violation.rule_code
        if rule_code in SAFE_RUN_RULES:
            kind = "run"
        elif rule_code in SAFE_PARA_RULES:
            kind = "para"
        else:
            skipped.append((violation, "unsupported_rule"))
            continue

        location = violation.location if isinstance(violation.location, dict) else {}
        paragraph_index = location.get("paragraph_index")
        if not isinstance(paragraph_index, int) or isinstance(paragraph_index, bool):
            skipped.append((violation, "no_location"))
            continue
        if not (0 <= paragraph_index < len(doc.paragraphs)):
            skipped.append((violation, "target_missing"))
            continue

        para = doc.paragraphs[paragraph_index]
        para_dict = paragraphs[paragraph_index]
        role = roles[paragraph_index] if paragraph_index < len(roles) else None

        if _unsafe_paragraph(para):
            skipped.append((violation, "unsafe_structure"))
            continue

        run_indexes: Optional[List[int]] = None
        run_dict: Optional[Dict] = None
        if kind == "run":
            if rule_code == "REFERENCE_HEADING_BOLD":
                run_indexes, _ = _reference_heading_bold_targets(audit, doc, paragraphs, violation)
            else:
                raw = location.get("run_indexes")
                if isinstance(raw, list):
                    run_indexes = [i for i in raw if isinstance(i, int) and not isinstance(i, bool)]
                else:
                    raw = location.get("run_index")
                    if isinstance(raw, int) and not isinstance(raw, bool):
                        run_indexes = [raw]
            if not run_indexes:
                skipped.append((violation, "no_location"))
                continue
            valid = [i for i in run_indexes if 0 <= i < len(para.runs)]
            if not valid:
                skipped.append((violation, "target_missing"))
                continue
            run_indexes = valid

            runs = [para.runs[i] for i in run_indexes]
            if any(_unsafe_run(r) for r in runs):
                skipped.append((violation, "unsafe_structure"))
                continue
            run_dict = paragraphs[paragraph_index]["runs"][run_indexes[0]]

        # Document-changed gate: recomputed current vs persisted actual.
        if _current_changed_from_actual(rule_code, para_dict, run_dict, violation.actual_value, para=para):
            skipped.append((violation, "document_changed"))
            continue

        # Resolve the expected value (never guessed).
        expected = _resolve_expected(violation, audit, doc, role, para_dict)
        if expected is None:
            skipped.append((violation, "expected_unresolved"))
            continue

        # Already compliant → no mutation.
        if _current_matches_expected(rule_code, expected, para_dict, run_dict, para=para):
            skipped.append((violation, "already_compliant"))
            continue

        # Apply.
        try:
            if rule_code == "FONT_SIZE":
                ok = all(_fix_font_size(r, expected) for r in runs)
            elif rule_code == "FONT_CONSISTENCY":
                ok = all(_fix_font_family(r, expected) for r in runs)
            elif rule_code == "ALIGNMENT":
                ok = _fix_alignment(para, expected)
            elif rule_code == "LINE_SPACING":
                ok = _fix_line_spacing(para, expected)
            elif rule_code == "SPACE_BEFORE":
                ok = _fix_space_before(para, expected)
            elif rule_code == "REFERENCE_HEADING_ALIGNMENT":
                ok = _fix_alignment(para, expected)
            elif rule_code == "REFERENCE_HEADING_BOLD":
                ok = all(_fix_reference_heading_bold(r, expected) for r in runs)
            else:  # SPACE_AFTER
                ok = _fix_space_after(para, expected)
        except Exception:
            failed.append(violation.id)
            break
        if not ok:
            failed.append(violation.id)
            break
        applied.append(violation)

    if failed or skipped or len(applied) != len(violations):
        raise GuidedFixFailed("not all selected findings could be applied")

    # Serialize to memory; the input bytes object is never touched.
    output = io.BytesIO()
    doc.save(output)
    corrected_bytes = output.getvalue()

    if any(violation.rule_code in ("LINE_SPACING", "SPACE_BEFORE", "SPACE_AFTER") for violation in applied):
        _validate_output(corrected_bytes, doc, applied, roles, audit, reviewed_targets,
                         original_source_bytes=source_bytes)
    else:
        _validate_output(corrected_bytes, doc, applied, roles, audit, reviewed_targets)
    return corrected_bytes, applied, []


def _validate_output(corrected_bytes: bytes, source_doc, applied: List[Violation], roles: List[str], audit, reviewed_targets, *, original_source_bytes: Optional[bytes] = None) -> None:
    """Reopen and verify every target's effective value, including grouped runs."""
    doc = parse_document(corrected_bytes)
    _preflight_font_targets(doc, applied)
    _preflight_line_spacing_targets(doc, applied)
    paragraphs = extract_paragraphs(doc)
    if len(doc.paragraphs) != len(source_doc.paragraphs):
        raise GuidedFixFailed("paragraph count changed")
    line_targets = [violation for violation in applied if violation.rule_code == "LINE_SPACING"]
    point_targets = [violation for violation in applied if violation.rule_code in ("SPACE_BEFORE", "SPACE_AFTER")]
    point_keys = set()
    for violation in point_targets:
        location = violation.location
        index = location.get("paragraph_index") if isinstance(location, dict) else None
        if type(index) is not int or not 0 <= index < len(doc.paragraphs):
            raise GuidedFixFailed("invalid paragraph")
        point_keys.add((index, None, violation.rule_code))
    reviewed_points = {key for key in reviewed_targets if key[2] in ("SPACE_BEFORE", "SPACE_AFTER")}
    if len(point_keys) != len(point_targets) or point_keys != reviewed_points:
        raise GuidedFixFailed("incomplete reviewed point targets")
    if line_targets or point_targets:
        from app.services.role_classifier import classify_paragraphs

        if original_source_bytes is None:
            reason = ("original line-spacing evidence unavailable" if not point_targets
                      else "original point-spacing evidence unavailable")
            raise GuidedFixFailed(reason)
        if point_targets and hashlib.sha256(original_source_bytes).hexdigest() != audit.document_sha256:
            raise GuidedFixFailed("original point-spacing source differs from audit")
        original = parse_document(original_source_bytes)
        _preflight_line_spacing_targets(original, line_targets)
        if len(original.paragraphs) != len(doc.paragraphs):
            raise GuidedFixFailed("original paragraph count changed")
        original_paragraphs = extract_paragraphs(original)
        original_roles = classify_paragraphs(original, original_paragraphs)
        output_roles = classify_paragraphs(doc, paragraphs)

        def run_boundaries(para):
            return [
                [(child.tag, child.text, tuple(sorted(child.attrib.items())))
                 for child in run._r if child.tag != _W + "rPr"]
                for run in para.runs
            ]

        def point_content(para):
            def node_content(node):
                return (node.tag, node.text, node.tail, tuple(sorted(node.attrib.items())),
                        tuple(node_content(child) for child in node))

            content = []
            for child in para._p:
                if child.tag == _W + "pPr":
                    continue
                if child.tag == _W + "r":
                    content.append((child.tag, child.text, child.tail, tuple(sorted(child.attrib.items())),
                                    tuple(node_content(node) for node in child if node.tag != _W + "rPr")))
                else:
                    content.append(node_content(child))
            return tuple(sorted(para._p.attrib.items())), content

        for violation in line_targets:
            index = violation.location["paragraph_index"]
            before, after = original.paragraphs[index], doc.paragraphs[index]
            if (before.text != after.text or run_boundaries(before) != run_boundaries(after) or
                    original_roles[index] != output_roles[index]):
                raise GuidedFixFailed("line-spacing target identity changed")
            expected = _resolve_line_spacing_expected(violation, audit, doc, output_roles[index])
            original_expected = _resolve_line_spacing_expected(
                violation, audit, original, original_roles[index],
            )
            if (expected is None or expected != original_expected or
                    _direct_line_spacing(after) != _line_spacing_encoded(violation.expected_value)):
                raise GuidedFixFailed("line-spacing target differs after save")

        line_indexes = {v.location["paragraph_index"] for v in line_targets}
        for violation in point_targets:
            index = violation.location["paragraph_index"]
            before, after = original.paragraphs[index], doc.paragraphs[index]
            if (before.text != after.text or point_content(before) != point_content(after) or
                    original_roles[index] != output_roles[index]):
                raise GuidedFixFailed("point-spacing target identity changed")
            original_expected = _resolve_point_spacing_expected(
                violation, audit, original, original_roles[index], original_paragraphs[index],
            )
            expected = _resolve_point_spacing_expected(
                violation, audit, doc, output_roles[index], paragraphs[index],
            )
            actual = _point_twips(violation.actual_value, requirement=False)
            source_twips = _direct_spacing_twips(before, violation.rule_code)
            _direct_spacing_twips(after, violation.rule_code)
            side = "before" if violation.rule_code == "SPACE_BEFORE" else "after"
            raw = after._p.pPr.find(_W + "spacing").get(_W + side)
            if (expected is None or expected != original_expected or actual != source_twips or
                    abs(source_twips - expected) <= 20 or raw != str(expected)):
                raise GuidedFixFailed("point-spacing target differs after save")

        for index in {key[0] for key in point_keys}:
            before_layers = _spacing_property_layers(original.paragraphs[index])
            after_layers = _spacing_property_layers(doc.paragraphs[index])
            if len(before_layers) != len(after_layers):
                raise GuidedFixFailed("point-spacing controls changed")
            allowed = {_W + side for side, code in (
                ("before", "SPACE_BEFORE"), ("after", "SPACE_AFTER"),
            ) if (index, None, code) in point_keys}
            if (index, None, "LINE_SPACING") in reviewed_targets:
                if index not in line_indexes:
                    raise GuidedFixFailed("incomplete reviewed line target")
                allowed.update({_W + "line", _W + "lineRule"})
            for layer_index, (before_layer, after_layer) in enumerate(zip(before_layers, after_layers)):
                before_spacing = before_layer.find(_W + "spacing")
                after_spacing = after_layer.find(_W + "spacing")
                if (before_spacing is None) != (after_spacing is None):
                    raise GuidedFixFailed("point-spacing controls changed")
                if before_spacing is not None:
                    before_attributes = dict(before_spacing.attrib)
                    after_attributes = dict(after_spacing.attrib)
                    if layer_index == 0:
                        for name in allowed:
                            before_attributes.pop(name, None)
                            after_attributes.pop(name, None)
                    if before_attributes != after_attributes:
                        raise GuidedFixFailed("point-spacing controls changed")
                before_context = before_layer.find(_W + "contextualSpacing")
                after_context = after_layer.find(_W + "contextualSpacing")
                if ((before_context is None) != (after_context is None) or
                        (before_context is not None and dict(before_context.attrib) != dict(after_context.attrib))):
                    raise GuidedFixFailed("point-spacing controls changed")
    for violation in applied:
        location = violation.location
        index = location["paragraph_index"]
        if violation.rule_code in ("LINE_SPACING", "SPACE_BEFORE", "SPACE_AFTER"):
            continue  # Independently certified against pristine bytes above.
        expected = _resolve_expected(
            violation, audit, doc, roles[index] if index < len(roles) else None, paragraphs[index],
        )
        if expected is None:
            raise GuidedFixFailed("requirement changed")
        if violation.rule_code == "REFERENCE_HEADING_BOLD":
            indexes = _reference_heading_bold_run_indexes(doc, paragraphs, index)
            reviewed = sorted(run for paragraph, run, rule in reviewed_targets
                              if paragraph == index and rule == "REFERENCE_HEADING_BOLD")
            before = source_doc.paragraphs[index]
            after = doc.paragraphs[index]
            if (indexes != reviewed or before.text != after.text or
                    [run.text for run in before.runs] != [run.text for run in after.runs] or
                    any(after.runs[i].font.bold is not expected for i in indexes)):
                raise GuidedFixFailed("incomplete Reference heading bold output")
            continue
        indexes = location.get("run_indexes", [location.get("run_index")]) if violation.rule_code in SAFE_RUN_RULES else [None]
        for run_index in indexes:
            if run_index is not None and (type(run_index) is not int or run_index >= len(paragraphs[index]["runs"])):
                raise GuidedFixFailed("run missing after save")
            if violation.rule_code in ("FONT_SIZE", "FONT_CONSISTENCY"):
                para = doc.paragraphs[index]
                raw_run = para.runs[run_index]
                if _unsupported_font_slots(raw_run, para, violation.rule_code):
                    raise GuidedFixFailed("unresolved alternate font slot after save")
                if violation.rule_code == "FONT_SIZE":
                    if int(raw_run._r.rPr.find(_W + "sz").get(_W + "val")) != _font_half_points(f"{expected}pt"):
                        raise GuidedFixFailed("target value differs after save")
                elif any(raw_run._r.rPr.rFonts.get(_W + name) != expected
                         for name in ("ascii", "hAnsi", "eastAsia")):
                    raise GuidedFixFailed("target value differs after save")
                continue
            run = paragraphs[index]["runs"][run_index] if run_index is not None else None
            if not _current_matches_expected(violation.rule_code, expected, paragraphs[index], run, para=doc.paragraphs[index]):
                raise GuidedFixFailed("target value differs after save")
