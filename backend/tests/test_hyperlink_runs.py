"""Visible hyperlink runs in deterministic formatting evidence.

Proves the canonical visible-run iterator (`iter_visible_runs`) yields
every visible text-bearing run exactly once in document order — direct
``w:r`` children plus ``w:r`` descendants of ``w:hyperlink`` — and that
existing font/size/bold/mixed, Reference, and citation paths consume the
added runs without changing URLs, relationships, indexes, severity, or
profiles. Hyperlink targets are metadata only and never visible text.
"""
import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.services.citation_sensor import run_citation_sensor, scan_citations
from app.services.document_parser import (
    extract_paragraphs,
    iter_visible_runs,
    resolve_effective_family,
    resolve_effective_size,
)
from app.services.formatting_evidence import resolve_bold, resolve_paragraph_format
from app.services.guided_fix import _unsafe_paragraph
from app.services.layout_engine import check_font_consistency, check_font_size
from app.services.citation_matching import parse_reference_entry
from app.services.reference_rules import run_reference_rules
from app.services.role_classifier import classify_paragraphs
from app.services.profile_preset_adapter import EffectiveProfileConfig
from app.services.profile_registry import SUC_PROFILE_ID, get_builtin_profile
from app.services.profile_schema import new_custom_profile
from app.services.profile_snapshot import resolve_snapshot

HLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def _hyperlink(parent, target=None, runs=(), anchor=None):
    """Append a w:hyperlink with the given (text, size_pt, family, bold) runs."""
    hl = OxmlElement("w:hyperlink")
    if target is not None:
        r_id = parent.part.relate_to(target, HLINK_REL, is_external=True)
        hl.set(qn("r:id"), r_id)
    else:
        assert anchor is not None
        hl.set(qn("w:anchor"), anchor)
    for text, size, family, bold in runs:
        r = OxmlElement("w:r")
        if size is not None or family is not None or bold is not None:
            rPr = OxmlElement("w:rPr")
            if size is not None:
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), str(int(size * 2)))
                rPr.append(sz)
            if family is not None:
                rf = OxmlElement("w:rFonts")
                rf.set(qn("w:ascii"), family)
                rPr.append(rf)
            if bold is not None:
                b = OxmlElement("w:b")
                if bold is not True:
                    b.set(qn("w:val"), "false")
                rPr.append(b)
            r.append(rPr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        hl.append(r)
    parent._p.append(hl)
    return hl


def _bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _reopen(data):
    return Document(io.BytesIO(data))


def _normal_12_tnr(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


class TNR12Preset:
    profile_label = '"test" profile'
    FONT_FAMILY = "Times New Roman"
    FONT_SIZE_BODY = 12.0
    FONT_SIZE_H1 = None
    FONT_SIZE_H2 = None
    FONT_SIZE_H3 = None
    HEADING_FONT_FAMILY = "Times New Roman"
    HEADING_FONT_SIZE_PT = 12.0
    BODY_ALLOWED_FONT_COMBOS = ()
    HEADING_ALLOWED_FONT_COMBOS = ()

    def heading_expected_size(self, level):
        return self.HEADING_FONT_SIZE_PT

    def is_font_pair_allowed(self, family, size, heading=False):
        if family is None or size is None:
            return True
        return family.lower() == "times new roman" and abs(size - 12.0) < 0.01


def _ref_profile(**refs):
    profile = new_custom_profile(
        "Hyperlink Tests", base=get_builtin_profile(SUC_PROFILE_ID),
    )
    for name, value in refs.items():
        setattr(profile.references, name, value)
    return EffectiveProfileConfig(resolve_snapshot(profile))


def _parsed(data):
    doc = _reopen(data)
    return doc, extract_paragraphs(doc)


# --- Phase 5: run extraction -----------------------------------------------------


def test_01_ordinary_run_only():
    doc = Document()
    doc.add_paragraph("Plain text here.")
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["Plain text here."]
    assert [r.text for r in para.runs] == ["Plain text here."]


def test_02_hyperlink_run_only():
    doc = Document()
    _hyperlink(doc.add_paragraph(), "https://example.invalid/only", [("Only link", 12, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert para.runs == []
    seen = list(iter_visible_runs(para))
    assert [r.text for r in seen] == ["Only link"]
    assert para.text == "Only link"


def test_03_ordinary_then_hyperlink_order():
    doc = Document()
    p = doc.add_paragraph("First. ")
    _hyperlink(p, "https://example.invalid/a", [("Link A", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["First. ", "Link A"]
    assert para.text == "First. Link A"


def test_04_hyperlink_then_ordinary_order():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/a", [("Link A", None, None, None)])
    p.add_run(" after.")
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["Link A", " after."]


def test_05_multiple_hyperlinks_distinct_relationships():
    doc = Document()
    p = doc.add_paragraph()
    h1 = _hyperlink(p, "https://example.invalid/a", [("Alpha", None, None, None)])
    h2 = _hyperlink(p, "https://example.invalid/b", [("Beta", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["Alpha", "Beta"]
    assert h1.get(qn("r:id")) != h2.get(qn("r:id"))


def test_06_one_hyperlink_split_into_several_runs():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/a", [("Part one ", None, None, None), ("part two", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["Part one ", "part two"]


def test_07_interleaved_runs_preserve_order():
    doc = Document()
    p = doc.add_paragraph("Start ")
    _hyperlink(p, "https://example.invalid/a", [("link", None, None, None)])
    p.add_run(" middle ")
    _hyperlink(p, "https://example.invalid/b", [("tail", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["Start ", "link", " middle ", "tail"]


def test_08_internal_bookmark_hyperlink_included():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, None, [("Jump target", None, None, None)], anchor="bkmk1")
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert [r.text for r in iter_visible_runs(para)] == ["Jump target"]
    assert "bkmk1" not in para.text


def test_09_label_differs_from_target_url_excluded():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/document", [("Official guidance", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert para.text == "Official guidance"
    assert "example.invalid" not in para.text
    assert all("example.invalid" not in (r.text or "") for r in iter_visible_runs(para))


def test_10_hyperlink_without_visible_text_contributes_nothing():
    doc = Document()
    p = doc.add_paragraph("Lead.")
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), "rId9")
    hl.append(OxmlElement("w:r"))
    p._p.append(hl)
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert para.text == "Lead."
    # Membership parity with Paragraph.runs: the contentless run element is
    # still yielded; it contributes no text and no finding.
    assert len(list(iter_visible_runs(para))) == 2
    assert [r.text for r in iter_visible_runs(para) if (r.text or "")] == ["Lead."]


def test_11_tabs_and_breaks_follow_existing_conventions():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/a", [("tabbed", None, None, None)])
    # Append a tab + break inside the same hyperlink run element.
    hl = p._p.findall(qn("w:hyperlink"))[0]
    r_el = hl.findall(qn("w:r"))[0]
    for tag in ("w:tab", "w:br"):
        r_el.append(OxmlElement(tag))
    plain = Document()
    q = plain.add_paragraph()
    rq = q.add_run("tabbed")
    rq._r.append(OxmlElement("w:tab"))
    rq._r.append(OxmlElement("w:br"))
    (hpara,) = _reopen(_bytes(doc)).paragraphs
    (opara,) = _reopen(_bytes(plain)).paragraphs
    (hrun,) = [r for r in iter_visible_runs(hpara) if (r.text or "")]
    (orun,) = list(opara.runs)
    assert hrun.text == orun.text


def test_12_field_instructions_remain_excluded():
    doc = Document()
    p = doc.add_paragraph()
    r = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " SEQ Table \\* ARABIC "
    r.append(instr)
    p._p.append(r)
    _hyperlink(p, "https://example.invalid/a", [("caption link", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert "SEQ" not in para.text
    assert "caption link" in para.text


def test_13_no_run_appears_twice():
    doc = Document()
    p = doc.add_paragraph("A ")
    _hyperlink(p, "https://example.invalid/a", [("B", None, None, None), ("C", None, None, None)])
    p.add_run(" D ")
    _hyperlink(p, "https://example.invalid/b", [("E", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    seen = list(iter_visible_runs(para))
    assert len({id(r._element) for r in seen}) == len(seen) == 5
    direct = {id(r._element) for r in para.runs}
    assert direct.issubset({id(r._element) for r in seen})


def test_14_iterator_falls_back_for_test_doubles():
    class Fake:
        runs = ["sentinel"]

    assert list(iter_visible_runs(Fake())) == ["sentinel"]
    assert list(iter_visible_runs(object())) == []


# --- Phase 6: formatting -----------------------------------------------------------


def _font_doc(link_runs, ordinary_size=12, ordinary_family=None):
    doc = Document()
    _normal_12_tnr(doc)
    p = doc.add_paragraph()
    r = p.add_run("Ordinary text. ")
    r.font.size = Pt(ordinary_size)
    if ordinary_family:
        r.font.name = ordinary_family
    _hyperlink(p, "https://example.invalid/x", link_runs)
    return _bytes(doc)


def _font_violations(data, code):
    doc = _reopen(data)
    paras = extract_paragraphs(doc)
    roles = classify_paragraphs(doc, paras)
    size = check_font_size(paras, TNR12Preset(), roles=roles)
    cons = check_font_consistency(paras, TNR12Preset(), roles=roles)
    return [v for v in size + cons if v.rule_code == code]


def test_f1_hyperlink_only_font_size_mismatch():
    doc = Document()
    _normal_12_tnr(doc)
    _hyperlink(doc.add_paragraph(), "https://example.invalid/x", [("Small link", 10, "Times New Roman", None)])
    data = _bytes(doc)
    assert any(v.rule_code == "FONT_SIZE" for v in _font_violations(data, "FONT_SIZE"))


def test_f2_hyperlink_only_font_family_mismatch():
    doc = Document()
    _normal_12_tnr(doc)
    _hyperlink(doc.add_paragraph(), "https://example.invalid/x", [("Calibri link", 12, "Calibri", None)])
    data = _bytes(doc)
    assert any(v.rule_code == "FONT_CONSISTENCY" for v in _font_violations(data, "FONT_CONSISTENCY"))


def test_f3_hyperlink_bold_evidence_consumed_where_supported():
    # A References heading must stay exactly the alias text, so rule-level
    # heading coverage lives in existing tests; here the shared consumer
    # (resolve_bold, used by REFERENCE_HEADING_BOLD) is proven directly.
    doc = Document()
    _hyperlink(doc.add_paragraph(), "https://example.invalid/x", [("Bold link", None, None, True)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert resolve_bold(para).effective_bold is True

    mixed = Document()
    mp = mixed.add_paragraph()
    off = mp.add_run("Plain. ")
    off.font.bold = False
    _hyperlink(mp, "https://example.invalid/x", [("Bold link", None, None, True)])
    (mpara,) = _reopen(_bytes(mixed)).paragraphs
    evidence = resolve_bold(mpara)
    assert evidence.has_mixed is True
    assert evidence.effective_bold is None


def test_f4_correct_hyperlink_produces_no_finding():
    doc = Document()
    _normal_12_tnr(doc)
    p = doc.add_paragraph()
    r = p.add_run("Ordinary text. ")
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"
    _hyperlink(p, "https://example.invalid/x", [("Good link", 12, "Times New Roman", None)])
    data = _bytes(doc)
    assert _font_violations(data, "FONT_SIZE") == []
    assert _font_violations(data, "FONT_CONSISTENCY") == []


def test_f5_mixed_sizes_detected_once_each():
    data = _font_doc([("Small link", 10, "Times New Roman", None)])
    doc, paras = _parsed(data)
    evidence = resolve_paragraph_format(doc.paragraphs[0])
    assert evidence.has_mixed_font_size is True
    assert _font_violations(data, "FONT_SIZE") != []


def test_f6_mixed_families_detected_once_each():
    data = _font_doc([("Calibri link", 12, "Calibri", None)], ordinary_family="Times New Roman")
    doc, paras = _parsed(data)
    evidence = resolve_paragraph_format(doc.paragraphs[0])
    assert evidence.has_mixed_font_family is True
    assert _font_violations(data, "FONT_CONSISTENCY") != []


def test_f7_inherited_formatting_resolves_without_fabrication():
    doc = Document()
    _normal_12_tnr(doc)
    p = doc.add_paragraph()
    r = p.add_run("Ordinary text. ")
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"
    _hyperlink(p, "https://example.invalid/x", [("Plain link", None, None, None)])
    data = _bytes(doc)
    assert _font_violations(data, "FONT_SIZE") == []
    assert _font_violations(data, "FONT_CONSISTENCY") == []


def test_f8_character_style_font_resolves_by_precedence():
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    _normal_12_tnr(doc)
    style = doc.styles.add_style("LinkChar", WD_STYLE_TYPE.CHARACTER)
    style.font.name = "Calibri"
    p = doc.add_paragraph()
    hl = _hyperlink(p, "https://example.invalid/x", [("Styled link", None, None, None)])
    r_el = hl.findall(qn("w:r"))[0]
    rPr = r_el.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        r_el.insert(0, rPr)
    rPr.append(_rstyle("LinkChar"))
    (para,) = _reopen(_bytes(doc)).paragraphs
    (run,) = [r for r in iter_visible_runs(para) if (r.text or "").strip()]
    family, source = resolve_effective_family(run, para, doc.styles.element)
    assert family == "Calibri"
    assert source == "character_style"


def _rstyle(name):
    el = OxmlElement("w:rStyle")
    el.set(qn("w:val"), name)
    return el


def test_f9_inherited_formatting_resolves_without_fabrication():
    doc = Document()
    _normal_12_tnr(doc)
    p = doc.add_paragraph()
    r = p.add_run("Ordinary text. ")
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"
    _hyperlink(p, "https://example.invalid/x", [("Bare link", None, None, None)])
    data = _bytes(doc)
    # The bare link inherits Normal 12 pt TNR: no finding, and crucially no
    # finding is fabricated from hyperlink membership alone.
    assert _font_violations(data, "FONT_SIZE") == []
    assert _font_violations(data, "FONT_CONSISTENCY") == []
    (para,) = _reopen(data).paragraphs
    (run,) = [x for x in iter_visible_runs(para) if (x.text or "").strip().startswith("Bare")]
    family, _ = resolve_effective_family(run, para, doc.styles.element)
    size, _ = resolve_effective_size(run, para, doc.styles.element)
    assert (family, size) == ("Times New Roman", 12.0)


def test_f9b_no_usable_source_stays_unresolved():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/x", [("Bare link", None, None, None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    (run,) = [x for x in iter_visible_runs(para) if (x.text or "")]
    family, _ = resolve_effective_family(run, para, doc.styles.element)
    # Default template declares no usable family: unresolved, never made up —
    # and no font finding is fabricated from it.
    assert family is None
    data = _bytes(doc)
    assert _font_violations(data, "FONT_CONSISTENCY") == []


def test_f10_run_indexes_count_each_physical_run_once():
    data = _font_doc([("A", 10, "Times New Roman", None), ("B", 10, "Times New Roman", None)])
    _, paras = _parsed(data)
    runs = paras[0]["runs"]
    assert [r["index"] for r in runs] == list(range(len(runs)))
    assert len(runs) == 3


def test_f11_ordinary_formatting_unchanged():
    doc = Document()
    _normal_12_tnr(doc)
    p = doc.add_paragraph()
    r = p.add_run("Ordinary text. ")
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"
    data = _bytes(doc)
    assert _font_violations(data, "FONT_SIZE") == []
    assert _font_violations(data, "FONT_CONSISTENCY") == []


# --- Phase 7: Reference + citation regressions ---------------------------------------


def _ref_profile():
    profile = new_custom_profile(
        "Hyperlink Tests", base=get_builtin_profile(SUC_PROFILE_ID),
    )
    profile.references.check_hanging_indent = True
    profile.references.hanging_indent_in = 0.5
    profile.references.check_heading_bold = True
    profile.references.check_font_family = True
    profile.references.font_family = "Times New Roman"
    profile.references.check_font_size = True
    profile.references.font_size_pt = 12.0
    return EffectiveProfileConfig(resolve_snapshot(profile))


def _entry_doc(entry_runs, entry_text_extra=" Test Journal."):
    doc = Document()
    _normal_12_tnr(doc)
    doc.add_paragraph("References")
    p = doc.add_paragraph()
    for text, size, family, bold in entry_runs:
        if text == "<LINK>":
            _hyperlink(p, "https://doi.org/10.0000/x", [("doi:10.0000/x", size, family, bold)])
        else:
            r = p.add_run(text)
            if size:
                r.font.size = Pt(size)
            if family:
                r.font.name = family
    return _bytes(doc)


def test_r1_doi_hyperlink_participates_in_reference_checks():
    data = _entry_doc([
        ("Smith, J. (2020). Title. ", 12, "Times New Roman", None),
        ("<LINK>", 10, "Times New Roman", None),
    ])
    from app.services.citation_matching import parse_reference_entry
    doc, paras = _parsed(data)
    assert "doi:10.0000/x" in paras[1]["text"]
    assert parse_reference_entry(paras[1]["text"], 1) is not None
    violations = run_reference_rules(doc, paras, _ref_profile())
    assert any(v.rule_code == "REFERENCE_FONT_SIZE" for v in violations)


def test_r2_human_label_participates_target_excluded():
    doc = Document()
    _normal_12_tnr(doc)
    doc.add_paragraph("References")
    p = doc.add_paragraph("Smith, J. (2020). See ")
    _hyperlink(p, "https://example.invalid/very-long-target-url", [("Official guidance", 12, "Calibri", None)])
    data = _bytes(doc)
    _, paras = _parsed(data)
    assert "very-long-target-url" not in paras[1]["text"]
    assert "Official guidance" in paras[1]["text"]
    violations = run_reference_rules(doc, paras, _ref_profile())
    assert any(v.rule_code == "REFERENCE_FONT_FAMILY" for v in violations)


def test_r3_reference_matching_stable_with_hyperlinks():
    from app.services.citation_sensor import run_citation_sensor
    from app.services.role_classifier import classify_paragraphs as _roles
    doc = Document()
    _normal_12_tnr(doc)
    doc.add_paragraph("Smith (2020) noted this.")
    doc.add_paragraph("References")
    p = doc.add_paragraph("Smith, J. (2020). Title. ")
    _hyperlink(p, "https://doi.org/10.0000/x", [("doi:10.0000/x", None, None, None)])
    data = _bytes(doc)
    redoc = _reopen(data)
    paras = extract_paragraphs(redoc)
    violations = run_citation_sensor(redoc, paras, config=None, roles=_roles(redoc, paras))
    assert [v.rule_code for v in violations if "CITATION" in v.rule_code or "REFERENCE" in v.rule_code] == []


def test_r4_hyperlink_only_citation_paragraph_detected():
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/s", [("(Smith, 2020)", None, None, None)])
    doc.add_paragraph("References")
    doc.add_paragraph("Smith, J. (2020). Title. Press.")
    data = _bytes(doc)
    redoc = _reopen(data)
    paras = extract_paragraphs(redoc)
    from app.services.role_classifier import classify_paragraphs as _roles
    roles = _roles(redoc, paras)
    assert roles[0] != "EMPTY"
    scan = scan_citations(paras[0]["text"], 0)
    assert len(scan.citations) == 1
    assert scan.citations[0].snippet == "(Smith, 2020)"
    assert "example.invalid" not in scan.citations[0].snippet


def test_r5_institutional_and_boundary_paths_unaffected():
    from app.services.citation_sensor import run_citation_sensor
    from app.services.role_classifier import classify_paragraphs as _roles
    from app.services.reference_section import reference_entry_indexes
    doc = Document()
    doc.add_paragraph("World Health Organization (2022) states this.")
    doc.add_paragraph("References")
    p = doc.add_paragraph("World Health Organization. (2022). Report. ")
    _hyperlink(p, "https://example.invalid/who", [("who.int", None, None, None)])
    data = _bytes(doc)
    redoc = _reopen(data)
    paras = extract_paragraphs(redoc)
    assert 2 in reference_entry_indexes(paras)
    violations = run_citation_sensor(redoc, paras, config=None, roles=_roles(redoc, paras))
    assert [v.rule_code for v in violations if "CITATION" in v.rule_code] == []


# --- Phase 10: API-level functional validation -----------------------------------------


def _api_doc(link_runs, ordinary_size=12, ordinary_family="Times New Roman"):
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Ordinary text. ")
    r.font.size = Pt(ordinary_size)
    r.font.name = ordinary_family
    _hyperlink(p, "https://example.invalid/x", link_runs)
    return _bytes(doc)


def _post(client, data, name="functional.docx"):
    resp = client.post(
        "/api/audit",
        files={"file": (name, data, "application/octet-stream")},
    )
    assert resp.status_code == 200
    return resp.json()


def _codes(payload):
    return [v["rule_code"] for v in payload["physical_layout_errors"]]


def test_api_A_hyperlink_size_mismatch_detected(client):
    payload = _post(client, _api_doc([("Small link", 10, "Times New Roman", None)]))
    assert "FONT_SIZE" in _codes(payload)


def test_api_B_hyperlink_family_difference_detected(client):
    payload = _post(client, _api_doc([("Calibri link", 12, "Calibri", None)]))
    assert "FONT_CONSISTENCY" in _codes(payload)


def test_api_C_correct_hyperlink_silent(client):
    payload = _post(client, _api_doc([("Good link", 12, "Times New Roman", None)]))
    assert "FONT_SIZE" not in _codes(payload)
    assert "FONT_CONSISTENCY" not in _codes(payload)


def test_api_D_reference_doi_hyperlink_stable(client):
    doc = Document()
    doc.add_paragraph("Smith (2020) noted this.")
    doc.add_paragraph("References")
    p = doc.add_paragraph("Smith, J. (2020). Title. ")
    _hyperlink(p, "https://doi.org/10.0000/x", [("doi:10.0000/x", None, None, None)])
    payload = _post(client, _bytes(doc))
    assert "CITATION_MISMATCH" not in _codes(payload)
    export = client.get(f"/api/audit/{payload['audit_id']}/export-pdf")
    assert export.status_code == 200
    assert export.headers["content-type"] == "application/pdf"
    got = client.get(f"/api/audit/{payload['audit_id']}")
    assert got.status_code == 200
    assert "doi:10.0000/x" not in str(got.json())


def test_api_E_hyperlink_only_citation_detected(client):
    doc = Document()
    p = doc.add_paragraph()
    _hyperlink(p, "https://example.invalid/s", [("(Smith, 2020)", None, None, None)])
    doc.add_paragraph("References")
    doc.add_paragraph("Smith, J. (2020). Title. Press.")
    payload = _post(client, _bytes(doc))
    assert "CITATION_MISMATCH" not in _codes(payload)
    # Without the entry the same hyperlink citation must be detected as a
    # mismatch at paragraph 0 with the visible label as evidence — proving
    # the paragraph is not empty and the URL never enters evidence.
    doc2 = Document()
    q = doc2.add_paragraph()
    _hyperlink(q, "https://example.invalid/s", [("(Smith, 2020)", None, None, None)])
    doc2.add_paragraph("References")
    payload2 = _post(client, _bytes(doc2))
    mismatches = [v for v in payload2["physical_layout_errors"] if v["rule_code"] == "CITATION_MISMATCH"]
    assert len(mismatches) == 1
    assert mismatches[0]["location"]["paragraph_index"] == 0
    assert mismatches[0]["actual_value"] == "(Smith, 2020)"
    history = client.get("/api/audits")
    assert history.status_code == 200


# --- cross-surface -------------------------------------------------------------------


def test_x_guided_fixes_skips_hyperlink_paragraphs():
    from app.services.guided_fix import _unsafe_paragraph
    doc = Document()
    p = doc.add_paragraph("Text ")
    _hyperlink(p, "https://example.invalid/x", [("link", 10, "Calibri", None)])
    (para,) = _reopen(_bytes(doc)).paragraphs
    assert _unsafe_paragraph(para) is True


def test_x_blocks_text_source_faithful_without_targets():
    from app.services.document_parser import extract_document_blocks
    doc = Document()
    p = doc.add_paragraph("See ")
    _hyperlink(p, "https://example.invalid/doc", [("Official guidance", None, None, None)])
    (block,) = [b for b in extract_document_blocks(_reopen(_bytes(doc))) if b["index"] == 0]
    assert block["text"] == "See Official guidance"
    assert "example.invalid" not in block["text"]
