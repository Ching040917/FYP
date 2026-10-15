"""Stored-story reviews: absence, positive witnesses, uncertainty and isolation."""
from copy import deepcopy
from io import BytesIO
from itertools import product
from unittest.mock import patch

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START

from app.services.document_parser import extract_document_blocks
from app.services.layout_engine import run_static_rules_engine
from app.services.page_elements import extract_page_elements
from app.services.page_element_rules import check_page_element_requirements, PAGE_ELEMENT_RULES
from app.services.profile_schema import new_custom_profile
from app.services.profile_snapshot import resolve_snapshot
from app.services.profile_preset_adapter import EffectiveProfileConfig
from app.services.scoring import calculate_weighted_score_detailed
from tests.test_page_elements import simple, complex_field, marker, instruction, xml, package


def config(header=False, footer=False, page=False, placement="either"):
    profile = new_custom_profile("Structural requirements")
    profile.headers.check_header_requirements = header
    profile.footers.check_footer_requirements = footer
    profile.page_numbers.check_page_number_requirements = page
    profile.page_numbers.placement = placement
    return EffectiveProfileConfig(resolve_snapshot(profile))


def saved(doc):
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def check(doc, **kwargs):
    return check_page_element_requirements(extract_page_elements(doc), config(**kwargs))


@pytest.mark.parametrize("source", ["none", "custom_off", "legacy", "built_in"])
def test_off_legacy_and_builtin_never_extract(source):
    from app.config import settings
    cfg = None if source == "none" else settings.PRESET if source == "legacy" else config()
    if source == "built_in":
        cfg.profile_source = "built_in"
        cfg.HEADER_CHECK_REQUIREMENTS = True
    with patch("app.services.layout_engine.extract_page_elements", side_effect=AssertionError("unexpected traversal")):
        findings = run_static_rules_engine(saved(Document()), config=cfg)
    assert not any(row.rule_code in PAGE_ELEMENT_RULES for row in findings)


@pytest.mark.parametrize("toggles", list(product([False, True], repeat=3))[1:])
def test_any_enabled_combination_extracts_once(toggles):
    with patch("app.services.layout_engine.extract_page_elements", wraps=extract_page_elements) as extract:
        run_static_rules_engine(saved(Document()), config=config(*toggles))
    assert extract.call_count == 1


@pytest.mark.parametrize("kind", ["header", "footer"])
@pytest.mark.parametrize("state,expected", [("absent", "missing"), ("empty", "empty"),
    ("text", None), ("drawing", "unresolved"), ("table", "unresolved")])
def test_content_truth_table(kind, state, expected):
    doc = Document()
    if state != "absent":
        target = getattr(doc.sections[0], kind)
        target.is_linked_to_previous = False
        if state == "text":
            target.paragraphs[0].text = "Stored content"
        elif state == "drawing":
            target.paragraphs[0]._p.append(xml("drawing"))
        elif state == "table":
            target._element.append(xml("tbl"))
    findings = check(doc, **{kind: True})
    assert [row.location["evidence_state"] for row in findings] == ([] if expected is None else [expected])
    if expected:
        assert findings[0].severity == "REVIEW"
        assert findings[0].rule_code == f"{kind.upper()}_CONTENT_REQUIRED"
        assert ("no text or supported field" in findings[0].actual_value) is (expected == "empty")


@pytest.mark.parametrize("kind", ["header", "footer"])
@pytest.mark.parametrize("field", ["PAGE", "NUMPAGES", "SECTION", "SECTIONPAGES"])
def test_supported_fields_are_content_witnesses_despite_font_uncertainty(kind, field):
    doc = Document()
    simple(getattr(doc.sections[0], kind).paragraphs[0], field, "")
    assert check(doc, **{kind: True}) == []


@pytest.mark.parametrize("kind", ["header", "footer"])
def test_inherited_shared_and_explicit_empty_consumers(kind):
    doc = Document()
    getattr(doc.sections[0], kind).paragraphs[0].text = "Inherited"
    doc.add_section(WD_SECTION_START.CONTINUOUS)
    doc.add_section()
    assert check(doc, **{kind: True}) == []
    target = getattr(doc.sections[1], kind)
    target.is_linked_to_previous = False
    target.paragraphs[0].text = ""
    findings = check(doc, **{kind: True})
    assert [(row.location["section_index"], row.location["evidence_state"]) for row in findings] == [(1, "empty"), (2, "empty")]


@pytest.mark.parametrize("placement", ["header", "footer", "either"])
@pytest.mark.parametrize("kind", ["header", "footer"])
@pytest.mark.parametrize("form", ["simple", "complex"])
def test_page_placement_identity_and_empty_stale_cache(kind, placement, form):
    doc = Document()
    p = getattr(doc.sections[0], kind).paragraphs[0]
    if form == "simple":
        field = simple(p, " PAGE \\* MERGEFORMAT ", "")
        field.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldLock", "true")
    else:
        complex_field(p, result="99")
    findings = check(doc, page=True, placement=placement)
    assert [row.location["evidence_state"] for row in findings] == ([] if placement in (kind, "either") else ["missing"])


@pytest.mark.parametrize("field", [None, "NUMPAGES", "SECTION", "SECTIONPAGES"])
def test_digits_and_total_fields_do_not_satisfy_page(field):
    doc = Document()
    p = doc.sections[0].footer.paragraphs[0]
    if field:
        simple(p, field, "1")
    else:
        p.text = "1"
    # Real direct fonts make absence provable; their values are not checked.
    from docx.shared import Pt
    for run in p._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
        from docx.text.run import Run
        r = Run(run, p)
        r.font.name = "Arial"
        r.font.size = Pt(11)
    findings = check(doc, page=True, placement="footer")
    assert [row.location["evidence_state"] for row in findings] == ["missing"]


@pytest.mark.parametrize("state", ["malformed", "unsupported", "font_incomplete"])
def test_page_absence_requires_complete_supported_evidence(state):
    doc = Document()
    p = doc.sections[0].footer.paragraphs[0]
    if state == "malformed":
        marker(p, "begin")
        instruction(p, " PAGE ")
    elif state == "unsupported":
        p._p.append(xml("drawing"))
    else:
        p.text = "1"
    findings = check(doc, page=True)
    assert [row.location["evidence_state"] for row in findings] == ["unresolved"]


@pytest.mark.parametrize("variant,enabled", product(["first_page_", "even_page_"], [False, True]))
def test_other_variants_never_satisfy_normal_pages(variant, enabled):
    doc = Document()
    doc.sections[0].different_first_page_header_footer = enabled
    doc.settings.odd_and_even_pages_header_footer = enabled
    for kind in ("header", "footer"):
        simple(getattr(doc.sections[0], variant + kind).paragraphs[0])
    findings = check(doc, header=True, footer=True, page=True)
    assert [row.location["evidence_state"] for row in findings] == ["missing"] * 3


def test_either_truth_table_and_unresolved_section_precedence():
    elements = extract_page_elements(Document())
    for header, footer in product(["positive", "missing", "unresolved"], repeat=2):
        current = deepcopy(elements)
        current["parts"] = []
        for key, (kind, state) in enumerate((("header", header), ("footer", footer))):
            story = next(s for s in current["sections"][0]["stories"] if s["kind"] == kind and s["variant"] == "default")
            story["status"] = "resolved" if state == "positive" else "absent" if state == "missing" else "unresolved"
            story["part_key"] = key if state == "positive" else None
            if state == "positive":
                current["parts"].append(dict(part_key=key, completeness="unresolved", paragraphs=[dict(literal_text="", fields=[dict(kind="PAGE", status="resolved")])]))
        findings = check_page_element_requirements(current, config(page=True))
        expected = [] if "positive" in (header, footer) else ["missing"] if header == footer == "missing" else ["unresolved"]
        assert [row.location["evidence_state"] for row in findings] == expected
        current["sections"][0]["status"] = "unresolved"
        assert check_page_element_requirements(current, config(page=True))[0].location["evidence_state"] == "unresolved"


def test_order_minimal_locations_score_and_package_body_immutability():
    doc = Document()
    doc.add_paragraph("Body alpha")
    doc.add_section(WD_SECTION_START.CONTINUOUS)
    doc.add_paragraph("Body beta")
    before = package(doc)
    blocks = extract_document_blocks(doc)
    findings = check(doc, header=True, footer=True, page=True)
    assert [(v.location["section_index"], v.rule_code) for v in findings] == [(i, code) for i in range(2) for code in ("HEADER_CONTENT_REQUIRED", "FOOTER_CONTENT_REQUIRED", "PAGE_NUMBER_FIELD_REQUIRED")]
    assert all(set(v.location) == {"structural_story", "section_index", "story_type", "story_variant", "evidence_state"} for v in findings)
    baseline = calculate_weighted_score_detailed([])
    actual = calculate_weighted_score_detailed(findings)
    assert actual == baseline
    assert actual.total == 100 and actual.major_count == actual.minor_count == 0
    assert package(doc) == before
    assert extract_document_blocks(doc) == blocks


@pytest.mark.parametrize("kind", ["header", "footer"])
@pytest.mark.parametrize("state", ["absent", "empty", "unresolved"])
def test_deficiency_and_uncertainty_have_manual_actions_and_no_score(kind, state):
    doc = Document()
    if state != "absent":
        target = getattr(doc.sections[0], kind)
        target.is_linked_to_previous = False
        if state == "unresolved":
            target.paragraphs[0]._p.append(xml("drawing"))
    row = check(doc, **{kind: True})[0]
    assert calculate_weighted_score_detailed([row]) == calculate_weighted_score_detailed([])
    assert f"Section 1 · {kind.title()} · Normal pages" in row.message
    assert ("ACA cannot determine" in row.message) is (state == "unresolved")
    assert ("Verify the rendered document manually." in row.message) is (state != "unresolved")


@pytest.mark.parametrize("state", ["absent", "header_page", "malformed"])
def test_structural_reviews_persist_immutably_without_ai_or_fixes(client, state, monkeypatch):
    import json
    from pypdf import PdfReader
    from types import SimpleNamespace
    from app.api import routes
    from app.services.ai_citation import AiCitationResult
    from app.services.ai_citation_eligibility import is_local_ai_citation_eligible
    from app.services.guided_fix import apply_guided_fixes

    doc = Document()
    doc.add_paragraph("Body remains unchanged.")
    if state == "header_page":
        simple(doc.sections[0].header.paragraphs[0], " PAGE ", "99")
    elif state == "malformed":
        p = doc.sections[0].header.paragraphs[0]
        marker(p, "begin")
        instruction(p, " PAGE ")
    source = saved(doc)
    original_package = package(doc)
    blocks = extract_document_blocks(doc)
    profile = new_custom_profile("Stored structural review")
    profile.headers.check_header_requirements = True
    profile.footers.check_footer_requirements = True
    profile.page_numbers.check_page_number_requirements = True
    ai_inputs = []

    async def capture_ai(*args, **kwargs):
        ai_inputs.extend(kwargs.get("citation_findings", []))
        return AiCitationResult(status="COMPLETED_NO_SUGGESTIONS", provider="LOCAL_OLLAMA", suggestions=[])

    monkeypatch.setattr(routes, "async_ai_citation_task", capture_ai)
    response = client.post("/api/audit", files={"file": ("structural.docx", source)},
                           data={"custom_profile": json.dumps(profile.to_dict())})
    assert response.status_code == 200
    submitted = response.json()
    rows = submitted["physical_layout_errors"]
    expected = [("FOOTER_CONTENT_REQUIRED", "missing")] if state == "header_page" else [
        ("HEADER_CONTENT_REQUIRED", "unresolved" if state == "malformed" else "missing"),
        ("FOOTER_CONTENT_REQUIRED", "missing"),
        ("PAGE_NUMBER_FIELD_REQUIRED", "unresolved" if state == "malformed" else "missing"),
    ]
    assert [(r["rule_code"], r["location"]["evidence_state"]) for r in rows] == expected
    assert submitted["major_count"] == submitted["minor_count"] == 0
    assert submitted["weighted_compliance_score"] == 100
    assert submitted["profile_snapshot"]["schema_version"] == 6
    assert ai_inputs == []
    assert all(not is_local_ai_citation_eligible(SimpleNamespace(**r)) for r in rows)
    _, applied, skipped = apply_guided_fixes(
        SimpleNamespace(profile_snapshot=submitted["profile_snapshot"]), source,
        [SimpleNamespace(**row) for row in rows])
    assert applied == []
    assert [reason for _, reason in skipped] == ["unsupported_rule"] * len(rows)
    audit_id = submitted["audit_id"]
    profile.headers.check_header_requirements = False
    profile.page_numbers.placement = "footer"
    with patch("app.services.layout_engine.extract_page_elements", side_effect=AssertionError("history re-audit")):
        restored = client.get(f"/api/audit/{audit_id}").json()
        history = client.get("/api/audits").json()
        pdf = client.get(f"/api/audit/{audit_id}/export-pdf")
    assert restored["profile_snapshot"] == submitted["profile_snapshot"]
    assert restored["violations"] == rows
    assert restored["score_breakdown"] == submitted["score_breakdown"]
    assert any(entry["id"] == audit_id for entry in history)
    assert client.get(f"/api/audit/{audit_id}/document-blocks").json()["blocks"] == blocks
    assert pdf.status_code == 200
    text = " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "Manual review" in text
    assert package(doc) == original_package
    assert extract_document_blocks(doc) == blocks
