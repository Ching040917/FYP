"""The correction is downloadable only after independent verification."""
import hashlib
import io
import json
import copy
import zipfile
from decimal import Decimal
from dataclasses import replace

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

import pytest
from sqlalchemy.orm import Session
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

from app.models.audit import AuditRecord, Violation
from app.services.guided_fix import GuidedFixFailed, review_targets
from types import SimpleNamespace
from app.services.guided_fix_verification import verify_reaudit
from app.services.layout_engine import run_static_rules_engine
from app.services.profile_preset_adapter import EffectiveProfileConfig
from app.services.profile_schema import new_custom_profile
from app.services.profile_snapshot import resolve_snapshot
from app.services.scoring import calculate_weighted_score_detailed
from app.services import guided_fix_storage
from tests.test_guided_fix_reference_heading import (
    BOLD_RULE, _bold_audit, _bold_bytes, _bold_doc,
)
from tests.test_guided_fixes import _clean_font_fixture_context


def _doc():
    doc = Document()
    _clean_font_fixture_context(doc)
    p = doc.add_paragraph('Body paragraph.')
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.name = 'Times New Roman'
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload(client, path, data, ids, review_id=None):
    fields = {'selection': json.dumps({'violation_ids': ids})}
    if review_id is not None:
        fields['review_id'] = review_id
    return client.post(path, data=fields, files={'file': ('original.docx', data)})


def test_review_apply_download_once(client, monkeypatch, tmp_path):
    monkeypatch.setattr(guided_fix_storage.settings, 'GUIDED_FIX_STORAGE_DIR', str(tmp_path))
    source = _doc()
    audit = client.post('/api/audit', files={'file': ('original.docx', source)})
    assert audit.status_code == 200, audit.text
    audit_id = audit.json()['audit_id']
    ids = [v['id'] for v in audit.json()['physical_layout_errors'] if v['rule_code'] == 'FONT_SIZE']
    assert ids
    root = f'/api/audit/{audit_id}/guided-fixes'
    assert _upload(client, root, source, ids).status_code == 422
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    assert review.json()['items'][0]['locations'] == ['Paragraph 1']
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    payload = applied.json()
    assert payload['status'] == 'passed'
    assert payload['resolved'] == ids[:1]
    assert payload['new_findings'] == []
    assert _upload(client, root, source, ids, review.json()['review_id']).status_code == 409
    download = client.get(root + '/download/' + payload['download_token'])
    assert download.status_code == 200
    assert download.headers['cache-control'] == 'private, no-store'
    assert Document(io.BytesIO(download.content)).paragraphs[0].runs[0].font.size.pt == 12
    assert client.get(root + '/download/' + payload['download_token']).status_code == 410
    assert hashlib.sha256(source).digest() != hashlib.sha256(download.content).digest()
    assert audit.json()['weighted_compliance_score'] == client.get(f'/api/audit/{audit_id}').json()['weighted_score']


def test_cancel_review_prevents_apply(client):
    source = _doc()
    audit = client.post('/api/audit', files={'file': ('original.docx', source)})
    audit_id = audit.json()['audit_id']
    ids = [v['id'] for v in audit.json()['physical_layout_errors'] if v['rule_code'] == 'FONT_SIZE']
    root = f'/api/audit/{audit_id}/guided-fixes'
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    token = review.json()['review_id']
    assert client.delete(root + '/review/' + token).status_code == 204
    assert client.delete(root + '/review/' + token).status_code == 204
    assert _upload(client, root, source, ids, token).status_code == 409

@pytest.mark.parametrize('rule', [
    'FONT_SIZE', 'FONT_CONSISTENCY', 'ALIGNMENT',
    'LINE_SPACING', 'SPACE_BEFORE', 'SPACE_AFTER',
])
def test_every_rule_rejects_missing_paragraph_target(rule):
    audit = AuditRecord(id='source', filename='original.docx', status='completed')
    finding = Violation(
        id='finding', audit_id='source', rule_code=rule, severity='MINOR',
        location={'paragraph_index': 19, **({'run_index': 0} if rule.startswith('FONT') else {})},
        message='m', actual_value='11pt', expected_value='12pt',
    )
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _doc(), [finding])


@pytest.mark.parametrize('run_indexes', [[0, 1], [0, 0], [True], []])
def test_run_group_requires_complete_distinct_indexes(run_indexes):
    audit = AuditRecord(id='source', filename='original.docx', status='completed')
    finding = Violation(
        id='finding', audit_id='source', rule_code='FONT_SIZE', severity='MINOR',
        location={'paragraph_index': 0, 'run_indexes': run_indexes},
        message='m', actual_value='11pt', expected_value='12pt',
    )
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _doc(), [finding])

def test_reference_page_check_never_passes_without_rendered_evidence(monkeypatch):
    profile = new_custom_profile('Reference page evidence')
    profile.references.check_start_on_new_page = True
    snapshot = resolve_snapshot(profile)
    doc = Document()
    doc.add_paragraph('The body text preceding the Reference List.')
    doc.add_paragraph('References')
    doc.add_paragraph('Smith, A. (2020). A reference. Press.')
    buffer = io.BytesIO()
    doc.save(buffer)
    source = buffer.getvalue()
    findings = run_static_rules_engine(source, config=EffectiveProfileConfig(snapshot))
    audit = SimpleNamespace(
        profile_snapshot=snapshot.to_dict(), paragraph_page_mapping=None,
        violations=findings, weighted_score=calculate_weighted_score_detailed(findings).total,
    )
    from app.services import docx_pdf_converter
    def unavailable(_data):
        raise RuntimeError('renderer unavailable')
    monkeypatch.setattr(docx_pdf_converter, 'convert_docx_to_pdf', unavailable)
    with pytest.raises(GuidedFixFailed):
        verify_reaudit(audit, source, source, [])

@pytest.mark.parametrize('rule,attribute', [
    ('FONT_SIZE', 'szCs'), ('FONT_CONSISTENCY', 'asciiTheme'),
])
def test_alternate_font_slots_require_manual_correction(rule, attribute):
    document = Document(io.BytesIO(_doc()))
    run = document.paragraphs[0].runs[0]
    if rule == 'FONT_SIZE':
        from docx.oxml import OxmlElement
        alternate = OxmlElement('w:szCs')
        alternate.set(qn('w:val'), '22')
        run._r.rPr.append(alternate)
    else:
        run._r.rPr.rFonts.set(qn('w:asciiTheme'), 'minorHAnsi')
    buffer = io.BytesIO()
    document.save(buffer)
    finding = Violation(
        id='finding', audit_id='source', rule_code=rule, severity='MINOR',
        location={'paragraph_index': 0, 'run_index': 0}, message='m',
        actual_value='11pt' if rule == 'FONT_SIZE' else 'Times New Roman',
        expected_value='12pt' if rule == 'FONT_SIZE' else 'Arial',
    )
    audit = AuditRecord(id='source', filename='original.docx', status='completed')
    with pytest.raises(GuidedFixFailed, match='alternate font slot'):
        review_targets(audit, buffer.getvalue(), [finding])


def test_deleting_audit_revokes_unconsumed_output(client, monkeypatch, tmp_path):
    monkeypatch.setattr(guided_fix_storage.settings, 'GUIDED_FIX_STORAGE_DIR', str(tmp_path))
    source = _doc()
    audit = client.post('/api/audit', files={'file': ('original.docx', source)})
    audit_id = audit.json()['audit_id']
    ids = [v['id'] for v in audit.json()['physical_layout_errors'] if v['rule_code'] == 'FONT_SIZE']
    root = f'/api/audit/{audit_id}/guided-fixes'
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    assert len(list(tmp_path.glob('*.docx'))) == 1
    assert client.delete(f'/api/audit/{audit_id}').status_code == 200
    assert client.get(root + '/download/' + applied.json()['download_token']).status_code == 404
    assert list(tmp_path.glob('*.docx')) == []


@pytest.fixture
def bold_boundary_storage(monkeypatch, tmp_path):
    """Keep new candidate proofs offline and away from session/owner storage."""
    from app.api import routes
    from app.services import preview_storage

    monkeypatch.setattr(guided_fix_storage.settings, 'GUIDED_FIX_STORAGE_DIR', str(tmp_path / 'outputs'))
    monkeypatch.setattr(preview_storage.settings, 'PREVIEW_STORAGE_DIR', str(tmp_path / 'previews'))
    monkeypatch.setattr(guided_fix_storage, '_outputs', {})
    monkeypatch.setattr(routes, '_reviews', {})

    def no_renderer(_source):
        raise RuntimeError('Offline candidate test: rendering deliberately unavailable')

    monkeypatch.setattr(routes, 'convert_docx_to_pdf', no_renderer)
    return tmp_path / 'outputs'


def _bold_selection(audit):
    findings = [v for v in audit['physical_layout_errors'] if v['rule_code'] == BOLD_RULE]
    assert len(findings) == 1
    return f'/api/audit/{audit["audit_id"]}/guided-fixes', [findings[0]['id']]


def _saved_source_state(engine, audit_id):
    """Include every persisted Audit, Finding and citation field, not just score."""
    with Session(engine) as db:
        audit = db.get(AuditRecord, audit_id)

        def columns(row):
            return {column.name: copy.deepcopy(getattr(row, column.name))
                    for column in row.__table__.columns}

        return {
            'audit': columns(audit),
            'findings': sorted((columns(v) for v in audit.violations), key=lambda v: v['id']),
            'citations': sorted((columns(c) for c in audit.citation_issues), key=lambda c: c['id']),
        }


def _assert_no_bold_output(response, root, audit_id):
    assert 'download_token' not in response.json()
    assert 'review_id' not in response.json()
    assert not list(root.glob('*.docx'))
    assert not any(output.audit_id == audit_id for output in guided_fix_storage._outputs.values())


@pytest.mark.parametrize('mutation_fails', [False, True])
def test_reference_bold_and_alignment_selection_is_atomic_and_source_views_stay_saved(
    client, test_engine, monkeypatch, bold_boundary_storage, mutation_fails,
):
    from app.services import guided_fix
    from pypdf import PdfReader

    source = _bold_bytes(_bold_doc(segmented=True))
    audit = _bold_audit(client, source, alignment=True)
    root, bold_ids = _bold_selection(audit)
    alignment_ids = [v['id'] for v in audit['physical_layout_errors']
                     if v['rule_code'] == 'REFERENCE_HEADING_ALIGNMENT']
    assert len(alignment_ids) == 1
    ids = alignment_ids + bold_ids
    audit_path = f'/api/audit/{audit["audit_id"]}'
    saved = _saved_source_state(test_engine, audit['audit_id'])
    detail = client.get(audit_path).json()
    history = client.get('/api/audits').json()

    def report_text():
        response = client.get(audit_path + '/export-pdf')
        assert response.status_code == 200, response.text
        return '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(response.content)).pages)

    pdf_text = report_text()
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    assert {item['violation_id'] for item in review.json()['items']} == set(ids)
    token = review.json()['review_id']
    subset = _upload(client, root, source, bold_ids, token)
    assert subset.status_code == 422
    _assert_no_bold_output(subset, bold_boundary_storage, audit['audit_id'])
    if mutation_fails:
        monkeypatch.setattr(guided_fix, '_fix_reference_heading_bold', lambda *_args: False)
    applied = _upload(client, root, source, ids, token)
    if mutation_fails:
        assert applied.status_code == 422, applied.text
        _assert_no_bold_output(applied, bold_boundary_storage, audit['audit_id'])
    else:
        assert applied.status_code == 200, applied.text
        assert set(applied.json()['resolved']) == set(ids)
        assert applied.json()['new_findings'] == []
        download = client.get(root + '/download/' + applied.json()['download_token'])
        assert download.status_code == 200
        heading = Document(io.BytesIO(download.content)).paragraphs[1]
        assert heading.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert all(run.font.bold is True for run in heading.runs if run.text != '')
        assert heading.text == Document(io.BytesIO(source)).paragraphs[1].text
    assert _saved_source_state(test_engine, audit['audit_id']) == saved
    assert client.get(audit_path).json() == detail
    assert client.get('/api/audits').json() == history
    assert report_text() == pdf_text
    original_heading = Document(io.BytesIO(source)).paragraphs[1]
    assert original_heading.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert all(run.font.bold is False for run in original_heading.runs if run.text != '')


@pytest.mark.parametrize('variant,reason', [
    ('duplicate', 'ambiguous_heading'),
    ('inherited', 'unverifiable_bold'),
    ('unresolved', 'unverifiable_bold'),
    ('signed', 'unsupported_package'),
    ('protected', 'unsupported_package'),
    ('macro', 'unsupported_package'),
])
def test_reference_bold_review_safe_reasons_require_genuine_saved_source(
    client, test_engine, bold_boundary_storage, variant, reason,
):
    from app.services.guided_fix import ReferenceHeadingBoldUnavailable

    doc = _bold_doc(segmented=True)
    if variant == 'duplicate':
        doc.add_paragraph('References').runs[0].font.bold = False
    elif variant == 'inherited':
        doc.styles['Normal'].font.bold = False
        for run in doc.paragraphs[1].runs:
            run.font.bold = None
    elif variant == 'unresolved':
        doc.paragraphs[1].runs[1].font.bold = None
    elif variant == 'protected':
        doc.settings.element.append(OxmlElement('w:documentProtection'))
    source = _bold_bytes(doc)
    if variant in {'signed', 'macro'}:
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(output, 'w') as after:
            for item in before.infolist():
                after.writestr(item, before.read(item.filename))
            after.writestr('_xmlsignatures/sig1.xml' if variant == 'signed'
                           else 'word/vbaProject.bin', b'<marker/>')
        source = output.getvalue()
    audit = _bold_audit(client, source)
    with Session(test_engine) as db:
        saved = db.get(AuditRecord, audit['audit_id'])
        assert saved.document_sha256 == hashlib.sha256(source).hexdigest()
        finding = next((v for v in saved.violations if v.rule_code == BOLD_RULE), None)
        if finding is None:
            # Detector silence is not candidate safety evidence. Seed only
            # this test's persisted Finding, never fabricate the source hash.
            finding = Violation(
                audit_id=saved.id, rule_code=BOLD_RULE, severity='MINOR',
                location={'paragraph_index': 1}, message='Persisted heading bold finding',
                actual_value='Not bold', expected_value='Bold',
            )
            db.add(finding)
            db.commit()
        finding_id = finding.id
    before = _saved_source_state(test_engine, audit['audit_id'])
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    response = _upload(client, root + '/review', source, [finding_id])
    assert response.status_code == 422, response.text
    assert response.json() == {'detail': ReferenceHeadingBoldUnavailable.DETAILS[reason]}
    _assert_no_bold_output(response, bold_boundary_storage, audit['audit_id'])
    assert _saved_source_state(test_engine, audit['audit_id']) == before


def test_reference_bold_internal_review_exception_is_not_a_public_reason(
    client, monkeypatch, bold_boundary_storage,
):
    from app.api import routes

    source = _bold_bytes(_bold_doc())
    audit = _bold_audit(client, source)
    root, ids = _bold_selection(audit)

    def internal_error(*_args):
        raise RuntimeError('word/document.xml C:/private/source.docx SQL secret heading text')

    monkeypatch.setattr(routes, 'review_targets', internal_error)
    response = _upload(client, root + '/review', source, ids)
    assert response.status_code == 422
    assert response.json() == {'detail': 'Fix cannot be applied safely. Manual correction required.'}
    _assert_no_bold_output(response, bold_boundary_storage, audit['audit_id'])


@pytest.mark.parametrize('change,status', [
    ('wrong_original', 409), ('stale_actual', 422), ('stale_expected', 422),
    ('disabled', 422), ('future_snapshot', 409), ('no_snapshot', 409), ('no_hash', 409),
])
@pytest.mark.parametrize('phase', ['review', 'apply'])
def test_reference_bold_saved_candidate_gates_fail_closed(
    client, test_engine, bold_boundary_storage, change, status, phase,
):
    source = _bold_bytes(_bold_doc())
    audit = _bold_audit(client, source)
    root, ids = _bold_selection(audit)
    review_id = None
    if phase == 'apply':
        review = _upload(client, root + '/review', source, ids)
        assert review.status_code == 200, review.text
        review_id = review.json()['review_id']
    with Session(test_engine) as db:
        saved = db.get(AuditRecord, audit['audit_id'])
        finding = db.get(Violation, ids[0])
        if change == 'stale_actual':
            finding.actual_value = 'Bold'
        elif change == 'stale_expected':
            finding.expected_value = 'Not bold'
        elif change == 'disabled':
            profile = new_custom_profile('Disabled saved bold requirement')
            profile.references.check_heading_bold = False
            profile.references.heading_bold = True
            saved.profile_snapshot = resolve_snapshot(profile).to_dict()
        elif change == 'future_snapshot':
            snapshot = copy.deepcopy(saved.profile_snapshot)
            snapshot['schema_version'] = 999
            saved.profile_snapshot = snapshot
        elif change == 'no_snapshot':
            saved.profile_snapshot = None
        elif change == 'no_hash':
            saved.document_sha256 = None
        db.commit()
    before = _saved_source_state(test_engine, audit['audit_id'])
    if change == 'wrong_original':
        wrong = _bold_doc()
        wrong.paragraphs[0].add_run(' Different original.')
        source = _bold_bytes(wrong)
    response = _upload(client, root + ('/review' if phase == 'review' else ''), source, ids, review_id)
    assert response.status_code == status, response.text
    _assert_no_bold_output(response, bold_boundary_storage, audit['audit_id'])
    assert _saved_source_state(test_engine, audit['audit_id']) == before


def test_reference_bold_uses_saved_snapshot_after_current_profile_changes(
    client, test_engine, monkeypatch, bold_boundary_storage,
):
    from app.api import routes
    from app.services import profile_registry

    profile = new_custom_profile('Mutable current profile')
    profile.references.check_heading_bold = True
    profile.references.heading_bold = True
    # Replace the catalog with a task-owned profile, not an owner/global object.
    # Later edits must not affect this audit's saved formatting requirement.
    monkeypatch.setattr(profile_registry, 'BUILTIN_PROFILES', {profile.profile_id: profile})
    source = _bold_bytes(_bold_doc())
    response = client.post('/api/audit', params={'profile_id': profile.profile_id},
                           files={'file': ('original.docx', source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    root, ids = _bold_selection(audit)
    saved = _saved_source_state(test_engine, audit['audit_id'])
    profile.references.heading_bold = False
    profile.references.check_heading_bold = False
    profile.profile_name = 'Changed current profile'

    def no_resolution(**_kwargs):
        pytest.fail('Guided correction must not re-resolve a current profile')

    monkeypatch.setattr(routes, 'resolve_request_profile', no_resolution)
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    assert review.json()['items'][0]['after'] == 'Bold'
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    output = client.get(root + '/download/' + applied.json()['download_token'])
    assert output.status_code == 200
    assert Document(io.BytesIO(output.content)).paragraphs[1].runs[0].font.bold is True
    assert _saved_source_state(test_engine, audit['audit_id']) == saved


@pytest.mark.parametrize('failure', [
    'italic_mutation', 'reaudit_unavailable', 'baseline_drift',
    'new_finding', 'category_increase', 'storage_failure',
])
def test_reference_bold_never_publishes_unverified_or_unstored_copy(
    client, test_engine, monkeypatch, bold_boundary_storage, failure,
):
    from app.api import routes
    from app.services import guided_fix, guided_fix_verification
    from app.services.layout_violation import LayoutViolation

    source = _bold_bytes(_bold_doc(segmented=True))
    audit = _bold_audit(client, source)
    root, ids = _bold_selection(audit)
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    injected = []
    verification_errors = []
    original_verification = routes.verify_reaudit

    def observe_verification(*args):
        try:
            return original_verification(*args)
        except GuidedFixFailed as exc:
            verification_errors.append(str(exc))
            raise

    monkeypatch.setattr(routes, 'verify_reaudit', observe_verification)
    if failure == 'italic_mutation':
        original_fix = guided_fix._fix_reference_heading_bold

        def change_italic(run, expected):
            injected.append('italic')
            result = original_fix(run, expected)
            run.font.italic = not run.font.italic
            return result

        monkeypatch.setattr(guided_fix, '_fix_reference_heading_bold', change_italic)
    elif failure == 'baseline_drift':
        with Session(test_engine) as db:
            saved = db.get(AuditRecord, audit['audit_id'])
            saved.weighted_score -= 1
            db.commit()
        injected.append('persisted baseline score')
    elif failure in {'reaudit_unavailable', 'new_finding'}:
        original_engine = guided_fix_verification.run_static_rules_engine

        def changed_engine(data, **kwargs):
            injected.append(data)
            if failure == 'reaudit_unavailable':
                raise RuntimeError('C:/private/converter.docx internal re-audit failure')
            findings = original_engine(data, **kwargs)
            if data != source:
                findings.append(LayoutViolation(
                    rule_code='FONT_SIZE', severity='MINOR', location={'paragraph_index': 0, 'run_index': 0},
                    message='New unrelated finding', expected_value='12pt', actual_value='9pt',
                ))
            return findings

        monkeypatch.setattr(guided_fix_verification, 'run_static_rules_engine', changed_engine)
    elif failure == 'category_increase':
        original_score = guided_fix_verification.calculate_weighted_score_detailed

        def increased_category(findings):
            score = original_score(findings)
            injected.append(score)
            if len(injected) == 2:
                breakdown = [replace(row, deduction=row.deduction + 1)
                             if row.category == 'page_margins' else row for row in score.breakdown]
                return replace(score, total=score.total - 1, breakdown=breakdown)
            return score

        monkeypatch.setattr(guided_fix_verification, 'calculate_weighted_score_detailed', increased_category)
    else:
        def disk_failure(_fd):
            injected.append('fsync')
            raise OSError('C:/private/outputs disk write failed')

        monkeypatch.setattr(guided_fix_storage.os, 'fsync', disk_failure)
    if failure != 'storage_failure':
        def forbidden_publication(*_args):
            pytest.fail('Output storage must not be called before all verification gates pass')

        monkeypatch.setattr(guided_fix_storage, 'store', forbidden_publication)
    before = _saved_source_state(test_engine, audit['audit_id'])
    response = _upload(client, root, source, ids, review.json()['review_id'])
    assert response.status_code == (503 if failure == 'storage_failure' else 422), response.text
    assert injected
    if failure == 'category_increase':
        assert len(injected) == 2
    if failure == 'new_finding':
        assert len(injected) == 2
        assert injected[0] == source and injected[1] != source
    expected_error = {
        'reaudit_unavailable': 'deterministic re-audit unavailable',
        'baseline_drift': 'persisted score differs from baseline',
        'new_finding': 'new or worsened finding after re-audit',
        'category_increase': 'category deduction increased',
    }.get(failure)
    if expected_error is not None:
        assert verification_errors == [expected_error]
    if failure in {'reaudit_unavailable', 'baseline_drift', 'new_finding', 'category_increase'}:
        assert response.json() == {
            'detail': 'Fix applied; verification failed. No file was produced.',
        }
    elif failure == 'storage_failure':
        assert response.json() == {'detail': 'Verified copy could not be stored. Try again.'}
    _assert_no_bold_output(response, bold_boundary_storage, audit['audit_id'])
    assert _saved_source_state(test_engine, audit['audit_id']) == before


@pytest.fixture
def spacing_boundary_storage(bold_boundary_storage):
    return bold_boundary_storage


def _spacing_document():
    document = Document()
    paragraph = document.add_paragraph('A body paragraph with independently checked spacing.')
    paragraph.paragraph_format.space_before = Pt(20)
    paragraph.paragraph_format.space_after = Pt(20)
    paragraph.runs[0].font.name = 'Times New Roman'
    paragraph.runs[0].font.size = Pt(11)
    return document


def _spacing_layer(document, layer):
    from docx.enum.style import WD_STYLE_TYPE

    paragraph = document.paragraphs[0]
    if layer == 'direct':
        return paragraph._p.get_or_add_pPr()
    if layer in {'style', 'base'}:
        base = document.styles.add_style('Spacing proof base', WD_STYLE_TYPE.PARAGRAPH)
        used = document.styles.add_style('Spacing proof used', WD_STYLE_TYPE.PARAGRAPH)
        used.base_style = base
        paragraph.style = used
        return (used if layer == 'style' else base)._element.get_or_add_pPr()
    assert layer == 'default'
    styles = document.styles.element
    defaults = styles.find(qn('w:docDefaults'))
    if defaults is None:
        defaults = OxmlElement('w:docDefaults')
        styles.insert(0, defaults)
    wrapper = defaults.find(qn('w:pPrDefault'))
    if wrapper is None:
        wrapper = OxmlElement('w:pPrDefault')
        defaults.append(wrapper)
    properties = wrapper.find(qn('w:pPr'))
    if properties is None:
        properties = OxmlElement('w:pPr')
        wrapper.append(properties)
    return properties


def _spacing_control(properties, side, control, value):
    if control == 'contextualSpacing':
        node = OxmlElement('w:contextualSpacing')
        if value is not None:
            node.set(qn('w:val'), value)
        properties.append(node)
    else:
        spacing = properties.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            properties.append(spacing)
        spacing.set(qn('w:' + side + control), value)


def _spacing_audit(client, source, *, before=6, after=6):
    profile = new_custom_profile('Direct point spacing proof')
    profile.body.space_before_pt = before
    profile.body.space_after_pt = after
    profile.lists.space_after_pt = after
    response = client.post('/api/audit', data={'custom_profile': json.dumps(profile.to_dict())},
                           files={'file': ('spacing.docx', source)})
    assert response.status_code == 200, response.text
    return response.json()


def _spacing_selection(audit, sides=('before', 'after')):
    codes = {'SPACE_' + side.upper() for side in sides}
    findings = [finding for finding in audit['physical_layout_errors']
                if finding['rule_code'] in codes]
    assert {finding['rule_code'] for finding in findings} == codes
    assert all(finding['location']['paragraph_index'] == 0 for finding in findings)
    return f'/api/audit/{audit["audit_id"]}/guided-fixes', [finding['id'] for finding in findings]


def _assert_no_spacing_output(response, output_root, audit_id):
    from app.api import routes

    _assert_no_bold_output(response, output_root, audit_id)
    assert not any(review[0] == audit_id for review in routes._reviews.values())


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('layer', ['direct', 'style', 'base', 'default'])
@pytest.mark.parametrize('control,value', [
    ('Lines', '100'), ('Lines', '0'), ('Autospacing', '1'),
    ('contextualSpacing', '1'), ('contextualSpacing', None),
])
def test_spacing_override_never_publishes_a_copy(
    client, spacing_boundary_storage, side, layer, control, value,
):
    document = _spacing_document()
    _spacing_control(_spacing_layer(document, layer), side, control, value)
    source = _bold_bytes(document)
    digest = hashlib.sha256(source).digest()
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    response = _upload(client, root + '/review', source, ids)
    assert response.status_code == 422, response.text
    assert response.json() == {'detail': 'Fix cannot be applied safely. Manual correction required.'}
    _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])
    assert hashlib.sha256(source).digest() == digest


def _spacing_finding_key(finding):
    return (finding.rule_code, finding.severity, json.dumps(finding.location, sort_keys=True),
            finding.actual_value, finding.expected_value)


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('required', [0, 0.05, 0.1, 0.5, 1, 6, 6.05, 6.1, 12, 239.95, 240])
def test_point_spacing_review_apply_download_and_reaudit(
    client, test_engine, spacing_boundary_storage, side, required,
):
    from pypdf import PdfReader
    from app.services.profile_snapshot import snapshot_from_dict
    from tests.test_guided_fix_preservation_storage import (
        _assert_spacing_package_unchanged, _spacing_preservation_document,
    )

    document = _spacing_preservation_document()
    opposite = 'after' if side == 'before' else 'before'
    # Opposite-side overrides are irrelevant to the selected side and protected.
    _spacing_control(document.paragraphs[0]._p.pPr, opposite, 'Lines', '100')
    _spacing_control(document.paragraphs[0]._p.pPr, opposite, 'Autospacing', '1')
    source = _bold_bytes(document)
    original = bytes(source)
    digest = hashlib.sha256(source).hexdigest()
    required_twips = int(Decimal(str(required)) * 20)
    audit = _spacing_audit(client, source, **{side: required})
    root, ids = _spacing_selection(audit, (side,))
    path = f'/api/audit/{audit["audit_id"]}'
    saved = _saved_source_state(test_engine, audit['audit_id'])
    detail, history = client.get(path).json(), client.get('/api/audits').json()

    def report_text():
        response = client.get(path + '/export-pdf')
        assert response.status_code == 200, response.text
        return '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(response.content)).pages)

    report = report_text()
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    assert [item['violation_id'] for item in review.json()['items']] == ids
    assert review.json()['items'][0]['locations'] == ['Paragraph 1']
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    assert applied.json()['status'] == 'passed'
    assert applied.json()['resolved'] == ids
    assert applied.json()['new_findings'] == []
    assert _upload(client, root, source, ids, review.json()['review_id']).status_code == 409
    download_path = root + '/download/' + applied.json()['download_token']
    download = client.get(download_path)
    assert download.status_code == 200
    assert download.headers['cache-control'] == 'private, no-store'
    assert client.get(download_path).status_code == 410
    assert download.content != source
    assert hashlib.sha256(download.content).hexdigest() != digest
    _assert_spacing_package_unchanged(source, download.content, {0: {side: required_twips}})
    output = Document(io.BytesIO(download.content))
    spacing = output.paragraphs[0]._p.pPr.find(qn('w:spacing'))
    assert spacing.get(qn('w:' + side)) == str(required_twips)
    assert spacing.get(qn('w:' + opposite + 'Lines')) == '100'
    assert spacing.get(qn('w:' + opposite + 'Autospacing')) == '1'
    config = EffectiveProfileConfig(snapshot_from_dict(saved['audit']['profile_snapshot']))
    before = run_static_rules_engine(source, config=config)
    after = run_static_rules_engine(download.content, config=config)
    code = 'SPACE_' + side.upper()
    selected = [finding for finding in before if finding.rule_code == code
                and finding.location == {'paragraph_index': 0}]
    assert len(selected) == 1
    assert not [finding for finding in after if finding.rule_code == code
                and finding.location == {'paragraph_index': 0}]
    assert sorted(map(_spacing_finding_key, after)) == sorted(
        _spacing_finding_key(finding) for finding in before if finding not in selected
    )
    before_score, after_score = [calculate_weighted_score_detailed(findings) for findings in (before, after)]
    assert after_score.total >= before_score.total
    deductions = {row.category: row.deduction for row in before_score.breakdown}
    assert all(row.deduction <= deductions[row.category] for row in after_score.breakdown)
    assert source == original and hashlib.sha256(source).hexdigest() == digest
    assert saved['audit']['document_sha256'] == digest
    assert _saved_source_state(test_engine, audit['audit_id']) == saved
    assert client.get(path).json() == detail
    assert client.get('/api/audits').json() == history
    assert report_text() == report


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('required', [6.00000000001, 6.025, 6.001, 0.01, 0.025, 1.333])
def test_spacing_unrepresentable_requirement_remains_manual(
    client, test_engine, spacing_boundary_storage, side, required,
):
    source = _bold_bytes(_spacing_document())
    digest = hashlib.sha256(source).hexdigest()
    audit = _spacing_audit(client, source, **{side: required})
    root, ids = _spacing_selection(audit, (side,))
    saved = _saved_source_state(test_engine, audit['audit_id'])
    detail = client.get(f'/api/audit/{audit["audit_id"]}').json()
    history = client.get('/api/audits').json()
    response = _upload(client, root + '/review', source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])
    assert hashlib.sha256(source).hexdigest() == digest
    assert _saved_source_state(test_engine, audit['audit_id']) == saved
    assert client.get(f'/api/audit/{audit["audit_id"]}').json() == detail
    assert client.get('/api/audits').json() == history


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('control', ['Autospacing', 'contextualSpacing'])
@pytest.mark.parametrize('false_value', ['0', 'false', 'off'])
@pytest.mark.parametrize('inherited_layer', ['style', 'base', 'default'])
def test_spacing_valid_false_overrides_inherited_true(
    client, spacing_boundary_storage, side, control, false_value, inherited_layer,
):
    document = _spacing_document()
    _spacing_control(_spacing_layer(document, inherited_layer), side, control, 'true')
    _spacing_control(_spacing_layer(document, 'direct'), side, control, false_value)
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    output = Document(io.BytesIO(download.content))
    properties = output.paragraphs[0]._p.pPr
    assert properties.find(qn('w:spacing')).get(qn('w:' + side)) == '120'
    if control == 'contextualSpacing':
        assert properties.find(qn('w:contextualSpacing')).get(qn('w:val')) == false_value
    else:
        assert properties.find(qn('w:spacing')).get(qn('w:' + side + control)) == false_value


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('control', ['Autospacing', 'contextualSpacing'])
def test_spacing_false_does_not_hide_malformed_inherited_control(
    client, spacing_boundary_storage, side, control,
):
    document = _spacing_document()
    _spacing_control(_spacing_layer(document, 'base'), side, control, 'invalid')
    _spacing_control(_spacing_layer(document, 'direct'), side, control, '0')
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    response = _upload(client, root + '/review', source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])


@pytest.mark.parametrize('unsafe', [False, True])
def test_spacing_before_after_selection_is_atomic(client, spacing_boundary_storage, unsafe):
    document = _spacing_document()
    if unsafe:
        _spacing_control(_spacing_layer(document, 'direct'), 'after', 'Autospacing', 'on')
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source, before=6, after=0)
    root, ids = _spacing_selection(audit)
    response = _upload(client, root + '/review', source, ids)
    if unsafe:
        assert response.status_code == 422, response.text
        _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])
        return
    assert response.status_code == 200, response.text
    assert {item['violation_id'] for item in response.json()['items']} == set(ids)
    applied = _upload(client, root, source, ids, response.json()['review_id'])
    assert applied.status_code == 200, applied.text
    assert set(applied.json()['resolved']) == set(ids)
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    spacing = Document(io.BytesIO(download.content)).paragraphs[0]._p.pPr.find(qn('w:spacing'))
    assert spacing.get(qn('w:before')) == '120'
    assert spacing.get(qn('w:after')) == '0'


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('fault', [
    'auto', 'lines', 'contextual', 'one_twip', 'minus_twip', 'rounded', 'missing',
    'sibling', 'line', 'lineRule', 'text', 'run', 'role', 'bookmark', 'partial',
])
def test_spacing_output_override_is_rejected(
    client, monkeypatch, spacing_boundary_storage, side, fault,
):
    from app.services import guided_fix
    from app.services.profile_snapshot import snapshot_from_dict

    document = _spacing_document()
    if fault == 'bookmark':
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        start.set(qn('w:id'), '7')
        start.set(qn('w:name'), 'OriginalBookmark')
        end.set(qn('w:id'), '7')
        document.paragraphs[0]._p.extend((start, end))
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source, **({side: 6.05} if fault == 'rounded' else {}))
    root, ids = _spacing_selection(audit, ('before', 'after') if fault == 'partial' else (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    original_mutator = getattr(guided_fix, '_fix_space_' + side)
    original_validator = guided_fix._validate_output
    errors, silent_outputs = [], []

    def faulty_mutator(paragraph, required):
        if fault == 'partial':
            return True
        assert original_mutator(paragraph, required)
        spacing = paragraph._p.pPr.find(qn('w:spacing'))
        if fault in {'one_twip', 'minus_twip'}:
            spacing.set(qn('w:' + side), str(required + (1 if fault == 'one_twip' else -1)))
        elif fault == 'rounded':
            spacing.set(qn('w:' + side), '120')  # A 6.025pt rounded result is not the required 121 twips.
        elif fault == 'missing':
            del spacing.attrib[qn('w:' + side)]
        elif fault == 'sibling':
            spacing.set(qn('w:' + ('after' if side == 'before' else 'before')), '399')
        elif fault in {'line', 'lineRule'}:
            spacing.set(qn('w:' + fault), '240' if fault == 'line' else 'exact')
        elif fault == 'text':
            paragraph.runs[0].text += ' changed'
        elif fault == 'run':
            paragraph.add_run('extra run')
        elif fault == 'role':
            paragraph.style = 'Heading 1'
        elif fault == 'bookmark':
            paragraph._p.find(qn('w:bookmarkStart')).set(qn('w:name'), 'ChangedBookmark')
        else:
            control = {'auto': 'Autospacing', 'lines': 'Lines', 'contextual': 'contextualSpacing'}[fault]
            _spacing_control(paragraph._p.pPr, side, control, '1')
        return True

    def inspected_validator(data, *args, **kwargs):
        if fault in {'auto', 'lines', 'contextual', 'one_twip', 'minus_twip', 'rounded'}:
            with zipfile.ZipFile(io.BytesIO(data)) as package:
                assert package.read('word/document.xml')
            # The real detector's silence must not substitute for raw output proof.
            snapshot = snapshot_from_dict(args[3].profile_snapshot)
            findings = run_static_rules_engine(data, config=EffectiveProfileConfig(snapshot))
            assert not [finding for finding in findings if finding.rule_code == 'SPACE_' + side.upper()]
            silent_outputs.append(data)
        try:
            return original_validator(data, *args, **kwargs)
        except GuidedFixFailed as exc:
            errors.append(str(exc))
            raise

    def forbidden_publication(*_args):
        pytest.fail('Unsafe spacing output reached storage')

    monkeypatch.setattr(guided_fix, '_fix_space_' + side, faulty_mutator)
    monkeypatch.setattr(guided_fix, '_validate_output', inspected_validator)
    monkeypatch.setattr(guided_fix_storage, 'store', forbidden_publication)
    response = _upload(client, root, source, ids, review.json()['review_id'])
    assert response.status_code == 422, response.text
    assert errors
    assert bool(silent_outputs) is (fault in {'auto', 'lines', 'contextual', 'one_twip', 'minus_twip', 'rounded'})
    _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('variant', [
    'paired', 'field', 'section', 'invalid_id', 'negative_id', 'unicode_id', 'duplicate_start',
    'duplicate_end', 'cross_paragraph', 'reversed',
])
def test_point_spacing_structure_and_bookmark_review(
    client, spacing_boundary_storage, side, variant,
):
    document = _spacing_document()
    paragraph = document.paragraphs[0]
    if variant == 'field':
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'DATE')
        paragraph._p.append(field)
    elif variant == 'section':
        paragraph._p.pPr.append(OxmlElement('w:sectPr'))
    else:
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        start.set(qn('w:id'), {'invalid_id': 'bad', 'negative_id': '-1', 'unicode_id': '١'}.get(variant, '7'))
        start.set(qn('w:name'), 'SpacingBookmark')
        end.set(qn('w:id'), '7')
        if variant == 'reversed':
            paragraph._p.append(end)
            paragraph._p.append(start)
        else:
            paragraph._p.append(start)
            if variant == 'duplicate_start':
                paragraph._p.append(copy.deepcopy(start))
            if variant == 'cross_paragraph':
                document.add_paragraph('Bookmark continuation.')._p.append(end)
            else:
                paragraph._p.append(end)
            if variant == 'duplicate_end':
                paragraph._p.append(copy.deepcopy(end))
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    if variant != 'paired':
        assert review.status_code == 422, review.text
        _assert_no_spacing_output(review, spacing_boundary_storage, audit['audit_id'])
        return
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    output = Document(io.BytesIO(download.content)).paragraphs[0]
    assert output._p.pPr.find(qn('w:spacing')).get(qn('w:' + side)) == '120'
    assert output._p.find(qn('w:bookmarkStart')).get(qn('w:id')) == '7'
    assert output._p.find(qn('w:bookmarkEnd')).get(qn('w:id')) == '7'



@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('difference', [19, 20, 21])
def test_point_spacing_exact_detector_threshold(side, difference):
    from app.services.guided_fix import _current_matches_expected

    paragraph = _spacing_document().paragraphs[0]
    paragraph._p.pPr.find(qn('w:spacing')).set(qn('w:' + side), str(68 + difference))
    assert _current_matches_expected(
        'SPACE_' + side.upper(), 68, {}, None, paragraph,
    ) is (difference <= 20)


@pytest.mark.parametrize('side', ['before', 'after'])
def test_point_spacing_large_actual_and_float_projection(side):
    from app.services.guided_fix import _current_changed_from_actual

    paragraph = _spacing_document().paragraphs[0]
    spacing = paragraph._p.pPr.find(qn('w:spacing'))
    code = 'SPACE_' + side.upper()
    spacing.set(qn('w:' + side), '6000')
    assert not _current_changed_from_actual(code, {}, None, '300pt', paragraph)
    spacing.set(qn('w:' + side), '9007199254740993')
    assert _current_changed_from_actual(code, {}, None, '450359962737049.6pt', paragraph)


@pytest.mark.parametrize('value', [
    '6.00000000001pt', '6.025pt', '6.001pt', '0.01pt', '0.025pt',
    '1.333pt', '-1pt', '-0pt', '-0.0pt', '1e-1pt', '240.05pt', 'NaNpt', 'infpt', 'nonsense', '',
    '6.00000000000000000000000000000000000001pt',
])
def test_point_spacing_requirement_parser_refuses_inexact_or_invalid(value):
    from app.services.guided_fix import _point_twips

    assert _point_twips(value, requirement=True) is None


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('spelling', ['6pt', '6.0pt', '6.00pt'])
def test_point_spacing_expected_numeric_equivalence_at_review(
    client, test_engine, spacing_boundary_storage, side, spelling,
):
    source = _bold_bytes(_spacing_document())
    audit = _spacing_audit(client, source, **{side: 6})
    root, ids = _spacing_selection(audit, (side,))
    with Session(test_engine) as db:
        db.get(Violation, ids[0]).expected_value = spelling
        db.commit()
    before = _saved_source_state(test_engine, audit['audit_id'])
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    assert _saved_source_state(test_engine, audit['audit_id']) == before
    with Session(test_engine) as db:
        assert db.get(Violation, ids[0]).expected_value == spelling


def _point_snapshot_audit(side, version=6):
    profile = new_custom_profile('Saved point spacing evidence')
    profile.body.space_before_pt = 6
    profile.body.space_after_pt = 6
    snapshot = replace(resolve_snapshot(profile), schema_version=version, fingerprint='').with_fingerprint()
    source = _bold_bytes(_spacing_document())
    findings = run_static_rules_engine(source, config=EffectiveProfileConfig(snapshot))
    for index, finding in enumerate(findings):
        finding.id = f'point-{index}'
    selected = [finding for finding in findings if finding.rule_code == 'SPACE_' + side.upper()]
    assert len(selected) == 1
    return source, SimpleNamespace(
        profile_snapshot=snapshot.to_dict(), paragraph_page_mapping=None, violations=findings,
        weighted_score=calculate_weighted_score_detailed(findings).total,
        document_sha256=hashlib.sha256(source).hexdigest(),
    ), selected


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('version', [1, 2, 3, 4, 5, 6])
def test_point_spacing_saved_snapshot_versions(side, version):
    source, audit, selected = _point_snapshot_audit(side, version)
    assert review_targets(audit, source, selected)[0]


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('value', [None, True, -1, 240.05, float('inf'), float('nan'), '6'])
def test_point_spacing_invalid_saved_requirement_is_manual(side, value):
    from app.services.profile_snapshot import snapshot_from_dict

    source, audit, selected = _point_snapshot_audit(side)
    snapshot = snapshot_from_dict(audit.profile_snapshot)
    audit.profile_snapshot = replace(
        snapshot, **{'body_space_' + side + '_pt': value}, fingerprint='',
    ).with_fingerprint().to_dict()
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('failure', ['preservation', 'reaudit_unavailable', 'new_finding', 'category_increase', 'storage_failure'])
def test_point_spacing_publication_barriers_preserve_saved_state(
    client, test_engine, monkeypatch, spacing_boundary_storage, side, failure,
):
    from app.api import routes
    from app.services import guided_fix_verification
    from app.services.layout_violation import LayoutViolation

    source = _bold_bytes(_spacing_document())
    digest = hashlib.sha256(source).hexdigest()
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    saved = _saved_source_state(test_engine, audit['audit_id'])
    history = client.get('/api/audits').json()
    called = []
    if failure == 'preservation':
        def fail_preservation(*_args):
            called.append(True)
            raise ValueError('injected package-preservation failure')

        monkeypatch.setattr(routes.docx_preservation, 'compare_preservation', fail_preservation)
    elif failure in {'reaudit_unavailable', 'new_finding'}:
        original_engine = guided_fix_verification.run_static_rules_engine

        def changed_reaudit(data, **kwargs):
            called.append(data)
            if failure == 'reaudit_unavailable':
                raise RuntimeError('injected re-audit failure')
            findings = original_engine(data, **kwargs)
            if data != source:
                findings.append(LayoutViolation(
                    rule_code='FONT_SIZE', severity='MINOR',
                    location={'paragraph_index': 0, 'run_index': 0},
                    message='New unrelated finding', expected_value='12pt', actual_value='9pt',
                ))
            return findings

        monkeypatch.setattr(guided_fix_verification, 'run_static_rules_engine', changed_reaudit)
    elif failure == 'category_increase':
        original_score = guided_fix_verification.calculate_weighted_score_detailed

        def increased_category(findings):
            score = original_score(findings)
            called.append(score)
            if len(called) == 2:
                breakdown = [replace(row, deduction=row.deduction + 1)
                             if row.category == 'page_margins' else row for row in score.breakdown]
                return replace(score, total=score.total - 1, breakdown=breakdown)
            return score

        monkeypatch.setattr(guided_fix_verification, 'calculate_weighted_score_detailed', increased_category)
    else:
        def fail_fsync(_fd):
            called.append(True)
            raise OSError('injected storage failure')

        monkeypatch.setattr(guided_fix_storage.os, 'fsync', fail_fsync)
    response = _upload(client, root, source, ids, review.json()['review_id'])
    assert response.status_code == (503 if failure == 'storage_failure' else 422), response.text
    assert called
    _assert_no_bold_output(response, spacing_boundary_storage, audit['audit_id'])
    assert hashlib.sha256(source).hexdigest() == digest
    assert _saved_source_state(test_engine, audit['audit_id']) == saved
    assert client.get('/api/audits').json() == history


@pytest.mark.parametrize('side', ['before', 'after'])
def test_point_spacing_missing_page_evidence_never_publishes(
    client, test_engine, monkeypatch, spacing_boundary_storage, side,
):
    from app.services import docx_pdf_converter

    document = _spacing_document()
    document.add_paragraph('References', style='Heading 1')
    document.add_paragraph('Smith, A. (2024). Synthetic reference. Academic Press.')
    source = _bold_bytes(document)
    profile = new_custom_profile('Point spacing with reference page evidence')
    setattr(profile.body, 'space_' + side + '_pt', 6)
    profile.references.check_start_on_new_page = True
    response = client.post('/api/audit', data={'custom_profile': json.dumps(profile.to_dict())},
                           files={'file': ('spacing.docx', source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    saved = _saved_source_state(test_engine, audit['audit_id'])
    assert saved['audit']['paragraph_page_mapping'] is None
    attempts = []

    def unavailable_renderer(_source):
        attempts.append(True)
        raise RuntimeError('synthetic page evidence unavailable')

    monkeypatch.setattr(docx_pdf_converter, 'convert_docx_to_pdf', unavailable_renderer)
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 422, applied.text
    assert attempts
    _assert_no_bold_output(applied, spacing_boundary_storage, audit['audit_id'])
    assert _saved_source_state(test_engine, audit['audit_id']) == saved


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('fault', [None, 'unchanged', 'wrong'])
def test_point_spacing_grouped_paragraphs_all_or_nothing(
    client, monkeypatch, spacing_boundary_storage, side, fault,
):
    from app.services import guided_fix

    document = _spacing_document()
    second = document.add_paragraph('Another body paragraph requiring the same spacing.')
    second.paragraph_format.space_before = Pt(20)
    second.paragraph_format.space_after = Pt(20)
    second.runs[0].font.name = 'Times New Roman'
    second.runs[0].font.size = Pt(11)
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source)
    findings = [finding for finding in audit['physical_layout_errors']
                if finding['rule_code'] == 'SPACE_' + side.upper()]
    assert {finding['location']['paragraph_index'] for finding in findings} == {0, 1}
    ids = [finding['id'] for finding in findings]
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    if fault:
        original = getattr(guided_fix, '_fix_space_' + side)
        calls = []

        def omit_last(paragraph, required):
            calls.append(paragraph)
            if len(calls) == 1:
                return original(paragraph, required)
            if fault == 'unchanged':
                return True
            result = original(paragraph, required)
            paragraph._p.pPr.find(qn('w:spacing')).set(qn('w:' + side), str(required + 1))
            return result

        monkeypatch.setattr(guided_fix, '_fix_space_' + side, omit_last)
    response = _upload(client, root, source, ids, review.json()['review_id'])
    if fault:
        assert len(calls) == 2
        assert response.status_code == 422, response.text
        _assert_no_bold_output(response, spacing_boundary_storage, audit['audit_id'])
        return
    assert response.status_code == 200, response.text
    download = client.get(root + '/download/' + response.json()['download_token'])
    assert download.status_code == 200
    output = Document(io.BytesIO(download.content))
    assert [p._p.pPr.find(qn('w:spacing')).get(qn('w:' + side))
            for p in output.paragraphs] == ['120', '120']


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('field,value', [('actual_value', '19.95pt'), ('expected_value', '7pt')])
def test_point_spacing_stale_saved_finding_has_no_review(
    client, test_engine, spacing_boundary_storage, side, field, value,
):
    source = _bold_bytes(_spacing_document())
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    with Session(test_engine) as db:
        setattr(db.get(Violation, ids[0]), field, value)
        db.commit()
    saved = _saved_source_state(test_engine, audit['audit_id'])
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 422, review.text
    _assert_no_spacing_output(review, spacing_boundary_storage, audit['audit_id'])
    assert _saved_source_state(test_engine, audit['audit_id']) == saved


@pytest.mark.parametrize('side', ['before', 'after'])
def test_point_spacing_actual_above_requirement_range_applies(
    client, spacing_boundary_storage, side,
):
    document = _spacing_document()
    setattr(document.paragraphs[0].paragraph_format, 'space_' + side, Pt(300))
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    output = Document(io.BytesIO(download.content))
    assert output.paragraphs[0]._p.pPr.find(qn('w:spacing')).get(qn('w:' + side)) == '120'


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('difference', [19, 20, 21])
def test_point_spacing_detector_float_boundary_is_conservative(
    client, spacing_boundary_storage, side, difference,
):
    document = _spacing_document()
    document.paragraphs[0]._p.pPr.find(qn('w:spacing')).set(
        qn('w:' + side), str(68 + difference),
    )
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source, **{side: 3.4})
    findings = [finding for finding in audit['physical_layout_errors']
                if finding['rule_code'] == 'SPACE_' + side.upper()]
    if difference == 19:
        assert not findings
        return
    assert len(findings) == 1
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    review = _upload(client, root + '/review', source, [findings[0]['id']])
    if difference == 20:
        assert review.status_code == 422, review.text
        _assert_no_spacing_output(review, spacing_boundary_storage, audit['audit_id'])
    else:
        assert review.status_code == 200, review.text


def _point_profile_audit(document, profile):
    snapshot = resolve_snapshot(profile)
    source = _bold_bytes(document)
    findings = run_static_rules_engine(source, config=EffectiveProfileConfig(snapshot))
    for index, finding in enumerate(findings):
        finding.id = f'point-role-{index}'
    return source, SimpleNamespace(
        profile_snapshot=snapshot.to_dict(), paragraph_page_mapping=None, violations=findings,
        weighted_score=calculate_weighted_score_detailed(findings).total,
        document_sha256=hashlib.sha256(source).hexdigest(),
    )


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('kind,role,required', [
    ('body', 'BODY', 6), ('heading', 'HEADING_1', 6.05),
    ('caption', 'CAPTION_FIGURE', 6.1), ('appendix', 'APPENDIX_BODY', 6),
])
def test_point_spacing_uses_saved_requirement_for_real_role(side, kind, role, required):
    from app.services.document_parser import extract_paragraphs
    from app.services.guided_fix import apply_guided_fixes
    from app.services.role_classifier import classify_paragraphs
    from app.services.docx_preservation import compare_preservation

    document = Document()
    if kind == 'appendix':
        document.add_paragraph('Main body prose that is long enough to be classified.')
        document.add_paragraph('Appendix A', style='Heading 1')
    text = 'Figure 1: Synthetic evidence diagram' if kind == 'caption' else (
        'Synthetic academic prose with enough text to qualify for this role.'
    )
    paragraph = document.add_paragraph(text, style='Heading 1' if kind == 'heading' else None)
    paragraph.paragraph_format.space_before = Pt(20)
    paragraph.paragraph_format.space_after = Pt(20)
    paragraph.runs[0].font.name = 'Times New Roman'
    paragraph.runs[0].font.size = Pt(11)
    index = len(document.paragraphs) - 1
    profile = new_custom_profile('Saved role-specific point requirements')
    for name in ('space_before_pt', 'space_after_pt'):
        setattr(profile.body, name, 6)
        setattr(profile.heading, name, 6.05)
        setattr(profile.captions, name, 6.1)
    source, audit = _point_profile_audit(document, profile)
    original = Document(io.BytesIO(source))
    assert classify_paragraphs(original, extract_paragraphs(original))[index] == role
    selected = [finding for finding in audit.violations
                if finding.rule_code == 'SPACE_' + side.upper()
                and finding.location == {'paragraph_index': index}]
    assert len(selected) == 1 and selected[0].expected_value == f'{required}pt'
    _, targets = review_targets(audit, source, selected)
    corrected, applied, skipped = apply_guided_fixes(audit, source, selected)
    assert applied == selected and not skipped
    raw = Document(io.BytesIO(corrected)).paragraphs[index]._p.pPr.find(qn('w:spacing'))
    assert raw.get(qn('w:' + side)) == str(int(Decimal(str(required)) * 20))
    compare_preservation(source, corrected, targets)
    verify_reaudit(audit, source, corrected, selected)


@pytest.mark.parametrize('kind,required', [('heading', 6.05), ('caption', 6.1), ('list', 12)])
def test_point_spacing_follows_detector_fallback_precedence(kind, required):
    from app.services.document_parser import extract_paragraphs
    from app.services.guided_fix import _resolve_point_spacing_expected
    from app.services.layout_engine import check_paragraph_typography

    document = _spacing_document()
    paragraph = document.paragraphs[0]
    if kind == 'heading':
        paragraph.style = 'Heading 1'
    elif kind == 'caption':
        paragraph.text = 'Figure 1: Synthetic evidence diagram'
    else:
        paragraph.style = 'List Bullet'
    profile = new_custom_profile('Detector fallback precedence')
    profile.body.space_after_pt = 6
    profile.heading.space_after_pt = 6.05
    profile.captions.space_after_pt = 6.1
    profile.lists.space_after_pt = 12
    snapshot = resolve_snapshot(profile)
    paragraphs = extract_paragraphs(document)
    findings = check_paragraph_typography(paragraphs, EffectiveProfileConfig(snapshot), document, ['BODY'])
    finding = next(f for f in findings if f.rule_code == 'SPACE_AFTER')
    assert finding.expected_value == f'{required}pt'
    audit = SimpleNamespace(profile_snapshot=snapshot.to_dict())
    assert _resolve_point_spacing_expected(finding, audit, document, 'BODY', paragraphs[0]) == int(
        Decimal(str(required)) * 20
    )


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('fault', ['source_hash', 'reviewed_target'])
def test_point_spacing_output_oracle_requires_source_and_review_binding(side, fault):
    from app.services.document_parser import extract_paragraphs, parse_document
    from app.services.guided_fix import _validate_output, apply_guided_fixes
    from app.services.role_classifier import classify_paragraphs

    source, audit, selected = _point_snapshot_audit(side)
    _, targets = review_targets(audit, source, selected)
    corrected, applied, _ = apply_guided_fixes(audit, source, selected)
    original = parse_document(source)
    roles = classify_paragraphs(original, extract_paragraphs(original))
    if fault == 'source_hash':
        audit.document_sha256 = '0' * 64
    else:
        targets = set()
    with pytest.raises(GuidedFixFailed):
        _validate_output(corrected, original, applied, roles, audit, targets, original_source_bytes=source)


def test_point_spacing_mixed_font_and_alignment_are_independently_verified():
    from app.services.docx_preservation import compare_preservation
    from app.services.guided_fix import apply_guided_fixes

    document = _clean_font_fixture_context(_spacing_document())
    paragraph = document.paragraphs[0]
    paragraph.runs[0].font.name = 'Calibri'
    paragraph.runs[0].font.size = Pt(11)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    profile = new_custom_profile('Mixed point font and alignment')
    profile.body.space_before_pt = 6.05
    profile.body.font_family = 'Times New Roman'
    profile.body.font_size_pt = 12
    profile.body.alignment = 'justify'
    source, audit = _point_profile_audit(document, profile)
    codes = {'SPACE_BEFORE', 'FONT_SIZE', 'FONT_CONSISTENCY', 'ALIGNMENT'}
    selected = [finding for finding in audit.violations if finding.rule_code in codes]
    assert {finding.rule_code for finding in selected} == codes
    _, targets = review_targets(audit, source, selected)
    corrected, applied, skipped = apply_guided_fixes(audit, source, selected)
    assert applied == selected and not skipped
    paragraph = Document(io.BytesIO(corrected)).paragraphs[0]
    assert paragraph._p.pPr.find(qn('w:spacing')).get(qn('w:before')) == '121'
    assert paragraph.runs[0].font.size.pt == 12
    assert paragraph.runs[0].font.name == 'Times New Roman'
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    compare_preservation(source, corrected, targets)
    verify_reaudit(audit, source, corrected, selected)


@pytest.mark.parametrize('side', ['before', 'after'])
def test_point_spacing_accepts_leading_zero_source_twips(client, spacing_boundary_storage, side):
    document = _spacing_document()
    document.paragraphs[0]._p.pPr.find(qn('w:spacing')).set(qn('w:' + side), '000400')
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source, **{side: 6.05})
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    spacing = Document(io.BytesIO(download.content)).paragraphs[0]._p.pPr.find(qn('w:spacing'))
    assert spacing.get(qn('w:' + side)) == '121'



def _ambiguous_spacing_document(side, variant, layer='direct'):
    document = _spacing_document()
    properties = _spacing_layer(document, layer)
    spacing = properties.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        properties.append(spacing)
    if variant == 'duplicate_spacing':
        properties.append(copy.deepcopy(spacing))
    elif variant == 'duplicate_pPr':
        properties.getparent().append(copy.deepcopy(properties))
    elif variant == 'duplicate_contextual':
        _spacing_control(properties, side, 'contextualSpacing', '0')
        _spacing_control(properties, side, 'contextualSpacing', '0')
    elif variant == 'unknown_attribute':
        spacing.set(qn('w:unknownSpacingControl'), '1')
    elif variant == 'spacing_child':
        spacing.append(OxmlElement('w:before'))
    elif variant == 'spacing_text':
        spacing.text = 'not a leaf'
    elif variant == 'contextual_child':
        _spacing_control(properties, side, 'contextualSpacing', '0')
        properties.find(qn('w:contextualSpacing')).append(OxmlElement('w:val'))
    elif variant == 'numbering':
        num = OxmlElement('w:numPr')
        identifier = OxmlElement('w:numId')
        identifier.set(qn('w:val'), '1')
        num.append(identifier)
        properties.append(num)
    elif variant == 'revision':
        properties.append(OxmlElement('w:pPrChange'))
    elif variant == 'alternate_content':
        from lxml import etree
        properties.append(etree.Element('{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent'))
    elif variant in {'missing_style', 'wrong_kind_style', 'cycle', 'duplicate_style',
                     'duplicate_pStyle', 'missing_base', 'duplicate_basedOn'}:
        _spacing_layer(document, 'style')
        paragraph = document.paragraphs[0]
        style = paragraph.style
        if variant == 'missing_style':
            paragraph._p.pPr.find(qn('w:pStyle')).set(qn('w:val'), 'MissingSpacingStyle')
        elif variant == 'wrong_kind_style':
            style._element.set(qn('w:type'), 'character')
        elif variant == 'cycle':
            style.base_style.base_style = style
        elif variant == 'duplicate_style':
            document.styles.element.append(copy.deepcopy(style._element))
        elif variant == 'duplicate_pStyle':
            paragraph._p.pPr.append(copy.deepcopy(paragraph._p.pPr.find(qn('w:pStyle'))))
        elif variant == 'missing_base':
            style._element.find(qn('w:basedOn')).set(qn('w:val'), 'MissingSpacingBase')
        else:
            style._element.append(copy.deepcopy(style._element.find(qn('w:basedOn'))))
    elif variant == 'duplicate_default':
        normal = document.styles['Normal']._element
        duplicate = copy.deepcopy(normal)
        duplicate.set(qn('w:styleId'), 'AnotherDefaultParagraph')
        document.styles.element.append(duplicate)
    else:
        raise AssertionError('Unknown spacing fixture')
    return document


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('variant', [
    'duplicate_spacing', 'duplicate_pPr', 'duplicate_contextual', 'unknown_attribute',
    'missing_style', 'wrong_kind_style', 'cycle', 'duplicate_style', 'duplicate_pStyle',
    'missing_base', 'duplicate_basedOn', 'duplicate_default',
])
def test_spacing_structural_and_used_style_ambiguity_never_publishes(
    client, spacing_boundary_storage, side, variant,
):
    source = _bold_bytes(_ambiguous_spacing_document(side, variant))
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    response = _upload(client, root + '/review', source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])


@pytest.mark.parametrize('layer', ['direct', 'style', 'base', 'default'])
def test_spacing_numbering_evidence_is_manual_at_every_used_layer(
    client, spacing_boundary_storage, layer,
):
    source = _bold_bytes(_ambiguous_spacing_document('after', 'numbering', layer))
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, ('after',))
    response = _upload(client, root + '/review', source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, spacing_boundary_storage, audit['audit_id'])


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('layer', ['direct', 'style', 'base', 'default'])
@pytest.mark.parametrize('variant', [
    'spacing_child', 'spacing_text', 'contextual_child', 'revision', 'alternate_content',
])
def test_spacing_raw_evidence_rejects_unreadable_used_containers(side, layer, variant):
    from app.services.guided_fix import _direct_spacing_points

    document = _ambiguous_spacing_document(side, variant, layer)
    source = _bold_bytes(document)
    reopened = Document(io.BytesIO(source))
    before = (reopened._element.xml, reopened.styles.element.xml)
    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(reopened.paragraphs[0], 'SPACE_' + side.upper())
    assert (reopened._element.xml, reopened.styles.element.xml) == before


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('value', ['-1', '1.5', 'NaN', 'inf', '', ' 400', '+400'])
def test_spacing_raw_points_require_nonnegative_integral_twips(side, value):
    from app.services.guided_fix import _direct_spacing_points

    document = _spacing_document()
    document.paragraphs[0]._p.pPr.find(qn('w:spacing')).set(qn('w:' + side), value)
    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(document.paragraphs[0], 'SPACE_' + side.upper())


@pytest.mark.parametrize('side', ['before', 'after'])
def test_spacing_missing_direct_points_cannot_fall_back_to_inheritance(side):
    from app.services.guided_fix import _direct_spacing_points

    document = _spacing_document()
    properties = _spacing_layer(document, 'style')
    inherited = OxmlElement('w:spacing')
    inherited.set(qn('w:' + side), '400')
    properties.append(inherited)
    del document.paragraphs[0]._p.pPr.find(qn('w:spacing')).attrib[qn('w:' + side)]
    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(document.paragraphs[0], 'SPACE_' + side.upper())


@pytest.mark.parametrize('side', ['before', 'after'])
def test_spacing_unused_broken_styles_do_not_reduce_safe_direct_point_availability(
    client, spacing_boundary_storage, side,
):
    from docx.enum.style import WD_STYLE_TYPE
    from tests.test_guided_fix_preservation_storage import _assert_spacing_package_unchanged

    document = _spacing_document()
    unused = document.styles.add_style('Unused broken spacing', WD_STYLE_TYPE.PARAGRAPH)
    unused.base_style = unused
    _spacing_control(unused._element.get_or_add_pPr(), side, 'Autospacing', 'invalid')
    source = _bold_bytes(document)
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    _assert_spacing_package_unchanged(source, download.content, {0: {side: 120}})


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('change,status', [
    ('wrong_original', 409), ('stale_actual', 422), ('stale_expected', 422),
    ('no_snapshot', 409), ('future_snapshot', 409), ('no_hash', 409),
    ('corrupt_snapshot', 409),
])
def test_spacing_saved_evidence_gates_leave_original_history_unchanged(
    client, test_engine, spacing_boundary_storage, side, change, status,
):
    source = _bold_bytes(_spacing_document())
    audit = _spacing_audit(client, source)
    root, ids = _spacing_selection(audit, (side,))
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    with Session(test_engine) as db:
        saved = db.get(AuditRecord, audit['audit_id'])
        finding = db.get(Violation, ids[0])
        if change == 'stale_actual':
            finding.actual_value = '10pt'
        elif change == 'stale_expected':
            finding.expected_value = '8pt'
        elif change == 'no_snapshot':
            saved.profile_snapshot = None
        elif change == 'future_snapshot':
            snapshot = copy.deepcopy(saved.profile_snapshot)
            snapshot['schema_version'] = 999
            saved.profile_snapshot = snapshot
        elif change == 'corrupt_snapshot':
            snapshot = copy.deepcopy(saved.profile_snapshot)
            snapshot['fingerprint'] = '0' * 64
            saved.profile_snapshot = snapshot
        elif change == 'no_hash':
            saved.document_sha256 = None
        db.commit()
    before = _saved_source_state(test_engine, audit['audit_id'])
    if change == 'wrong_original':
        changed = _spacing_document()
        changed.paragraphs[0].add_run(' A different original.')
        source = _bold_bytes(changed)
    response = _upload(client, root, source, ids, review.json()['review_id'])
    assert response.status_code == status, response.text
    _assert_no_bold_output(response, spacing_boundary_storage, audit['audit_id'])
    assert _saved_source_state(test_engine, audit['audit_id']) == before


@pytest.mark.parametrize('rule', ['LINE_SPACING', '', None])
def test_spacing_raw_oracle_does_not_expand_to_other_operations(rule):
    from app.services.guided_fix import _direct_spacing_points

    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(_spacing_document().paragraphs[0], rule)


@pytest.mark.parametrize('side', ['before', 'after'])
def test_spacing_raw_oracle_requires_main_body_context(side):
    from app.services.guided_fix import _direct_spacing_points

    document = _spacing_document()
    table_paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    table_paragraph.paragraph_format.space_before = Pt(20)
    table_paragraph.paragraph_format.space_after = Pt(20)
    header_paragraph = document.sections[0].header.paragraphs[0]
    header_paragraph.paragraph_format.space_before = Pt(20)
    header_paragraph.paragraph_format.space_after = Pt(20)
    for paragraph in (table_paragraph, header_paragraph, None):
        with pytest.raises(GuidedFixFailed):
            _direct_spacing_points(paragraph, 'SPACE_' + side.upper())


@pytest.mark.parametrize('side', ['before', 'after'])
def test_spacing_saved_requirement_survives_current_profile_change(
    client, test_engine, monkeypatch, spacing_boundary_storage, side,
):
    from datetime import datetime
    from app.services import profile_registry
    from tests.test_guided_fix_preservation_storage import _assert_spacing_package_unchanged

    profile = new_custom_profile('Historical direct point spacing')
    profile.body.space_before_pt = 6
    profile.body.space_after_pt = 6
    monkeypatch.setattr(profile_registry, 'BUILTIN_PROFILES', {profile.profile_id: profile})
    source = _bold_bytes(_spacing_document())
    response = client.post('/api/audit', params={'profile_id': profile.profile_id},
                           files={'file': ('historical.docx', source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    root, ids = _spacing_selection(audit, (side,))
    with Session(test_engine) as db:
        db.get(AuditRecord, audit['audit_id']).created_at = datetime(2024, 1, 1)
        db.commit()
    saved = _saved_source_state(test_engine, audit['audit_id'])
    profile.body.space_before_pt = 12
    profile.body.space_after_pt = 12
    review = _upload(client, root + '/review', source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()['review_id'])
    assert applied.status_code == 200, applied.text
    download = client.get(root + '/download/' + applied.json()['download_token'])
    assert download.status_code == 200
    _assert_spacing_package_unchanged(source, download.content, {0: {side: 120}})
    assert _saved_source_state(test_engine, audit['audit_id']) == saved


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('layer', ['direct', 'style', 'base', 'default'])
@pytest.mark.parametrize('variant', ['duplicate_spacing', 'duplicate_pPr', 'duplicate_contextual'])
def test_spacing_duplicate_containers_reject_in_all_used_layers(side, layer, variant):
    from app.services.guided_fix import _direct_spacing_points

    document = _ambiguous_spacing_document(side, variant, layer)
    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(document.paragraphs[0], 'SPACE_' + side.upper())


@pytest.mark.parametrize('side', ['before', 'after'])
@pytest.mark.parametrize('control', ['Autospacing', 'contextualSpacing'])
@pytest.mark.parametrize('literal', ['TRUE', 'False', 'yes', '2', '', ' false'])
def test_spacing_control_literals_are_exact_not_python_truthiness(side, control, literal):
    from app.services.guided_fix import _direct_spacing_points

    document = _spacing_document()
    _spacing_control(_spacing_layer(document, 'direct'), side, control, literal)
    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(document.paragraphs[0], 'SPACE_' + side.upper())


@pytest.mark.parametrize('variant', [
    'missing_styles_relationship', 'duplicate_docDefaults', 'duplicate_pPrDefault',
    'duplicate_default_pPr', 'duplicate_base_reference', 'missing_default_style',
])
def test_spacing_style_and_default_context_must_be_complete(variant):
    from app.services.guided_fix import _direct_spacing_points

    document = _spacing_document()
    defaults = _spacing_layer(document, 'default')
    if variant == 'missing_styles_relationship':
        relationship = next(identifier for identifier, relationship in document.part.rels.items()
                            if relationship.reltype.endswith('/styles'))
        del document.part.rels[relationship]
    elif variant == 'duplicate_docDefaults':
        document.styles.element.append(copy.deepcopy(defaults.getparent().getparent()))
    elif variant == 'duplicate_pPrDefault':
        defaults.getparent().getparent().append(copy.deepcopy(defaults.getparent()))
    elif variant == 'duplicate_default_pPr':
        defaults.getparent().append(copy.deepcopy(defaults))
    elif variant == 'duplicate_base_reference':
        _spacing_layer(document, 'style')
        style = document.paragraphs[0].style._element
        style.append(copy.deepcopy(style.find(qn('w:basedOn'))))
    else:
        for style in document.styles.element.findall(qn('w:style')):
            if style.get(qn('w:type')) == 'paragraph':
                style.attrib.pop(qn('w:default'), None)
    with pytest.raises(GuidedFixFailed):
        _direct_spacing_points(document.paragraphs[0], 'SPACE_BEFORE')


def _font_document(grouped=False):
    document = _clean_font_fixture_context(Document())
    paragraph = document.add_paragraph()
    for text in (("Ordinary ASCII ", "academic prose.") if grouped else ("Ordinary ASCII academic prose.",)):
        run = paragraph.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.italic = True
    sibling = document.add_paragraph("Unselected sibling.")
    sibling.runs[0].font.name = "Times New Roman"
    sibling.runs[0].font.size = Pt(12)
    return document


def _font_audit(document, required=12, family="Times New Roman", profile=None, version=6):
    profile = profile or new_custom_profile("Strict direct font evidence")
    profile.body.font_family = family
    profile.body.font_size_pt = required
    snapshot = replace(resolve_snapshot(profile), schema_version=version, fingerprint="").with_fingerprint()
    source = _bold_bytes(document)
    findings = run_static_rules_engine(source, config=EffectiveProfileConfig(snapshot))
    for index, finding in enumerate(findings):
        finding.id = f"font-{index}"
    audit = SimpleNamespace(
        profile_snapshot=snapshot.to_dict(), paragraph_page_mapping=None, violations=findings,
        weighted_score=calculate_weighted_score_detailed(findings).total,
        document_sha256=hashlib.sha256(source).hexdigest(),
    )
    return source, audit


def _font_findings(audit, *codes):
    findings = [finding for finding in audit.violations if finding.rule_code in codes
                and finding.location.get("paragraph_index") == 0]
    assert findings, "A regression must select genuine detector findings"
    return findings


def _font_layer(document, layer):
    from docx.enum.style import WD_STYLE_TYPE

    paragraph = document.paragraphs[0]
    if layer == "direct":
        return paragraph.runs[0]._r.get_or_add_rPr()
    if layer == "defaults":
        return document.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    if layer == "default_character":
        return document.styles["Default Paragraph Font"].element.get_or_add_rPr()
    character = layer in ("character", "character_base")
    kind = WD_STYLE_TYPE.CHARACTER if character else WD_STYLE_TYPE.PARAGRAPH
    style = document.styles.add_style("Font evidence " + layer, kind)
    if character:
        paragraph.runs[0].style = style
    else:
        paragraph.style = style
    if layer in ("character_base", "paragraph_base"):
        base = document.styles.add_style("Font evidence base " + layer, kind)
        style.base_style = base
        style = base
    return style.element.get_or_add_rPr()


@pytest.mark.parametrize("codes", [("FONT_CONSISTENCY",), ("FONT_SIZE",), ("FONT_CONSISTENCY", "FONT_SIZE")])
@pytest.mark.parametrize("grouped", [False, True])
@pytest.mark.parametrize("required", [12, 12.5])
def test_font_supported_direct_composition_is_exact_and_preserved(codes, grouped, required):
    from app.services.guided_fix import apply_guided_fixes
    from tests.test_guided_fix_preservation_storage import _assert_font_package_unchanged

    source, audit = _font_audit(_font_document(grouped), required=required)
    original_hash = hashlib.sha256(source).hexdigest()
    selected = _font_findings(audit, *codes)
    items, targets = review_targets(audit, source, selected)
    assert {item["before"] for item in items} == {
        finding.actual_value for finding in selected
    }
    output, applied, skipped = apply_guided_fixes(audit, source, selected)
    assert skipped == []
    assert applied == selected
    _assert_font_package_unchanged(source, output, targets)
    after = Document(io.BytesIO(output))
    for run in after.paragraphs[0].runs:
        if "FONT_CONSISTENCY" in codes:
            assert all(run._r.rPr.rFonts.get(qn("w:" + name)) == "Times New Roman"
                       for name in ("ascii", "hAnsi", "eastAsia"))
        if "FONT_SIZE" in codes:
            assert run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == str(int(required * 2))
        assert run.italic is True
    summary = verify_reaudit(audit, source, output, selected)
    assert summary["new_findings"] == []
    assert summary["after_score"] >= summary["before_score"]
    assert hashlib.sha256(source).hexdigest() == original_hash


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("layer", ["character", "character_base", "default_character", "paragraph",
                                  "paragraph_base", "defaults"])
def test_font_inherited_active_complex_slots_are_manual(rule, layer):
    document = _font_document()
    props = _font_layer(document, layer)
    for tag, value in (("cs", "1"), ("szCs", "44")):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:val"), value)
        props.append(node)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:cs"), "Courier New")
    props.append(fonts)
    source, audit = _font_audit(document)
    selected = _font_findings(audit, rule)
    from app.services.guided_fix import apply_guided_fixes
    for operation in (review_targets, apply_guided_fixes):
        with pytest.raises(GuidedFixFailed):
            operation(audit, source, selected)
    assert fonts.get(qn("w:cs")) == "Courier New"
    assert props.find(qn("w:szCs")).get(qn("w:val")) == "44"


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("layer", ["direct", "character", "paragraph_base", "defaults"])
@pytest.mark.parametrize("tag,literal", [
    ("cs", "0"), ("cs", "false"), ("cs", "bad"), ("rtl", "0"), ("rtl", ""),
    ("szCs", "22"), ("script", "latin"), ("bidi", "0"),
])
def test_font_script_controls_are_never_cleared_or_interpreted(rule, layer, tag, literal):
    document = _font_document()
    node = OxmlElement("w:" + tag)
    node.set(qn("w:val"), literal)
    _font_layer(document, layer).append(node)
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, rule))


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("attribute", ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme",
                                     "csTheme", "cs", "hint", "unknown"])
def test_font_every_alternate_or_unknown_direct_font_attribute_is_manual(rule, attribute):
    document = _font_document()
    document.paragraphs[0].runs[0]._r.rPr.rFonts.set(qn("w:" + attribute), "minorBidi")
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, rule))


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("layer", ["character", "character_base", "default_character", "paragraph",
                                  "paragraph_base", "defaults"])
def test_font_inherited_literal_fonts_are_manual_even_when_shadowed(rule, layer):
    document = _font_document()
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    _font_layer(document, layer).append(fonts)
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, rule))


@pytest.mark.parametrize("layer", ["character", "character_base", "default_character", "paragraph",
                                  "paragraph_base", "defaults"])
def test_font_size_inherited_size_is_manual_but_family_does_not_interpret_it(layer):
    document = _font_document()
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    _font_layer(document, layer).append(size)
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, "FONT_SIZE"))
    assert review_targets(audit, source, _font_findings(audit, "FONT_CONSISTENCY"))[0]


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("text", ["", "café", "ASCII 中文", "ASCII العربية", "ASCII \u0301", "ASCII \x7f"])
def test_font_non_ascii_or_empty_selected_text_is_manual(rule, text):
    document = _font_document()
    source, audit = _font_audit(document)
    selected = _font_findings(audit, rule)
    document.paragraphs[0].runs[0].text = text
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), selected)


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("variant", [
    "no_rPr", "duplicate_rPr", "no_fonts", "duplicate_fonts", "no_ascii", "no_hAnsi",
    "empty_ascii", "mixed_hAnsi", "mixed_eastAsia", "font_child", "font_text",
    "duplicate_size", "no_size", "zero_size", "fraction_size", "negative_size", "size_attribute",
])
def test_font_malformed_or_ambiguous_direct_evidence_is_manual(rule, variant):
    document = _font_document()
    source, audit = _font_audit(document)
    selected = _font_findings(audit, rule)
    run = document.paragraphs[0].runs[0]
    props = run._r.rPr
    fonts = props.rFonts
    size = props.find(qn("w:sz"))
    if variant == "no_rPr":
        run._r.remove(props)
    elif variant == "duplicate_rPr":
        run._r.append(copy.deepcopy(props))
    elif variant == "no_fonts":
        props.remove(fonts)
    elif variant == "duplicate_fonts":
        props.append(copy.deepcopy(fonts))
    elif variant in ("no_ascii", "no_hAnsi"):
        del fonts.attrib[qn("w:" + variant.removeprefix("no_"))]
    elif variant == "empty_ascii":
        fonts.set(qn("w:ascii"), "")
    elif variant in ("mixed_hAnsi", "mixed_eastAsia"):
        fonts.set(qn("w:" + variant.removeprefix("mixed_")), "Arial")
    elif variant == "font_child":
        fonts.append(OxmlElement("w:cs"))
    elif variant == "font_text":
        fonts.text = "unexpected"
    elif variant == "duplicate_size":
        props.append(copy.deepcopy(size))
    elif variant == "no_size":
        props.remove(size)
    elif variant == "size_attribute":
        size.set(qn("w:unknown"), "1")
    else:
        size.set(qn("w:val"), {"zero_size": "0", "fraction_size": "22.5", "negative_size": "-22"}[variant])
    if rule == "FONT_CONSISTENCY" and variant == "no_size":
        assert review_targets(audit, _bold_bytes(document), selected)[0]
        return  # A family correction does not require a direct size.
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), selected)


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("variant", [
    "dangling_character", "dangling_paragraph", "duplicate_used_style", "wrong_style_type",
    "cycle", "duplicate_basedOn", "linked", "duplicate_docDefaults", "duplicate_rPrDefault",
    "duplicate_default_rPr", "missing_default_character", "duplicate_default_character",
    "missing_default_paragraph", "duplicate_default_paragraph", "duplicate_style_rPr",
    "duplicate_pPr", "duplicate_rStyle", "duplicate_pStyle", "direct_bidi", "style_bidi",
    "defaults_bidi", "fldSimple",
])
def test_font_used_style_graph_and_paragraph_context_are_unambiguous(rule, variant):
    document = _font_document()
    paragraph = document.paragraphs[0]
    props = _font_layer(document, "paragraph")
    style = props.getparent()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    if variant == "dangling_character":
        paragraph.runs[0]._r.get_or_add_rPr().append(OxmlElement("w:rStyle"))
        paragraph.runs[0]._r.rPr.find(qn("w:rStyle")).set(qn("w:val"), "Missing")
    elif variant == "dangling_paragraph":
        paragraph._p.pPr.find(qn("w:pStyle")).set(qn("w:val"), "Missing")
    elif variant == "duplicate_used_style":
        document.styles.element.append(copy.deepcopy(style))
    elif variant == "wrong_style_type":
        style.set(qn("w:type"), "character")
    elif variant in ("cycle", "duplicate_basedOn", "linked"):
        tag = "link" if variant == "linked" else "basedOn"
        node = OxmlElement("w:" + tag)
        node.set(qn("w:val"), style.get(qn("w:styleId")))
        style.append(node)
        if variant == "duplicate_basedOn":
            style.append(copy.deepcopy(node))
    elif variant == "duplicate_docDefaults":
        document.styles.element.append(copy.deepcopy(defaults))
    elif variant == "duplicate_rPrDefault":
        defaults.append(copy.deepcopy(defaults.find(qn("w:rPrDefault"))))
    elif variant == "duplicate_default_rPr":
        wrapper = defaults.find(qn("w:rPrDefault"))
        wrapper.append(copy.deepcopy(wrapper.find(qn("w:rPr"))))
    elif variant in ("missing_default_character", "duplicate_default_character",
                     "missing_default_paragraph", "duplicate_default_paragraph"):
        kind = "character" if variant.endswith("character") else "paragraph"
        if kind == "paragraph":
            paragraph._p.remove(paragraph._p.pPr)
        default = next(node for node in document.styles.element.findall(qn("w:style"))
                       if node.get(qn("w:type")) == kind and node.get(qn("w:default")) == "1")
        if variant.startswith("missing"):
            default.attrib.pop(qn("w:default"))
        else:
            extra = copy.deepcopy(default)
            extra.set(qn("w:styleId"), "OtherDefault")
            document.styles.element.append(extra)
    elif variant == "duplicate_style_rPr":
        style.append(copy.deepcopy(props))
    elif variant == "duplicate_pPr":
        paragraph._p.append(copy.deepcopy(paragraph._p.pPr))
    elif variant == "duplicate_rStyle":
        run_props = paragraph.runs[0]._r.rPr
        for _ in range(2):
            node = OxmlElement("w:rStyle")
            node.set(qn("w:val"), "DefaultParagraphFont")
            run_props.append(node)
    elif variant == "duplicate_pStyle":
        paragraph._p.pPr.append(copy.deepcopy(paragraph._p.pPr.find(qn("w:pStyle"))))
    elif variant in ("direct_bidi", "style_bidi", "defaults_bidi"):
        if variant == "direct_bidi":
            ppr = paragraph._p.pPr
        elif variant == "style_bidi":
            ppr = style.get_or_add_pPr()
        else:
            ppr = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr"))
        node = OxmlElement("w:bidi")
        node.set(qn("w:val"), "0")
        ppr.append(node)
    else:
        wrapper = OxmlElement("w:fldSimple")
        paragraph._p.append(wrapper)
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, rule))


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
def test_font_unused_broken_styles_do_not_reduce_supported_availability(rule):
    document = _font_document()
    from docx.enum.style import WD_STYLE_TYPE
    style = document.styles.add_style("Unused broken font evidence", WD_STYLE_TYPE.PARAGRAPH)
    base = OxmlElement("w:basedOn")
    base.set(qn("w:val"), "Missing")
    style.element.append(base)
    style.element.append(copy.deepcopy(base))
    source, audit = _font_audit(document)
    assert review_targets(audit, source, _font_findings(audit, rule))[0]


@pytest.mark.parametrize("required", [12.005, 12.25, 12.00000000001])
def test_font_non_halfpoint_saved_requirement_is_manual_before_review(required):
    source, audit = _font_audit(_font_document(), required=required)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, "FONT_SIZE"))


@pytest.mark.parametrize("actual", ["11.5pt", "11.005pt", "11.00000000000000000000000000000001pt"])
def test_font_current_size_requires_exact_direct_halfpoint(actual):
    source, audit = _font_audit(_font_document())
    selected = _font_findings(audit, "FONT_SIZE")
    selected[0].actual_value = actual
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("version", [1, 2, 3, 4, 5, 6])
def test_font_saved_snapshot_versions_bind_requirements_independently(rule, version):
    source, audit = _font_audit(_font_document(), version=version)
    selected = _font_findings(audit, rule)
    assert review_targets(audit, source, selected)[0]
    selected[0].expected_value = "Arial" if rule == "FONT_CONSISTENCY" else "13pt"
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("change", ["missing", "corrupt", "future"])
def test_font_missing_or_invalid_saved_snapshot_never_uses_persisted_scalar_alone(rule, change):
    source, audit = _font_audit(_font_document())
    selected = _font_findings(audit, rule)
    if change == "missing":
        audit.profile_snapshot = None
    elif change == "future":
        audit.profile_snapshot["schema_version"] = 999
    else:
        audit.profile_snapshot["body"]["font_size_pt"] = 13
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("fault", ["ascii", "hAnsi", "eastAsia", "sz", "missing_eastAsia",
                                  "inherited_cs", "inherited_fonts", "inherited_size", "paragraph_bidi"])
def test_font_faulty_authorized_output_is_discarded_before_publication(
    client, monkeypatch, test_engine, bold_boundary_storage, fault,
):
    from app.services import guided_fix
    source = _bold_bytes(_font_document(grouped=True))
    response = client.post("/api/audit", files={"file": ("font.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    rule = "FONT_SIZE" if fault in ("sz", "inherited_size") else "FONT_CONSISTENCY"
    selected = [f["id"] for f in audit["physical_layout_errors"]
                if f["rule_code"] == rule and f["location"].get("paragraph_index") == 0]
    assert selected
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    saved = _saved_source_state(test_engine, audit["audit_id"])
    review = _upload(client, root + "/review", source, selected)
    assert review.status_code == 200, review.text
    name = "_fix_font_size" if rule == "FONT_SIZE" else "_fix_font_family"
    real = getattr(guided_fix, name)
    calls = 0

    def wrong_last_run(run, value):
        nonlocal calls
        ok = real(run, value)
        calls += 1
        if calls == 2:
            if fault == "sz":
                run._r.rPr.find(qn("w:sz")).set(qn("w:val"), "25")
            elif fault == "missing_eastAsia":
                del run._r.rPr.rFonts.attrib[qn("w:eastAsia")]
            elif fault.startswith("inherited_"):
                styles = next(rel.target_part.element for rel in run.part.rels.values()
                              if rel.reltype.endswith("/styles"))
                props = styles.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
                tag = {"inherited_cs": "cs", "inherited_fonts": "rFonts", "inherited_size": "sz"}[fault]
                node = OxmlElement("w:" + tag)
                node.set(qn("w:ascii") if tag == "rFonts" else qn("w:val"), "Calibri" if tag == "rFonts" else "22")
                props.append(node)
            elif fault == "paragraph_bidi":
                node = OxmlElement("w:bidi")
                node.set(qn("w:val"), "0")
                run._parent._p.get_or_add_pPr().append(node)
            else:
                run._r.rPr.rFonts.set(qn("w:" + fault), "Courier New")
        return ok

    monkeypatch.setattr(guided_fix, name, wrong_last_run)
    applied = _upload(client, root, source, selected, review.json()["review_id"])
    assert applied.status_code == 422, applied.text
    assert calls == 2
    _assert_no_bold_output(applied, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved


@pytest.mark.parametrize("mixed", [False, True])
def test_font_manual_template_or_mixed_selection_has_no_review_or_output(
    client, test_engine, bold_boundary_storage, mixed,
):
    document = _font_document(grouped=True) if mixed else Document()
    if mixed:
        node = OxmlElement("w:cs")
        node.set(qn("w:val"), "0")
        document.paragraphs[0].runs[1]._r.rPr.append(node)
    else:
        run = document.add_paragraph("Ordinary template with inherited font controls.").runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    source = _bold_bytes(document)
    response = client.post("/api/audit", files={"file": ("manual.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    ids = [finding["id"] for finding in audit["physical_layout_errors"]
           if finding["rule_code"] in ("FONT_SIZE", "FONT_CONSISTENCY")]
    assert ids
    saved = _saved_source_state(test_engine, audit["audit_id"])
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 422, review.text
    _assert_no_bold_output(review, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved


@pytest.mark.parametrize("role", ["HEADING_1", "HEADING_2", "HEADING_3", "APPENDIX_HEADING"])
def test_font_heading_size_uses_authoritative_saved_adapter_requirement(role):
    from app.services.document_parser import extract_paragraphs
    from app.services.role_classifier import classify_paragraphs
    from app.services.guided_fix import apply_guided_fixes

    document = _font_document()
    paragraph = document.paragraphs[0]
    if role == "APPENDIX_HEADING":
        paragraph.runs[0].text = "Appendix A"
    else:
        paragraph.style = document.styles["Heading " + role[-1]]
        style = paragraph.style.element
        for child in list(style):
            if child.tag in (qn("w:rPr"), qn("w:link")):
                style.remove(child)
    profile = new_custom_profile("Authoritative heading font")
    profile.heading.font_family = "Times New Roman"
    profile.heading.font_size_pt = 14
    profile.heading.level_1 = {"font_size_pt": 16}
    profile.heading.level_2 = {"font_size_pt": 18}
    profile.heading.level_3 = {"font_size_pt": 20}
    source, audit = _font_audit(document, profile=profile)
    assert classify_paragraphs(document, extract_paragraphs(document))[0] == role
    selected = _font_findings(audit, "FONT_SIZE")
    assert selected[0].expected_value == "14pt"
    assert review_targets(audit, source, selected)[0]
    output, _, _ = apply_guided_fixes(audit, source, selected)
    assert Document(io.BytesIO(output)).paragraphs[0].runs[0]._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "28"
    assert verify_reaudit(audit, source, output, selected)["new_findings"] == []


@pytest.mark.parametrize("rule", ["FONT_SIZE", "FONT_CONSISTENCY"])
def test_font_ambiguous_saved_allowed_pairs_are_not_a_new_pair_selection(rule):
    profile = new_custom_profile("Ambiguous allowed font pairs")
    profile.body.allowed_font_combos = (("Arial", 12), ("Times New Roman", 14))
    source, audit = _font_audit(_font_document(), profile=profile, family="Arial")
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, rule))


@pytest.mark.parametrize("fault", ["none", "mutation", "preservation", "reaudit", "wrong_original"])
def test_font_combined_api_keeps_saved_history_pdf_and_publication_atomic(
    client, test_engine, monkeypatch, bold_boundary_storage, fault,
):
    from app.api import routes
    from app.services import guided_fix
    from tests.test_guided_fix_preservation_storage import _assert_font_package_unchanged
    from pypdf import PdfReader

    source = _bold_bytes(_font_document(grouped=True))
    response = client.post("/api/audit", files={"file": ("combined-font.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    ids = [f["id"] for f in audit["physical_layout_errors"]
           if f["rule_code"] in ("FONT_SIZE", "FONT_CONSISTENCY") and f["location"].get("paragraph_index") == 0]
    assert len(ids) == 2
    audit_path = f'/api/audit/{audit["audit_id"]}'
    root = audit_path + "/guided-fixes"
    saved = _saved_source_state(test_engine, audit["audit_id"])
    detail = client.get(audit_path).json()
    history = client.get("/api/audits").json()

    def pdf_text():
        response = client.get(audit_path + "/export-pdf")
        assert response.status_code == 200
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.content)).pages)

    pdf = pdf_text()
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    upload = source
    if fault == "mutation":
        monkeypatch.setattr(guided_fix, "_fix_font_size", lambda *_args: False)
    elif fault in ("preservation", "reaudit"):
        def refuse(*_args, **_kwargs):
            raise GuidedFixFailed("Controlled font publication failure")
        target = routes.docx_preservation if fault == "preservation" else routes
        monkeypatch.setattr(target, "compare_preservation" if fault == "preservation" else "verify_reaudit", refuse)
    elif fault == "wrong_original":
        different = _font_document(grouped=True)
        different.paragraphs[0].add_run(" Other original.")
        upload = _bold_bytes(different)
    applied = _upload(client, root, upload, ids, review.json()["review_id"])
    if fault == "none":
        assert applied.status_code == 200, applied.text
        assert set(applied.json()["resolved"]) == set(ids)
        assert applied.json()["new_findings"] == []
        download = client.get(root + "/download/" + applied.json()["download_token"])
        assert download.status_code == 200
        targets = {(0, i, rule) for i in (0, 1) for rule in ("FONT_SIZE", "FONT_CONSISTENCY")}
        _assert_font_package_unchanged(source, download.content, targets)
        runs = Document(io.BytesIO(download.content)).paragraphs[0].runs
        assert all(run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "24" for run in runs)
        assert all(run._r.rPr.rFonts.get(qn("w:" + attr)) == "Times New Roman"
                   for run in runs for attr in ("ascii", "hAnsi", "eastAsia"))
        assert client.get(root + "/download/" + applied.json()["download_token"]).status_code == 410
    else:
        assert applied.status_code == (409 if fault == "wrong_original" else 422), applied.text
        _assert_no_bold_output(applied, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved
    assert client.get(audit_path).json() == detail
    assert client.get("/api/audits").json() == history
    assert pdf_text() == pdf


@pytest.mark.parametrize("rule", ["FONT_SIZE", "FONT_CONSISTENCY"])
def test_font_saved_requirement_does_not_follow_current_profile_edits(
    client, test_engine, monkeypatch, bold_boundary_storage, rule,
):
    from app.services import profile_registry
    profile = new_custom_profile("Historical direct font")
    profile.body.font_family = "Times New Roman"
    profile.body.font_size_pt = 12.5
    monkeypatch.setattr(profile_registry, "BUILTIN_PROFILES", {profile.profile_id: profile})
    source = _bold_bytes(_font_document())
    response = client.post("/api/audit", params={"profile_id": profile.profile_id},
                           files={"file": ("saved-font.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    ids = [f["id"] for f in audit["physical_layout_errors"] if f["rule_code"] == rule
           and f["location"].get("paragraph_index") == 0]
    assert ids
    saved = _saved_source_state(test_engine, audit["audit_id"])
    profile.body.font_family = "Arial"
    profile.body.font_size_pt = 16
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    applied = _upload(client, root, source, ids, review.json()["review_id"])
    assert applied.status_code == 200, applied.text
    download = client.get(root + "/download/" + applied.json()["download_token"])
    assert download.status_code == 200
    run = Document(io.BytesIO(download.content)).paragraphs[0].runs[0]
    if rule == "FONT_SIZE":
        assert run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "25"
    else:
        assert all(run._r.rPr.rFonts.get(qn("w:" + attr)) == "Times New Roman"
                   for attr in ("ascii", "hAnsi", "eastAsia"))
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved


@pytest.mark.parametrize("text", ["ASCII\ttext\nwith XML whitespace.", "ASCII\rtext.", " "])
def test_font_selected_ascii_xml_whitespace_keeps_direct_run_identity(text):
    from app.services.guided_fix import apply_guided_fixes
    document = _font_document()
    document.paragraphs[0].runs[0].text = text
    source, audit = _font_audit(document)
    if not text.strip():
        # Whitespace-only paragraphs have no font candidate; do not invent one.
        assert not [f for f in audit.violations if f.rule_code.startswith("FONT")
                    and f.location.get("paragraph_index") == 0]
        return
    selected = _font_findings(audit, "FONT_SIZE", "FONT_CONSISTENCY")
    output, _, _ = apply_guided_fixes(audit, source, selected)
    assert Document(io.BytesIO(output)).paragraphs[0].runs[0].text == document.paragraphs[0].runs[0].text


@pytest.mark.parametrize("layer", ["character", "character_base", "default_character", "paragraph",
                                  "paragraph_base"])
@pytest.mark.parametrize("variant", ["duplicate", "cycle", "dangling", "linked", "nested_control"])
def test_font_each_reachable_style_chain_is_validated(layer, variant):
    document = _font_document()
    props = _font_layer(document, layer)
    style = props.getparent()
    if variant == "duplicate":
        document.styles.element.append(copy.deepcopy(style))
    elif variant == "nested_control":
        wrapper = OxmlElement("w:unknown")
        wrapper.append(OxmlElement("w:cs"))
        props.append(wrapper)
    else:
        node = OxmlElement("w:link" if variant == "linked" else "w:basedOn")
        node.set(qn("w:val"), "Missing" if variant == "dangling" else style.get(qn("w:styleId")))
        style.append(node)
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, "FONT_SIZE", "FONT_CONSISTENCY"))


@pytest.mark.parametrize("variant", ["missing_styles", "duplicate_defaults_wrapper", "nested_defaults",
                                    "unknown_run_property", "nested_text", "raw_text_tail"])
def test_font_malformed_reachable_wrappers_do_not_hide_font_controls(variant):
    document = _font_document()
    source, audit = _font_audit(document)
    selected = _font_findings(audit, "FONT_SIZE", "FONT_CONSISTENCY")
    run = document.paragraphs[0].runs[0]
    if variant == "missing_styles":
        key = next(key for key, rel in document.part.rels.items() if rel.reltype.endswith("/styles"))
        del document.part.rels[key]
    elif variant == "duplicate_defaults_wrapper":
        defaults = document.styles.element.find(qn("w:docDefaults"))
        defaults.append(copy.deepcopy(defaults.find(qn("w:pPrDefault"))))
    elif variant == "nested_defaults":
        wrapper = document.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault"))
        wrapper.append(OxmlElement("w:unknown"))
    elif variant == "unknown_run_property":
        run._r.rPr.append(OxmlElement("w:unknown"))
    elif variant == "nested_text":
        run._r.find(qn("w:t")).append(OxmlElement("w:cs"))
    else:
        run._r.find(qn("w:t")).tail = "Unenumerated text"
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), selected)


@pytest.mark.parametrize("variant", ["single_combined", "family_existing_allowed_size", "ambiguous_size_manual"])
def test_font_unique_allowed_family_preserves_pair_choice(variant):
    from app.services.guided_fix import apply_guided_fixes
    from tests.test_guided_fix_preservation_storage import _assert_font_package_unchanged
    profile = new_custom_profile("Saved allowed family")
    profile.body.allowed_font_combos = (
        (("Times New Roman", 12),) if variant == "single_combined"
        else (("Times New Roman", 12), ("Times New Roman", 14))
    )
    document = _font_document()
    if variant == "family_existing_allowed_size":
        document.paragraphs[0].runs[0].font.size = Pt(12)
    source, audit = _font_audit(document, profile=profile)
    if variant == "ambiguous_size_manual":
        with pytest.raises(GuidedFixFailed):
            review_targets(audit, source, _font_findings(audit, "FONT_SIZE"))
        return
    codes = ("FONT_CONSISTENCY", "FONT_SIZE") if variant == "single_combined" else ("FONT_CONSISTENCY",)
    selected = _font_findings(audit, *codes)
    _, targets = review_targets(audit, source, selected)
    output, _, _ = apply_guided_fixes(audit, source, selected)
    _assert_font_package_unchanged(source, output, targets)
    run = Document(io.BytesIO(output)).paragraphs[0].runs[0]
    assert run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == "24"
    assert all(run._r.rPr.rFonts.get(qn("w:" + attr)) == "Times New Roman"
               for attr in ("ascii", "hAnsi", "eastAsia"))
    assert verify_reaudit(audit, source, output, selected)["new_findings"] == []


@pytest.mark.parametrize("variant", [
    "12pt", "11.5pt", "11.005pt", "11.00000000000000000000000000000001pt",
    "missing_direct_size", "grouped_last_size",
])
def test_font_allowed_pair_current_family_size_is_exact_for_every_run(variant):
    profile = new_custom_profile("Exact allowed-pair Current")
    profile.body.allowed_font_combos = (("Times New Roman", 12),)
    document = _font_document(grouped=variant == "grouped_last_size")
    source, audit = _font_audit(document, profile=profile)
    selected = _font_findings(audit, "FONT_CONSISTENCY")
    assert review_targets(audit, source, selected)[0]
    if variant == "missing_direct_size":
        props = document.paragraphs[0].runs[0]._r.rPr
        props.remove(props.find(qn("w:sz")))
        source = _bold_bytes(document)
    elif variant == "grouped_last_size":
        document.paragraphs[0].runs[-1]._r.rPr.find(qn("w:sz")).set(qn("w:val"), "23")
        source = _bold_bytes(document)
    else:
        selected[0].actual_value = "Calibri " + variant
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


def test_font_stale_allowed_pair_current_has_no_actionable_api_review(
    client, test_engine, monkeypatch, bold_boundary_storage,
):
    from app.services import profile_registry
    profile = new_custom_profile("Saved allowed pair Current")
    profile.body.font_family = "Times New Roman"
    profile.body.font_size_pt = 12
    profile.body.allowed_font_combos = (("Times New Roman", 12),)
    monkeypatch.setattr(profile_registry, "BUILTIN_PROFILES", {profile.profile_id: profile})
    source = _bold_bytes(_font_document())
    response = client.post("/api/audit", params={"profile_id": profile.profile_id},
                           files={"file": ("combo-current.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    ids = [finding["id"] for finding in audit["physical_layout_errors"]
           if finding["rule_code"] == "FONT_CONSISTENCY" and finding["location"].get("paragraph_index") == 0]
    assert len(ids) == 1
    with Session(test_engine) as db:
        db.get(Violation, ids[0]).actual_value = "Calibri 12pt"
        db.commit()
    saved = _saved_source_state(test_engine, audit["audit_id"])
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 422, review.text
    _assert_no_bold_output(review, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved


@pytest.mark.parametrize("rule", ["FONT_CONSISTENCY", "FONT_SIZE"])
@pytest.mark.parametrize("layer", ["direct", "character", "default_character", "paragraph_base", "defaults"])
@pytest.mark.parametrize("tag,attribute,value", [
    ("eastAsianLayout", "combine", "0"),
    ("eastAsianLayout", "combine", "false"),
    ("eastAsianLayout", "unknown", "malformed"),
    ("oMath", "val", "0"),
    ("lang", "bidi", "ar-SA"),
    ("lang", "eastAsia", "zh-CN"),
    ("lang", "unknown", "en-US"),
    ("lang", "val", ""),
])
def test_font_script_layout_and_language_ambiguity_is_manual(rule, layer, tag, attribute, value):
    document = _font_document()
    props = _font_layer(document, layer)
    node = props.find(qn("w:" + tag))
    if node is None:
        node = OxmlElement("w:" + tag)
        props.append(node)
    if tag == "lang":
        node.set(qn("w:val"), "en-US")
    node.set(qn("w:" + attribute), value)
    source, audit = _font_audit(document)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, _font_findings(audit, rule))


def _line_spacing_document(raw_lines=(240,), role="BODY"):
    from tests.test_guided_fixes import _clean_line_spacing_fixture_context

    document = _clean_line_spacing_fixture_context(Document())
    if role == "APPENDIX_BODY":
        document.add_paragraph("Appendix A")
    elif role == "REFERENCE_ENTRY":
        document.add_paragraph("References")
    for index, line in enumerate(raw_lines):
        paragraph = document.add_paragraph()
        texts = (("Smith, A. (2024). ", f"Academic reference {index + 1}. Press.")
                 if role == "REFERENCE_ENTRY" else
                 ("Ordinary academic prose ", f"with an independently preserved ending {index + 1}."))
        for text in texts:
            paragraph.add_run(text)
        spacing = OxmlElement("w:spacing")
        for key, value in (("line", line), ("lineRule", "auto"), ("before", "120"), ("after", "160")):
            spacing.set(qn("w:" + key), str(value))
        paragraph._p.get_or_add_pPr().append(spacing)
    return document


def _line_spacing_audit(document, required=1.5, *, profile=None, version=6):
    profile = profile or new_custom_profile("Exact direct line spacing")
    profile.body.line_spacing = required
    snapshot = replace(resolve_snapshot(profile), schema_version=version, fingerprint="").with_fingerprint()
    source = _bold_bytes(document)
    findings = run_static_rules_engine(source, config=EffectiveProfileConfig(snapshot))
    for index, finding in enumerate(findings):
        finding.id = f"line-{index}"
    return source, SimpleNamespace(
        profile_snapshot=snapshot.to_dict(), paragraph_page_mapping=None, violations=findings,
        weighted_score=calculate_weighted_score_detailed(findings).total,
        document_sha256=hashlib.sha256(source).hexdigest(),
    )


def _line_spacing_findings(audit):
    findings = [finding for finding in audit.violations if finding.rule_code == "LINE_SPACING"]
    assert findings, "A line regression must start with genuine detector findings"
    return findings


def _line_spacing_api_audit(client, source, required=1.5, *, profile=None):
    profile = profile or new_custom_profile("Direct line spacing API")
    profile.body.line_spacing = required
    response = client.post("/api/audit", data={"custom_profile": json.dumps(profile.to_dict())},
                           files={"file": ("line-spacing.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    findings = [finding for finding in audit["physical_layout_errors"] if finding["rule_code"] == "LINE_SPACING"]
    assert findings, "API journeys must select real saved line findings"
    return audit, f'/api/audit/{audit["audit_id"]}/guided-fixes', [finding["id"] for finding in findings]


@pytest.mark.parametrize("raw_lines,required,encoded", [
    ((240,), 1.5, 360), ((360,), 2.0, 480), ((480,), 1.5, 360), ((276,), 1.5, 360),
    ((241,), 1.5, 360), ((480,), 1, 240), ((480,), 1.15, 276), ((480,), 1.0125, 243),
    ((240,), 4.0, 960), ((240, 480), 1.5, 360),
])
def test_line_spacing_exact_multiplier_review_apply_download_and_saved_views(
    client, test_engine, bold_boundary_storage, raw_lines, required, encoded,
):
    from pypdf import PdfReader
    from app.services.profile_snapshot import snapshot_from_dict
    from tests.test_guided_fix_preservation_storage import _assert_line_spacing_package_unchanged

    source = _bold_bytes(_line_spacing_document(raw_lines))
    original = bytes(source)
    digest = hashlib.sha256(source).hexdigest()
    audit, root, ids = _line_spacing_api_audit(client, source, required)
    path = f'/api/audit/{audit["audit_id"]}'
    saved = _saved_source_state(test_engine, audit["audit_id"])
    detail, history = client.get(path).json(), client.get("/api/audits").json()

    def pdf_text():
        response = client.get(path + "/export-pdf")
        assert response.status_code == 200
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.content)).pages)

    report = pdf_text()
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    assert [item["before"] for item in review.json()["items"]] == [str(raw / 240) for raw in raw_lines]
    assert [item["after"] for item in review.json()["items"]] == [str(required)] * len(raw_lines)
    applied = _upload(client, root, source, ids, review.json()["review_id"])
    assert applied.status_code == 200, applied.text
    assert applied.json()["status"] == "passed"
    assert applied.json()["resolved"] == ids
    assert applied.json()["new_findings"] == []
    download_path = root + "/download/" + applied.json()["download_token"]
    download = client.get(download_path)
    assert download.status_code == 200
    assert download.headers["cache-control"] == "private, no-store"
    assert client.get(download_path).status_code == 410
    _assert_line_spacing_package_unchanged(source, download.content, dict(enumerate([encoded] * len(raw_lines))))
    output = Document(io.BytesIO(download.content))
    for paragraph in output.paragraphs:
        spacing = paragraph._p.pPr.find(qn("w:spacing"))
        assert spacing.get(qn("w:line")) == str(encoded)
        assert spacing.get(qn("w:lineRule")) == "auto"
        assert spacing.get(qn("w:before")) == "120"
        assert spacing.get(qn("w:after")) == "160"
    config = EffectiveProfileConfig(snapshot_from_dict(saved["audit"]["profile_snapshot"]))
    before, after = [run_static_rules_engine(data, config=config) for data in (source, download.content)]
    assert not any(f.rule_code == "LINE_SPACING" for f in after)
    assert sorted(map(_spacing_finding_key, after)) == sorted(
        _spacing_finding_key(f) for f in before if f.rule_code != "LINE_SPACING"
    )
    scores = [calculate_weighted_score_detailed(findings) for findings in (before, after)]
    assert scores[1].total >= scores[0].total
    deductions = {row.category: row.deduction for row in scores[0].breakdown}
    assert all(row.deduction <= deductions[row.category] for row in scores[1].breakdown)
    assert source == original and hashlib.sha256(source).hexdigest() == digest
    assert hashlib.sha256(download.content).hexdigest() != digest
    assert saved["audit"]["document_sha256"] == digest
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved
    assert client.get(path).json() == detail
    assert client.get("/api/audits").json() == history
    assert pdf_text() == report


@pytest.mark.parametrize("role", ["BODY", "APPENDIX_BODY", "REFERENCE_ENTRY"])
@pytest.mark.parametrize("version", [1, 2, 3, 4, 5, 6])
def test_line_spacing_saved_versions_and_roles_resolve_independent_requirements(role, version):
    from app.services.guided_fix import apply_guided_fixes
    from app.services.document_parser import extract_paragraphs
    from app.services.role_classifier import classify_paragraphs
    from tests.test_guided_fix_preservation_storage import _assert_line_spacing_package_unchanged

    profile = new_custom_profile("Separate saved body and References multipliers")
    profile.references.line_spacing = 2.0
    profile.references.check_line_spacing = True
    document = _line_spacing_document(role=role)
    source, audit = _line_spacing_audit(document, profile=profile, version=version)
    # Legacy Reference payloads omit the toggle: retain a genuinely emitted
    # finding from the enabled current Snapshot, without fabricating an emitter.
    if role == "REFERENCE_ENTRY" and version < 3:
        _, current = _line_spacing_audit(document, profile=profile)
        audit.violations = current.violations
        with pytest.raises(GuidedFixFailed):
            review_targets(audit, source, _line_spacing_findings(audit))
        return
    selected = _line_spacing_findings(audit)
    index = selected[0].location["paragraph_index"]
    assert classify_paragraphs(document, extract_paragraphs(document))[index] == role
    required, encoded = (2.0, 480) if role == "REFERENCE_ENTRY" else (1.5, 360)
    assert selected[0].expected_value == str(required)
    items, _ = review_targets(audit, source, selected)
    assert items[0]["after"] == str(required)
    output, _, _ = apply_guided_fixes(audit, source, selected)
    _assert_line_spacing_package_unchanged(source, output, {index: encoded})
    assert verify_reaudit(audit, source, output, selected)["new_findings"] == []
    selected[0].expected_value = "1.75"
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("required", [1.501, 1.333, 1.0041666666666667])
def test_line_spacing_nonrepresentable_saved_decimal_never_reaches_review(
    client, test_engine, bold_boundary_storage, required,
):
    source = _bold_bytes(_line_spacing_document((480,)))
    audit, root, ids = _line_spacing_api_audit(client, source, required)
    before = _saved_source_state(test_engine, audit["audit_id"])
    response = _upload(client, root + "/review", source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == before


@pytest.mark.parametrize("actual", ["NaN", "nan", "inf", "-inf", "1", "1.00", " 1.0", "1.0x", "1.01", None])
def test_line_spacing_current_requires_literal_finite_detector_projection(actual):
    source, audit = _line_spacing_audit(_line_spacing_document())
    selected = _line_spacing_findings(audit)
    selected[0].actual_value = actual
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("value", [None, True, False, 0.5, 4.001, float("inf"), float("nan")])
def test_line_spacing_saved_requirement_must_be_finite_real_in_bounds(value):
    from app.services.profile_snapshot import snapshot_from_dict

    source, audit = _line_spacing_audit(_line_spacing_document())
    selected = _line_spacing_findings(audit)
    snapshot = snapshot_from_dict(audit.profile_snapshot)
    audit.profile_snapshot = replace(snapshot, body_line_spacing=value, fingerprint="").with_fingerprint().to_dict()
    selected[0].expected_value = str(value)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("change", ["missing", "corrupt", "future", "disabled_reference", "missing_reference"])
def test_line_spacing_missing_or_disabled_saved_evidence_cannot_use_finding_scalar(change):
    from app.services.profile_snapshot import snapshot_from_dict

    profile = new_custom_profile("Reference saved toggle")
    profile.references.line_spacing = 2.0
    profile.references.check_line_spacing = True
    source, audit = _line_spacing_audit(_line_spacing_document(role="REFERENCE_ENTRY"), profile=profile)
    selected = _line_spacing_findings(audit)
    if change == "missing":
        audit.profile_snapshot = None
    elif change == "corrupt":
        audit.profile_snapshot["body"]["line_spacing"] = 2.0
    elif change == "future":
        audit.profile_snapshot["schema_version"] = 999
    else:
        snapshot = snapshot_from_dict(audit.profile_snapshot)
        changes = ({"references_check_line_spacing": False} if change == "disabled_reference"
                   else {"references_line_spacing": None})
        audit.profile_snapshot = replace(snapshot, **changes, fingerprint="").with_fingerprint().to_dict()
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("location", [
    {}, {"paragraph_index": True}, {"paragraph_index": -1}, {"paragraph_index": 99},
    {"paragraph_index": 0, "run_index": 0}, {"paragraph_index": 0, "run_indexes": []},
    {"paragraph_index": 0, "paragraph_indexes": [0, 1], "occurrences": 2},
    {"paragraph_index": 0, "paragraph_indexes": [0], "occurrences": 2},
])
def test_line_spacing_requires_complete_main_story_paragraph_identity(location):
    source, audit = _line_spacing_audit(_line_spacing_document())
    selected = _line_spacing_findings(audit)
    selected[0].location = location
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


def test_line_spacing_duplicate_selected_target_is_not_a_complete_second_correction():
    source, audit = _line_spacing_audit(_line_spacing_document())
    selected = _line_spacing_findings(audit)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected + selected)


@pytest.mark.parametrize("rule", ["exact", "atLeast"])
def test_line_spacing_absolute_findings_are_genuine_but_never_confirmed(
    client, bold_boundary_storage, rule,
):
    document = _line_spacing_document((360,))
    document.paragraphs[0]._p.pPr.find(qn("w:spacing")).set(qn("w:lineRule"), rule)
    source = _bold_bytes(document)
    audit, root, ids = _line_spacing_api_audit(client, source)
    finding = next(f for f in audit["physical_layout_errors"] if f["id"] == ids[0])
    assert finding["actual_value"] == "228600"
    response = _upload(client, root + "/review", source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, bold_boundary_storage, audit["audit_id"])


@pytest.mark.parametrize("layer", ["style", "base", "default"])
@pytest.mark.parametrize("attribute,value", [
    ("line", "240"), ("lineRule", "auto"), ("beforeLines", "0"), ("afterLines", "100"),
    ("beforeAutospacing", "0"), ("afterAutospacing", "false"),
])
def test_line_spacing_shadowed_inherited_line_or_competing_control_is_manual(layer, attribute, value):
    document = _line_spacing_document()
    source, audit = _line_spacing_audit(document)
    properties = _spacing_layer(document, layer)
    spacing = properties.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        properties.append(spacing)
    spacing.set(qn("w:" + attribute), value)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


@pytest.mark.parametrize("layer", ["direct", "style", "base", "default"])
@pytest.mark.parametrize("tag", ["contextualSpacing", "snapToGrid", "numPr", "pPrChange", "unknownProperty"])
def test_line_spacing_used_controls_are_manual_even_when_false_or_shadowed(layer, tag):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    node = OxmlElement("w:" + tag)
    node.set(qn("w:val"), "0")
    _spacing_layer(document, layer).append(node)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


@pytest.mark.parametrize("attribute,value", [
    ("line", None), ("lineRule", None), ("lineRule", "exactly"), ("line", "0"),
    ("line", "-240"), ("line", "0240"), ("line", "+240"), ("line", " 240"),
    ("line", "240.0"), ("line", "NaN"), ("line", "2147483648"),
    ("before", "-1"), ("after", "1.5"), ("beforeLines", "0"),
    ("afterLines", "0"), ("beforeAutospacing", "false"), ("afterAutospacing", "0"),
    ("unknown", "1"),
])
def test_line_spacing_malformed_selected_raw_source_refuses_before_scalar_getters(attribute, value):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    selected = _line_spacing_findings(audit)
    spacing = document.paragraphs[0]._p.pPr.find(qn("w:spacing"))
    if value is None:
        spacing.attrib.pop(qn("w:" + attribute))
    else:
        spacing.set(qn("w:" + attribute), value)
    # Malformed source is deliberately introduced only after a genuine finding
    # is captured. This is service preflight evidence, not an upload journey.
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), selected)


@pytest.mark.parametrize("layer", ["direct", "style", "base", "default"])
@pytest.mark.parametrize("variant", [
    "duplicate_spacing", "duplicate_pPr", "duplicate_leaf", "nested_leaf",
    "property_text", "property_tail", "spacing_text", "spacing_child",
])
def test_line_spacing_every_used_property_container_is_plain_and_unambiguous(layer, variant):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    properties = _spacing_layer(document, layer)
    spacing = properties.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        properties.append(spacing)
    if variant == "duplicate_spacing":
        properties.append(copy.deepcopy(spacing))
    elif variant == "duplicate_pPr":
        properties.getparent().append(copy.deepcopy(properties))
    elif variant == "duplicate_leaf":
        for _ in range(2):
            node = OxmlElement("w:jc")
            node.set(qn("w:val"), "left")
            properties.append(node)
    elif variant == "nested_leaf":
        node = OxmlElement("w:keepLines")
        node.append(OxmlElement("w:keepNext"))
        properties.append(node)
    elif variant == "property_text":
        properties.text = "not property whitespace"
    elif variant == "property_tail":
        spacing.tail = "not property whitespace"
    elif variant == "spacing_text":
        spacing.text = "not a spacing leaf"
    else:
        spacing.append(OxmlElement("w:before"))
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


@pytest.mark.parametrize("variant", [
    "missing_style", "wrong_kind", "duplicate_style", "cycle", "missing_base",
    "duplicate_basedOn", "duplicate_pStyle", "linked_style", "linked_base",
    "missing_styles", "duplicate_docDefaults", "duplicate_pPrDefault", "duplicate_default_pPr",
    "missing_default", "duplicate_default", "nested_defaults",
])
def test_line_spacing_used_style_and_default_graph_must_be_complete(variant):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    paragraph = document.paragraphs[0]
    _spacing_layer(document, "style")
    style = paragraph.style
    if variant == "missing_style":
        paragraph._p.pPr.find(qn("w:pStyle")).set(qn("w:val"), "AbsentLineStyle")
    elif variant == "wrong_kind":
        style.element.set(qn("w:type"), "character")
    elif variant == "duplicate_style":
        document.styles.element.append(copy.deepcopy(style.element))
    elif variant == "cycle":
        style.base_style.base_style = style
    elif variant == "missing_base":
        style.element.find(qn("w:basedOn")).set(qn("w:val"), "AbsentLineBase")
    elif variant == "duplicate_basedOn":
        style.element.append(copy.deepcopy(style.element.find(qn("w:basedOn"))))
    elif variant == "duplicate_pStyle":
        paragraph._p.pPr.append(copy.deepcopy(paragraph._p.pPr.find(qn("w:pStyle"))))
    elif variant in {"linked_style", "linked_base"}:
        node = OxmlElement("w:link")
        node.set(qn("w:val"), "UninterpretedLinkedStyle")
        (style if variant == "linked_style" else style.base_style).element.append(node)
    elif variant == "missing_styles":
        for key, relationship in list(document.part.rels.items()):
            if relationship.reltype.endswith("/styles"):
                del document.part.rels[key]
    elif variant in {"missing_default", "duplicate_default"}:
        # Force default paragraph-style resolution rather than an explicit style.
        paragraph._p.pPr.remove(paragraph._p.pPr.find(qn("w:pStyle")))
        normal = document.styles["Normal"].element
        if variant == "missing_default":
            normal.attrib.pop(qn("w:default"))
        else:
            duplicate = copy.deepcopy(normal)
            duplicate.set(qn("w:styleId"), "SecondDefaultLineStyle")
            document.styles.element.append(duplicate)
    else:
        properties = _spacing_layer(document, "default")
        defaults = properties.getparent().getparent()
        if variant == "duplicate_docDefaults":
            defaults.getparent().append(copy.deepcopy(defaults))
        elif variant == "duplicate_pPrDefault":
            defaults.append(copy.deepcopy(properties.getparent()))
        elif variant == "duplicate_default_pPr":
            properties.getparent().append(copy.deepcopy(properties))
        else:
            defaults.append(OxmlElement("w:unknownDefaults"))
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


@pytest.mark.parametrize("variant", [
    "bookmark", "simple_field", "hyperlink", "tracked_insert", "content_control",
    "run_field", "run_change", "duplicate_rPr", "misplaced_rPr", "paragraph_text",
    "paragraph_tail", "section_boundary", "empty", "numbered_style",
])
def test_line_spacing_selected_structure_does_not_expand_plain_paragraph_envelope(variant):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    paragraph = document.paragraphs[0]
    if variant in {"bookmark", "simple_field", "hyperlink", "tracked_insert", "content_control"}:
        tag = {"bookmark": "bookmarkStart", "simple_field": "fldSimple", "hyperlink": "hyperlink",
               "tracked_insert": "ins", "content_control": "sdt"}[variant]
        node = OxmlElement("w:" + tag)
        if variant != "bookmark":
            node.append(copy.deepcopy(paragraph.runs[0]._r))
        paragraph._p.append(node)
    elif variant in {"run_field", "run_change"}:
        paragraph.runs[0]._r.append(OxmlElement("w:fldChar" if variant == "run_field" else "w:rPrChange"))
    elif variant == "duplicate_rPr":
        run = paragraph.runs[0]._r
        run.get_or_add_rPr()
        run.append(OxmlElement("w:rPr"))
    elif variant == "misplaced_rPr":
        paragraph.runs[0]._r.append(OxmlElement("w:rPr"))
    elif variant == "paragraph_text":
        from lxml import etree
        etree._Element.text.__set__(paragraph._p, "unrepresented raw paragraph text")
    elif variant == "paragraph_tail":
        paragraph.runs[0]._r.tail = "unrepresented raw run tail"
    elif variant == "section_boundary":
        paragraph._p.pPr.append(copy.deepcopy(document.sections[0]._sectPr))
    elif variant == "empty":
        for run in paragraph.runs:
            run.text = ""
    else:
        paragraph.style = "List Number"
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


@pytest.mark.parametrize("grid_type", [None, "default", "lines", "linesAndChars"])
def test_line_spacing_any_main_document_grid_is_manual_even_in_other_section(grid_type):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    document.add_section()
    grid = OxmlElement("w:docGrid")
    grid.set(qn("w:linePitch"), "360")
    if grid_type is not None:
        grid.set(qn("w:type"), grid_type)
    document.sections[-1]._sectPr.append(grid)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


@pytest.mark.parametrize("story", ["table", "header", "footer", "first_page_header", "even_page_footer"])
def test_line_spacing_non_main_story_cannot_be_selected_by_paragraph_alias(story):
    document = _line_spacing_document()
    _, audit = _line_spacing_audit(document)
    selected = _line_spacing_findings(audit)
    original = document.paragraphs[0]
    if story == "table":
        target = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    else:
        target = getattr(document.sections[0], story).paragraphs[0]
    target._p.getparent().replace(target._p, copy.deepcopy(original._p))
    original._p.getparent().remove(original._p)
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), selected)


@pytest.mark.parametrize("role", ["HEADING_1", "APPENDIX_HEADING", "CAPTION", "APPENDIX_FORM"])
def test_line_spacing_wrong_role_never_borrows_body_requirement(role):
    document = _line_spacing_document((480,))
    paragraph = document.paragraphs[0]
    if role == "HEADING_1":
        paragraph.style = "Heading 1"
        for node in list(paragraph.style.element.findall(qn("w:link"))):
            paragraph.style.element.remove(node)
    elif role == "APPENDIX_HEADING":
        paragraph.text = "Appendix A"
    # Headings below use a real heading detector finding; skipped-role forms
    # instead retain a genuine body finding before changing only its context.
    source, audit = _line_spacing_audit(document)
    selected = _line_spacing_findings(audit)
    if role == "CAPTION":
        paragraph.text = "Figure 1: Sample illustration"
    elif role == "APPENDIX_FORM":
        document._element.body.insert(0, OxmlElement("w:p"))
        document.paragraphs[0].text = "Appendix A"
        paragraph.text = "Name:"
        selected[0].location = {"paragraph_index": 1}
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), selected)


@pytest.mark.parametrize("variant", ["duplicate", "table"])
def test_line_spacing_reference_requires_unique_main_story_reference_span(variant):
    profile = new_custom_profile("Unique reference line requirement")
    profile.references.line_spacing = 2.0
    profile.references.check_line_spacing = True
    document = _line_spacing_document(role="REFERENCE_ENTRY")
    _, audit = _line_spacing_audit(document, profile=profile)
    if variant == "duplicate":
        document.add_paragraph("References")
    else:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "References"
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, _bold_bytes(document), _line_spacing_findings(audit))


def test_line_spacing_unused_broken_styles_and_ordinary_sibling_points_remain_available():
    from docx.enum.style import WD_STYLE_TYPE
    from app.services.guided_fix import apply_guided_fixes
    from tests.test_guided_fix_preservation_storage import _assert_line_spacing_package_unchanged

    document = _line_spacing_document()
    unused = document.styles.add_style("Unused broken line style", WD_STYLE_TYPE.PARAGRAPH)
    based = OxmlElement("w:basedOn")
    based.set(qn("w:val"), "AbsentUnusedBase")
    unused.element.append(based)
    unused.element.append(OxmlElement("w:link"))
    for layer in ("style", "default"):
        properties = _spacing_layer(document, layer)
        spacing = properties.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            properties.append(spacing)
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "240")
    source, audit = _line_spacing_audit(document)
    selected = _line_spacing_findings(audit)
    output, _, _ = apply_guided_fixes(audit, source, selected)
    _assert_line_spacing_package_unchanged(source, output, {0: 360})


def test_line_spacing_inherited_only_upload_has_no_invented_selectable_finding(client, bold_boundary_storage):
    document = _line_spacing_document()
    spacing = document.paragraphs[0]._p.pPr.find(qn("w:spacing"))
    for attribute in ("line", "lineRule"):
        del spacing.attrib[qn("w:" + attribute)]
    document.styles["Normal"].paragraph_format.line_spacing = 1.0
    profile = new_custom_profile("Inherited line spacing is not direct evidence")
    profile.body.line_spacing = 1.5
    response = client.post("/api/audit", data={"custom_profile": json.dumps(profile.to_dict())},
                           files={"file": ("inherited.docx", _bold_bytes(document))})
    assert response.status_code == 200, response.text
    assert not any(f["rule_code"] == "LINE_SPACING" for f in response.json()["physical_layout_errors"])


@pytest.mark.parametrize("unsafe", ["absolute_second", "stock_context"])
def test_line_spacing_safe_and_unsafe_source_selection_has_no_review_or_output(
    client, test_engine, bold_boundary_storage, unsafe,
):
    document = _line_spacing_document((240, 480))
    if unsafe == "absolute_second":
        document.paragraphs[1]._p.pPr.find(qn("w:spacing")).set(qn("w:lineRule"), "exact")
    else:
        document.styles["Normal"].paragraph_format.line_spacing = 1.15
    source = _bold_bytes(document)
    audit, root, ids = _line_spacing_api_audit(client, source)
    assert len(ids) == 2
    before = _saved_source_state(test_engine, audit["audit_id"])
    response = _upload(client, root + "/review", source, ids)
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == before


@pytest.mark.parametrize("fault", [
    "line_plus_one", "line_minus_one", "missing_rule", "exact", "atLeast", "absolute_alias",
    "absolute_alias_atLeast",
    "grouped_last_337", "unchanged_last", "text", "run_boundary", "run_attribute",
    "paragraph_count", "role", "before", "after", "contextual", "autospacing",
])
def test_line_spacing_faulty_output_is_discarded_before_token_or_storage(
    client, test_engine, monkeypatch, bold_boundary_storage, fault,
):
    from app.api import routes
    from app.services import guided_fix
    from pypdf import PdfReader

    document = _line_spacing_document((240, 240))
    heading = document.styles["Heading 1"].element
    for link in list(heading.findall(qn("w:link"))):
        heading.remove(link)
    required = 2.0 if fault in {"absolute_alias", "absolute_alias_atLeast"} else 1.5
    source = _bold_bytes(document)
    digest = hashlib.sha256(source).hexdigest()
    audit, root, ids = _line_spacing_api_audit(client, source, required)
    path = f'/api/audit/{audit["audit_id"]}'
    before = _saved_source_state(test_engine, audit["audit_id"])
    detail, history = client.get(path).json(), client.get("/api/audits").json()

    def pdf_text():
        response = client.get(path + "/export-pdf")
        assert response.status_code == 200
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.content)).pages)

    report = pdf_text()
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    fix = guided_fix._fix_line_spacing
    calls = []

    def corrupt_after_real_setter(paragraph, ratio):
        calls.append(paragraph.text)
        if fault == "unchanged_last" and len(calls) == 2:
            return True
        result = fix(paragraph, ratio)
        if fault == "unchanged_last":
            return result
        if len(calls) != (2 if fault == "grouped_last_337" else 1):
            return result
        spacing = paragraph._p.pPr.find(qn("w:spacing"))
        if fault in {"line_plus_one", "line_minus_one", "grouped_last_337"}:
            value = 337 if fault == "grouped_last_337" else int(spacing.get(qn("w:line"))) + (
                1 if fault == "line_plus_one" else -1)
            spacing.set(qn("w:line"), str(value))
        elif fault == "missing_rule":
            del spacing.attrib[qn("w:lineRule")]
        elif fault in {"exact", "atLeast", "absolute_alias", "absolute_alias_atLeast"}:
            rule = "atLeast" if fault == "absolute_alias_atLeast" else "exact" if fault == "absolute_alias" else fault
            spacing.set(qn("w:lineRule"), rule)
            if fault in {"absolute_alias", "absolute_alias_atLeast"}:
                spacing.set(qn("w:line"), "0.00015748pt")
        elif fault == "text":
            paragraph.runs[0].text += " Changed original text."
        elif fault == "run_boundary":
            text = paragraph.text
            paragraph.runs[0].text = text
            paragraph._p.remove(paragraph.runs[1]._r)
        elif fault == "run_attribute":
            paragraph.runs[0]._r.find(qn("w:t")).set(
                "{http://www.w3.org/XML/1998/namespace}space", "default")
        elif fault == "paragraph_count":
            paragraph._p.getparent().insert(0, copy.deepcopy(paragraph._p))
        elif fault == "role":
            paragraph.style = "Heading 1"
        elif fault in {"before", "after"}:
            spacing.set(qn("w:" + fault), "999")
        elif fault == "contextual":
            paragraph._p.pPr.append(OxmlElement("w:contextualSpacing"))
        else:
            spacing.set(qn("w:beforeAutospacing"), "0")
        return result

    monkeypatch.setattr(guided_fix, "_fix_line_spacing", corrupt_after_real_setter)
    if fault not in {"before", "after"}:
        def no_preservation_escape(*_args, **_kwargs):
            pytest.fail("The independent reopened line/identity oracle must reject this fault first")
        monkeypatch.setattr(routes.docx_preservation, "compare_preservation", no_preservation_escape)

    def no_publication(*_args, **_kwargs):
        pytest.fail("A faulty line copy must not reach output storage")

    monkeypatch.setattr(guided_fix_storage, "store", no_publication)
    response = _upload(client, root, source, ids, review.json()["review_id"])
    assert response.status_code == 422, response.text
    _assert_no_spacing_output(response, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == before
    assert client.get(path).json() == detail
    assert client.get("/api/audits").json() == history
    assert pdf_text() == report
    assert hashlib.sha256(source).hexdigest() == digest


def test_line_spacing_output_oracle_requires_pristine_original_not_mutated_source_doc():
    from app.services.guided_fix import _validate_output, apply_guided_fixes
    from app.services.document_parser import extract_paragraphs
    from app.services.role_classifier import classify_paragraphs

    document = _line_spacing_document()
    source, audit = _line_spacing_audit(document)
    selected = _line_spacing_findings(audit)
    _, targets = review_targets(audit, source, selected)
    roles = classify_paragraphs(document, extract_paragraphs(document))
    output, _, _ = apply_guided_fixes(audit, source, selected)
    already_mutated = Document(io.BytesIO(output))
    with pytest.raises(GuidedFixFailed):
        _validate_output(output, already_mutated, selected, roles, audit, targets)


@pytest.mark.parametrize("failure", [
    "preservation", "baseline_drift", "new_finding", "category_increase",
    "reaudit_unavailable", "page_mapping_unavailable", "storage_failure",
])
def test_line_spacing_publication_barriers_preserve_saved_audit_history_and_pdf(
    client, test_engine, monkeypatch, bold_boundary_storage, failure,
):
    from app.api import routes
    from app.services import guided_fix_verification, docx_pdf_converter
    from app.services.layout_violation import LayoutViolation
    from pypdf import PdfReader

    source = _bold_bytes(_line_spacing_document())
    audit, root, ids = _line_spacing_api_audit(client, source)
    path = f'/api/audit/{audit["audit_id"]}'
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    injected = []
    if failure in {"baseline_drift", "page_mapping_unavailable"}:
        with Session(test_engine) as db:
            saved = db.get(AuditRecord, audit["audit_id"])
            if failure == "baseline_drift":
                saved.weighted_score -= 1
            else:
                saved.paragraph_page_mapping = {"0": 1}
            db.commit()
        if failure == "page_mapping_unavailable":
            def no_page_renderer(_source):
                injected.append("required fresh renderer")
                raise RuntimeError("Controlled unavailable fresh page evidence")
            monkeypatch.setattr(docx_pdf_converter, "convert_docx_to_pdf", no_page_renderer)
        else:
            injected.append("saved score differs")
    elif failure == "preservation":
        def preservation_refusal(*_args):
            injected.append("package preservation")
            raise GuidedFixFailed("Controlled package gate refusal")
        monkeypatch.setattr(routes.docx_preservation, "compare_preservation", preservation_refusal)
    elif failure in {"reaudit_unavailable", "new_finding"}:
        engine = guided_fix_verification.run_static_rules_engine

        def failed_engine(data, **kwargs):
            injected.append(data)
            if failure == "reaudit_unavailable":
                raise RuntimeError("Controlled unavailable deterministic re-audit")
            findings = engine(data, **kwargs)
            if data != source:
                findings.append(LayoutViolation(
                    rule_code="FONT_SIZE", severity="MINOR",
                    location={"paragraph_index": 0, "run_index": 0},
                    message="An unrelated new finding", actual_value="9pt", expected_value="12pt",
                ))
            return findings

        monkeypatch.setattr(guided_fix_verification, "run_static_rules_engine", failed_engine)
    elif failure == "category_increase":
        score = guided_fix_verification.calculate_weighted_score_detailed

        def increased_score(findings):
            result = score(findings)
            injected.append(result)
            if len(injected) == 2:
                breakdown = [replace(row, deduction=row.deduction + 1)
                             if row.category == "page_margins" else row for row in result.breakdown]
                return replace(result, total=result.total - 1, breakdown=breakdown)
            return result

        monkeypatch.setattr(guided_fix_verification, "calculate_weighted_score_detailed", increased_score)
    else:
        def failed_fsync(_descriptor):
            injected.append("fsync")
            raise OSError("Controlled output disk write failure")
        monkeypatch.setattr(guided_fix_storage.os, "fsync", failed_fsync)
    if failure != "storage_failure":
        def no_publication(*_args):
            pytest.fail("Storage must not run after a failed line verification gate")
        monkeypatch.setattr(guided_fix_storage, "store", no_publication)
    saved = _saved_source_state(test_engine, audit["audit_id"])
    detail, history = client.get(path).json(), client.get("/api/audits").json()

    def pdf_text():
        response = client.get(path + "/export-pdf")
        assert response.status_code == 200
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.content)).pages)

    report = pdf_text()
    response = _upload(client, root, source, ids, review.json()["review_id"])
    assert response.status_code == (503 if failure == "storage_failure" else 422), response.text
    assert injected
    if failure in {"new_finding", "category_increase"}:
        assert len(injected) == 2
    _assert_no_spacing_output(response, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved
    assert client.get(path).json() == detail
    assert client.get("/api/audits").json() == history
    assert pdf_text() == report


@pytest.mark.parametrize("phase", ["review", "apply"])
@pytest.mark.parametrize("change,status", [
    ("wrong_original", 409), ("stale_actual", 422), ("stale_expected", 422),
    ("no_snapshot", 409), ("future_snapshot", 409), ("corrupt_snapshot", 409), ("no_hash", 409),
])
def test_line_spacing_saved_candidate_gates_reject_before_mutation(
    client, test_engine, monkeypatch, bold_boundary_storage, phase, change, status,
):
    from app.services import guided_fix

    source = _bold_bytes(_line_spacing_document())
    audit, root, ids = _line_spacing_api_audit(client, source)
    token = None
    if phase == "apply":
        review = _upload(client, root + "/review", source, ids)
        assert review.status_code == 200, review.text
        token = review.json()["review_id"]
    with Session(test_engine) as db:
        saved = db.get(AuditRecord, audit["audit_id"])
        finding = db.get(Violation, ids[0])
        if change == "stale_actual":
            finding.actual_value = "NaN"
        elif change == "stale_expected":
            finding.expected_value = "1.75"
        elif change == "no_snapshot":
            saved.profile_snapshot = None
        elif change == "future_snapshot":
            payload = copy.deepcopy(saved.profile_snapshot)
            payload["schema_version"] = 999
            saved.profile_snapshot = payload
        elif change == "corrupt_snapshot":
            payload = copy.deepcopy(saved.profile_snapshot)
            payload["body"]["line_spacing"] = 2.0
            saved.profile_snapshot = payload
        elif change == "no_hash":
            saved.document_sha256 = None
        db.commit()
    before = _saved_source_state(test_engine, audit["audit_id"])
    if change == "wrong_original":
        different = _line_spacing_document()
        different.paragraphs[0].runs[0].text += " A different original."
        source = _bold_bytes(different)

    def no_mutation(*_args):
        pytest.fail("Wrong source or stale saved evidence must be rejected before mutation")

    monkeypatch.setattr(guided_fix, "_fix_line_spacing", no_mutation)
    response = _upload(client, root + ("/review" if phase == "review" else ""), source, ids, token)
    assert response.status_code == status, response.text
    # _guided_request can reject before consuming a valid Review token; do not
    # prescribe token cleanup at that earlier, unchanged route boundary.
    _assert_no_bold_output(response, bold_boundary_storage, audit["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == before


def test_line_spacing_saved_requirement_survives_current_profile_edits(
    client, test_engine, monkeypatch, bold_boundary_storage,
):
    from app.services import profile_registry
    from tests.test_guided_fix_preservation_storage import _assert_line_spacing_package_unchanged

    profile = new_custom_profile("Historical exact direct line spacing")
    profile.body.line_spacing = 1.5
    monkeypatch.setattr(profile_registry, "BUILTIN_PROFILES", {profile.profile_id: profile})
    source = _bold_bytes(_line_spacing_document())
    response = client.post("/api/audit", params={"profile_id": profile.profile_id},
                           files={"file": ("historical-line.docx", source)})
    assert response.status_code == 200, response.text
    audit = response.json()
    ids = [f["id"] for f in audit["physical_layout_errors"] if f["rule_code"] == "LINE_SPACING"]
    assert ids
    saved = _saved_source_state(test_engine, audit["audit_id"])
    profile.body.line_spacing = 2.0
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    assert review.json()["items"][0]["after"] == "1.5"
    applied = _upload(client, root, source, ids, review.json()["review_id"])
    assert applied.status_code == 200, applied.text
    download = client.get(root + "/download/" + applied.json()["download_token"])
    assert download.status_code == 200
    _assert_line_spacing_package_unchanged(source, download.content, {0: 360})
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved


@pytest.mark.parametrize("expected", ["1.50", " 1.5", "1.5 ", "1.5x", "15e-1", "NaN", None])
def test_line_spacing_required_literal_cannot_be_reinterpreted_or_normalized(expected):
    source, audit = _line_spacing_audit(_line_spacing_document())
    selected = _line_spacing_findings(audit)
    selected[0].expected_value = expected
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, selected)


@pytest.mark.parametrize("raw", [1, 2147483647])
def test_line_spacing_canonical_positive_source_boundaries_are_not_saved_ratio_bounds(raw):
    from app.services.guided_fix import apply_guided_fixes
    from tests.test_guided_fix_preservation_storage import _assert_line_spacing_package_unchanged

    source, audit = _line_spacing_audit(_line_spacing_document((raw,)))
    selected = _line_spacing_findings(audit)
    assert selected[0].actual_value == str(raw / 240)
    assert review_targets(audit, source, selected)[0][0]["before"] == str(raw / 240)
    output, _, _ = apply_guided_fixes(audit, source, selected)
    _assert_line_spacing_package_unchanged(source, output, {0: 360})


def test_line_spacing_plain_tabs_breaks_and_multiple_run_boundaries_are_preserved():
    from app.services.guided_fix import apply_guided_fixes
    from tests.test_guided_fix_preservation_storage import _assert_line_spacing_package_unchanged

    document = _line_spacing_document()
    paragraph = document.paragraphs[0]
    paragraph.runs[0].text = "Ordinary\tacademic\nprose "
    paragraph.runs[1]._r.append(OxmlElement("w:cr"))
    paragraph.runs[1].add_text("with an unchanged trailing boundary.")
    source, audit = _line_spacing_audit(document)
    selected = _line_spacing_findings(audit)
    output, _, _ = apply_guided_fixes(audit, source, selected)
    _assert_line_spacing_package_unchanged(source, output, {0: 360})
    reopened = Document(io.BytesIO(output)).paragraphs[0]
    assert reopened.text == paragraph.text
    assert [run.text for run in reopened.runs] == [run.text for run in paragraph.runs]


@pytest.mark.parametrize("unsafe", [False, True])
def test_line_spacing_and_point_spacing_mixed_selection_is_atomic(
    client, test_engine, bold_boundary_storage, unsafe,
):
    source = _bold_bytes(_line_spacing_document())
    profile = new_custom_profile("Mixed direct line and point spacing")
    profile.body.space_before_pt = 0
    audit, root, line_ids = _line_spacing_api_audit(client, source, profile=profile)
    point_ids = [f["id"] for f in audit["physical_layout_errors"] if f["rule_code"] == "SPACE_BEFORE"]
    assert point_ids
    ids = line_ids + point_ids
    if unsafe:
        with Session(test_engine) as db:
            finding = db.get(Violation, line_ids[0])
            finding.actual_value = "NaN"
            db.commit()
    before = _saved_source_state(test_engine, audit["audit_id"])
    review = _upload(client, root + "/review", source, ids)
    if unsafe:
        assert review.status_code == 422, review.text
        _assert_no_spacing_output(review, bold_boundary_storage, audit["audit_id"])
    else:
        from app.services.docx_preservation import compare_preservation

        assert review.status_code == 200, review.text
        applied = _upload(client, root, source, ids, review.json()["review_id"])
        assert applied.status_code == 200, applied.text
        assert set(applied.json()["resolved"]) == set(ids)
        download = client.get(root + "/download/" + applied.json()["download_token"])
        assert download.status_code == 200
        spacing = Document(io.BytesIO(download.content)).paragraphs[0]._p.pPr.find(qn("w:spacing"))
        assert {key: spacing.get(qn("w:" + key)) for key in ("line", "lineRule", "before", "after")} == {
            "line": "360", "lineRule": "auto", "before": "0", "after": "160",
        }
        compare_preservation(source, download.content, {(0, None, "LINE_SPACING"), (0, None, "SPACE_BEFORE")})
    assert _saved_source_state(test_engine, audit["audit_id"]) == before


@pytest.mark.parametrize("boundary", ["foreign_finding", "foreign_review", "changed_selection", "incomplete_audit"])
def test_line_spacing_api_audit_ownership_and_review_binding_precede_mutation(
    client, test_engine, monkeypatch, bold_boundary_storage, boundary,
):
    from app.services import guided_fix

    source = _bold_bytes(_line_spacing_document((240, 240)))
    audit, root, ids = _line_spacing_api_audit(client, source)
    foreign, foreign_root, foreign_ids = _line_spacing_api_audit(client, source)
    review = _upload(client, root + "/review", source, ids)
    assert review.status_code == 200, review.text
    token, target, chosen = review.json()["review_id"], root, ids
    if boundary == "foreign_finding":
        target, chosen = root + "/review", foreign_ids
    elif boundary == "foreign_review":
        target, chosen = foreign_root, foreign_ids
    elif boundary == "changed_selection":
        chosen = ids[:1]
    else:
        with Session(test_engine) as db:
            db.get(AuditRecord, audit["audit_id"]).status = "pending"
            db.commit()
    saved = _saved_source_state(test_engine, audit["audit_id"])
    other_saved = _saved_source_state(test_engine, foreign["audit_id"])

    def no_mutation(*_args):
        pytest.fail("Audit ownership and Review binding must be checked before a line edit")

    monkeypatch.setattr(guided_fix, "_fix_line_spacing", no_mutation)
    response = _upload(client, target, source, chosen, token)
    assert response.status_code == (422 if boundary in {"foreign_finding", "changed_selection"} else 409)
    _assert_no_bold_output(response, bold_boundary_storage, audit["audit_id"])
    _assert_no_bold_output(response, bold_boundary_storage, foreign["audit_id"])
    assert _saved_source_state(test_engine, audit["audit_id"]) == saved
    assert _saved_source_state(test_engine, foreign["audit_id"]) == other_saved


def test_ambient_aggregate_positions_permits_unselected_truncation_but_strictly_refuses_selected():
    """Unselected ambient aggregate findings evaluate available positions; selected requires 1:1 completeness."""
    from app.services.guided_fix_verification import _positions
    from app.services.layout_violation import LayoutViolation

    finding = LayoutViolation(
        rule_code="REFERENCE_HANGING_INDENT",
        severity="MINOR",
        location={
            "paragraph_index": 2,
            "occurrences": 15,
            "paragraph_indexes": list(range(2, 12)),
        },
        message="Reference entry hanging indent does not match",
        expected_value="0.5in",
        actual_value="0in",
    )
    # Ambient / unselected: evaluates using available stored indexes
    positions = _positions(finding, strict=False)
    assert len(positions) == 10
    assert [p for p, _, _ in positions] == list(range(2, 12))

    # Selected: strictly requires 1:1 completeness
    with pytest.raises(GuidedFixFailed, match="incomplete finding locations"):
        _positions(finding, strict=True)

    # Repeated indexes do not constitute complete distinct selected targets.
    finding.location = {
        "paragraph_index": 2, "occurrences": 2, "paragraph_indexes": [2, 2],
    }
    with pytest.raises(GuidedFixFailed, match="incomplete finding locations"):
        _positions(finding, strict=True)

    # Corrupt location: negative index or occurrences less than stored set
    corrupt = LayoutViolation(
        rule_code="REFERENCE_HANGING_INDENT",
        severity="MINOR",
        location={
            "paragraph_index": 2,
            "occurrences": 5,
            "paragraph_indexes": list(range(2, 12)),
        },
        message="Reference entry hanging indent does not match",
        expected_value="0.5in",
        actual_value="0in",
    )
    with pytest.raises(GuidedFixFailed, match="incomplete finding locations"):
        _positions(corrupt, strict=False)


def test_ambient_aggregate_references_do_not_crash_valid_space_before_reaudit():
    """Unselected ambient aggregate findings (>10 references) do not crash re-audit of valid SPACE_BEFORE fix."""
    from app.services.guided_fix import apply_guided_fixes
    from app.services.docx_preservation import compare_preservation

    document = Document()
    _clean_font_fixture_context(document)
    p0 = document.add_paragraph(
        "Synthetic body prose cites Jones (2020) and Jones (2020) again."
    )
    p0.paragraph_format.space_before = Pt(20)
    document.add_paragraph("References", style="Heading 1")
    for i in range(15):
        document.add_paragraph(f"Smith, A. (2020). Title of paper {i}. Journal of Testing, 1(1), 10-20.")

    profile = new_custom_profile("Ambient 15 references hanging indent")
    profile.references.check_hanging_indent = True
    profile.references.hanging_indent_in = 0.5
    profile.body.space_before_pt = 6

    source, audit = _point_profile_audit(document, profile)
    ref_findings = [f for f in audit.violations if f.rule_code == "REFERENCE_HANGING_INDENT"]
    assert len(ref_findings) == 1
    assert ref_findings[0].location["occurrences"] == 15
    assert len(ref_findings[0].location["paragraph_indexes"]) == 10
    duplicates = [f for f in audit.violations if f.rule_code == "REFERENCE_DUPLICATE"]
    assert len(duplicates) == 1
    assert duplicates[0].location["entry_count"] == 15
    assert len(duplicates[0].location["paragraph_indexes"]) == 10
    citations = [f for f in audit.violations if f.rule_code == "CITATION_MISMATCH"]
    assert len(citations) == 1
    assert citations[0].location["occurrences"] == 2
    assert citations[0].location["paragraph_indexes"] == [0]

    selected = [f for f in audit.violations if f.rule_code == "SPACE_BEFORE"]
    assert len(selected) == 1

    _, targets = review_targets(audit, source, selected)
    corrected, applied, skipped = apply_guided_fixes(audit, source, selected)
    assert applied == selected and not skipped
    compare_preservation(source, corrected, targets)

    # verify_reaudit now succeeds despite 15 ambient reference entries
    result = verify_reaudit(audit, source, corrected, selected)
    assert result["before_score"] < result["after_score"]
    assert result["new_findings"] == []
    for side in ("before_findings", "after_findings"):
        summary = {item["rule_code"]: item for item in result[side]}
        assert summary["REFERENCE_HANGING_INDENT"]["count"] == 15
        assert summary["REFERENCE_DUPLICATE"]["count"] == 15
        assert summary["CITATION_MISMATCH"]["count"] == 2

    # Selecting the truncated aggregate itself continues to fail at Review
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, ref_findings)

    # Attempting to verify re-audit with the truncated aggregate selected fails strictly
    with pytest.raises(GuidedFixFailed, match="incomplete finding locations"):
        verify_reaudit(audit, source, corrected, ref_findings)


def test_ambient_aggregate_worsened_occurrence_count_is_detected():
    """Increasing occurrence count of an ambient aggregate finding triggers re-audit failure."""
    from app.services.guided_fix import apply_guided_fixes

    document = Document()
    _clean_font_fixture_context(document)
    p0 = document.add_paragraph("Synthetic body prose with enough text to qualify for the body role.")
    p0.paragraph_format.space_before = Pt(20)
    document.add_paragraph("References", style="Heading 1")
    for i in range(15):
        document.add_paragraph(f"Author{i:02d}, A. (2020). Title of paper {i}. Journal of Testing, 1(1), 10-20.")

    profile = new_custom_profile("Ambient 15 references hanging indent")
    profile.references.check_hanging_indent = True
    profile.references.hanging_indent_in = 0.5
    profile.body.space_before_pt = 6

    source, audit = _point_profile_audit(document, profile)
    selected = [f for f in audit.violations if f.rule_code == "SPACE_BEFORE"]
    corrected, applied, _ = apply_guided_fixes(audit, source, selected)

    # Create a corrupted output DOCX that adds a 16th non-compliant reference entry
    corrupted_doc = Document(io.BytesIO(corrected))
    corrupted_doc.add_paragraph("Author15, A. (2020). Another unindented paper. Journal of Testing, 1(1), 10-20.")
    corrupted_buf = io.BytesIO()
    corrupted_doc.save(corrupted_buf)
    corrupted_bytes = corrupted_buf.getvalue()

    with pytest.raises(GuidedFixFailed, match="new or worsened finding after re-audit"):
        verify_reaudit(audit, source, corrupted_bytes, selected)
