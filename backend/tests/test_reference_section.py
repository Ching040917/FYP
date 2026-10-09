"""Reference List structure + formatting-evidence foundation tests.

Proves the shared heading helper, entry selection, indentation evidence and
page-start evidence behave as documented — and that none of it changes any
existing consumer's behaviour.
"""

import io

import pytest
from docx import Document

from app.services import citation_sensor as cs
from app.services import role_classifier as rc
from app.services.document_parser import extract_paragraphs
from app.services.formatting_evidence import (
    SOURCE_DIRECT,
    SOURCE_NUMBERING,
    SOURCE_PARAGRAPH_STYLE,
    SOURCE_STYLE_CHAIN,
    SOURCE_UNRESOLVED,
    resolve_indentation,
    resolve_paragraph_format,
)
from app.services.reference_section import (
    PAGE_START_EXPLICIT_NEW_PAGE,
    PAGE_START_MAPPED_NEW_PAGE,
    PAGE_START_MAPPED_SAME_PAGE,
    PAGE_START_UNAVAILABLE,
    PAGE_START_UNCERTAIN,
    REFERENCE_ALIASES_NOT_CURRENTLY_MATCHED,
    classify_reference_page_start,
    find_reference_heading_index,
    is_reference_heading_text,
    normalise_reference_heading,
    reference_entry_indexes,
    reference_section_bounds,
)
from tests import reference_structure_fixtures as fx


def _paragraphs(docx_bytes: bytes):
    return extract_paragraphs(Document(io.BytesIO(docx_bytes)))


def _objects(docx_bytes: bytes):
    return {i: p for i, p in enumerate(Document(io.BytesIO(docx_bytes)).paragraphs)}


# --- Phase 1: shared heading helper -----------------------------------------


def test_accepted_aliases_are_recognised_case_insensitively():
    for alias in ("references", "References", "REFERENCES", "References "):
        assert is_reference_heading_text(alias) is True, alias


def test_every_documented_alias_variant_is_recognised():
    for name, data in fx.heading_variants().items():
        paragraphs = _paragraphs(data)
        found = find_reference_heading_index(paragraphs)
        if name.endswith("(singular, not an alias)") or name.startswith("Prose"):
            assert found is None, f"{name} must not be treated as a heading"
        else:
            assert found is not None, f"{name} must be recognised"


def test_numbered_section_prefix_is_stripped_before_matching():
    assert normalise_reference_heading("6 References") == "references"
    assert normalise_reference_heading("6.1 References") == "references"
    assert normalise_reference_heading("6.1.2 Bibliography") == "bibliography"


def test_prose_containing_the_word_is_not_a_heading():
    assert is_reference_heading_text("References are listed below") is False
    assert is_reference_heading_text("See References") is False
    assert is_reference_heading_text("") is False
    assert is_reference_heading_text(None) is False


def test_singular_reference_is_not_currently_an_alias():
    # Guarded explicitly so enabling it later is a deliberate profile change,
    # never an accidental behaviour drift.
    assert "reference" in REFERENCE_ALIASES_NOT_CURRENTLY_MATCHED
    assert is_reference_heading_text("Reference") is False


def test_alias_parity_with_previous_consumers():
    """The shared helper agrees with both implementations it replaced."""
    corpus = [
        "References", "references", "REFERENCES", " Bibliography ",
        "Works Cited", "works cited", "Reference List", "reference list",
        "6 References", "6. References", "6.1 References", "6.1. Bibliography",
        "Reference", "References are listed below", "", "   ", "refs",
        "Appendix A", "Reference lists are useful",
    ]
    for text in corpus:
        shared = is_reference_heading_text(text)
        sensor = cs._is_references_header(text)
        classifier = rc._is_references_heading_text(text)
        assert sensor == classifier == shared, text


def test_numeric_prefix_regex_is_unchanged():
    # Both modules previously carried a byte-identical copy of this regex.
    from app.services.reference_section import NUMERIC_SECTION_PREFIX

    assert NUMERIC_SECTION_PREFIX.pattern == r"^\d+(?:\.\d+)*\.?\s+"


# --- Phase 2: entry selection -----------------------------------------------


def test_entry_selection_stops_at_the_appendix_heading():
    paragraphs = _paragraphs(fx.appendix_terminates_section())
    bounds = reference_section_bounds(paragraphs)
    assert bounds is not None
    by_index = paragraphs_by_index(paragraphs)
    assert by_index[bounds.heading_index] == "References"
    # The Appendix heading and anything after it are excluded.
    entries = [by_index[i] for i in bounds.entry_indexes]
    assert all(not t.startswith("Appendix") for t in entries)
    assert len(entries) == 2
    assert bounds.end_index == 4  # the "Appendix A" paragraph
    assert bounds.first_entry_index == bounds.entry_indexes[0]


def paragraphs_by_index(paragraphs):
    return {p["index"]: p["text"] for p in paragraphs}


def test_empty_paragraphs_are_separators_not_entries():
    paragraphs = _paragraphs(fx.empty_paragraphs_inside_section())
    # Blank and whitespace-only paragraphs separate entries; they are never
    # themselves entries.
    assert reference_entry_indexes(paragraphs) == (2, 5)


def test_long_wrapped_entry_stays_one_paragraph():
    paragraphs = _paragraphs(fx.long_wrapped_entry())
    assert reference_entry_indexes(paragraphs) == (2,)


def test_manual_line_break_stays_one_paragraph():
    paragraphs = _paragraphs(fx.manual_line_break_entry())
    assert reference_entry_indexes(paragraphs) == (2,)


def test_no_heading_yields_no_section():
    docx_bytes = fx.body_with_heading("Not a heading")
    assert reference_section_bounds(_paragraphs(docx_bytes)) is None
    assert reference_entry_indexes(_paragraphs(docx_bytes)) == ()


def test_table_and_textbox_references_are_excluded():
    for builder in (fx.table_based_reference, fx.textbox_reference):
        paragraphs = _paragraphs(builder())
        indexes = reference_entry_indexes(paragraphs)
        # Neither structure is exposed as a parser paragraph, so neither
        # becomes an entry. The heading itself is also absent from entries.
        assert all(i != 0 for i in indexes)
        texts = [paragraphs_by_index(paragraphs)[i] for i in indexes]
        assert not any("Garcia" in t for t in texts), builder.__name__


# --- Phase 3: indentation evidence -----------------------------------------


def test_direct_hanging_indent_resolves_from_w_hanging():
    objs = _objects(fx.direct_hanging_indent(0.5))
    evidence = resolve_indentation(objs[1])
    assert evidence.left_indent_in == pytest.approx(0.5)
    # Verified: w:hanging surfaces as a NEGATIVE first-line indent.
    assert evidence.first_line_indent_in == pytest.approx(-0.5)
    assert evidence.hanging_indent_in == pytest.approx(0.5)
    assert evidence.source == SOURCE_DIRECT
    assert evidence.is_resolved


def test_positive_first_line_indent_is_not_a_hanging_indent():
    objs = _objects(fx.direct_first_line_indent(0.5))
    evidence = resolve_indentation(objs[1])
    assert evidence.first_line_indent_in == pytest.approx(0.5)
    assert evidence.hanging_indent_in is None
    assert evidence.left_indent_in is None


def test_wrong_hanging_indent_distance_is_reported_verbatim():
    objs = _objects(fx.direct_hanging_indent(1.0))
    assert resolve_indentation(objs[1]).hanging_indent_in == pytest.approx(1.0)


def test_no_explicit_indent_anywhere_is_unresolved_not_zero():
    objs = _objects(fx.no_explicit_indent())
    evidence = resolve_indentation(objs[1])
    assert evidence.source == SOURCE_UNRESOLVED
    assert evidence.left_indent_in is None
    assert evidence.hanging_indent_in is None
    # Critically: not silently reported as 0.
    assert evidence.left_indent_in != 0
    assert evidence.hanging_indent_in != 0


def test_style_inherited_hanging_indent_resolves_to_paragraph_style():
    objs = _objects(fx.style_inherited_hanging_indent(0.5))
    evidence = resolve_indentation(objs[1])
    assert evidence.source == SOURCE_PARAGRAPH_STYLE
    assert evidence.hanging_indent_in == pytest.approx(0.5)


def test_base_style_inheritance_resolves_to_style_chain():
    objs = _objects(fx.base_style_inherited_hanging_indent(0.5))
    evidence = resolve_indentation(objs[1])
    assert evidence.source == SOURCE_STYLE_CHAIN
    assert evidence.hanging_indent_in == pytest.approx(0.5)


def test_numbered_paragraph_is_numbering_sourced_and_never_zero():
    objs = _objects(fx.numbered_reference())
    evidence = resolve_indentation(objs[1])
    assert evidence.is_numbered is True
    assert evidence.source == SOURCE_NUMBERING
    # Indentation comes from the list definition, which this module does not
    # read — so no value is asserted and no mismatch can be implied.
    assert evidence.hanging_indent_in is None
    assert evidence.left_indent_in is None


def test_indentation_evidence_exposes_no_low_level_units():
    objs = _objects(fx.direct_hanging_indent(0.5))
    evidence = resolve_indentation(objs[1])
    for value in (evidence.left_indent_in, evidence.hanging_indent_in):
        assert isinstance(value, float)
        # Inches, not twips (720) and not EMU (457200).
        assert 0 < value < 10


# --- Phase 4: paragraph format evidence ------------------------------------


def test_direct_alignment_and_spacing_are_captured():
    objs = _objects(fx.direct_alignment_and_spacing())
    evidence = resolve_paragraph_format(objs[1])
    assert "JUSTIFY" in str(evidence.alignment)
    assert evidence.line_spacing == pytest.approx(2.0)
    assert evidence.space_before_pt == pytest.approx(0.0)
    assert evidence.space_after_pt == pytest.approx(0.0)
    assert evidence.alignment_source == SOURCE_DIRECT


def test_inherited_alignment_and_spacing_stay_unresolved():
    objs = _objects(fx.inherited_alignment_and_spacing())
    evidence = resolve_paragraph_format(objs[1])
    # python-docx resolves no inheritance here: unresolved, never substituted.
    assert evidence.line_spacing is None
    assert evidence.space_before_pt is None
    assert evidence.space_after_pt is None
    assert evidence.line_spacing_source == SOURCE_UNRESOLVED
    assert evidence.space_before_source == SOURCE_UNRESOLVED


def test_effective_font_family_and_size_are_resolved():
    objs = _objects(fx.effective_font_family_and_size("Times New Roman", 12.0))
    evidence = resolve_paragraph_format(objs[1])
    assert evidence.effective_font_family == "Times New Roman"
    assert evidence.effective_font_size_pt == pytest.approx(12.0)
    assert evidence.has_mixed_font_family is False
    assert evidence.has_mixed_font_size is False


def test_mixed_font_family_is_reported_explicitly():
    objs = _objects(fx.mixed_font_family())
    evidence = resolve_paragraph_format(objs[1])
    assert evidence.has_mixed_font_family is True


def test_mixed_font_size_is_reported_explicitly():
    objs = _objects(fx.mixed_font_size())
    evidence = resolve_paragraph_format(objs[1])
    assert evidence.has_mixed_font_size is True


# --- Phase 5: page-start evidence ------------------------------------------


def test_explicit_page_break_run_is_proven_from_ooxml():
    data = fx.explicit_page_break_run_before_heading()
    paragraphs, objs = _paragraphs(data), _objects(data)
    heading = find_reference_heading_index(paragraphs)
    # The break run lives in the heading paragraph; the text still matches.
    assert heading is not None
    evidence = classify_reference_page_start(
        paragraphs, heading, paragraph_objs=objs
    )
    assert evidence.state == PAGE_START_EXPLICIT_NEW_PAGE
    assert evidence.has_explicit_page_break is True


def test_page_break_before_property_is_proven_from_ooxml():
    data = fx.page_break_before_property()
    paragraphs, objs = _paragraphs(data), _objects(data)
    heading = find_reference_heading_index(paragraphs)
    evidence = classify_reference_page_start(
        paragraphs, heading, paragraph_objs=objs
    )
    assert evidence.state == PAGE_START_EXPLICIT_NEW_PAGE
    assert evidence.has_page_break_before is True


def test_explicit_break_outranks_a_conflicting_page_mapping():
    data = fx.explicit_page_break_run_before_heading()
    paragraphs, objs = _paragraphs(data), _objects(data)
    heading = find_reference_heading_index(paragraphs)
    # A mapping that claims SAME page must not override proven OOXML.
    evidence = classify_reference_page_start(
        paragraphs, heading,
        page_mapping={0: 1, heading: 1},
        paragraph_objs=objs,
    )
    assert evidence.state == PAGE_START_EXPLICIT_NEW_PAGE


def test_natural_pagination_with_mapping_is_distinct_from_explicit():
    data = fx.no_page_break()
    paragraphs = _paragraphs(data)
    heading = find_reference_heading_index(paragraphs)
    later = classify_reference_page_start(
        paragraphs, heading, page_mapping={0: 1, heading: 3}
    )
    same = classify_reference_page_start(
        paragraphs, heading, page_mapping={0: 1, heading: 1}
    )
    assert later.state == PAGE_START_MAPPED_NEW_PAGE
    assert same.state == PAGE_START_MAPPED_SAME_PAGE
    # Neither claims an explicit break.
    assert later.has_explicit_page_break is False
    assert later.has_page_break_before is False


def test_missing_mapping_is_unavailable_and_implies_nothing():
    paragraphs = _paragraphs(fx.no_page_break())
    heading = find_reference_heading_index(paragraphs)
    evidence = classify_reference_page_start(paragraphs, heading, page_mapping=None)
    assert evidence.state == PAGE_START_UNAVAILABLE
    assert evidence.heading_page is None


def test_partial_mapping_is_uncertain():
    paragraphs = _paragraphs(fx.no_page_break())
    heading = find_reference_heading_index(paragraphs)
    # Heading mapped, predecessor unknown.
    evidence = classify_reference_page_start(
        paragraphs, heading, page_mapping={heading: 2}
    )
    assert evidence.state == PAGE_START_UNCERTAIN


def test_page_start_helper_is_read_only():
    """It reports evidence only; it never raises a violation type."""
    import inspect

    sig = inspect.signature(classify_reference_page_start)
    assert "violation" not in " ".join(sig.parameters).lower()
