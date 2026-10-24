"""Reference-heading Guided Fix identity, direct evidence and verified copies."""
import hashlib
import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx.shared import Pt
from app.services.profile_schema import new_custom_profile
from app.services.guided_fix import GuidedFixFailed, review_targets
from app.services.profile_snapshot import resolve_snapshot
from tests.test_guided_fixes import _fix_post


RULE = "REFERENCE_HEADING_ALIGNMENT"


def _source(*, duplicate=False, table=False, field=False, inherited=False):
    doc = Document()
    doc.add_paragraph("A body paragraph.")
    heading = doc.add_paragraph("References")
    if inherited:
        heading.style = "Heading 1"
        doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if field:
        run = heading.add_run()
        run._r.append(OxmlElement("w:fldChar"))
        run._r[-1].set(qn("w:fldCharType"), "begin")
    doc.add_paragraph("Smith, A. (2024). An entry. Press.")
    if duplicate:
        doc.add_paragraph("References")
    if table:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "References"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _audit(client, source, *, enabled=True, required="center"):
    profile = new_custom_profile("Reference heading alignment")
    profile.references.check_heading_alignment = enabled
    profile.references.heading_alignment = required
    result = client.post("/api/audit", data={"custom_profile": json.dumps(profile.to_dict())},
                         files={"file": ("source.docx", source)})
    assert result.status_code == 200, result.text
    findings = [v for v in result.json()["physical_layout_errors"] if v["rule_code"] == RULE]
    return result.json(), findings


@pytest.mark.parametrize("inherited", [False, True])
def test_unique_heading_review_apply_download_and_reaudit(client, tmp_path, monkeypatch, inherited):
    from app.services import guided_fix_storage
    monkeypatch.setattr(guided_fix_storage.settings, "GUIDED_FIX_STORAGE_DIR", str(tmp_path))
    source = _source(inherited=inherited)
    original_digest = hashlib.sha256(source).digest()
    audit, findings = _audit(client, source)
    assert len(findings) == 1
    finding = findings[0]
    assert finding["location"] == {"paragraph_index": 1}
    assert finding["actual_value"] == "Left"
    assert finding["expected_value"] == "Center"
    base = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    selection = {"selection": json.dumps({"violation_ids": [finding["id"]]})}
    review = client.post(base + "/review", data=selection, files={"file": ("source.docx", source)})
    assert review.status_code == 200, review.text
    assert review.json()["items"][0]["locations"] == ["Paragraph 2"]
    assert review.json()["items"][0]["before"] == "Left"
    assert review.json()["items"][0]["after"] == "Center"
    applied = client.post(base, data={**selection, "review_id": review.json()["review_id"]},
                          files={"file": ("source.docx", source)})
    assert applied.status_code == 200, applied.text
    assert applied.json()["status"] == "passed"
    assert applied.json()["resolved"] == [finding["id"]]
    assert applied.json()["new_findings"] == []
    output = client.get(base + "/download/" + applied.json()["download_token"])
    assert output.status_code == 200
    doc = Document(io.BytesIO(output.content))
    assert doc.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert doc.paragraphs[1].text == "References"
    assert doc.paragraphs[0].text == "A body paragraph."
    assert hashlib.sha256(source).digest() == original_digest
    assert hashlib.sha256(output.content).digest() != original_digest
    assert RULE not in [v["rule_code"] for v in _audit(client, output.content)[1]]


@pytest.mark.parametrize("variant", ["duplicate", "table", "field"])
def test_ambiguous_or_unsafe_heading_has_no_download(client, variant):
    source = _source(**{variant: True})
    audit, findings = _audit(client, source)
    assert findings
    response = _fix_post(client, audit["audit_id"], source, [findings[0]["id"]])
    assert response.status_code == 422


def test_disabled_profile_has_no_heading_candidate(client):
    _, findings = _audit(client, _source(), enabled=False)
    assert findings == []


def test_already_matching_heading_has_no_candidate(client):
    _, findings = _audit(client, _source(), required="left")
    assert findings == []


@pytest.mark.parametrize("change", ["wrong_location", "extra_location", "wrong_required", "stale_actual", "disabled"])
def test_persisted_target_must_match_heading_source_and_saved_profile(change):
    source = _source()
    profile = new_custom_profile("Persisted Reference rule")
    profile.references.check_heading_alignment = change != "disabled"
    profile.references.heading_alignment = "center"
    audit = SimpleNamespace(profile_snapshot=resolve_snapshot(profile).to_dict())
    finding = SimpleNamespace(
        id="heading-1", rule_code=RULE, location={"paragraph_index": 1},
        actual_value="Left", expected_value="Center",
    )
    if change == "wrong_location":
        finding.location = {"paragraph_index": 0}
    elif change == "extra_location":
        finding.location["paragraph_indexes"] = [1]
    elif change == "wrong_required":
        finding.expected_value = "Right"
    elif change == "stale_actual":
        finding.actual_value = "Right"
    with pytest.raises(GuidedFixFailed):
        review_targets(audit, source, [finding])


BOLD_RULE = "REFERENCE_HEADING_BOLD"


def _bold_doc(current=False, *, segmented=False, rich=False):
    doc = Document()
    body = doc.add_paragraph("A body paragraph.")
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for text in (("Ref", "erences", " ") if segmented else ("References",)):
        run = heading.add_run(text)
        run.font.bold = current
        run.font.italic = True
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    if segmented:
        heading.add_run()  # No text: its absent bold must remain absent.
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "7")
        start.set(qn("w:name"), "ReferenceHeading")
        heading._p.insert(1, start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "7")
        heading._p.append(end)
    doc.add_paragraph("Smith, A. (2024). An entry. Press.")
    if rich:
        import base64
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI

        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "DATE")
        body._p.append(field)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Untouched table content."
        doc.add_paragraph("Untouched numbered item.", style="List Number")
        doc.sections[0].header.paragraphs[0].text = "Untouched header"
        doc.sections[0].footer.paragraphs[0].text = "Untouched footer"
        doc.add_picture(io.BytesIO(base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aA6sAAAAASUVORK5CYII="
        )))
        part = Part(PackURI("/custom/aca.bin"), "application/octet-stream",
                    b"Untouched opaque content", doc.part.package)
        doc.part.relate_to(part, "urn:aca:preservation-fixture")
    return doc


def _bold_bytes(doc):
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _bold_audit(client, source, *, required=True, alignment=False):
    profile = new_custom_profile("Reference heading bold")
    profile.references.check_heading_bold = True
    profile.references.heading_bold = required
    profile.references.check_heading_alignment = alignment
    profile.references.heading_alignment = "center" if alignment else None
    response = client.post("/api/audit", data={"custom_profile": json.dumps(profile.to_dict())},
                           files={"file": ("source.docx", source)})
    assert response.status_code == 200, response.text
    return response.json()


@pytest.mark.parametrize("required", [True, False])
@pytest.mark.parametrize("segmented", [False, True])
def test_reference_heading_bold_review_apply_download_and_reaudit(
    client, tmp_path, monkeypatch, required, segmented,
):
    import zipfile
    from app.services import guided_fix_storage
    from app.services.docx_preservation import compare_preservation

    monkeypatch.setattr(guided_fix_storage.settings, "GUIDED_FIX_STORAGE_DIR", str(tmp_path))
    source = _bold_bytes(_bold_doc(not required, segmented=segmented, rich=True))
    original_digest = hashlib.sha256(source).hexdigest()
    audit = _bold_audit(client, source, required=required)
    findings = [v for v in audit["physical_layout_errors"] if v["rule_code"] == BOLD_RULE]
    assert len(findings) == 1
    finding = findings[0]
    assert finding["location"] == {"paragraph_index": 1}
    before = client.get(f'/api/audit/{audit["audit_id"]}').json()
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    selection = {"selection": json.dumps({"violation_ids": [finding["id"]]})}
    reviewed = client.post(root + "/review", data=selection, files={"file": ("source.docx", source)})
    assert reviewed.status_code == 200, reviewed.text
    item = reviewed.json()["items"][0]
    heading_text = Document(io.BytesIO(source)).paragraphs[1].text
    assert item["before"] == ("Not bold" if required else "Bold")
    assert item["after"] == ("Bold" if required else "Not bold")
    assert item["property"] == "Reference List heading bold (all heading text)"
    assert item["affected_count"] == 1
    assert item["locations"] == [f'Reference List heading "{heading_text}" — Paragraph 2']
    applied = client.post(root, data={**selection, "review_id": reviewed.json()["review_id"]},
                          files={"file": ("source.docx", source)})
    assert applied.status_code == 200, applied.text
    payload = applied.json()
    assert payload["status"] == "passed"
    assert payload["resolved"] == [finding["id"]]
    assert payload["new_findings"] == []
    downloaded = client.get(root + "/download/" + payload["download_token"])
    assert downloaded.status_code == 200
    assert client.get(root + "/download/" + payload["download_token"]).status_code == 410
    assert client.post(root, data={**selection, "review_id": reviewed.json()["review_id"]},
                       files={"file": ("source.docx", source)}).status_code == 409
    original = Document(io.BytesIO(source))
    corrected = Document(io.BytesIO(downloaded.content))
    heading = corrected.paragraphs[1]
    expected_indexes = [i for i, run in enumerate(original.paragraphs[1].runs) if run.text != ""]
    assert all(heading.runs[i].font.bold is required for i in expected_indexes)
    assert [(run.text, run.font.italic, run.font.name, run.font.size) for run in heading.runs] == [
        (run.text, run.font.italic, run.font.name, run.font.size) for run in original.paragraphs[1].runs
    ]
    if segmented:
        assert heading.runs[-1].font.bold is None
    assert heading.text == original.paragraphs[1].text
    assert heading.alignment == WD_ALIGN_PARAGRAPH.LEFT
    targets = {(1, i, BOLD_RULE) for i in expected_indexes}
    compare_preservation(source, downloaded.content, targets)
    with zipfile.ZipFile(io.BytesIO(source)) as a, zipfile.ZipFile(io.BytesIO(downloaded.content)) as b:
        assert set(a.namelist()) == set(b.namelist())
        for name in a.namelist():
            if name not in {"word/document.xml", "docProps/core.xml", "docProps/app.xml"}:
                assert a.read(name) == b.read(name), name
    assert hashlib.sha256(source).hexdigest() == original_digest
    assert hashlib.sha256(downloaded.content).hexdigest() != original_digest
    assert client.get(f'/api/audit/{audit["audit_id"]}').json() == before
    reaudit = _bold_audit(client, downloaded.content, required=required)
    assert BOLD_RULE not in [v["rule_code"] for v in reaudit["physical_layout_errors"]]


def _bold_saved(required=True):
    profile = new_custom_profile("Saved bold requirement")
    profile.references.check_heading_bold = True
    profile.references.heading_bold = required
    return SimpleNamespace(profile_snapshot=resolve_snapshot(profile).to_dict())


def _bold_finding(required=True):
    return SimpleNamespace(
        id="bold-heading", rule_code=BOLD_RULE, location={"paragraph_index": 1},
        actual_value="Not bold" if required else "Bold",
        expected_value="Bold" if required else "Not bold",
    )


@pytest.mark.parametrize("variant", [
    "inherited", "partly_unresolved", "mixed", "whitespace_unresolved", "invalid_bool",
    "duplicate_bold", "duplicate_rpr", "bold_attribute", "bold_child", "bold_text",
    "bCs", "cs", "rtl", "bidi", "style_bCs", "default_bCs", "style_cycle", "missing_style",
    "hyperlink", "simple_field", "complex_field", "revision", "drawing", "textbox",
    "math", "sdt", "alternate", "unknown", "numbered", "inherited_numbered",
    "section", "table_heading", "duplicate_heading", "cross_bookmark",
    "missing_ppr", "malformed_based_on", "invalid_style_bold", "invalid_default_flag",
    "unknown_default_wrapper", "malformed_default_wrapper",
])
def test_reference_heading_bold_strict_evidence_unavailable(variant):
    from app.services.guided_fix import ReferenceHeadingBoldUnavailable, _reference_heading_bold_targets
    from app.services.document_parser import extract_paragraphs

    doc = _bold_doc(False, segmented=True)
    heading = doc.paragraphs[1]
    run = heading.runs[0]
    if variant == "inherited":
        for segment in heading.runs:
            segment.font.bold = None
        doc.styles["Normal"].font.bold = False
    elif variant == "partly_unresolved":
        heading.runs[1].font.bold = None
    elif variant == "whitespace_unresolved":
        heading.runs[2].font.bold = None
    elif variant == "mixed":
        heading.runs[1].font.bold = True
    elif variant == "invalid_bool":
        run._r.rPr.b.set(qn("w:val"), "maybe")
    elif variant == "duplicate_bold":
        run._r.rPr.append(OxmlElement("w:b"))
    elif variant == "duplicate_rpr":
        run._r.append(OxmlElement("w:rPr"))
    elif variant == "bold_attribute":
        run._r.rPr.b.set(qn("w:unknown"), "1")
    elif variant == "bold_child":
        run._r.rPr.b.append(OxmlElement("w:i"))
    elif variant == "bold_text":
        run._r.rPr.b.text = "unverifiable"
    elif variant in {"bCs", "cs", "rtl"}:
        run._r.rPr.append(OxmlElement("w:" + variant))
    elif variant == "bidi":
        heading._p.pPr.append(OxmlElement("w:bidi"))
    elif variant == "style_bCs":
        doc.styles["Normal"].element.get_or_add_rPr().append(OxmlElement("w:bCs"))
    elif variant == "default_bCs":
        doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).append(OxmlElement("w:bCs"))
    elif variant == "style_cycle":
        doc.styles["Normal"].base_style = doc.styles["Normal"]
    elif variant == "missing_style":
        heading._p.get_or_add_pPr().get_or_add_pStyle().val = "NoSuchStyle"
    elif variant == "missing_ppr":
        heading._p.remove(heading._p.pPr)
    elif variant == "malformed_based_on":
        from docx.enum.style import WD_STYLE_TYPE
        style = doc.styles.add_style("Direct bold source", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.element.find(qn("w:basedOn")).append(OxmlElement("w:unknown"))
        heading.style = style
    elif variant == "invalid_style_bold":
        doc.styles["Normal"].font.bold = True
        doc.styles["Normal"].element.rPr.b.set(qn("w:val"), "maybe")
    elif variant == "invalid_default_flag":
        doc.styles["Default Paragraph Font"].element.set(qn("w:default"), "maybe")
    elif variant == "unknown_default_wrapper":
        doc.styles.element.find(qn("w:docDefaults")).append(OxmlElement("w:unknown"))
    elif variant == "malformed_default_wrapper":
        doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).append(OxmlElement("w:unknown"))
    elif variant in {"numbered", "section"}:
        heading._p.pPr.append(OxmlElement("w:numPr" if variant == "numbered" else "w:sectPr"))
    elif variant == "inherited_numbered":
        heading.style = "List Number"
    elif variant == "table_heading":
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "References"
    elif variant == "duplicate_heading":
        doc.add_paragraph("References")
    elif variant == "cross_bookmark":
        end = heading._p.find(qn("w:bookmarkEnd"))
        heading._p.remove(end)
        doc.paragraphs[0]._p.append(end)
    elif variant == "alternate":
        from lxml import etree
        heading._p.append(etree.Element(
            "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
        ))
    else:
        tag = {
            "hyperlink": "w:hyperlink", "simple_field": "w:fldSimple",
            "complex_field": "w:fldChar", "revision": "w:ins", "drawing": "w:drawing",
            "textbox": "w:txbxContent", "math": "m:oMath", "sdt": "w:sdt",
            "unknown": "w:unknown",
        }[variant]
        heading._p.append(OxmlElement(tag))
    # Do not ask the detector to resolve malformed evidence. The candidate
    # independently inspects the live OOXML and cannot inherit its omissions.
    paragraphs = [{"index": i, "text": p.text, "style_name": p.style.name}
                  for i, p in enumerate(doc.paragraphs)]
    with pytest.raises(ReferenceHeadingBoldUnavailable):
        _reference_heading_bold_targets(_bold_saved(), doc, paragraphs, _bold_finding())


@pytest.mark.parametrize("change", [
    "wrong_index", "bool_index", "extra_location", "wrong_actual", "wrong_expected",
    "disabled", "missing_snapshot", "future_snapshot", "wrong_bool", "noop",
])
def test_reference_heading_bold_saved_requirement_and_identity_fail_closed(change):
    from app.services.guided_fix import ReferenceHeadingBoldUnavailable, _reference_heading_bold_targets
    from app.services.document_parser import extract_paragraphs

    doc = _bold_doc()
    saved, finding = _bold_saved(), _bold_finding()
    if change == "wrong_index":
        finding.location["paragraph_index"] = 0
    elif change == "bool_index":
        finding.location["paragraph_index"] = True
    elif change == "extra_location":
        finding.location["run_index"] = 0
    elif change == "wrong_actual":
        finding.actual_value = "Bold"
    elif change == "wrong_expected":
        finding.expected_value = "bold"
    elif change == "disabled":
        saved.profile_snapshot["references"]["check_heading_bold"] = False
    elif change == "missing_snapshot":
        saved.profile_snapshot = None
    elif change == "future_snapshot":
        saved.profile_snapshot["schema_version"] = 999
    elif change == "wrong_bool":
        saved.profile_snapshot["references"]["heading_bold"] = 1
    elif change == "noop":
        doc.paragraphs[1].runs[0].font.bold = True
        finding.actual_value = "Bold"
    with pytest.raises(ReferenceHeadingBoldUnavailable):
        _reference_heading_bold_targets(saved, doc, extract_paragraphs(doc), finding)


def test_reference_heading_bold_partial_mutation_never_publishes_token(
    client, monkeypatch, tmp_path,
):
    from app.services import guided_fix, guided_fix_storage
    from app.services.reference_rules import _check_heading_bold
    from app.services.reference_section import reference_section_bounds
    from app.services.document_parser import extract_paragraphs
    from app.services.profile_preset_adapter import EffectiveProfileConfig

    source = _bold_bytes(_bold_doc(False, segmented=True))
    audit = _bold_audit(client, source)
    finding = next(v for v in audit["physical_layout_errors"] if v["rule_code"] == BOLD_RULE)
    root = f'/api/audit/{audit["audit_id"]}/guided-fixes'
    selected = {"selection": json.dumps({"violation_ids": [finding["id"]]})}
    reviewed = client.post(root + "/review", data=selected, files={"file": ("source.docx", source)})
    assert reviewed.status_code == 200, reviewed.text
    calls = []

    def partial(run, required):
        calls.append(run.text)
        if len(calls) == 1:
            run.font.bold = required
        return True

    monkeypatch.setattr(guided_fix, "_fix_reference_heading_bold", partial)
    published = []
    monkeypatch.setattr(guided_fix_storage, "store", lambda *args: published.append(args))
    applied = client.post(root, data={**selected, "review_id": reviewed.json()["review_id"]},
                          files={"file": ("source.docx", source)})
    assert applied.status_code == 422
    assert "download_token" not in applied.json()
    assert published == []
    # This is the detector-silence hazard the all-run verifier must defeat.
    doc = Document(io.BytesIO(source))
    doc.paragraphs[1].runs[0].font.bold = True
    paragraphs = extract_paragraphs(doc)
    profile = new_custom_profile("Silent mixed")
    profile.references.check_heading_bold = True
    profile.references.heading_bold = True
    assert _check_heading_bold(doc.paragraphs[1], reference_section_bounds(paragraphs),
                               EffectiveProfileConfig(resolve_snapshot(profile))) == []
