"""Consumer contracts for transient, read-only section/story evidence."""

from copy import deepcopy
from io import BytesIO
import socket
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt
from lxml import etree

from app.services.document_parser import extract_document_blocks, extract_sections
from app.services.page_elements import extract_page_elements


KINDS = ("header", "footer")
VARIANTS = ("default", "first", "even")


def xml(name, **attributes):
    element = OxmlElement("w:" + name)
    for key, value in attributes.items():
        element.set(qn("w:" + key), value)
    return element


def story(document, kind="footer", variant="default", index=0):
    prefix = {"default": "", "first": "first_page_", "even": "even_page_"}[variant]
    return getattr(document.sections[index], prefix + kind)


def record(result, index=0, kind="footer", variant="default"):
    return next(s for s in result["sections"][index]["stories"]
                if s["kind"] == kind and s["variant"] == variant)


def paragraph(result, index=0, kind="footer", variant="default"):
    key = record(result, index, kind, variant)["part_key"]
    return result["parts"][key]["paragraphs"][0]


def simple(p, instruction=" PAGE ", result="99"):
    field = xml("fldSimple", instr=instruction)
    run = xml("r")
    text = xml("t")
    text.text = result
    run.append(text)
    field.append(run)
    p._p.append(field)
    return field


def marker(p, typ):
    run = xml("r")
    run.append(xml("fldChar", fldCharType=typ))
    p._p.append(run)


def instruction(p, text):
    run = xml("r")
    element = xml("instrText")
    element.text = text
    run.append(element)
    p._p.append(run)


def complex_field(p, pieces=(" PAGE ",), result="99"):
    marker(p, "begin")
    for text in pieces:
        instruction(p, text)
    marker(p, "separate")
    p.add_run(result)
    marker(p, "end")


def reopened(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream)


def package(document):
    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_absent_stories_do_not_materialize_and_body_identity_survives():
    document = Document()
    document.add_paragraph("Body alpha")
    document.add_paragraph("Body beta")
    before = package(document)
    blocks = extract_document_blocks(document)
    result = extract_page_elements(document)
    assert result["parts"] == []
    section = result["sections"][0]
    assert section["first_page_enabled"] is False
    assert section["even_pages_enabled"] is False
    assert {(s["kind"], s["variant"]) for s in section["stories"]} == {
        (kind, variant) for kind in KINDS for variant in VARIANTS}
    assert all(s["status"] == "absent" and s["part_key"] is None
               and not s["linked_to_previous"] for s in section["stories"])
    assert package(document) == before
    assert extract_document_blocks(document) == blocks
    assert [(b["index"], b["text"]) for b in blocks] == [(0, "Body alpha"), (1, "Body beta")]


@pytest.mark.parametrize("break_type", ["continuous", "nextPage", "oddPage", "evenPage", "nextColumn"])
def test_sections_preserve_break_declarations_and_body_indexes(break_type):
    document = Document()
    document.add_paragraph("one")
    document.add_section()
    document.add_paragraph("two")
    sect = document.sections[0]._sectPr
    sect.append(xml("type", val=break_type))
    before = extract_document_blocks(document)
    result = extract_page_elements(reopened(document))
    assert [s["section_index"] for s in result["sections"]] == [0, 1]
    assert [s["break_type"] for s in result["sections"]] == [break_type, "nextPage"]
    assert [s["section_index"] for s in result["sections"]] == [s["section_index"] for s in extract_sections(document)]
    assert all(len(s["stories"]) == 6 for s in result["sections"])
    assert extract_document_blocks(document) == before


@pytest.mark.parametrize("value", ["unknown", "", None])
def test_unknown_and_malformed_breaks_are_not_next_page(value):
    document = Document()
    element = xml("type") if value is None else xml("type", val=value)
    document.sections[0]._sectPr.append(element)
    section = extract_page_elements(document)["sections"][0]
    assert section["break_type"] is None
    assert section["status"] == "unresolved"
    assert section["raw_break_type"] == value


def test_duplicate_break_and_missing_boundary_are_unresolved():
    document = Document()
    sect = document.sections[0]._sectPr
    sect.extend([xml("type", val="continuous"), xml("type", val="nextPage")])
    assert extract_page_elements(document)["sections"][0]["break_type"] is None
    document.element.body.remove(sect)
    section = extract_page_elements(document)["sections"][0]
    assert section["status"] == "unresolved"
    assert all(s["status"] == "unresolved" for s in section["stories"])


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("variant", VARIANTS)
def test_each_story_variant_is_independent_and_ordered(kind, variant):
    document = Document()
    target = story(document, kind, variant)
    target.paragraphs[0].text = "first"
    target.add_paragraph("second")
    result = extract_page_elements(reopened(document))
    selected = record(result, kind=kind, variant=variant)
    assert selected["reference_state"] == "explicit"
    assert selected["source_section_index"] == 0
    assert selected["status"] == "resolved"
    part = result["parts"][selected["part_key"]]
    assert part["kind"] == kind
    assert part["consumers"] == [{"section_index": 0, "kind": kind, "variant": variant}]
    assert [(p["paragraph_index"], p["literal_text"]) for p in part["paragraphs"]] == [(0, "first"), (1, "second")]
    assert all(s["status"] == "absent" for s in result["sections"][0]["stories"] if s != selected)
    assert selected["activation"] == ("enabled" if variant == "default" else "disabled")


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("variant", VARIANTS)
def test_inheritance_empty_unlinked_and_shared_parts(kind, variant):
    document = Document()
    story(document, kind, variant).paragraphs[0].text = "shared"
    document.add_section()
    document.add_section()
    third = story(document, kind, variant, 2)
    third.is_linked_to_previous = False
    document.add_section()
    document = reopened(document)
    result = extract_page_elements(document)
    records = [record(result, i, kind, variant) for i in range(4)]
    assert [r["linked_to_previous"] for r in records] == [False, True, False, True]
    assert [r["source_section_index"] for r in records] == [0, 0, 2, 2]
    assert records[0]["part_key"] == records[1]["part_key"]
    assert records[2]["part_key"] == records[3]["part_key"] != records[0]["part_key"]
    assert paragraph(result, 2, kind, variant)["literal_text"] == ""
    assert [p["consumers"] for p in result["parts"]] == [
        [{"section_index": i, "kind": kind, "variant": variant} for i in indexes]
        for indexes in ((0, 1), (2, 3))]
    story(document, kind, variant, 2).paragraphs[0].text = "independent"
    changed = extract_page_elements(document)
    assert paragraph(changed, 0, kind, variant)["literal_text"] == "shared"
    assert paragraph(changed, 3, kind, variant)["literal_text"] == "independent"


@pytest.mark.parametrize("kind", KINDS)
def test_explicit_shared_relationship_and_first_reference_key_order(kind):
    document = Document()
    story(document, kind, "even").paragraphs[0].text = "earliest reference"
    story(document, kind, "default").paragraphs[0].text = "later reference"
    document.add_section()
    original = document.sections[0]._sectPr.find(qn("w:" + kind + "Reference"))
    document.sections[1]._sectPr.append(deepcopy(original))
    result = extract_page_elements(document)
    assert record(result, kind=kind, variant="even")["part_key"] == 0
    assert record(result, kind=kind, variant="default")["part_key"] == 1
    second = record(result, 1, kind, "even")
    assert second["reference_state"] == "explicit" and second["linked_to_previous"] is False
    assert second["part_key"] == 0 and second["source_section_index"] == 1
    assert len(result["parts"]) == 2
    assert result["parts"][0]["consumers"] == [
        {"section_index": i, "kind": kind, "variant": "even"} for i in (0, 1)]


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("failure", ["duplicate", "dangling", "external", "wrong_type", "wrong_target", "unknown_variant"])
def test_bad_relationships_are_unresolved_and_inheritance_remains_unresolved(kind, failure, monkeypatch):
    document = Document()
    target = story(document, kind)
    target.paragraphs[0].text = "text"
    sect = document.sections[0]._sectPr
    ref = sect.find(qn("w:" + kind + "Reference"))
    if failure == "duplicate":
        sect.append(deepcopy(ref))
    elif failure == "dangling":
        ref.set(qn("r:id"), "missing")
    elif failure == "external":
        rid = document.part.rels.get_or_add_ext_rel(getattr(RT, kind.upper()), "https://never-fetch.invalid/story")
        ref.set(qn("r:id"), rid)
    elif failure == "wrong_type":
        rid = document.part.relate_to(target.part, RT.IMAGE)
        ref.set(qn("r:id"), rid)
    elif failure == "wrong_target":
        rid = document.part.relate_to(document.part, getattr(RT, kind.upper()))
        ref.set(qn("r:id"), rid)
    else:
        ref.set(qn("w:type"), "unknown")
    document.add_section()
    def deny(*args, **kwargs):
        pytest.fail("Extraction attempted network access")
    monkeypatch.setattr(socket, "create_connection", deny)
    result = extract_page_elements(document)
    assert record(result, kind=kind)["reference_state"] == "unresolved"
    assert record(result, kind=kind)["status"] == "unresolved"
    inherited = record(result, 1, kind)
    assert inherited["linked_to_previous"] is True
    assert inherited["status"] == "unresolved"
    assert "unresolved_predecessor" in inherited["reasons"]
    assert inherited["part_key"] is None


@pytest.mark.parametrize("kind", KINDS)
def test_hyperlink_text_is_ordered_without_target_and_external_flag(kind):
    document = Document()
    target = story(document, kind)
    p = target.paragraphs[0]
    p.add_run("before ")
    link = xml("hyperlink")
    rid = target.part.rels.get_or_add_ext_rel(RT.HYPERLINK, "https://never-fetch.invalid/secret")
    link.set(qn("r:id"), rid)
    for text in ("link", " label"):
        run, element = xml("r"), xml("t")
        element.text = text
        run.append(element)
        link.append(run)
    p._p.append(link)
    p.add_run(" after")
    result = extract_page_elements(reopened(document))
    observed = paragraph(result, kind=kind)
    assert observed["literal_text"] == "before link label after"
    assert [r["run_index"] for r in observed["formatting"]["runs"]] == [0, 1, 2, 3]
    assert "never-fetch" not in str(result)
    assert result["parts"][0]["unsupported_content"]["external_relationships"] is True
    assert result["parts"][0]["completeness"] == "unsupported"


@pytest.mark.parametrize("kind", KINDS)
@pytest.mark.parametrize("shape,flag", [
    ("tbl", "tables"), ("drawing", "drawings"), ("pict", "shapes"),
    ("txbxContent", "text_boxes"), ("ins", "tracked_revisions"),
    ("sdt", "unsupported_containers"), ("alternate", "alternate_content")])
def test_unsupported_content_is_not_flattened_or_proven_absent(kind, shape, flag):
    document = Document()
    target = story(document, kind)
    p = target.paragraphs[0]
    p.add_run("supported")
    if shape == "alternate":
        wrapper = etree.Element("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
    else:
        wrapper = xml(shape)
    inner = xml("p")
    fake_para = type(p)(inner, p._parent)
    simple(fake_para)
    wrapper.append(inner)
    (target._element if shape == "tbl" else p._p).append(wrapper)
    result = extract_page_elements(reopened(document))
    part = result["parts"][0]
    assert [(x["paragraph_index"], x["literal_text"]) for x in part["paragraphs"]] == [(0, "supported")]
    assert paragraph(result, kind=kind)["fields"] == []
    assert part["unsupported_content"][flag] is True
    assert part["unsupported_content"]["unsupported_fields"] is True
    assert part["completeness"] == "unsupported"
    if shape != "tbl":
        assert paragraph(result, kind=kind)["formatting"]["completeness"] == "unsupported"


@pytest.mark.parametrize("kind", KINDS)
def test_table_paragraph_formatting_does_not_enter_direct_evidence(kind):
    document = Document()
    target = story(document, kind)
    target.paragraphs[0].text = "direct"
    table = target.add_table(1, 1, Inches(2))
    cell = table.cell(0, 0).paragraphs[0]
    cell.add_run("table font").font.name = "TableOnly"
    result = extract_page_elements(reopened(document))
    part = result["parts"][0]
    assert [p["literal_text"] for p in part["paragraphs"]] == ["direct"]
    assert all(r["font_family"] != "TableOnly" for p in part["paragraphs"] for r in p["formatting"]["runs"])
    assert part["unsupported_content"]["tables"] is True


@pytest.mark.parametrize("form", ["simple", "complex", "split"])
@pytest.mark.parametrize("token,expected", [("PAGE", "PAGE"), ("page", "PAGE"), ("NUMPAGES", "NUMPAGES"),
    ("SECTION", "SECTION"), ("SECTIONPAGES", "SECTIONPAGES"), ("PAGEREF", "other"), ("NOTPAGE", "other")])
def test_field_identity_uses_exact_instruction_token_not_cached_value(form, token, expected):
    document = Document()
    p = story(document).paragraphs[0]
    p.add_run("Page ")
    text = " " + token + " \\* MERGEFORMAT "
    if form == "simple":
        simple(p, text)
    else:
        pieces = (text,) if form == "complex" else (text[:3], text[3:])
        complex_field(p, pieces)
    p.add_run(" trailing")
    result = extract_page_elements(reopened(document))
    observed = paragraph(result)
    assert observed["literal_text"] == "Page  trailing"
    assert observed["cached_text"] == "99"
    assert observed["visible_text"] == "Page 99 trailing"
    assert observed["fields"] == [dict(kind=expected, form="simple" if form == "simple" else "complex",
        instruction=text, cached_result="99", paragraph_index=0, status="resolved", reasons=[])]


def test_literal_digits_and_page_of_numpages_are_distinct():
    document = Document()
    p = story(document).paragraphs[0]
    p.add_run("1")
    literal = paragraph(extract_page_elements(document))
    assert literal["fields"] == [] and literal["literal_text"] == "1" and literal["cached_text"] == ""
    p.text = "Page "
    simple(p, " PAGE ", "1")
    p.add_run(" of ")
    complex_field(p, (" NUMPAGES ",), "8")
    observed = paragraph(extract_page_elements(document))
    assert [f["kind"] for f in observed["fields"]] == ["PAGE", "NUMPAGES"]
    assert observed["literal_text"] == "Page  of " and observed["cached_text"] == "18"
    assert observed["visible_text"] == "Page 1 of 8"


@pytest.mark.parametrize("failure", ["nested_complex", "nested_simple", "missing_end", "missing_begin", "duplicate_separator",
    "missing_separator", "orphan_instruction", "instruction_after_separator", "invalid_marker", "empty_instruction", "wrapper_crossing"])
def test_malformed_fields_are_unresolved_not_page_proof(failure):
    document = Document()
    p = story(document).paragraphs[0]
    if failure == "nested_simple":
        outer = simple(p)
        inner = simple(p)
        outer.append(inner)
    elif failure == "orphan_instruction":
        instruction(p, " PAGE ")
    elif failure == "missing_begin":
        marker(p, "end")
    else:
        marker(p, "begin")
        instruction(p, "" if failure == "empty_instruction" else " PAGE ")
        if failure == "nested_complex":
            complex_field(p)
        if failure != "missing_separator":
            marker(p, "separate")
        if failure == "duplicate_separator":
            marker(p, "separate")
        if failure == "instruction_after_separator":
            instruction(p, " NUMPAGES ")
        if failure == "invalid_marker":
            marker(p, "bogus")
        if failure == "wrapper_crossing":
            wrapper = xml("sdt")
            wrapper.append(xml("r"))
            p._p.append(wrapper)
        p.add_run("99")
        if failure != "missing_end":
            marker(p, "end")
    result = extract_page_elements(reopened(document))
    fields = paragraph(result)["fields"]
    assert fields and all(f["kind"] == "unresolved" and f["status"] == "unresolved" and f["reasons"] for f in fields)
    assert result["parts"][0]["unsupported_content"]["unsupported_fields"] is True
    assert paragraph(result)["formatting"]["completeness"] != "complete"


def test_complex_field_state_spans_paragraph_boundaries():
    document = Document()
    target = story(document)
    p = target.paragraphs[0]
    p.add_run("before ")
    marker(p, "begin")
    instruction(p, " PA")
    second = target.add_paragraph()
    instruction(second, "GE ")
    marker(second, "separate")
    second.add_run("99")
    third = target.add_paragraph()
    third.add_run("100")
    marker(third, "end")
    third.add_run(" after")
    result = extract_page_elements(reopened(document))
    paragraphs = result["parts"][0]["paragraphs"]
    field = paragraphs[0]["fields"][0]
    assert field["kind"] == "PAGE" and field["instruction"] == " PAGE "
    assert field["cached_result"] == "99100" and field["paragraph_index"] == 0
    assert [p["cached_text"] for p in paragraphs] == ["", "99", "100"]
    assert [p["literal_text"] for p in paragraphs] == ["before ", "", " after"]
    assert paragraphs[1]["fields"] == paragraphs[2]["fields"] == []


@pytest.mark.parametrize("fmt", [None, "decimal", "lowerRoman", "upperRoman", "lowerLetter", "upperLetter"])
@pytest.mark.parametrize("start", [None, "0", "1", "12"])
def test_numbering_raw_declarations_do_not_compute_sequence(fmt, start):
    document = Document()
    attrs = {"chapStyle": "2", "chapSep": "hyphen"}
    if fmt is not None:
        attrs["fmt"] = fmt
    if start is not None:
        attrs["start"] = start
    document.sections[0]._sectPr.append(xml("pgNumType", **attrs))
    numbering = extract_page_elements(reopened(document))["sections"][0]["numbering"]
    assert numbering == dict(format=fmt, start=start, chapter_style="2", chapter_separator="hyphen",
                             status="resolved", reasons=[])


def test_absent_numbering_is_not_restart_at_one_and_distances_are_declarations():
    document = Document()
    document.sections[0].header_distance = Inches(0.5)
    document.sections[0].footer_distance = Inches(0.75)
    section = extract_page_elements(document)["sections"][0]
    assert section["numbering"] == dict(format=None, start=None, chapter_style=None,
                                       chapter_separator=None, status="resolved", reasons=[])
    assert section["distances"] == {"header": 0.5, "footer": 0.75}
    margin = document.sections[0]._sectPr.find(qn("w:pgMar"))
    margin.set(qn("w:header"), "bad")
    assert extract_page_elements(document)["sections"][0]["distances"]["header"] is None


@pytest.mark.parametrize("attrs", [{"start": "bad"}, {"start": "-1"}, {"fmt": ""}, {"fmt": "notAFormat"},
                                    {"chapStyle": "10"}, {"chapSep": "bad"}])
def test_malformed_numbering_retains_raw_values_but_is_unresolved(attrs):
    document = Document()
    document.sections[0]._sectPr.append(xml("pgNumType", **attrs))
    numbering = extract_page_elements(document)["sections"][0]["numbering"]
    assert numbering["status"] == "unresolved" and numbering["reasons"]
    mapping = {"start": "start", "fmt": "format", "chapStyle": "chapter_style", "chapSep": "chapter_separator"}
    assert all(numbering[mapping[k]] == v for k, v in attrs.items())


def test_duplicate_numbering_has_no_fabricated_selected_declaration():
    document = Document()
    document.sections[0]._sectPr.extend([xml("pgNumType", start="1"), xml("pgNumType", start="9")])
    numbering = extract_page_elements(document)["sections"][0]["numbering"]
    assert numbering["status"] == "unresolved"
    assert all(numbering[k] is None for k in ("format", "start", "chapter_style", "chapter_separator"))


@pytest.mark.parametrize("name", ["titlePg", "evenAndOddHeaders"])
@pytest.mark.parametrize("value,expected", [("absent", False), (None, True), ("1", True), ("true", True), ("on", True),
                                          ("0", False), ("false", False), ("off", False), ("invalid", None), ("TRUE", None), ("duplicate", None)])
def test_on_off_activation_preserves_false_and_uncertainty(name, value, expected):
    document = Document()
    parent = document.sections[0]._sectPr if name == "titlePg" else document.settings.element
    for old in list(parent.findall(qn("w:" + name))):
        parent.remove(old)
    if value != "absent":
        element = xml(name) if value is None else xml(name, val="true" if value == "duplicate" else value)
        parent.append(element)
        if value == "duplicate":
            parent.append(xml(name, val="false"))
    result = extract_page_elements(document)
    section = result["sections"][0]
    assert section["first_page_enabled" if name == "titlePg" else "even_pages_enabled"] is expected
    variant = "first" if name == "titlePg" else "even"
    assert record(result, variant=variant)["activation"] == ("unresolved" if expected is None else "enabled" if expected else "disabled")
    assert record(result)["activation"] == "enabled"


@pytest.mark.parametrize("kind", KINDS)
def test_direct_and_inherited_formatting_keep_provenance_and_mixed_runs(kind):
    document = Document()
    style = document.styles.add_style("StoryEvidence", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    p = story(document, kind).paragraphs[0]
    p.style = style
    p._p.get_or_add_pPr().append(xml("jc", val="right"))
    p.add_run("inherited ")
    direct = p.add_run("direct")
    direct.font.name = "Arial"
    direct.font.size = Pt(14)
    direct.bold = True
    direct.italic = False
    italic = p.add_run(" italic")
    italic.italic = True
    runs = paragraph(extract_page_elements(reopened(document)), kind=kind)["formatting"]
    assert runs["alignment"] == "right" and runs["completeness"] == "complete"
    first, second, third = runs["runs"]
    assert (first["font_family"], first["font_size_pt"], first["bold"], first["italic"]) == (None, None, None, None)
    assert (first["effective_family"], first["effective_family_source"], first["effective_size_pt"], first["effective_size_source"]) == ("Georgia", "paragraph_style", 11, "paragraph_style")
    assert (second["font_family"], second["font_size_pt"], second["bold"], second["italic"]) == ("Arial", 14, True, False)
    assert (second["effective_family"], second["effective_family_source"], second["effective_size_pt"], second["effective_size_source"]) == ("Arial", "direct", 14, "direct")
    assert third["italic"] is True and third["effective_family"] == "Georgia"


def test_unresolved_run_does_not_prove_uniform_formatting():
    document = Document()
    # Remove all source fonts/sizes so neither default nor inheritance can resolve.
    for element in list(document.styles.element.iter()):
        if element.tag in (qn("w:rFonts"), qn("w:sz")):
            element.getparent().remove(element)
    p = story(document).paragraphs[0]
    run = p.add_run("resolved")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    p.add_run(" unresolved")
    observed = paragraph(extract_page_elements(document))["formatting"]
    assert observed["runs"][0]["effective_family"] == "Arial"
    assert observed["runs"][1]["effective_family"] is None
    assert observed["runs"][1]["effective_size_pt"] is None
    assert observed["completeness"] == "unresolved"


def test_hidden_text_is_not_complete_visible_content():
    document = Document()
    p = story(document).paragraphs[0]
    p.add_run("visible")
    p.add_run("hidden").font.hidden = True
    result = extract_page_elements(document)
    assert paragraph(result)["literal_text"] == "visible"
    assert paragraph(result)["formatting"]["completeness"] == "unsupported"
    assert result["parts"][0]["unsupported_content"]["hidden_text"] is True


def test_package_fields_relationships_and_body_are_immutable_and_repeatable(monkeypatch):
    document = Document()
    document.add_paragraph("body zero")
    header = story(document, "header")
    header.paragraphs[0].text = "header"
    footer = story(document)
    simple(footer.paragraphs[0])
    complex_field(footer.add_paragraph(), (" PA", "GE "))
    document.add_section()
    document.add_paragraph("body three")
    document.add_section()
    first_footer = story(document, "footer", "first", 2)
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = "first only"
    document = reopened(document)
    before = package(document)
    body_before = extract_document_blocks(document)
    def deny(*args, **kwargs):
        pytest.fail("Extractor must not perform network IO")
    monkeypatch.setattr(socket.socket, "connect", deny)
    monkeypatch.setattr(socket, "create_connection", deny)
    first = extract_page_elements(document)
    second = extract_page_elements(document)
    assert first == second
    assert len(first["parts"]) == 3
    assert record(first, 0, "header")["part_key"] == record(first, 2, "header")["part_key"]
    assert record(first, 0, "footer")["part_key"] == record(first, 2, "footer")["part_key"]
    assert record(first, 0, "footer", "first")["part_key"] is None
    assert extract_document_blocks(document) == body_before
    assert package(document) == before
    assert all(len(section["stories"]) == 6 for section in first["sections"])


def test_missing_settings_and_styles_are_not_materialized():
    document = Document()
    story(document).paragraphs[0].text = "no styles"
    for rid, rel in list(document.part.rels.items()):
        if rel.reltype in (RT.SETTINGS, RT.STYLES):
            del document.part.rels[rid]
    before = package(document)
    result = extract_page_elements(document)
    assert result["sections"][0]["even_pages_enabled"] is False
    assert paragraph(result)["literal_text"] == "no styles"
    assert paragraph(result)["formatting"]["completeness"] == "unresolved"
    assert package(document) == before


def test_inherited_hidden_content_is_not_visible_evidence():
    document = Document()
    style = document.styles.add_style("HiddenStory", WD_STYLE_TYPE.PARAGRAPH)
    style.font.hidden = True
    p = story(document).paragraphs[0]
    p.style = style
    p.add_run("inherited hidden")
    p.add_run("explicit visible").font.hidden = False
    result = extract_page_elements(document)
    assert paragraph(result)["literal_text"] == "explicit visible"
    assert result["parts"][0]["unsupported_content"]["hidden_text"] is True


def test_theme_override_does_not_fall_through_to_inherited_concrete_family():
    document = Document()
    style = document.styles.add_style("ConcreteFamily", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Georgia"
    p = story(document).paragraphs[0]
    p.style = style
    run = p.add_run("theme")
    run._r.get_or_add_rPr().append(xml("rFonts", asciiTheme="minorHAnsi"))
    formatting = paragraph(extract_page_elements(document))["formatting"]
    assert formatting["runs"][0]["effective_family"] is None
    assert formatting["completeness"] == "unresolved"


@pytest.mark.parametrize("kind", KINDS)
def test_real_image_relationship_is_flagged_without_mutation(kind):
    import base64
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII=")
    document = Document()
    p = story(document, kind).paragraphs[0]
    p.add_run().add_picture(BytesIO(png), width=Inches(0.1))
    document = reopened(document)
    before = package(document)
    result = extract_page_elements(document)
    assert paragraph(result, kind=kind)["literal_text"] == ""
    assert result["parts"][0]["unsupported_content"]["drawings"] is True
    assert result["parts"][0]["completeness"] == "unsupported"
    assert package(document) == before


def test_duplicate_paragraph_property_boundaries_remain_unresolved():
    document = Document()
    p = document.add_paragraph("boundary")
    for _ in range(2):
        properties = xml("pPr")
        properties.append(xml("sectPr"))
        p._p.append(properties)
    result = extract_page_elements(document)
    assert [s["status"] for s in result["sections"][:2]] == ["unresolved", "unresolved"]
    assert all(story["status"] == "unresolved" for section in result["sections"][:2]
               for story in section["stories"])


@pytest.mark.parametrize("shape", ["invalid_bold", "duplicate_size", "missing_style"])
def test_malformed_formatting_does_not_prove_complete_resolution(shape):
    document = Document()
    p = story(document).paragraphs[0]
    run = p.add_run("formatting")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    if shape == "invalid_bold":
        run._r.get_or_add_rPr().append(xml("b", val="invalid"))
    elif shape == "duplicate_size":
        run._r.get_or_add_rPr().append(xml("sz", val="28"))
    else:
        run._r.get_or_add_rPr().append(xml("rStyle", val="MissingStyle"))
    assert paragraph(extract_page_elements(document))["formatting"]["completeness"] == "unresolved"


def test_style_false_toggle_does_not_prove_inherited_hidden_text_visible():
    document = Document()
    base = document.styles.add_style("HiddenBase", WD_STYLE_TYPE.PARAGRAPH)
    base.font.hidden = True
    base.font.name = "Arial"
    base.font.size = Pt(12)
    derived = document.styles.add_style("HiddenDerived", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    derived.font.hidden = False
    p = story(document).paragraphs[0]
    p.style = derived
    p.add_run("still hidden")
    result = extract_page_elements(document)
    assert paragraph(result)["literal_text"] == ""
    assert result["parts"][0]["unsupported_content"]["hidden_text"] is True
    assert result["parts"][0]["completeness"] == "unsupported"


@pytest.mark.parametrize("inherited", [False, True])
def test_complex_script_formatting_is_not_resolved_as_latin(inherited):
    document = Document()
    p = story(document).paragraphs[0]
    run = p.add_run("complex script formatting")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    properties = run._r.get_or_add_rPr()
    if inherited:
        style = document.styles.add_style("ComplexStory", WD_STYLE_TYPE.PARAGRAPH)
        p.style = style
        properties = style.element.get_or_add_rPr()
    properties.append(xml("cs"))
    properties.append(xml("szCs", val="36"))
    result = paragraph(extract_page_elements(document))["formatting"]
    assert result["runs"][0]["font_family"] == "Arial"
    assert result["runs"][0]["font_size_pt"] == 12
    assert result["runs"][0]["effective_family"] is None
    assert result["runs"][0]["effective_size_pt"] is None
    assert result["completeness"] == "unresolved"


@pytest.mark.parametrize("inherited", [False, True])
def test_numbering_markers_are_unsupported_not_omitted_complete_text(inherited):
    document = Document()
    p = story(document).paragraphs[0]
    run = p.add_run("numbered item")
    run.font.name = "Arial"
    run.font.size = Pt(12)
    properties = p._p.get_or_add_pPr()
    if inherited:
        style = document.styles.add_style("NumberedStory", WD_STYLE_TYPE.PARAGRAPH)
        p.style = style
        properties = style.element.get_or_add_pPr()
    numbering = xml("numPr")
    numbering.append(xml("numId", val="1"))
    properties.append(numbering)
    result = extract_page_elements(document)
    assert paragraph(result)["literal_text"] == "numbered item"
    assert paragraph(result)["formatting"]["completeness"] == "unsupported"
    assert result["parts"][0]["unsupported_content"]["automatic_numbering"] is True
    assert result["parts"][0]["completeness"] == "unsupported"
