"""Guided Fixes backend — identity, request security, supported fixes,
structural safety, preservation, output contract, migration.

Backend-only: no frontend, no network, AI mocked by conftest fixtures.
"""
import hashlib
import io
import json
import logging
import sqlite3
import urllib.parse
import uuid
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from sqlalchemy.orm import sessionmaker

from app.models.audit import AuditRecord, Violation
from app.services import guided_fix
from app.services.guided_fix import (
    LEGACY_DETAIL,
    MISMATCH_DETAIL,
    GuidedFixFailed,
    apply_guided_fixes,
    build_corrected_filename,
)

# conftest is importable: tests/ is a package.
from tests import conftest as testkit

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
HLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _clean_font_fixture_context(document):
    """Synthetic direct-font fixture, not a cleanup operation on uploaded documents.

    The stock Word template contains inherited theme fonts and szCs. Positive
    font tests deliberately construct a smaller supported context instead.
    """
    containers = [document.styles.element.find(qn("w:docDefaults"))]
    containers += [document.styles[name].element for name in ("Normal", "Default Paragraph Font")]
    for container in containers:
        if container is None:
            continue
        for props in container.iter(qn("w:rPr")):
            for child in list(props):
                if child.tag in {qn("w:" + tag) for tag in ("rFonts", "sz", "szCs", "cs", "rtl")}:
                    props.remove(child)
                elif child.tag == qn("w:lang"):
                    child.attrib.pop(qn("w:bidi"), None)
                    child.attrib.pop(qn("w:eastAsia"), None)
    return document


def _clean_line_spacing_fixture_context(document):
    """Construct synthetic direct-auto evidence, never normalize an upload."""
    for properties in document.styles.element.iter(qn("w:pPr")):
        for spacing in properties.findall(qn("w:spacing")):
            spacing.attrib.pop(qn("w:line"), None)
            spacing.attrib.pop(qn("w:lineRule"), None)
    for grid in list(document.element.iter(qn("w:docGrid"))):
        grid.getparent().remove(grid)
    return document


def _font_fixture_bytes(data):
    document = _clean_font_fixture_context(Document(io.BytesIO(data)))
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                run.font.name = "Times New Roman"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()

def _post_audit(client, docx_factory, paragraphs=("Body paragraph.",), **kw):
    return client.post(
        "/api/audit",
        files={"file": ("audit.docx", docx_factory(paragraphs=list(paragraphs), **kw), "application/octet-stream")},
    )


def _post_audit_bytes(client, file_bytes, filename="audit.docx", query=""):
    return client.post(
        "/api/audit" + query,
        files={"file": (filename, file_bytes, "application/octet-stream")},
    )


def _rule_ids(resp, rule_code):
    return [
        v["id"] for v in resp.json()["physical_layout_errors"]
        if v["rule_code"] == rule_code
    ]


def _fix_post(client, audit_id, file_bytes, violation_ids, filename="audit.docx"):
    """Exercise the public mandatory Review → Apply → Download flow."""
    path = f"/api/audit/{audit_id}/guided-fixes"
    fields = {"selection": json.dumps({"violation_ids": violation_ids})}
    review = client.post(path + "/review", data=fields, files={"file": (filename, file_bytes, DOCX_MIME)})
    if review.status_code != 200:
        return review
    applied = client.post(path, data={**fields, "review_id": review.json()["review_id"]},
                          files={"file": (filename, file_bytes, DOCX_MIME)})
    if applied.status_code != 200:
        return applied
    assert applied.json()["status"] == "passed"
    return client.get(path + "/download/" + applied.json()["download_token"])


def _make_docx(text="Body paragraph.", size=11.0, family="Calibri"):
    """Minimal single-paragraph docx with explicit font size 11pt."""
    doc = Document()
    _clean_font_fixture_context(doc)
    p = doc.add_paragraph(text)
    run = p.runs[0]
    run.font.name = family
    run.font.size = Pt(size)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _rezip(docx_bytes):
    """Preserve document contents but deterministically change ZIP metadata."""
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as src, zipfile.ZipFile(buf, "w") as dst:
        dst.comment = src.comment
        for index, item in enumerate(src.infolist()):
            if index == 0:
                marker = b"ACA identity resave"
                item.comment = b"" if item.comment == marker else marker
            dst.writestr(item, src.read(item.filename))
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 1. Identity (items 1-5)
# ---------------------------------------------------------------------------

class TestIdentity:
    def test_audit_persists_document_sha256(self, client, docx_factory, test_engine):
        file_bytes = docx_factory(paragraphs=["Body paragraph."])
        resp = client.post(
            "/api/audit",
            files={"file": ("audit.docx", file_bytes, "application/octet-stream")},
        )
        assert resp.status_code == 200
        audit_id = resp.json()["audit_id"]
        Session = sessionmaker(bind=test_engine)
        s = Session()
        row = s.query(AuditRecord).filter(AuditRecord.id == audit_id).first()
        assert row.document_sha256 == hashlib.sha256(file_bytes).hexdigest()
        s.close()

    def test_exact_bytes_fix_returns_docx(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        # The SAME bytes are audited and re-uploaded: "exact bytes" is the
        # claim under test. Building a second document from identical
        # arguments is NOT byte-identical — make_docx_bytes stamps the
        # current time into the ZIP entries — so that variant would assert
        # the fixture's clock instead of Guided Fixes.
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(DOCX_MIME)
        assert "ACA-fixed.docx" in r.headers["content-disposition"]
        d = Document(io.BytesIO(r.content))
        assert d.paragraphs[0].runs[0].font.size.pt == 12.0

    def test_different_docx_same_filename_409(self, client, docx_factory):
        resp = _post_audit(client, docx_factory, ["Original text."])
        audit_id = resp.json()["audit_id"]
        other = docx_factory(paragraphs=["Different text."])
        r = _fix_post(client, audit_id, other, ["nonexistent"])
        assert r.status_code == 409
        assert r.json()["detail"] == MISMATCH_DETAIL

    def test_resave_changes_hash(self, client, docx_factory):
        # A content-preserving resave still fails byte identity, regardless
        # of when it occurs or whether its ZIP timestamps change.
        original = docx_factory(paragraphs=["Body paragraph."], body_size=11)
        resp = _post_audit_bytes(client, original)
        assert resp.status_code == 200
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert ids

        resaved = _rezip(original)
        assert resaved != original
        assert hashlib.sha256(resaved).hexdigest() != hashlib.sha256(original).hexdigest()
        # Content is semantically identical — the endpoint still refuses.
        r = _fix_post(client, audit_id, resaved, [ids[0]])
        assert r.status_code == 409
        assert r.json()["detail"] == MISMATCH_DETAIL
        assert _post_audit_bytes(client, resaved).status_code == 200

    @pytest.mark.parametrize("member_comment", [b"", b"ACA identity resave"])
    def test_rezip_preserves_members_and_is_deterministic(self, docx_factory, member_comment):
        source = docx_factory(paragraphs=["Body paragraph."])
        stream = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as src, zipfile.ZipFile(stream, "w") as dst:
            dst.comment = b"Preserved archive comment"
            for index, item in enumerate(src.infolist()):
                if index == 0:
                    item.comment = member_comment
                dst.writestr(item, src.read(item.filename))
        original = stream.getvalue()

        resaved = _rezip(original)
        assert resaved != original
        assert _rezip(original) == resaved
        assert zipfile.is_zipfile(io.BytesIO(resaved))
        with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(resaved)) as after:
            assert after.testzip() is None
            assert after.namelist() == before.namelist()
            assert after.comment == before.comment
            assert after.infolist()[0].comment != before.infolist()[0].comment
            for name in before.namelist():
                assert after.read(name) == before.read(name)
        assert Document(io.BytesIO(resaved)).paragraphs[0].text == "Body paragraph."

    def test_null_hash_409_legacy(self, client, test_engine, docx_factory):
        audit_id = str(uuid.uuid4())
        Session = sessionmaker(bind=test_engine)
        s = Session()
        s.add(AuditRecord(id=audit_id, filename="old.docx", file_size=100,
                          weighted_score=90, deploy_mode="LOCAL", status="completed"))
        s.add(Violation(id="legacy-v1", audit_id=audit_id, rule_code="FONT_SIZE",
                        severity="MINOR", location={"paragraph_index": 0, "run_index": 0},
                        message="m", expected_value="12pt", actual_value="11pt"))
        s.commit()
        s.close()
        r = _fix_post(client, audit_id, docx_factory(paragraphs=["x"]), ["legacy-v1"])
        assert r.status_code == 409
        assert r.json()["detail"] == LEGACY_DETAIL


# ---------------------------------------------------------------------------
# 2. Request security (items 6-13)
# ---------------------------------------------------------------------------

class TestRequestSecurity:
    def test_extra_field_422(self, client, docx_factory):
        file_bytes = docx_factory(paragraphs=["Body paragraph."])
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        payload = {"violation_ids": [str(uuid.uuid4())], "run_indexes": [0]}
        r = client.post(
            f"/api/audit/{audit_id}/guided-fixes/review",
            data={"selection": json.dumps(payload)},
            files={"file": ("audit.docx", file_bytes, DOCX_MIME)},
        )
        assert r.status_code == 422
        assert r.json()["detail"] == "Invalid selection payload."

    def test_foreign_violation_id_422(self, client, docx_factory):
        file_bytes_1 = docx_factory(paragraphs=["First audit."])
        # Second audit must carry a real finding so its ids are extractable.
        file_bytes_2 = docx_factory(paragraphs=["Second audit."], body_size=11)
        resp1 = _post_audit_bytes(client, file_bytes_1)
        resp2 = _post_audit_bytes(client, file_bytes_2)
        ids2 = _rule_ids(resp2, "FONT_SIZE")
        assert ids2
        r = _fix_post(client, resp1.json()["audit_id"], file_bytes_1, [ids2[0]])
        assert r.status_code == 422
        assert r.json()["detail"] == "A selected finding does not belong to this audit."

    def test_duplicate_selection_rejected(self, client, docx_factory):
        file_bytes = docx_factory(paragraphs=["Body paragraph."], body_size=11)
        resp = _post_audit_bytes(client, file_bytes)
        ids = _rule_ids(resp, "FONT_SIZE")
        r = _fix_post(client, resp.json()["audit_id"], file_bytes, [ids[0], ids[0]])
        assert r.status_code == 422

    def test_empty_selection_422(self, client, docx_factory):
        file_bytes = docx_factory(paragraphs=["Body paragraph."])
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        r = client.post(
            f"/api/audit/{audit_id}/guided-fixes/review",
            data={"selection": json.dumps({"violation_ids": []})},
            files={"file": ("audit.docx", file_bytes, DOCX_MIME)},
        )
        assert r.status_code == 422

    def test_21_ids_422(self, client, docx_factory):
        file_bytes = docx_factory(paragraphs=["Body paragraph."])
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = [str(uuid.uuid4()) for _ in range(21)]
        r = client.post(
            f"/api/audit/{audit_id}/guided-fixes/review",
            data={"selection": json.dumps({"violation_ids": ids})},
            files={"file": ("audit.docx", file_bytes, DOCX_MIME)},
        )
        assert r.status_code == 422

    def test_garbage_docx_400(self, client, docx_factory):
        file_bytes = docx_factory(paragraphs=["Body paragraph."])
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        r = client.post(
            f"/api/audit/{audit_id}/guided-fixes/review",
            data={"selection": json.dumps({"violation_ids": ["x"]})},
            files={"file": ("garbage.docx", b"not a zip at all", DOCX_MIME)},
        )
        assert r.status_code == 400
        assert r.json()["detail"] == "The uploaded file is not a valid DOCX."

    def test_oversize_400(self, client_with_small_cap):
        # Validation order: size gate precedes audit lookup — a random UUID
        # still yields the oversize 400 (no audit creation needed under the
        # tiny cap).
        audit_id = str(uuid.uuid4())
        r = client_with_small_cap.post(
            f"/api/audit/{audit_id}/guided-fixes/review",
            data={"selection": json.dumps({"violation_ids": ["x"]})},
            files={"file": ("big.docx", b"x" * 2048, DOCX_MIME)},
        )
        assert r.status_code == 400
        assert "10MB" in r.json()["detail"]

    def test_crlf_dotdot_filename_sanitized(self, client, docx_factory):
        hostile = "..\\report\r\n.docx"
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]], filename=hostile)
        assert r.status_code == 200, r.text
        cd = r.headers["content-disposition"]
        assert "\r" not in cd and "\n" not in cd
        assert ".." not in cd
        assert "\\" not in cd
        # Sanitizer unit contract: control chars and traversal dots removed.
        name = build_corrected_filename(hostile)
        assert "\r" not in name and "\n" not in name
        assert ".." not in name and "\\" not in name
        assert name.endswith(".docx")


# ---------------------------------------------------------------------------
# 3. Supported fixes (items 14-23)
# ---------------------------------------------------------------------------

class TestSupportedFixes:
    def test_font_size_fix(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d = Document(io.BytesIO(r.content))
        assert abs(d.paragraphs[0].runs[0].font.size.pt - 12.0) < 0.01

    def test_font_consistency_fix(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], font_name="Calibri"))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_CONSISTENCY")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d = Document(io.BytesIO(r.content))
        run = d.paragraphs[0].runs[0]
        assert run.font.name == "Times New Roman"
        rfonts = run._element.rPr.rFonts
        assert rfonts.get(qn("w:ascii")) == "Times New Roman"
        assert rfonts.get(qn("w:hAnsi")) == "Times New Roman"
        assert rfonts.get(qn("w:eastAsia")) == "Times New Roman"

    def test_alignment_fix(self, client, docx_factory):
        file_bytes = docx_factory(paragraphs=["Body paragraph."], alignment=WD_ALIGN_PARAGRAPH.LEFT)
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "ALIGNMENT")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d = Document(io.BytesIO(r.content))
        assert d.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    def test_line_spacing_fix(self, client, docx_factory):
        document = Document(io.BytesIO(docx_factory(paragraphs=["Body paragraph."], line_spacing=1.0)))
        _clean_line_spacing_fixture_context(document)
        stream = io.BytesIO()
        document.save(stream)
        file_bytes = stream.getvalue()
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "LINE_SPACING")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d = Document(io.BytesIO(r.content))
        spacing = d.paragraphs[0]._p.pPr.find(qn("w:spacing"))
        assert spacing.get(qn("w:line")) == "360"
        assert spacing.get(qn("w:lineRule")) == "auto"

    def test_space_before_fix(self, client, docx_factory):
        base = docx_factory(paragraphs=["Body paragraph."])
        d = Document(io.BytesIO(base))
        d.paragraphs[0].paragraph_format.space_before = Pt(20)
        buf = io.BytesIO()
        d.save(buf)
        file_bytes = buf.getvalue()
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "SPACE_BEFORE")
        assert ids, resp.json()["physical_layout_errors"]
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d2 = Document(io.BytesIO(r.content))
        assert abs(d2.paragraphs[0].paragraph_format.space_before.pt - 0.0) < 0.01

    def test_space_after_fix(self, client, docx_factory):
        base = docx_factory(paragraphs=["Body paragraph."])
        d = Document(io.BytesIO(base))
        d.paragraphs[0].paragraph_format.space_after = Pt(20)
        buf = io.BytesIO()
        d.save(buf)
        file_bytes = buf.getvalue()
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "SPACE_AFTER")
        assert ids, resp.json()["physical_layout_errors"]
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d2 = Document(io.BytesIO(r.content))
        assert abs(d2.paragraphs[0].paragraph_format.space_after.pt - 6.0) < 0.01

    def test_sibling_untouched(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(
            paragraphs=["Body paragraph.", "Sibling intact."], body_size=11
        ))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert len(ids) >= 2
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d = Document(io.BytesIO(r.content))
        assert d.paragraphs[0].runs[0].font.size.pt == 12.0
        assert d.paragraphs[1].runs[0].font.size.pt == 11.0
        assert d.paragraphs[1].text == "Sibling intact."

    def test_highlight_survives(self, client, docx_factory):
        doc = Document()
        _clean_font_fixture_context(doc)
        p = doc.add_paragraph("Body paragraph.")
        run = p.runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        run._element.rPr.append(hl)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_CONSISTENCY")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d2 = Document(io.BytesIO(r.content))
        rpr = d2.paragraphs[0].runs[0]._element.rPr
        hl2 = rpr.find(qn("w:highlight"))
        assert hl2 is not None and hl2.get(qn("w:val")) == "yellow"

    def test_multi_run_grouped_fix(self, client, docx_factory):
        doc = Document()
        _clean_font_fixture_context(doc)
        p = doc.add_paragraph()
        p.add_run("One ").font.size = Pt(11)
        p.add_run("two").font.size = Pt(11)
        for run in p.runs:
            run.font.name = "Times New Roman"
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert ids
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200, r.text
        d2 = Document(io.BytesIO(r.content))
        assert d2.paragraphs[0].runs[0].font.size.pt == 12.0
        assert d2.paragraphs[0].runs[1].font.size.pt == 12.0

    def test_ambiguous_family_never_guessed(self, client, docx_factory):
        # Allowed combinations span several families: no family may be guessed.
        doc = Document()
        _clean_font_fixture_context(doc)
        p = doc.add_paragraph("Body paragraph.")
        run = p.runs[0]
        run.font.name = "Comic Sans MS"
        run.font.size = Pt(13)
        p.paragraph_format.line_spacing = 2.0
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        resp = _post_audit_bytes(client, file_bytes, query="?profile_id=apa7-student-paper")
        assert resp.status_code == 200, resp.text
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_CONSISTENCY")
        assert ids, resp.json()["physical_layout_errors"]
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 422
        assert r.json()["detail"] == "Fix cannot be applied safely. Manual correction required."


# ---------------------------------------------------------------------------
# 4. Safety (items 24-32)
# ---------------------------------------------------------------------------

class TestSafety:
    def test_document_changed_skip(self):
        audit = AuditRecord(id="gf-unit", filename="gf.docx", file_size=1, status="completed")
        audit.profile_snapshot = None
        v = Violation(
            id="v1", audit_id="gf-unit", rule_code="FONT_SIZE", severity="MINOR",
            location={"paragraph_index": 0, "run_index": 0}, message="m",
            expected_value="12pt", actual_value="9pt",
        )
        with pytest.raises(GuidedFixFailed):
            apply_guided_fixes(audit, _make_docx(size=11.0), [v])

    def test_hyperlink_unsafe_structure(self):
        doc = Document()
        p = doc.add_paragraph("text ")
        r_id = doc.part.relate_to("https://example.com", HLINK_REL, is_external=True)
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "22")  # 11pt
        rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = "link"
        r.append(t)
        hl.append(r)
        p._p.append(hl)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()

        audit = AuditRecord(id="gf-unit", filename="gf.docx", file_size=1, status="completed")
        audit.profile_snapshot = None
        v = Violation(
            id="v2", audit_id="gf-unit", rule_code="ALIGNMENT", severity="MINOR",
            location={"paragraph_index": 0}, message="m",
            expected_value="justify", actual_value="left",
        )
        with pytest.raises(GuidedFixFailed):
            apply_guided_fixes(audit, data, [v])

    def test_unsupported_rule_skip(self):
        audit = AuditRecord(id="gf-unit", filename="gf.docx", file_size=1, status="completed")
        audit.profile_snapshot = None
        v = Violation(
            id="v3", audit_id="gf-unit", rule_code="MARGIN_LEFT", severity="MAJOR",
            location={"paragraph_index": 0}, message="m",
            expected_value="1in", actual_value="0.5in",
        )
        with pytest.raises(GuidedFixFailed):
            apply_guided_fixes(audit, _make_docx(), [v])

    def test_fixer_failure_422_verify(self, client, docx_factory, monkeypatch):
        file_bytes = docx_factory(paragraphs=["Body paragraph."], alignment=WD_ALIGN_PARAGRAPH.LEFT)
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "ALIGNMENT")
        assert ids

        def boom(para, name):
            raise RuntimeError("boom")

        monkeypatch.setattr(guided_fix, "_fix_alignment", boom)
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 422
        assert r.json()["detail"] == "Fix not applied. No file was produced."
        assert not r.content.startswith(b"PK")

    def test_nothing_persisted(self, client_file_db, docx_factory, tmp_path):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client_file_db, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        assert ids
        before_files = set(p.name for p in tmp_path.iterdir())
        before_get = client_file_db.get(f"/api/audit/{audit_id}").json()
        r = _fix_post(client_file_db, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200
        after_get = client_file_db.get(f"/api/audit/{audit_id}").json()
        assert after_get["weighted_score"] == before_get["weighted_score"]
        assert after_get["violations"] == before_get["violations"]
        assert after_get["status"] == "completed"
        after_files = set(p.name for p in tmp_path.iterdir())
        assert after_files == before_files, f"new files: {after_files - before_files}"

    def test_logs_do_not_leak_content(self, client, docx_factory, caplog):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        marker = "GFLOGLEAKMARKER12345"
        with caplog.at_level(logging.DEBUG):
            r = _fix_post(client, audit_id, file_bytes, [ids[0]])
            assert r.status_code == 200
        assert marker not in caplog.text
        assert "Body paragraph" not in caplog.text


# ---------------------------------------------------------------------------
# 5. Preservation (items 33-42)
# ---------------------------------------------------------------------------

def _rich_docx_bytes():
    """One fixable body paragraph + hyperlink, bookmark, tracked insert,
    numbered list, SEQ caption, table, image, header, Unicode text."""
    doc = Document()
    _clean_font_fixture_context(doc)
    body = doc.add_paragraph("Plain body paragraph to fix.")
    for run in body.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
    pf = body.paragraph_format
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)

    uni = doc.add_paragraph("论文测试 中文 body")
    for run in uni.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)

    link_para = doc.add_paragraph()
    link_para.add_run("Visit ")
    r_id = doc.part.relate_to("https://example.com", HLINK_REL, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = "example"
    new_run.append(t)
    hl.append(new_run)
    link_para._p.append(hl)

    bm_para = doc.add_paragraph("Bookmarked text")
    bm = OxmlElement("w:bookmarkStart")
    bm.set(qn("w:id"), "7")
    bm.set(qn("w:name"), "bm1")
    bm_para._p.insert(0, bm)

    ins_para = doc.add_paragraph()
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "1")
    ins.set(qn("w:author"), "ACA")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "inserted text"
    r.append(t)
    ins.append(r)
    ins_para._p.append(ins)

    list_para = doc.add_paragraph("List item")
    ppr = list_para._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "1")
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)

    testkit._add_word_caption(doc, "Table", ": Test table")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    table.rows[1].cells[0].text = "C"
    table.rows[1].cells[1].text = "D"

    pic_para = doc.add_paragraph()
    pic_para.add_run().add_picture(io.BytesIO(testkit._PNG_1PX), width=doc.sections[0].page_width * 0.1)

    doc.sections[0].header.paragraphs[0].text = "Header text"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestPreservation:
    def test_rich_document_preserved(self, client, docx_factory):
        data = _rich_docx_bytes()
        resp = _post_audit_bytes(client, data)
        assert resp.status_code == 200, resp.text
        audit_id = resp.json()["audit_id"]
        ids = [
            v["id"] for v in resp.json()["physical_layout_errors"]
            if v["rule_code"] == "FONT_SIZE" and v["location"].get("paragraph_index") == 0
        ]
        assert ids, resp.json()["physical_layout_errors"]
        r = _fix_post(client, audit_id, data, [ids[0]])
        assert r.status_code == 200, r.text

        src = Document(io.BytesIO(data))
        out = Document(io.BytesIO(r.content))

        # Paragraph + table text identical.
        assert [p.text for p in out.paragraphs] == [p.text for p in src.paragraphs]
        assert [[[c.text for c in row.cells] for row in t.rows] for t in out.tables] == [
            [[c.text for c in row.cells] for row in t.rows] for t in src.tables
        ]
        # Hyperlink rel count.
        def hyperlinks(d):
            return sum(1 for rel in d.part.rels.values() if "hyperlink" in rel.reltype)
        assert hyperlinks(out) == hyperlinks(src)
        # Bookmark start count.
        assert len(out.element.body.findall(".//" + qn("w:bookmarkStart"))) == \
            len(src.element.body.findall(".//" + qn("w:bookmarkStart")))
        # Drawing count.
        assert len(out.element.body.findall(".//" + qn("w:drawing"))) == \
            len(src.element.body.findall(".//" + qn("w:drawing")))
        # Table count.
        assert len(out.tables) == len(src.tables)
        # Numbering XML unchanged.
        assert out.part.numbering_part.blob == src.part.numbering_part.blob
        # SEQ field instrText content unchanged.
        def instr_texts(d):
            return [t.text for t in d.element.body.findall(".//" + qn("w:instrText"))]
        assert instr_texts(out) == instr_texts(src)
        # Tracked insert count unchanged.
        assert len(out.element.body.findall(".//" + qn("w:ins"))) == \
            len(src.element.body.findall(".//" + qn("w:ins")))
        # Header text unchanged.
        assert out.sections[0].header.paragraphs[0].text == \
            src.sections[0].header.paragraphs[0].text
        # Unicode paragraph intact.
        assert any("论文测试 中文" in p.text for p in out.paragraphs)
        # Target paragraph actually fixed (11pt → 12pt).
        assert out.paragraphs[0].runs[0].font.size.pt == 12.0


# ---------------------------------------------------------------------------
# 6. Output contract (items 43-48)
# ---------------------------------------------------------------------------

class TestOutputContract:
    def test_content_type_and_filename(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(DOCX_MIME)
        cd = r.headers["content-disposition"]
        assert "audit_ACA-fixed.docx" in cd
        assert "\r" not in cd and "\n" not in cd


    def test_corrected_bytes_reopenable(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200
        d = Document(io.BytesIO(r.content))
        assert d.paragraphs

    def test_export_pdf_and_history_intact(self, client, docx_factory):
        file_bytes = _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11))
        resp = _post_audit_bytes(client, file_bytes)
        audit_id = resp.json()["audit_id"]
        ids = _rule_ids(resp, "FONT_SIZE")
        r = _fix_post(client, audit_id, file_bytes, [ids[0]])
        assert r.status_code == 200
        pdf = client.get(f"/api/audit/{audit_id}/export-pdf")
        assert pdf.status_code == 200
        history = client.get("/api/audits")
        assert history.status_code == 200
        assert any(a["id"] == audit_id for a in history.json())


# ---------------------------------------------------------------------------
# 7. Migration (items 49-52)
# ---------------------------------------------------------------------------

class TestMigration:
    def _run_alembic(self, tmp_path, target):
        from pathlib import Path
        import os as _os
        from alembic import command
        from alembic.config import Config
        from app.config import settings as _settings

        db = tmp_path / "gf-migration.db"
        url = f"sqlite:///{db}"
        backend_dir = Path(__file__).resolve().parents[1]
        cfg = Config(str(backend_dir / "alembic.ini"))
        cfg.set_main_option("script_location", str(backend_dir / "alembic"))
        orig = _settings.DATABASE_URL
        orig_env = _os.environ.get("DATABASE_URL")
        try:
            _os.environ["DATABASE_URL"] = url
            _settings.DATABASE_URL = url
            cfg.set_main_option("sqlalchemy.url", url)
            command.upgrade(cfg, target)
        finally:
            _settings.DATABASE_URL = orig
            if orig_env is None:
                _os.environ.pop("DATABASE_URL", None)
            else:
                _os.environ["DATABASE_URL"] = orig_env
        return db

    def test_upgrade_90fc_to_head_preserves_rows(self, tmp_path):
        db = self._run_alembic(tmp_path, "90fc17718e11")
        con = sqlite3.connect(str(db))
        con.execute(
            "INSERT INTO audit_records (id, filename, file_size, deploy_mode, status) "
            "VALUES (?,?,?,?,?)",
            ("mig-audit-1", "mig.docx", 100, "LOCAL", "completed"),
        )
        con.commit()
        con.close()
        self._run_alembic(tmp_path, "head")
        con = sqlite3.connect(str(db))
        cols = {r[1] for r in con.execute("PRAGMA table_info(audit_records)")}
        assert "document_sha256" in cols
        row = con.execute(
            "SELECT filename, document_sha256 FROM audit_records WHERE id='mig-audit-1'"
        ).fetchone()
        assert row[0] == "mig.docx"
        assert row[1] is None
        con.close()

    def test_single_head(self):
        from pathlib import Path
        from alembic.config import Config
        from alembic.script import ScriptDirectory

        backend_dir = Path(__file__).resolve().parents[1]
        cfg = Config(str(backend_dir / "alembic.ini"))
        cfg.set_main_option("script_location", str(backend_dir / "alembic"))
        sd = ScriptDirectory.from_config(cfg)
        assert sd.get_heads() == ["b3e8d19e2b3f"]


# ---------------------------------------------------------------------------
# Citation ↔ Reference List matching findings are manual-only
# ---------------------------------------------------------------------------

def test_new_citation_codes_are_never_selectable_guided_fixes(client, docx_factory):
    """Citation and semantic reference findings cannot be auto-applied."""
    file_bytes = docx_factory(
        paragraphs=["Orphan (Garcia, 2018) text.", "Another (Lee, 2019) line."],
        references=["Other, X. (2020). Unrelated entry. Press."],
    )
    resp = _post_audit_bytes(client, file_bytes)
    audit_id = resp.json()["audit_id"]
    ids = _rule_ids(resp, "CITATION_MISMATCH")
    assert ids, "fixture must produce citation findings"

    r = _fix_post(client, audit_id, file_bytes, ids)
    assert r.status_code == 422
    assert r.json()["detail"] == "Fix cannot be applied safely. Manual correction required."
    assert r.content != file_bytes or r.headers.get("content-type", "").startswith("application/json")


def _policy_fixture(rule, docx_factory):
    """Build a genuine candidate and its profile before adding package evidence."""
    from app.services.profile_schema import new_custom_profile
    from tests.test_guided_fix_reference_heading import _source, _bold_doc, _bold_bytes
    if rule == "FONT_CONSISTENCY":
        return _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], font_name="Calibri")), None
    if rule == "FONT_SIZE":
        return _font_fixture_bytes(docx_factory(paragraphs=["Body paragraph."], body_size=11)), None
    if rule == "REFERENCE_HEADING_ALIGNMENT":
        profile = new_custom_profile("Package policy heading alignment")
        profile.references.check_heading_alignment = True
        profile.references.heading_alignment = "center"
        return _source(), profile
    if rule == "REFERENCE_HEADING_BOLD":
        profile = new_custom_profile("Package policy heading bold")
        profile.references.check_heading_bold = True
        profile.references.heading_bold = True
        return _bold_bytes(_bold_doc()), profile
    kwargs = {"alignment": WD_ALIGN_PARAGRAPH.LEFT} if rule == "ALIGNMENT" else {}
    if rule == "LINE_SPACING":
        kwargs["line_spacing"] = 1.0
    document = Document(io.BytesIO(docx_factory(paragraphs=["Body paragraph."], **kwargs)))
    if rule == "LINE_SPACING":
        _clean_line_spacing_fixture_context(document)
    if rule in ("SPACE_BEFORE", "SPACE_AFTER"):
        setattr(document.paragraphs[0].paragraph_format, rule.lower(), Pt(20))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue(), None


def _policy_marker(data, marker):
    from lxml import etree
    from tests.test_guided_fix_preservation_storage import _add_parts, _change_part
    if marker == "signed":
        return _add_parts(data, {"_xmlsignatures/alternate.xml": b"<marker/>"})
    if marker == "macro_part":
        return _add_parts(data, {"word/vbaProject.bin": b"inert non-executable test data"})
    if marker in ("documentProtection", "writeProtection"):
        return _change_part(data, "word/settings.xml",
                            lambda root: etree.SubElement(root, qn("w:" + marker)))
    data = _add_parts(data, {"word/alternate.bin": b"inert non-executable test data"})
    if marker == "macro_type":
        def add_type(root):
            etree.SubElement(root, root.tag.removesuffix("Types") + "Override",
                             PartName="/word/alternate.bin",
                             ContentType="application/vnd.ms-office.vbaProject")
        return _change_part(data, "[Content_Types].xml", add_type)
    if marker == "macro_relationship":
        def add_binary_type(root):
            etree.SubElement(root, root.tag.removesuffix("Types") + "Override",
                             PartName="/word/alternate.bin", ContentType="application/octet-stream")
        data = _change_part(data, "[Content_Types].xml", add_binary_type)
        def add_relation(root):
            etree.SubElement(root, root.tag.removesuffix("Relationships") + "Relationship",
                             Id="policy-marker", Target="alternate.bin",
                             Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject")
        return _change_part(data, "word/_rels/document.xml.rels", add_relation)
    raise AssertionError(marker)


@pytest.mark.parametrize("rule", [
    "FONT_CONSISTENCY", "FONT_SIZE", "ALIGNMENT", "LINE_SPACING",
    "SPACE_BEFORE", "SPACE_AFTER", "REFERENCE_HEADING_ALIGNMENT", "REFERENCE_HEADING_BOLD",
])
@pytest.mark.parametrize("marker", [
    "signed", "documentProtection", "writeProtection", "macro_part", "macro_type", "macro_relationship",
])
def test_package_policy_refuses_each_genuine_candidate(client, docx_factory, monkeypatch, tmp_path, rule, marker):
    from app.api import routes
    from app.config import settings
    from app.services.docx_pdf_converter import DocxConversionError

    def no_conversion(_data):
        raise DocxConversionError("libreoffice_unavailable")

    monkeypatch.setattr(routes, "convert_docx_to_pdf", no_conversion)
    monkeypatch.setattr(settings, "GUIDED_FIX_STORAGE_DIR", str(tmp_path))
    ordinary, profile = _policy_fixture(rule, docx_factory)
    data = {"custom_profile": json.dumps(profile.to_dict())} if profile else {}

    def audit_and_review(source):
        response = client.post("/api/audit", data=data, files={"file": ("renamed.docx", source, DOCX_MIME)})
        assert response.status_code == 200, response.text
        findings = _rule_ids(response, rule)
        assert findings, (rule, response.json()["physical_layout_errors"])
        audit_id = response.json()["audit_id"]
        path = f"/api/audit/{audit_id}/guided-fixes"
        fields = {"selection": json.dumps({"violation_ids": [findings[0]]})}
        return audit_id, path, fields

    control_id, control_path, control_fields = audit_and_review(ordinary)
    control = client.post(control_path + "/review", data=control_fields,
                          files={"file": ("renamed.docx", ordinary, DOCX_MIME)})
    assert control.status_code == 200, (rule, control.text)
    assert client.delete(control_path + "/review/" + control.json()["review_id"]).status_code == 204

    source = _policy_marker(ordinary, marker)
    digest = hashlib.sha256(source).digest()
    audit_id, path, fields = audit_and_review(source)
    before = client.get(f"/api/audit/{audit_id}").json()
    history_before = client.get("/api/audits").json()
    response = client.post(path + "/review", data=fields,
                           files={"file": ("renamed.docx", source, DOCX_MIME)})
    assert response.status_code == 422
    assert response.json() == {"detail": (
        "This correction is unavailable for signed, protected or macro-enabled documents. "
        "Correct the selected findings manually in Word; your original document is unchanged."
    )}
    assert not any(review[0] == audit_id for review in routes._reviews.values())
    assert not list(tmp_path.rglob("*.docx"))
    assert hashlib.sha256(source).digest() == digest
    assert client.get(f"/api/audit/{audit_id}").json() == before
    assert client.get("/api/audits").json() == history_before


def test_package_policy_fails_closed_on_unverifiable_xml(docx_factory):
    from tests.test_guided_fix_preservation_storage import _replace_part
    source = _replace_part(docx_factory(paragraphs=["Body paragraph."]), "word/settings.xml",
                           lambda _: b"<broken")
    with pytest.raises(GuidedFixFailed, match="unverifiable DOCX package"):
        guided_fix._validate_guided_fix_package(source)


def test_package_refusal_preserves_report_content(client, docx_factory, monkeypatch, tmp_path):
    from app.api import routes
    from app.config import settings
    from app.services.docx_pdf_converter import DocxConversionError
    from pypdf import PdfReader

    def no_conversion(_data):
        raise DocxConversionError("libreoffice_unavailable")

    monkeypatch.setattr(routes, "convert_docx_to_pdf", no_conversion)
    monkeypatch.setattr(settings, "GUIDED_FIX_STORAGE_DIR", str(tmp_path))
    ordinary, _ = _policy_fixture("FONT_SIZE", docx_factory)
    source = _policy_marker(ordinary, "signed")
    audit = _post_audit_bytes(client, source)
    assert audit.status_code == 200
    audit_id = audit.json()["audit_id"]
    ids = _rule_ids(audit, "FONT_SIZE")
    assert ids
    report_path = f"/api/audit/{audit_id}/export-pdf"

    def report_text():
        response = client.get(report_path)
        assert response.status_code == 200
        return [page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages]

    before = report_text()
    response = client.post(f"/api/audit/{audit_id}/guided-fixes/review",
                           data={"selection": json.dumps({"violation_ids": [ids[0]]})},
                           files={"file": ("renamed.docx", source, DOCX_MIME)})
    assert response.status_code == 422
    assert report_text() == before
