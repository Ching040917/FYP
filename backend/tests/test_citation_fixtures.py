"""Synthetic citation/reference fixtures driven through the deterministic
sensor and the full static rules engine.

The fixture table (tests/citation_fixtures.py) is the support-matrix
contract: every row states whether the sensor reports, stays silent, or
flags manual review — and regression rows assert that the profile toggles
can switch the new behaviour off completely.
"""
import io

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.services.citation_sensor import APPENDIX_HEADER_PATTERN, run_citation_sensor
from app.services.citation_sensor import scan_citations as scan_citations_for
from app.services.document_parser import extract_paragraphs, parse_document
from app.services.layout_engine import run_static_rules_engine
from app.services.profile_preset_adapter import EffectiveProfileConfig
from app.services.profile_registry import get_builtin_profile, SUC_PROFILE_ID
from app.services.profile_snapshot import resolve_snapshot
from app.services.role_classifier import classify_paragraphs
from app.services.scoring import calculate_weighted_score_detailed
from tests.citation_fixtures import (
    CLEAN_FIXTURES,
    CLEAN_UNDER_ALL_RULES,
    UNSUPPORTED_OCCURRENCE_FIXTURES,
    CitationDoc,
    FIXTURES,
    scan,
)

CITATION_CODES = (
    "CITATION_MISMATCH",
    "CITATION_REFERENCE_MISMATCH",
    "CITATION_YEAR_SUFFIX_MISMATCH",
    "CITATION_REFERENCE_AMBIGUOUS",
    "REFERENCE_NOT_CITED",
    "REFERENCE_DUPLICATE",
)


def _all_off_config(base=None):
    """Real snapshot-derived config with every citation toggle disabled."""
    profile = get_builtin_profile(SUC_PROFILE_ID)
    if base is not None:
        profile = base
    for name in type(profile.citation_matching).__dataclass_fields__:
        setattr(profile.citation_matching, name, False)
    return EffectiveProfileConfig(resolve_snapshot(profile))


def _all_on_config():
    profile = get_builtin_profile(SUC_PROFILE_ID)
    for name in type(profile.citation_matching).__dataclass_fields__:
        setattr(profile.citation_matching, name, True)
    return EffectiveProfileConfig(resolve_snapshot(profile))


# ---------------------------------------------------------------------------
# Support matrix — silent cases
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", CLEAN_FIXTURES)
def test_silent_fixtures(name, docx_factory):
    assert scan(FIXTURES[name], docx_factory) == []


@pytest.mark.parametrize("name", CLEAN_UNDER_ALL_RULES)
def test_silent_fixtures_stay_silent_with_every_rule_enabled(name, docx_factory):
    """Even with all six rules armed, the guarded cases emit nothing."""
    assert scan(FIXTURES[name], docx_factory, config=_all_on_config()) == []


@pytest.mark.parametrize("name", [
    "numeric_citation",
    "malformed_citation",
    "personal_communication",
    "as_cited_in",
    "year_suffixes",
])
def test_unsupported_or_unmatched_occurrences_produce_no_citation_finding(name, docx_factory):
    """No supported CITATION_MISMATCH arises for these documents: their
    citations are out of scope, unsupported occurrences (personal
    communication / "as cited in"), or they match an entry exactly. With
    every rule armed, the genuinely uncited WORK is still reported —
    cited state is work-level, never author-level."""
    assert scan(FIXTURES[name], docx_factory) == []
    viols = scan(FIXTURES[name], docx_factory, config=_all_on_config())
    assert [v.rule_code for v in viols] == ["REFERENCE_NOT_CITED"]
    # The reported work is the one nothing points at, not merely "the author".
    assert "2024a" in viols[0].actual_value or viols[0].actual_value in FIXTURES[
        name
    ].references[0]


# ---------------------------------------------------------------------------
# Support matrix — reporting cases
# ---------------------------------------------------------------------------

def test_missing_reference_reports_one_consolidated_finding(docx_factory):
    viols = scan(FIXTURES["missing_reference"], docx_factory)
    assert len(viols) == 1
    v = viols[0]
    assert v.rule_code == "CITATION_MISMATCH"
    assert v.severity == "MAJOR"
    assert v.location["paragraph_index"] == 0
    assert v.location["occurrences"] == 1
    assert v.expected_value == "Reference entry for Smith"


def test_uncited_reference_is_opt_in(docx_factory):
    fixture = FIXTURES["uncited_reference"]
    assert scan(fixture, docx_factory) == []
    viols = scan(fixture, docx_factory, config=_all_on_config())
    assert [v.rule_code for v in viols] == ["REFERENCE_NOT_CITED"]
    v = viols[0]
    assert v.severity == "MINOR"
    assert "Jones" in v.actual_value
    # Entry paragraph index (header + first entry are paragraphs 1 and 2).
    assert v.location["paragraph_index"] == 3


def test_grouped_parenthetical_reports_each_missing_item(docx_factory):
    viols = scan(FIXTURES["grouped_missing"], docx_factory)
    assert len(viols) == 2
    assert {v.actual_value for v in viols} == {"Smith, 2024", "Lee, 2023"}
    assert all(v.rule_code == "CITATION_MISMATCH" for v in viols)


def test_repeated_missing_citation_consolidates(docx_factory):
    viols = scan(FIXTURES["repeated_missing"], docx_factory)
    assert len(viols) == 1
    v = viols[0]
    assert v.location["occurrences"] == 3
    assert v.location["paragraph_index"] == 0
    assert v.location["paragraph_indexes"] == [0, 1, 2]


def test_consolidated_finding_points_at_the_first_occurrence(docx_factory):
    """Page mapping is keyed by paragraph_index, so a consolidated finding
    must carry the primary occurrence's index."""
    v = scan(FIXTURES["repeated_missing"], docx_factory)[0]
    assert v.location["paragraph_index"] == min(v.location["paragraph_indexes"])


def test_duplicate_reference_reports_the_later_entry(docx_factory):
    viols = scan(FIXTURES["duplicate_reference"], docx_factory)
    assert [v.rule_code for v in viols] == ["REFERENCE_DUPLICATE"]
    v = viols[0]
    assert v.severity == "MINOR"
    assert v.location["entry_count"] == 2
    assert v.location["paragraph_index"] == v.location["paragraph_indexes"][-1]


def test_missing_section_reports_one_root_finding(docx_factory):
    viols = scan(FIXTURES["no_references_heading"], docx_factory)
    assert len(viols) == 1
    v = viols[0]
    assert v.rule_code == "CITATION_MISMATCH"
    assert v.severity == "MAJOR"
    assert v.location == {
        "paragraph_index": 0,
        "citation_count": 2,
        "missing_references_section": True,
    }
    assert "2 in-text citation(s)" in v.message


def test_ambiguous_author_is_review_only(docx_factory):
    viols = scan(FIXTURES["same_surname_different_initials"], docx_factory)
    assert [v.rule_code for v in viols] == ["CITATION_REFERENCE_AMBIGUOUS"]
    assert viols[0].severity == "REVIEW"
    result = calculate_weighted_score_detailed(viols)
    assert result.total == 100
    assert result.major_count == 0 and result.minor_count == 0


def test_entries_after_an_appendix_are_not_indexed(docx_factory):
    viols = scan(FIXTURES["references_in_appendix"], docx_factory)
    assert len(viols) == 1
    assert "Smith" in viols[0].message
    assert viols[0].rule_code == "CITATION_MISMATCH"


# ---------------------------------------------------------------------------
# Regression guarantees
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("name", sorted(FIXTURES))
def test_every_toggle_off_removes_all_citation_and_reference_findings(name, docx_factory):
    """Turning every citation control off restores the pre-build output for
    citation content: no CITATION_*/REFERENCE_* finding may appear."""
    viols = scan(FIXTURES[name], docx_factory, config=_all_off_config())
    assert [v for v in viols if v.rule_code in CITATION_CODES] == []


def test_profile_toggles_do_not_change_non_citation_findings(docx_factory):
    """A citation-free document audits identically with every citation
    control on and off — the new rules never touch other findings or score."""
    data = docx_factory(paragraphs=["Body text only."])
    with_on = run_static_rules_engine(data, config=_all_on_config())
    with_off = run_static_rules_engine(data, config=_all_off_config())
    assert with_on == with_off
    assert (
        calculate_weighted_score_detailed(with_on).total
        == calculate_weighted_score_detailed(with_off).total
    )


def test_score_deduction_comes_only_from_the_emitted_findings(docx_factory):
    """A MAJOR citation mismatch deducts exactly its category weight; a
    REVIEW nothing at all."""
    missing = scan(FIXTURES["missing_reference"], docx_factory)
    assert calculate_weighted_score_detailed(missing).total == 95  # 5 = citation major weight
    ambiguous = scan(FIXTURES["same_surname_different_initials"], docx_factory)
    assert calculate_weighted_score_detailed(ambiguous).total == 100


def test_sensor_runs_without_ai_and_with_roles(docx_factory):
    """The sensor is pure: no network, no AI, and role filtering only
    removes citation scanning for bibliography/furniture paragraphs."""
    fixture = FIXTURES["missing_reference"]
    assert scan(fixture, docx_factory) != []
    assert scan(fixture, docx_factory, roles=["REFERENCE_ENTRY"]) == []


def test_real_classifier_roles_keep_detecting_body_citations(docx_factory):
    fixture = FIXTURES["missing_reference"]
    data = fixture.build(docx_factory)
    doc = parse_document(data)
    paragraphs = extract_paragraphs(doc)
    roles = classify_paragraphs(doc, paragraphs)
    viols = run_citation_sensor(doc, paragraphs, roles=roles)
    assert [v.rule_code for v in viols] == ["CITATION_MISMATCH"]


def test_appendix_heading_boundary_is_unchanged(docx_factory):
    """layout_engine/role_classifier import this literal — it must stay."""
    assert APPENDIX_HEADER_PATTERN.match("Appendix A: Material")


# ---------------------------------------------------------------------------
# Observed OOXML behaviour (fixture 25/26): tracked changes and text boxes
# ---------------------------------------------------------------------------

def _reopen(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def _save(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _scan_bytes(data: bytes):
    doc = parse_document(data)
    return run_citation_sensor(doc, extract_paragraphs(doc))


def test_citation_inside_a_tracked_insertion_is_not_detected(docx_factory):
    """Observed python-docx behaviour: text wrapped in w:ins is not part of
    paragraph.text, so a citation inside a tracked insertion is invisible to
    the sensor. Pinned so a future parser change is a conscious decision."""
    data = docx_factory(paragraphs=["Body text.", "References", "Smith, J. (2020). T. Press."])
    doc = _reopen(data)
    para = doc.paragraphs[0]
    ins = parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="A" w:date="2024-01-01T00:00:00Z">'
        f'<w:r><w:t>Smith (2020) wrote about it.</w:t></w:r></w:ins>'
    )
    para._p.append(ins)
    assert _scan_bytes(_save(doc)) == []
    # Control: the same sentence in a normal run IS detected.
    data2 = docx_factory(
        paragraphs=["Smith (2020) wrote about it.", "References", "Jones, A. (2019). T. Press."]
    )
    assert [v.rule_code for v in _scan_bytes(data2)] == ["CITATION_MISMATCH"]


def test_citation_inside_a_text_box_is_not_detected(docx_factory):
    """Observed behaviour: text boxes are not body paragraphs, so their text
    is never extracted — a citation there produces no finding."""
    data = docx_factory(paragraphs=["Body text.", "References", "Jones, A. (2019). T. Press."])
    doc = _reopen(data)
    para = doc.paragraphs[0]
    pict = parse_xml(
        f'<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<v:shape><v:textbox><w:txbxContent>'
        f'<w:p><w:r><w:t>Garcia (2018) is mentioned here.</w:t></w:r></w:p>'
        f'</w:txbxContent></v:textbox></v:shape></w:pict>'
    )
    para._p.append(pict)
    assert _scan_bytes(_save(doc)) == []


# ---------------------------------------------------------------------------
# Navigable missing-References location
# ---------------------------------------------------------------------------

def test_missing_section_root_finding_points_at_the_first_citation(docx_factory):
    """Two supported citations, no References section: ONE finding whose
    paragraph_index is the earliest supported citation paragraph."""
    fixture = CitationDoc(
        "two_citations_no_section",
        ["Body line.", "Smith (2020) reported it.", "Lee (2021) agreed."],
        include_header=False,
    )
    viols = scan(fixture, docx_factory)
    assert len(viols) == 1
    v = viols[0]
    assert v.rule_code == "CITATION_MISMATCH"
    assert v.severity == "MAJOR"
    assert v.location == {
        "paragraph_index": 1,
        "citation_count": 2,
        "missing_references_section": True,
    }
def test_missing_section_root_finding_counts_every_occurrence_truthfully(docx_factory):
    """Repeated citations of one source stay a single finding whose count is
    the number of detected citation occurrences."""
    fixture = CitationDoc(
        "repeated_no_section",
        ["Smith (2020) first.", "Smith (2020) again.", "Smith (2020) once more."],
        include_header=False,
    )
    viols = scan(fixture, docx_factory)
    assert len(viols) == 1
    assert viols[0].location == {
        "paragraph_index": 0,
        "citation_count": 3,
        "missing_references_section": True,
    }


def test_no_supported_citation_means_no_missing_section_finding(docx_factory):
    """Document-level null location is never fabricated: with no supported
    citation the sensor emits no missing-References finding at all."""
    fixture = CitationDoc(
        "no_citations_no_section",
        ["Body text only.", "Per (J. Smith, personal communication, 2024) it holds."],
        include_header=False,
    )
    assert scan(fixture, docx_factory) == []


def test_null_primary_location_is_built_safely_when_no_citation_has_one(docx_factory):
    """The defensive null case is reachable and well formed: a citation whose
    paragraph identity is unknown yields a null paragraph_index plus the
    explicit flag — never a fabricated page anchor."""
    from app.services.citation_sensor import _first_citation_paragraph, _missing_section_violation
    from app.services.citation_matching import CitationRef, make_key

    citation = CitationRef(
        author_display="Smith",
        key=make_key("Smith", "2020"),
        paragraph_index=None,
        snippet="(Smith, 2020)",
    )
    assert _first_citation_paragraph([citation]) is None
    violation = _missing_section_violation(1, _first_citation_paragraph([citation]))
    assert violation.location == {
        "paragraph_index": None,
        "citation_count": 1,
        "missing_references_section": True,
    }

    # An unsupported occurrence is never a citation key, so it cannot anchor
    # the root finding either.
    scan = scan_citations_for("Per (J. Smith, personal communication, 2024) it holds.")
    assert scan.citations == []
    assert scan.unsupported == 1



def test_personal_communication_alone_produces_no_missing_reference(docx_factory):
    """The personal communication itself never needs a Reference List entry."""
    fixture = CitationDoc(
        "pc_only",
        [
            "Per (J. Smith, personal communication, October 5, 2024) the value holds.",
            "Jones (2018) is cited here.",
        ],
        ["Other, X. (2019). Unrelated. Press.", "Jones, A. (2018). Cited above. Press."],
    )
    assert scan(fixture, docx_factory) == []


def test_personal_communication_does_not_exempt_the_same_author_entry(docx_factory):
    """An unrelated Reference List work by the personal-communication author
    is STILL reported as uncited — the exclusion is per occurrence."""
    fixture = CitationDoc(
        "pc_and_entry",
        ["Per (J. Smith, personal communication, October 5, 2024) the value holds."],
        ["Smith, J. (2019). Unrelated paper. Press."],
    )
    assert scan(fixture, docx_factory) == []
    viols = scan(fixture, docx_factory, config=_all_on_config())
    assert [v.rule_code for v in viols] == ["REFERENCE_NOT_CITED"]
    assert "Smith" in viols[0].actual_value


def test_normal_citation_by_the_same_author_is_matched_independently(docx_factory):
    """A personal communication never suppresses a separate normal citation by
    the same author: the normal citation is compared against the list."""
    fixture = CitationDoc(
        "pc_plus_normal",
        [
            "Per (J. Smith, personal communication, October 5, 2024) it holds.",
            "An earlier study (Smith, 2019) agrees.",
        ],
        ["Smith, J. (2019). The earlier paper. Press."],
    )
    assert scan(fixture, docx_factory) == []


def test_as_cited_in_occurrence_generates_no_deterministic_match(docx_factory):
    """A secondary citation is unsupported material: no key, no finding."""
    fixture = CitationDoc(
        "as_cited_in_only",
        ["Reported by (Smith, 2020, as cited in Jones, 2019)."],
        ["Other, X. (2017). Unrelated. Press."],
    )
    viols = scan(fixture, docx_factory)
    assert [v.rule_code for v in viols] == []


def test_as_cited_in_does_not_exempt_either_named_author(docx_factory):
    """Neither author mentioned in "as cited in" is exempt: their unrelated
    entries remain eligible for REFERENCE_NOT_CITED."""
    fixture = CitationDoc(
        "as_cited_in_entries",
        ["Reported by (Smith, 2020, as cited in Jones, 2019)."],
        [
            "Smith, J. (2020). Unrelated Smith work. Press.",
            "Jones, A. (2019). Unrelated Jones work. Press.",
        ],
    )
    assert scan(fixture, docx_factory) == []
    viols = scan(fixture, docx_factory, config=_all_on_config())
    assert [v.rule_code for v in viols] == ["REFERENCE_NOT_CITED", "REFERENCE_NOT_CITED"]


def test_matching_pipeline_has_no_author_exemption_surface():
    """Structural guarantee: neither module exposes an author-wide exemption."""
    import inspect

    from app.services import citation_matching, citation_sensor

    for module in (citation_matching, citation_sensor):
        source = inspect.getsource(module).lower()
        for banned in ("exempt_author_keys", "exempt_authors", "exempt_names"):
            assert banned not in source, f"{module.__name__} still references {banned}"
    assert "exempt" not in inspect.signature(citation_matching.match).parameters
    # The scanner reports unsupported occurrences as a count, not a name set.
    scan_fields = set(citation_sensor.CitationScan.__dataclass_fields__)
