"""Institutional-author citation matching — conservative full-identity support.

Proves a dual-sided contract: narrative citations resolve a bounded
multi-token author run against same-year Reference candidates, and
Reference entries classify institutional author phrases, while every
personal-author path stays byte-identical. No final-token alias exists:
personal surname keys contain no spaces, so institutional phrase keys can
never collide with them.
"""
from app.services.citation_matching import (
    build_reference_index,
    match,
    normalise_institutional_author,
    parse_institutional_author,
    parse_reference_entry,
)
from app.services.citation_sensor import (
    author_resolution_context,
    extend_institutional_run,
    resolve_institutional_citations,
    scan_citations,
)


def _scan(text, idx=2):
    return scan_citations(text, idx).citations


def _resolve(cites, texts, entries):
    index = build_reference_index(entries)
    return resolve_institutional_citations(cites, texts, index), index


def _texts(*paras):
    return {i: t for i, t in enumerate(paras)}


# --- personal regressions (must stay byte-identical) --------------------------


def test_1_personal_narrative_unchanged():
    (c,) = _scan("Smith (2020)")
    assert (c.author_display, c.key.author_key, c.institutional_phrase) == ("Smith", "smith", "")
    assert c.snippet == "Smith (2020)"


def test_2_personal_and_pair_unchanged():
    (c,) = _scan("Smith and Lee (2020)")
    assert c.author_display == "Smith and Lee"
    assert c.key.author_key == "smith"
    assert c.institutional_phrase == ""


def test_3_personal_ampersand_pair_unchanged():
    (c,) = _scan("Smith & Lee (2020)")
    assert c.author_display == "Smith & Lee"
    assert c.institutional_phrase == ""


def test_4_personal_et_al_unchanged():
    (c,) = _scan("Smith et al. (2020) reported it.")
    assert c.author_display == "Smith et al."
    assert c.institutional_phrase == ""


def test_5_curly_apostrophe_surname_unchanged():
    (c,) = _scan("O'Connor (2020) replied.")
    assert c.key.author_key == "o'connor"
    assert c.institutional_phrase == ""


def test_6_hyphenated_surname_unchanged():
    (c,) = _scan("Smith-Jones (2020) reported it.")
    assert c.key.author_key == "smith-jones"
    assert c.institutional_phrase == ""


def test_7_particle_behavior_unchanged():
    (c,) = _scan("van der Meer (2020)")
    assert (c.author_display, c.key.author_key) == ("Meer", "meer")
    assert c.institutional_phrase == ""


# --- institutional positives ----------------------------------------------------


def test_8_full_institution_matches_full_entry():
    cites = _scan("Institutional Standards Board (2021)")
    assert cites[0].institutional_phrase == "Institutional Standards Board"
    resolved, index = _resolve(
        cites, _texts("Institutional Standards Board (2021)"),
        [(9, "Institutional Standards Board. (2021). Academic standard.")],
    )
    assert resolved[0].author_display == "Institutional Standards Board"
    assert resolved[0].key.author_key == "institutional standards board"
    assert resolved[0].snippet == "Institutional Standards Board (2021)"
    assert match(resolved, index) == []


def test_9_world_health_organization_matches():
    cites = _scan("World Health Organization (2022) states this.")
    resolved, index = _resolve(
        cites, _texts("World Health Organization (2022) states this."),
        [(9, "World Health Organization. (2022). Report. WHO Press.")],
    )
    assert match(resolved, index) == []
    assert resolved[0].key.author_key == "world health organization"


def test_10_ministry_of_education_matches():
    cites = _scan("Ministry of Education (2023) requires it.")
    resolved, index = _resolve(
        cites, _texts("Ministry of Education (2023) requires it."),
        [(9, "Ministry of Education. (2023). Policy.")],
    )
    assert match(resolved, index) == []


def test_11_national_research_council_matches():
    cites = _scan("Guidance from the National Research Council (2020) helps.")
    assert cites[0].institutional_phrase == "National Research Council"
    resolved, index = _resolve(
        cites, _texts("Guidance from the National Research Council (2020) helps."),
        [(9, "National Research Council. (2020). Guidance.")],
    )
    assert match(resolved, index) == []


def test_12_and_institution_not_split():
    cites = _scan("Research and Standards Council (2020) met.")
    assert cites[0].institutional_phrase == "Research and Standards Council"
    resolved, index = _resolve(
        cites, _texts("Research and Standards Council (2020) met."),
        [(9, "Research and Standards Council. (2020). Minutes.")],
    )
    assert match(resolved, index) == []
    assert resolved[0].key.author_key == "research and standards council"


def test_parenthetical_institution_matches():
    cites = _scan("See (Institutional Standards Board, 2021) here.")
    assert len(cites) == 1
    assert cites[0].institutional_phrase == "Institutional Standards Board"
    assert cites[0].snippet == "(Institutional Standards Board, 2021)"
    resolved, index = _resolve(
        cites, _texts("See (Institutional Standards Board, 2021) here."),
        [(9, "Institutional Standards Board. (2021). X.")],
    )
    assert match(resolved, index) == []


def test_parenthetical_institution_without_entry_stays_silent():
    cites = _scan("See (Institutional Standards Board, 2021) here.")
    resolved, index = _resolve(
        cites, _texts("See (Institutional Standards Board, 2021) here."),
        [(9, "Smith, J. (2021). Other.")],
    )
    assert resolved == []
    assert match(resolved, index) == []


# --- collision and ambiguity ------------------------------------------------------


def test_13_full_institution_never_matches_personal_board():
    cites = _scan("Institutional Standards Board (2021)")
    resolved, index = _resolve(
        cites, _texts("Institutional Standards Board (2021)"),
        [(10, "Board, J. (2021). Other.")],
    )
    assert resolved[0].key.author_key == "institutional standards board"
    violations = match(resolved, index)
    assert [v.rule_code for v in violations] == ["CITATION_MISMATCH"]
    assert "Institutional Standards Board (2021)" in violations[0].message
    assert violations[0].actual_value == "Institutional Standards Board (2021)"


def test_14_short_board_still_matches_personal_board():
    cites = _scan("Board (2021) cited.")
    assert cites[0].institutional_phrase == ""
    resolved, index = _resolve(
        cites, _texts("Board (2021) cited."),
        [(10, "Board, J. (2021). Other.")],
    )
    assert match(resolved, index) == []


def test_15_full_institution_wins_over_personal_collision():
    cites = _scan("Institutional Standards Board (2021)")
    resolved, index = _resolve(
        cites, _texts("Institutional Standards Board (2021)"),
        [
            (9, "Institutional Standards Board. (2021). Academic standard."),
            (10, "Board, J. (2021). Other."),
        ],
    )
    assert match(resolved, index) == []


def test_16_duplicate_institutional_entries_stay_ambiguous():
    cites = _scan("Institutional Standards Board (2021)")
    resolved, index = _resolve(
        cites, _texts("Institutional Standards Board (2021)"),
        [
            (9, "Institutional Standards Board. (2021). First."),
            (10, "Institutional Standards Board. (2021). Second."),
        ],
    )
    # Identical full keys: exact direction-1 match holds; the duplication
    # itself is reported deterministically — never a guessed match.
    violations = match(resolved, index)
    assert [v.rule_code for v in violations] == ["REFERENCE_DUPLICATE"]
    # Suffixed variants sharing the author key stay REVIEW-ambiguous.
    resolved2, index2 = _resolve(
        cites, _texts("Institutional Standards Board (2021)"),
        [
            (9, "Institutional Standards Board. (2021a). First."),
            (10, "Institutional Standards Board. (2021b). Second."),
        ],
    )
    violations2 = match(resolved2, index2)
    assert [v.rule_code for v in violations2] == ["CITATION_REFERENCE_AMBIGUOUS"]
    assert violations2[0].severity == "REVIEW"


def test_17_two_institutions_sharing_final_token_stay_distinct():
    cites = _scan("National Standards Board (2021) met.")
    resolved, index = _resolve(
        cites, _texts("National Standards Board (2021) met."),
        [
            (9, "Institutional Standards Board. (2021). Other work."),
            (10, "National Standards Board. (2021). Minutes."),
        ],
    )
    assert match(resolved, index) == []
    assert resolved[0].key.author_key == "national standards board"


# --- context safety -----------------------------------------------------------------


def test_capitalized_prose_without_keyword_keeps_final_token():
    # "Later" is capitalized but carries no organization noun: the existing
    # personal path (including dedupe grouping) must survive untouched.
    cites = _scan("Smith (2020) said one thing. Later Smith (2020) said another.")
    assert len(cites) == 2
    assert all(c.institutional_phrase == "" for c in cites)
    assert all(c.key.author_key == "smith" for c in cites)
    resolved, index = _resolve(
        cites, _texts("Smith (2020) said one thing. Later Smith (2020) said another."),
        [(9, "Smith, J. (2020). Title.")],
    )
    assert match(resolved, index) == []


def test_keywordless_institution_keeps_safe_fallback():
    (c,) = _scan("Open Consortium (2021) met.")
    assert c.institutional_phrase == ""
    assert c.key.author_key == "consortium"
    resolved, index = _resolve(
        [c], _texts("Open Consortium (2021) met."),
        [(9, "Smith, J. (2021). Other.")],
    )
    violations = match(resolved, index)
    assert [v.rule_code for v in violations] == ["CITATION_MISMATCH"]
    assert violations[0].message.startswith("Citation 'Consortium (2021)'")


def test_18_prose_before_personal_smith_not_absorbed():
    (c,) = _scan("According to Smith (2020), the result is significant.")
    assert (c.author_display, c.institutional_phrase) == ("Smith", "")


def test_19_prose_before_institution_resolves_without_prose():
    (c,) = _scan("According to the Institutional Standards Board (2021), it applies.")
    assert c.institutional_phrase == "Institutional Standards Board"


def test_20_capitalized_prose_without_proof_never_matches():
    (c,) = _scan("Expected Result Board (2021) shows this.")
    assert c.institutional_phrase == "Expected Result Board"
    resolved, index = _resolve(
        [c], _texts("Expected Result Board (2021) shows this."),
        [(9, "Smith, J. (2021). Other.")],
    )
    violations = match(resolved, index)
    assert [v.rule_code for v in violations] == ["CITATION_MISMATCH"]
    assert violations[0].message.startswith("Citation 'Expected Result Board (2021)'")
    assert violations[0].message != "Citation 'Board (2021)' was found in text, but no matching entry was found in the References bibliography."


def test_21_title_like_text_not_institutional():
    assert parse_institutional_author("A Study of Academic Standards.") is None
    assert parse_reference_entry("A Study of Academic Standards. (2020). X.", 1) is None


def test_22_and_organization_not_split_without_proof():
    resolved, index = _resolve(
        _scan("Research and Standards Council (2020) met."),
        _texts("Research and Standards Council (2020) met."),
        [(9, "Smith, J. (2020). Other.")],
    )
    violations = match(resolved, index)
    assert [v.rule_code for v in violations] == ["CITATION_MISMATCH"]
    assert "Research and Standards Council (2020)" in violations[0].message


# --- classification negatives ----------------------------------------------------------


def test_personal_comma_form_untouched_by_institutional_branch():
    entry = parse_reference_entry("Board, J. (2021). Other.", 10)
    assert entry is not None
    assert entry.author_display == "Board"
    assert entry.key.author_key == "board"


def test_direct_order_personal_name_stays_unkeyable():
    assert parse_institutional_author("John Michael Smith.") is None
    assert parse_reference_entry("John Michael Smith. (2020). Notes.", 1) is None


def test_single_token_keeps_existing_behavior():
    entry = parse_reference_entry("Organization. (2022). Notes.", 1)
    assert entry is not None
    assert entry.key.author_key == "organization"


def test_digits_and_punctuation_reject_institutional():
    assert parse_institutional_author("2020 Report.") is None
    assert parse_institutional_author("Board; Notes.") is None
    assert parse_institutional_author("Standards Board extra words here now today.") is None


def test_normalise_institutional_author_collapses_and_casefolds():
    assert normalise_institutional_author("  Institutional   Standards Board. ") == "institutional standards board"


# --- context helper contracts --------------------------------------------------------------


def test_context_stops_at_clause_boundaries():
    text = "First sentence. According to the Board (2021)"
    chain_at = text.index("Board")
    assert author_resolution_context(text, chain_at) == " According to the "
    assert author_resolution_context("Institutional Standards Board (2021)", 0) == ""


def test_extend_personal_multi_is_identity():
    assert extend_institutional_run("text ", "Smith and Lee") == ("Smith and Lee", 5)
    assert extend_institutional_run("text ", "Smith et al.") == ("Smith et al.", 5)


# --- cross-surface integrity ------------------------------------------------------------------


def test_resolved_location_and_paragraph_unchanged():
    cites = _scan("Institutional Standards Board (2021)", 4)
    resolved, index = _resolve(
        cites, {4: "Institutional Standards Board (2021)"},
        [(9, "Institutional Standards Board. (2021). X.")],
    )
    assert resolved[0].paragraph_index == 4


def test_missing_references_path_untouched():
    from app.services.citation_sensor import run_citation_sensor
    from docx import Document
    import io as _io

    doc = Document()
    doc.add_paragraph("Institutional Standards Board (2021) matters.")
    from app.services.document_parser import extract_paragraphs
    paragraphs = extract_paragraphs(doc)
    violations = run_citation_sensor(doc, paragraphs, config=None, roles=None)
    assert len(violations) == 1
    assert violations[0].rule_code == "CITATION_MISMATCH"
    assert violations[0].location.get("missing_references_section") is True


def test_full_sensor_run_matches_institution_end_to_end():
    from app.services.citation_sensor import run_citation_sensor
    from docx import Document
    from app.services.document_parser import extract_paragraphs

    doc = Document()
    doc.add_paragraph("Institutional Standards Board (2021) matters.")
    doc.add_paragraph("References")
    doc.add_paragraph("Institutional Standards Board. (2021). Academic standard.")
    violations = run_citation_sensor(doc, extract_paragraphs(doc), config=None, roles=None)
    assert violations == []


# --- API-level functional audits -------------------------------------------------------


def _post(client, docx_factory, paragraphs, references):
    file_bytes = docx_factory(paragraphs=paragraphs, references=references)
    resp = client.post(
        "/api/audit",
        files={"file": ("test.docx", file_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200
    return resp.json()


def _citation_codes(payload):
    return [v["rule_code"] for v in payload["physical_layout_errors"]]


def test_api_A_full_institution_matches(client, docx_factory):
    payload = _post(
        client, docx_factory,
        ["Institutional Standards Board (2021) matters."],
        ["Institutional Standards Board. (2021). Academic standard."],
    )
    assert "CITATION_MISMATCH" not in _citation_codes(payload)


def test_api_B_full_institution_vs_personal_collision_mismatches(client, docx_factory):
    payload = _post(
        client, docx_factory,
        ["Institutional Standards Board (2021) matters."],
        ["Board, J. (2021). Other work. Press."],
    )
    mismatches = [v for v in payload["physical_layout_errors"] if v["rule_code"] == "CITATION_MISMATCH"]
    assert len(mismatches) == 1
    assert "Institutional Standards Board (2021)" in mismatches[0]["message"]


def test_api_C_personal_board_preserved(client, docx_factory):
    payload = _post(
        client, docx_factory,
        ["Board (2021) cited here."],
        ["Board, J. (2021). Other work. Press."],
    )
    assert "CITATION_MISMATCH" not in _citation_codes(payload)


def test_api_D_and_institution_matches(client, docx_factory):
    payload = _post(
        client, docx_factory,
        ["Research and Standards Council (2020) met."],
        ["Research and Standards Council. (2020). Minutes."],
    )
    assert "CITATION_MISMATCH" not in _citation_codes(payload)


def test_api_E_missing_references_unchanged(client, docx_factory):
    payload = _post(
        client, docx_factory,
        ["Institutional Standards Board (2021) matters."],
        [],
    )
    mismatches = [v for v in payload["physical_layout_errors"] if v["rule_code"] == "CITATION_MISMATCH"]
    assert len(mismatches) == 1
    assert mismatches[0]["location"].get("missing_references_section") is True
