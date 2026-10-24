"""Focused preservation and ephemeral output-storage contract tests."""

import io
import os
import zipfile
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree

from app.config import settings
from app.services import guided_fix_storage
from app.services.docx_preservation import (
    ReferenceHeadingBoldPackageUnsupported,
    compare_preservation,
    validate_package,
    validate_reference_heading_bold_package,
)


def _docx(size=11):
    doc = Document()
    run = doc.add_paragraph("Original content").runs[0]
    run.font.size = Pt(size)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _replace_part(data, name, transform):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as old, zipfile.ZipFile(out, "w") as new:
        for info in old.infolist():
            content = old.read(info.filename)
            new.writestr(info, transform(content) if info.filename == name else content)
    return out.getvalue()


def _change_part(data, name, change):
    def transform(content):
        root = etree.fromstring(content)
        change(root)
        return etree.tostring(root)
    return _replace_part(data, name, transform)


def _add_parts(data, parts):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as old, zipfile.ZipFile(out, "w") as new:
        for info in old.infolist():
            new.writestr(info, old.read(info.filename))
        for name, content in parts.items():
            new.writestr(name, content)
    return out.getvalue()


def _bold_heading(bold):
    document = Document()
    paragraph = document.add_paragraph()
    for text in ("Refer", " ", "ences", ""):
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
    paragraph.paragraph_format.space_before = Pt(6)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("before_value,after_value", [(False, True), (True, False)])
def test_reference_heading_bold_multirun_changes_only_selected_bold(before_value, after_value):
    before = _bold_heading(before_value)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def change(root):
        for run in root.findall(".//" + w + "body/" + w + "p/" + w + "r")[:3]:
            run.find(w + "rPr/" + w + "b").set(w + "val", "1" if after_value else "0")

    after = _change_part(before, "word/document.xml", change)
    targets = {(0, index, "REFERENCE_HEADING_BOLD") for index in range(3)}
    compare_preservation(before, after, targets)
    with pytest.raises(ValueError):
        compare_preservation(before, after, {(0, 0, "REFERENCE_HEADING_BOLD")})
    with pytest.raises(ValueError):
        compare_preservation(before, after, {(0, None, "REFERENCE_HEADING_BOLD")})


@pytest.mark.parametrize("property_name,attribute,value", [
    ("i", "val", "0"),
    ("rFonts", "ascii", "Arial"),
    ("sz", "val", "28"),
    ("spacing", "val", "20"),
    ("b", "other", "changed"),
    ("bCs", "val", "1"),
])
def test_reference_heading_bold_does_not_mask_run_siblings(property_name, attribute, value):
    before = _bold_heading(False)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def change(root):
        properties = root.find(".//" + w + "rPr")
        properties.find(w + "b").set(w + "val", "1")
        prop = properties.find(w + property_name)
        if prop is None:
            prop = etree.SubElement(properties, w + property_name)
        prop.set(w + attribute, value)

    with pytest.raises(ValueError):
        compare_preservation(before, _change_part(before, "word/document.xml", change),
                             {(0, 0, "REFERENCE_HEADING_BOLD")})


def test_reference_heading_bold_does_not_mask_paragraph_spacing_or_empty_run():
    before = _bold_heading(False)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def change_spacing(root):
        root.find(".//" + w + "pPr/" + w + "spacing").set(w + "before", "240")

    def change_empty(root):
        root.findall(".//" + w + "rPr/" + w + "b")[-1].set(w + "val", "1")

    targets = {(0, index, "REFERENCE_HEADING_BOLD") for index in range(3)}
    for change in (change_spacing, change_empty):
        with pytest.raises(ValueError):
            compare_preservation(before, _change_part(before, "word/document.xml", change), targets)


def _excluded_package(indicator, value):
    data = _docx()
    if indicator == "part":
        return _add_parts(data, {value: b"<marker/>"})
    if indicator == "relationship":
        data = _add_parts(data, {"word/alternate.bin": b"<marker/>"})

        def add_relationship(root):
            etree.SubElement(root, root.tag.removesuffix("Relationships") + "Relationship",
                             Id="excluded", Type=value, Target="alternate.bin")

        return _change_part(data, "word/_rels/document.xml.rels", add_relationship)
    if indicator == "content_type":
        data = _add_parts(data, {"word/alternate.bin": b"<marker/>"})

        def add_content_type(root):
            etree.SubElement(root, root.tag.removesuffix("Types") + "Override",
                             PartName="/word/alternate.bin", ContentType=value)

        return _change_part(data, "[Content_Types].xml", add_content_type)
    if indicator == "protection":
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        return _change_part(data, "word/settings.xml",
                            lambda root: etree.SubElement(root, w + value))
    raise AssertionError("unknown fixture indicator")


@pytest.mark.parametrize("indicator,value", [
    ("part", "_xmlsignatures/alternate.xml"),
    ("part", "word/vbaProject.bin"),
    ("part", "word/vbaData.xml"),
    ("part", "word/vbaProjectSignature.bin"),
    ("relationship", "http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/origin"),
    ("relationship", "http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/signature"),
    ("relationship", "http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/certificate"),
    ("relationship", "http://schemas.microsoft.com/office/2006/relationships/vbaProject"),
    ("relationship", "http://schemas.microsoft.com/office/2006/relationships/vbaData"),
    ("relationship", "http://schemas.microsoft.com/office/2006/relationships/vbaProjectSignature"),
    ("content_type", "application/vnd.openxmlformats-package.digital-signature-origin"),
    ("content_type", "application/vnd.openxmlformats-package.digital-signature-xmlsignature+xml"),
    ("content_type", "application/vnd.openxmlformats-package.digital-signature-certificate"),
    ("content_type", "application/vnd.ms-office.vbaProject"),
    ("content_type", "application/vnd.ms-office.vbaData+xml"),
    ("content_type", "application/vnd.ms-office.vbaProjectSignature"),
    ("content_type", "application/vnd.ms-word.document.macroEnabled.main+xml"),
    ("content_type", "application/vnd.ms-word.template.macroEnabledTemplate.main+xml"),
    ("protection", "documentProtection"),
    ("protection", "writeProtection"),
])
def test_reference_heading_bold_excludes_proven_package_features_only_for_candidate(indicator, value):
    data = _excluded_package(indicator, value)
    validate_package(data)
    with pytest.raises(ReferenceHeadingBoldPackageUnsupported):
        validate_reference_heading_bold_package(data)


def test_reference_heading_bold_ordinary_package_is_eligible():
    data = _docx()
    validate_package(data)
    validate_reference_heading_bold_package(data)


@pytest.mark.parametrize("content_type", [
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "application/xml",
    "",
])
def test_reference_heading_bold_requires_normal_docx_main_type(content_type):
    def change(root):
        main = next(child for child in root if child.get("PartName") == "/word/document.xml")
        main.set("ContentType", content_type)

    data = _change_part(_docx(), "[Content_Types].xml", change)
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


def test_reference_heading_bold_macro_main_type_is_typed_exclusion():
    def change(root):
        main = next(child for child in root if child.get("PartName") == "/word/document.xml")
        main.set("ContentType", "application/vnd.ms-word.document.macroEnabled.main+xml")

    data = _change_part(_docx(), "[Content_Types].xml", change)
    with pytest.raises(ReferenceHeadingBoldPackageUnsupported):
        validate_reference_heading_bold_package(data)


@pytest.mark.parametrize("part,content", [
    ("word/styles.xml", b"<broken"),
    ("word/settings.xml", b"<broken"),
    ("word/document.xml", b"<broken"),
    ("word/_rels/document.xml.rels", b"<broken"),
    ("[Content_Types].xml", b"<broken"),
    ("word/settings.xml", b"<!DOCTYPE x [<!ENTITY x 'unsafe'>]><x>&x;</x>"),
    ("word/settings.xml", b"<settings/>"),
    ("[Content_Types].xml", b"<Types/>"),
])
def test_reference_heading_bold_malformed_xml_is_not_feature_exclusion(part, content):
    data = _replace_part(_docx(), part, lambda old: content)
    validate_package(data)
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


def test_reference_heading_bold_relationship_validation_is_not_feature_exclusion():
    data = _change_part(_docx(), "_rels/.rels",
                        lambda root: root[0].set("Target", "missing.xml"))
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


def test_reference_heading_bold_checks_declared_xml_with_alternative_part_name():
    data = _excluded_package("content_type", "application/xml")
    validate_reference_heading_bold_package(data)
    data = _replace_part(data, "word/alternate.bin", lambda content: b"<broken")
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


@pytest.mark.parametrize("attribute,value", [("Id", ""), ("TargetMode", "Unknown")])
def test_reference_heading_bold_malformed_relationship_is_plain_value_error(attribute, value):
    data = _change_part(_docx(), "_rels/.rels",
                        lambda root: root[0].set(attribute, value))
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


@pytest.mark.parametrize("change", ["missing_main", "duplicate", "wrong_entry"])
def test_reference_heading_bold_malformed_content_types_are_plain_value_error(change):
    def mutate(root):
        if change == "missing_main":
            root.remove(next(child for child in root if child.get("PartName") == "/word/document.xml"))
        elif change == "duplicate":
            etree.SubElement(root, root[0].tag, **dict(root[0].attrib))
        else:
            etree.SubElement(root, "unknown")

    data = _change_part(_docx(), "[Content_Types].xml", mutate)
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


@pytest.mark.parametrize("invalid", ["zip", "duplicate", "path", "missing", "expanded", "total", "count"])
def test_reference_heading_bold_keeps_existing_package_boundaries(invalid, monkeypatch):
    from app.services import docx_preservation

    data = _docx()
    if invalid == "zip":
        data = b"not a zip"
    elif invalid in ("duplicate", "path"):
        name = "word/document.xml" if invalid == "duplicate" else "../outside.xml"
        data = _add_parts(data, {name: b"<x/>"})
    elif invalid == "missing":
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("word/document.xml", b"<x/>")
        data = stream.getvalue()
    else:
        boundary = {"expanded": "_MAX_EXPANDED_PART", "total": "_MAX_EXPANDED_TOTAL",
                    "count": "_MAX_PARTS"}[invalid]
        monkeypatch.setattr(docx_preservation, boundary, 1)
    with pytest.raises(ValueError) as caught:
        validate_reference_heading_bold_package(data)
    assert type(caught.value) is ValueError


def test_declared_run_size_and_repacked_package_pass():
    before = _docx(11)
    after = _docx(12)
    compare_preservation(before, after, {(0, 0, "FONT_SIZE")})
    compare_preservation(before, _replace_part(before, "word/document.xml", lambda part: part), set())

@pytest.mark.parametrize("rule,property_name,attribute,value,run_index", [
    ("FONT_CONSISTENCY", "rFonts", "ascii", "Times New Roman", 0),
    ("ALIGNMENT", "jc", "val", "center", None),
    ("REFERENCE_HEADING_ALIGNMENT", "jc", "val", "center", None),
    ("REFERENCE_HEADING_BOLD", "b", "val", "1", 0),
    ("LINE_SPACING", "spacing", "line", "360", None),
    ("SPACE_BEFORE", "spacing", "before", "120", None),
    ("SPACE_AFTER", "spacing", "after", "120", None),
])
def test_declared_property_only(rule, property_name, attribute, value, run_index):
    before = _docx()
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def change(root):
        paragraph = root.find(".//" + w + "body/" + w + "p")
        target = paragraph.find(w + "r") if run_index is not None else paragraph
        parent_name = "rPr" if run_index is not None else "pPr"
        parent = target.find(w + parent_name)
        if parent is None:
            parent = etree.Element(w + parent_name)
            target.insert(0, parent)
        prop = parent.find(w + property_name)
        if prop is None:
            prop = etree.SubElement(parent, w + property_name)
        prop.set(w + attribute, value)

    after = _change_part(before, "word/document.xml", change)
    compare_preservation(before, after, {(0, run_index, rule)})
    with pytest.raises(ValueError):
        compare_preservation(before, after, set())


def test_other_attribute_in_selected_property_is_protected():
    before = _docx()
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def change(root):
        run_properties = root.find(".//" + w + "rPr")
        etree.SubElement(run_properties, w + "rFonts").set(w + "cs", "Changed")

    with pytest.raises(ValueError):
        compare_preservation(before, _change_part(before, "word/document.xml", change),
                             {(0, 0, "FONT_CONSISTENCY")})


def test_only_save_metadata_is_ignored():
    before = _docx()
    dcterms = "{http://purl.org/dc/terms/}"

    def change_modified(root):
        element = root.find(dcterms + "modified")
        assert element is not None
        element.text = "2026-10-18T12:00:00Z"

    saved = _change_part(before, "docProps/core.xml", change_modified)
    compare_preservation(before, saved, set())
    changed_author = _change_part(before, "docProps/core.xml",
                                  lambda root: root.find("{http://purl.org/dc/elements/1.1/}creator").__setattr__("text", "Other author"))
    with pytest.raises(ValueError):
        compare_preservation(before, changed_author, set())


def test_unselected_and_wrong_property_changes_fail():
    before, after = _docx(11), _docx(12)
    with pytest.raises(ValueError):
        compare_preservation(before, after, set())
    with pytest.raises(ValueError):
        compare_preservation(before, after, {(0, 0, "FONT_CONSISTENCY")})
    with pytest.raises(ValueError):
        compare_preservation(before, after, {(0, 1, "FONT_SIZE")})
    with pytest.raises(ValueError):
        compare_preservation(before, after, {(0, None, "FONT_SIZE")})

def test_every_grouped_run_must_be_declared():
    document = Document()
    paragraph = document.add_paragraph()
    for text in ("one", "two"):
        paragraph.add_run(text).font.size = Pt(11)
    stream = io.BytesIO()
    document.save(stream)
    before = stream.getvalue()
    for run in paragraph.runs:
        run.font.size = Pt(12)
    stream = io.BytesIO()
    document.save(stream)
    after = stream.getvalue()
    with pytest.raises(ValueError):
        compare_preservation(before, after, {(0, 0, "FONT_SIZE")})
    compare_preservation(before, after, {(0, 0, "FONT_SIZE"), (0, 1, "FONT_SIZE")})


def test_binary_media_is_preserved():
    from tests.conftest import _PNG_1PX

    document = Document(io.BytesIO(_docx()))
    document.paragraphs[0].add_run().add_picture(io.BytesIO(_PNG_1PX))
    stream = io.BytesIO()
    document.save(stream)
    before = stream.getvalue()
    with zipfile.ZipFile(io.BytesIO(before)) as archive:
        media = next(name for name in archive.namelist() if name.startswith("word/media/"))
    with pytest.raises(ValueError, match="binary"):
        compare_preservation(before, _replace_part(before, media, lambda data: data + b"x"), set())


def test_text_media_relationship_and_unknown_xml_change_fail():
    before = _docx()
    edited = _change_part(before, "word/document.xml", lambda root: root.xpath("//*[local-name()='t']")[0].__setattr__("text", "changed"))
    with pytest.raises(ValueError):
        compare_preservation(before, edited, {(0, 0, "FONT_SIZE")})
    edited = _change_part(before, "word/styles.xml", lambda root: root.set("other", "value"))
    with pytest.raises(ValueError):
        compare_preservation(before, edited, set())
    edited = _change_part(before, "_rels/.rels", lambda root: root[0].set("Target", "missing.xml"))
    with pytest.raises(ValueError):
        compare_preservation(before, edited, set())


def test_duplicate_zip_entry_and_missing_part_fail():
    before = _docx()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("word/document.xml", b"<document/>")
        archive.writestr("word/document.xml", b"<document/>")
    with pytest.raises(ValueError, match="duplicate"):
        compare_preservation(before, out.getvalue(), set())
    missing = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(missing, "w") as new:
        for info in old.infolist():
            if info.filename != "word/styles.xml":
                new.writestr(info, old.read(info.filename))
    with pytest.raises(ValueError, match="missing required DOCX part"):
        compare_preservation(before, missing.getvalue(), set())

def test_overexpanded_unreferenced_member_rejected_before_parsing():
    before = _docx()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(before)) as source, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, source.read(entry.filename))
        target.writestr('word/unused.bin', b'0' * (33 * 1024 * 1024))
    assert len(out.getvalue()) < 10 * 1024 * 1024
    with pytest.raises(ValueError, match='expanded-size'):
        validate_package(out.getvalue())


@pytest.mark.parametrize("rule", ["FONT_SIZE", "REFERENCE_HEADING_BOLD"])
def test_cannot_select_indirect_runs(rule):
    before = _docx()
    edited = _change_part(before, "word/document.xml", _wrap_run_in_hyperlink)
    with pytest.raises(ValueError):
        compare_preservation(edited, edited, {(0, 0, rule)})


def _wrap_run_in_hyperlink(root):
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph = root.find(".//" + w + "body/" + w + "p")
    run = paragraph.find(w + "r")
    paragraph.remove(run)
    etree.SubElement(paragraph, w + "hyperlink").append(run)


@pytest.fixture
def output_storage(tmp_path, monkeypatch):
    path = tmp_path / "outputs"
    monkeypatch.setattr(settings, "GUIDED_FIX_STORAGE_DIR", str(path))
    yield path
    with guided_fix_storage._LOCK:
        guided_fix_storage._outputs.clear()


def test_output_token_one_use_and_audit_bound(output_storage):
    data = _docx()
    token, expires = guided_fix_storage.store("audit-one", data, "fixed.docx")
    assert datetime.fromisoformat(expires) > datetime.now(timezone.utc)
    with pytest.raises(guided_fix_storage.DownloadWrongAuditError):
        guided_fix_storage.retrieve("audit-two", token)
    assert guided_fix_storage.retrieve("audit-one", token) == (data, "fixed.docx")
    with pytest.raises(guided_fix_storage.DownloadGoneError):
        guided_fix_storage.retrieve("audit-one", token)
    assert list(output_storage.glob("*.docx")) == []


def test_expired_and_crashed_output_removed(output_storage):
    data = _docx()
    token, _ = guided_fix_storage.store("audit-one", data, "fixed.docx")
    with guided_fix_storage._LOCK:
        record = guided_fix_storage._outputs[token]
        from dataclasses import replace
        guided_fix_storage._outputs[token] = replace(record, expiry=datetime.now(timezone.utc) - timedelta(seconds=1))
    guided_fix_storage.cleanup()
    with pytest.raises(guided_fix_storage.DownloadGoneError):
        guided_fix_storage.retrieve("audit-one", token)
    stale = output_storage / ("x" * 43 + ".docx")
    stale.write_bytes(b"orphaned")
    old = datetime.now(timezone.utc).timestamp() - 1000
    os.utime(stale, (old, old))
    guided_fix_storage.cleanup()
    assert not stale.exists()


def test_modified_output_cannot_be_downloaded(output_storage):
    token, _ = guided_fix_storage.store("audit-one", _docx(), "fixed.docx")
    (output_storage / f"{token}.docx").write_bytes(b"tampered")
    with pytest.raises(OSError, match="hash"):
        guided_fix_storage.retrieve("audit-one", token)
    assert not (output_storage / f"{token}.docx").exists()


def test_bad_output_filename_never_creates_storage(output_storage):
    with pytest.raises(ValueError):
        guided_fix_storage.store("audit-one", _docx(), "../other.docx")
    assert not output_storage.exists()


def test_revocation_and_exclusive_creation(output_storage):
    data = _docx()
    first, _ = guided_fix_storage.store("audit-one", data, "fixed.docx")
    second, _ = guided_fix_storage.store("audit-one", data, "fixed.docx")
    assert first != second
    guided_fix_storage.revoke("other-audit", first)
    assert (output_storage / f"{first}.docx").exists()
    guided_fix_storage.revoke("audit-one", first)
    guided_fix_storage.revoke("audit-one", first)
    with pytest.raises(guided_fix_storage.DownloadGoneError):
        guided_fix_storage.retrieve("audit-one", first)
    assert guided_fix_storage.retrieve("audit-one", second) == (data, "fixed.docx")


def test_store_failure_does_not_publish_token(output_storage, monkeypatch):
    data = _docx()
    def reject_write(*args, **kwargs):
        raise OSError("disk unavailable")

    monkeypatch.setattr(guided_fix_storage.os, "open", reject_write)
    with pytest.raises(OSError, match="disk unavailable"):
        guided_fix_storage.store("audit-one", data, "fixed.docx")
    assert list(output_storage.iterdir()) == []
