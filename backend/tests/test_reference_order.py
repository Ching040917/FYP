"""Reference List alphabetical-order review — author identities only.

Covers the single REFERENCE_ORDER rule:

  * adjacent reliable pairs compare author keys (personal first-surname,
    institutional full phrase), both directions;
  * unreliable entries (unparseable, title-first, lowercase particles,
    typed or styled numbering) break adjacency — never compared across;
  * numbered lists are excluded entirely;
  * equal author keys stay silent (same author, coauthor variants,
    suffixes, n.d. vs dated, exact duplicates);
  * years, suffixes, coauthor tails, titles and DOI/URL text are never
    compared;
  * the toggle defaults off; disabled emits nothing and scores nothing;
  * findings are MINOR under citation_apa with profile-scoped wording and
    no normalized-key leakage;
  * sensor, matching, formatting, AI and Guided Fixes behaviour is
    unchanged; no network use.
"""

import io
import json
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.document_parser import extract_paragraphs
from app.services.layout_engine import run_static_rules_engine
from app.services.layout_violation import LayoutViolation
from app.services.pdf_report import (
    _REQUIRED_ACTIONS,
    _RULE_NAME,
    _pdf_finding_action,
)
from app.services.profile_preset_adapter import EffectiveProfileConfig
from app.services.profile_registry import SUC_PROFILE_ID, get_builtin_profile
from app.services.profile_schema import new_custom_profile
from app.services.profile_snapshot import resolve_snapshot
from app.services.reference_order import (
    RULE_REFERENCE_ORDER,
    author_sort_identity,
    run_reference_order,
)
from app.services.reference_rules import run_reference_rules
from app.services.scoring import (
    CATEGORY_META,
    calculate_weighted_score_detailed,
    rule_code_to_category,
)
from tests.conftest import make_docx_bytes

HLINK_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/"
    "relationships/hyperlink"
)


def _profile(order_on: bool = True) -> EffectiveProfileConfig:
    profile = new_custom_profile(
        "Order Tests", base=get_builtin_profile(SUC_PROFILE_ID),
    )
    profile.references.check_reference_order = order_on
    return EffectiveProfileConfig(resolve_snapshot(profile))


def _doc_bytes(references) -> bytes:
    return make_docx_bytes(
        paragraphs=["Prior work motivates this study."],
        references=references,
    )


def _run(references, order_on: bool = True):
    data = _doc_bytes(references)
    doc = Document(io.BytesIO(data))
    return run_reference_rules(
        doc, extract_paragraphs(doc), _profile(order_on),
    )


def _order(violations):
    return [v for v in violations if v.rule_code == RULE_REFERENCE_ORDER]


def _styled_numbered_docx(entries) -> bytes:
    """Entries carrying direct <w:numPr> (style-inherited numbering shape)."""
    doc = Document()
    doc.add_paragraph("Body text.")
    doc.add_paragraph("References")
    for text in entries:
        para = doc.add_paragraph(text)
        pPr = para._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        numId = OxmlElement("w:numId")
        numId.set(qn("w:val"), "1")
        numPr.append(ilvl)
        numPr.append(numId)
        pPr.append(numPr)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _hyperlink_docx() -> bytes:
    """Inverted pair where the first entry ends in a DOI hyperlink whose
    hidden target differs from its visible label."""
    doc = Document()
    doc.add_paragraph("Body text.")
    doc.add_paragraph("References")
    first = doc.add_paragraph("Smith, J. (2020). Title. ")
    hl = OxmlElement("w:hyperlink")
    r_id = first.part.relate_to(
        "https://hidden-target.invalid/secret-entry", HLINK_REL,
        is_external=True,
    )
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = "https://doi.org/10.1000/xyz"
    run.append(text)
    hl.append(run)
    first._p.append(hl)
    doc.add_paragraph("Adams, A. (2021). Title.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


ADAMS = "Adams, A. (2020). First title. Publisher."
BROWN = "Brown, B. (2020). Second title. Publisher."
SMITH = "Smith, J. (2020). Third title. Publisher."


# --- 1/2: personal-author order ---------------------------------------------


def test_correct_personal_order_is_silent():
    assert _order(_run([ADAMS, BROWN, SMITH])) == []


def test_inverted_personal_order_flags_the_later_entry():
    violations = _order(_run([SMITH, ADAMS, BROWN]))
    assert len(violations) == 1
    finding = violations[0]
    assert finding.severity == "MINOR"
    # Later physical entry: Adams is the second entry (body 0, heading 1).
    assert finding.location == {"paragraph_index": 3}
    assert "Adams" in (finding.actual_value or "")


def test_middle_inversion_flags_only_the_later_entry():
    violations = _order(_run([ADAMS, SMITH, BROWN]))
    assert len(violations) == 1
    assert violations[0].location == {"paragraph_index": 4}
    assert "Brown" in (violations[0].actual_value or "")


# --- 3/4: institutional-author order ----------------------------------------


INSTITUTIONAL = "Institutional Standards Board. (2021). Title."
MINISTRY = "Ministry of Education. (2023). Title."


def test_correct_institutional_order_is_silent():
    assert _order(_run([INSTITUTIONAL, MINISTRY])) == []


def test_inverted_institutional_order_flags_the_later_entry():
    violations = _order(_run([MINISTRY, INSTITUTIONAL]))
    assert len(violations) == 1
    assert violations[0].location == {"paragraph_index": 3}
    assert "Institutional" in (violations[0].actual_value or "")


# --- 5/6: mixed personal and institutional ----------------------------------


def test_mixed_personal_institutional_correct_order_is_silent():
    # adams < institutional standards board < smith on normalized keys.
    assert _order(_run([ADAMS, INSTITUTIONAL, SMITH])) == []


def test_mixed_personal_institutional_inversion_is_flagged():
    violations = _order(_run([SMITH, INSTITUTIONAL]))
    assert len(violations) == 1
    assert violations[0].location == {"paragraph_index": 3}


# --- 7/8: apostrophe and hyphen normalization -------------------------------


def test_straight_apostrophe_surname_has_a_reliable_identity():
    assert author_sort_identity("O'Connor, P. (2020). Title.") == "o'connor"


def test_curly_apostrophe_surname_is_unsupported_and_breaks_adjacency():
    # Existing-parser limitation (the entry-head pattern accepts only the
    # straight apostrophe): a curly O’Connor is unkeyable, so it breaks
    # adjacency instead of being guessed.
    assert author_sort_identity("O\u2019Connor, P. (2020). Title.") is None
    assert _order(_run([
        SMITH,
        "O\u2019Connor, P. (2020). Title.",
        ADAMS,
    ])) == []


def test_apostrophe_surname_orders_before_plain_surname():
    assert _order(_run([
        "O'Connor, P. (2020). Title.",
        SMITH,
    ])) == []
    violations = _order(_run([
        SMITH,
        "O'Connor, P. (2020). Title.",
    ]))
    assert len(violations) == 1


def test_hyphenated_surname_orders_as_one_token():
    assert author_sort_identity("Smith-Jones, A. (2020). Title.") == "smith-jones"
    assert _order(_run([
        ADAMS,
        "Smith-Jones, A. (2020). Title.",
    ])) == []
    violations = _order(_run([
        "Smith-Jones, A. (2020). Title.",
        ADAMS,
    ]))
    assert len(violations) == 1


# --- 9/10/11: unreliable entries break adjacency ----------------------------


def test_lowercase_particle_entry_breaks_adjacency():
    assert author_sort_identity("van der Meer, J. (2020). Title.") is None
    assert _order(_run([
        SMITH,
        "van der Meer, J. (2020). Title.",
        ADAMS,
    ])) == []


def test_title_first_entry_breaks_adjacency():
    assert author_sort_identity("A study of things. (2020). Title.") is None
    assert _order(_run([
        SMITH,
        "A study of things. (2020). Title.",
        ADAMS,
    ])) == []


def test_unparseable_entry_breaks_adjacency():
    assert author_sort_identity("Smith, J. Title without year.") is None
    assert _order(_run([
        SMITH,
        "Smith, J. Title without year.",
        ADAMS,
    ])) == []


def test_comparison_resumes_after_a_gap():
    # Smith|gap breaks; Brown before Adams is a fresh adjacent reliable
    # pair, so the genuine inversion is still reported on Adams (index 5).
    violations = _order(_run([
        SMITH,
        "Notes without a year token.",
        BROWN,
        ADAMS,
    ]))
    assert len(violations) == 1
    assert violations[0].location == {"paragraph_index": 5}


# --- 12: numbered lists excluded --------------------------------------------


def test_typed_numbering_is_excluded_entirely():
    assert author_sort_identity("1. Smith, J. (2020). Title.") is None
    assert _order(_run([
        "1. Smith, J. (2020). Title.",
        "2. Adams, A. (2020). Title.",
    ])) == []


def test_styled_numbering_is_excluded_entirely():
    data = _styled_numbered_docx([SMITH, ADAMS])
    doc = Document(io.BytesIO(data))
    violations = run_reference_order(
        doc, extract_paragraphs(doc), _profile(True),
    )
    assert violations == []


# --- 13/14/15: same-author silence ------------------------------------------


def test_same_first_surname_with_different_coauthors_is_silent():
    assert _order(_run([
        "Smith, J. (2019). Solo title.",
        "Smith, J., & Lee, K. (2020). Joint title.",
    ])) == []
    assert _order(_run([
        "Smith, J., & Lee, K. (2020). Joint title.",
        "Smith, J. (2019). Solo title.",
    ])) == []


def test_same_author_with_different_years_is_silent():
    assert _order(_run([
        "Smith, J. (2019). Early title.",
        "Smith, J. (2020). Later title.",
    ])) == []
    assert _order(_run([
        "Smith, J. (2020). Later title.",
        "Smith, J. (2019). Early title.",
    ])) == []


def test_same_author_with_suffixes_is_silent():
    assert _order(_run([
        "Smith, J. (2020a). First title.",
        "Smith, J. (2020b). Second title.",
    ])) == []
    assert _order(_run([
        "Smith, J. (2020b). Second title.",
        "Smith, J. (2020a). First title.",
    ])) == []


# --- 16/17: n.d. behaviour --------------------------------------------------


def test_reliable_nd_author_compares_by_author_only():
    assert _order(_run([
        "Adams, A. (n.d.). Title.",
        "Brown, B. (2020). Title.",
    ])) == []
    violations = _order(_run([
        "Brown, B. (2020). Title.",
        "Adams, A. (n.d.). Title.",
    ]))
    assert len(violations) == 1
    assert violations[0].location == {"paragraph_index": 3}


def test_same_author_nd_and_dated_is_silent():
    assert _order(_run([
        "Smith, J. (n.d.). Title.",
        "Smith, J. (2020). Title.",
    ])) == []
    assert _order(_run([
        "Smith, J. (2020). Title.",
        "Smith, J. (n.d.). Title.",
    ])) == []


# --- 18: duplicates belong to REFERENCE_DUPLICATE ----------------------------


def test_duplicate_full_keys_are_silent_for_order():
    assert _order(_run([SMITH, SMITH])) == []


# --- 19/20: toggle ----------------------------------------------------------


def test_disabled_toggle_emits_nothing():
    assert _run([SMITH, ADAMS], order_on=False) == []


def test_enabled_toggle_emits_the_inversion():
    assert len(_order(_run([SMITH, ADAMS], order_on=True))) == 1


# --- 21/22/23: location and wording -----------------------------------------


def test_location_is_the_later_physical_entry_only():
    violations = _order(_run([SMITH, ADAMS]))
    assert violations[0].location == {"paragraph_index": 3}


def test_current_required_and_action_are_profile_scoped():
    violations = _order(_run([SMITH, ADAMS]))
    finding = violations[0]
    assert finding.actual_value and len(finding.actual_value) <= 120
    assert "Adams" in finding.actual_value
    assert finding.expected_value == \
        "Arrange Reference List entries alphabetically by author"
    assert "Order Tests" in finding.message
    assert "APA 7" not in finding.message


def test_no_normalized_key_or_index_array_is_exposed():
    violations = _order(_run([SMITH, ADAMS]))
    finding = violations[0]
    assert set(finding.location.keys()) == {"paragraph_index"}
    blob = json.dumps(finding.to_dict())
    assert "author_key" not in blob
    assert "paragraph_indexes" not in blob


# --- 24/25/26: existing behaviour unchanged ---------------------------------


def test_formatting_findings_unchanged_apart_from_order():
    data = _doc_bytes([SMITH, ADAMS])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    off = run_reference_rules(doc, paragraphs, _profile(False))
    on = run_reference_rules(doc, paragraphs, _profile(True))
    off_codes = sorted(v.rule_code for v in off)
    on_codes = sorted(v.rule_code for v in on)
    assert on_codes == sorted(off_codes + [RULE_REFERENCE_ORDER])


def test_citation_sensor_output_ignores_the_new_toggle():
    from app.services.citation_sensor import run_citation_sensor

    data = _doc_bytes([SMITH, SMITH])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    off = run_citation_sensor(doc, paragraphs, config=_profile(False))
    on = run_citation_sensor(doc, paragraphs, config=_profile(True))
    assert [(v.rule_code, v.severity) for v in off] == \
        [(v.rule_code, v.severity) for v in on]
    assert any(v.rule_code == "REFERENCE_DUPLICATE" for v in on)
    assert not any(v.rule_code == RULE_REFERENCE_ORDER for v in on)


def test_institutional_matching_ignores_the_new_toggle():
    from app.services.citation_sensor import run_citation_sensor

    data = make_docx_bytes(
        paragraphs=["The board reports growth (Institutional Standards Board, 2021)."],
        references=[INSTITUTIONAL],
    )
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    off = run_citation_sensor(doc, paragraphs, config=_profile(False))
    on = run_citation_sensor(doc, paragraphs, config=_profile(True))
    assert [(v.rule_code, v.severity) for v in off] == \
        [(v.rule_code, v.severity) for v in on]


# --- 27/28: hyperlinks and network ------------------------------------------


def test_visible_doi_tail_does_not_change_identity():
    with_doi = "Smith, J. (2020). Title. https://doi.org/10.1000/xyz"
    assert author_sort_identity(with_doi) == "smith"
    assert _order(_run([with_doi, ADAMS])) != []


def test_hidden_hyperlink_target_is_never_exposed():
    data = _hyperlink_docx()
    doc = Document(io.BytesIO(data))
    violations = run_reference_order(
        doc, extract_paragraphs(doc), _profile(True),
    )
    # Smith (with visible DOI label) before Adams: one inversion on Adams.
    assert len(violations) == 1
    assert violations[0].location == {"paragraph_index": 3}
    blob = json.dumps([v.to_dict() for v in violations])
    assert "hidden-target" not in blob


def test_visible_hyperlink_label_keeps_a_reliable_identity():
    data = _hyperlink_docx()
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    identities = [
        author_sort_identity(
            p.get("text") or "",
            p,
            {i: q for i, q in enumerate(doc.paragraphs)}.get(p.get("index")),
        )
        for p in paragraphs if (p.get("text") or "").strip().startswith(("Smith", "Adams"))
    ]
    assert identities == ["smith", "adams"]


def test_no_network_use_in_the_order_module():
    import app.services.reference_order as mod

    source = open(mod.__file__, encoding="utf-8").read()
    for token in ("socket", "urllib", "http.client", "requests", "httpx", "urlopen"):
        assert token not in source


# --- engine wiring ----------------------------------------------------------


def test_static_engine_wires_the_rule_through():
    data = _doc_bytes([SMITH, ADAMS])
    off = run_static_rules_engine(data, config=_profile(False))
    on = run_static_rules_engine(data, config=_profile(True))
    assert RULE_REFERENCE_ORDER not in [v.rule_code for v in off]
    assert [v.rule_code for v in on].count(RULE_REFERENCE_ORDER) == 1


# --- scoring ----------------------------------------------------------------


def test_order_maps_to_citation_apa_in_both_contracts():
    assert rule_code_to_category("REFERENCE_ORDER") == "citation_apa"
    assert rule_code_to_category("REFERENCE_ORDER", legacy_spacing=True) == "citation_apa"


def test_category_weights_and_caps_are_unchanged():
    meta = CATEGORY_META["citation_apa"]
    assert (meta.major_weight, meta.minor_weight, meta.cap) == (5, 2, 25)


def test_single_minor_deducts_once_and_reconciles():
    violation = LayoutViolation(
        rule_code=RULE_REFERENCE_ORDER, severity="MINOR",
        location={"paragraph_index": 3},
        message="m", expected_value="e", actual_value="a",
    )
    result = calculate_weighted_score_detailed([violation])
    assert result.total == 98
    assert result.minor_count == 1
    assert result.major_count == 0
    row = next(b for b in result.breakdown if b.category == "citation_apa")
    assert (row.minor, row.deduction) == (1, 2)
    assert sum(b.deduction for b in result.breakdown) == 100 - result.total


# --- PDF copy ---------------------------------------------------------------


def test_pdf_title_and_action_match_the_approved_wording():
    assert _RULE_NAME["REFERENCE_ORDER"] == "Reference List is not in alphabetical order"
    assert _REQUIRED_ACTIONS["REFERENCE_ORDER"] == (
        "Move this Reference List entry to the correct alphabetical position "
        "according to the selected profile."
    )
    violation = LayoutViolation(
        rule_code=RULE_REFERENCE_ORDER, severity="MINOR",
        location={"paragraph_index": 3},
        message="Deterministic message.",
        expected_value="Arrange Reference List entries alphabetically by author",
        actual_value="Adams, A. (2021). Title.",
    )
    action = _pdf_finding_action(violation)
    assert action == _REQUIRED_ACTIONS["REFERENCE_ORDER"]
    assert "APA 7" not in action


# --- API / History / PDF parity ---------------------------------------------


def _api_profile(order_on: bool) -> dict:
    profile = new_custom_profile(
        "Order API", base=get_builtin_profile(SUC_PROFILE_ID),
    )
    profile.references.check_reference_order = order_on
    return profile.to_dict()


def _post_audit(client, references, order_on: bool):
    return client.post(
        "/api/audit",
        files={"file": ("order.docx", _doc_bytes(references), "application/octet-stream")},
        data={"custom_profile": json.dumps(_api_profile(order_on))},
    )


def test_api_end_to_end_with_order_enabled(client):
    post = _post_audit(client, [SMITH, ADAMS], True)
    assert post.status_code == 200, post.text
    payload = post.json()
    codes = [v["rule_code"] for v in payload["physical_layout_errors"]]
    assert codes.count(RULE_REFERENCE_ORDER) == 1
    finding = next(
        v for v in payload["physical_layout_errors"]
        if v["rule_code"] == RULE_REFERENCE_ORDER
    )
    assert finding["severity"] == "MINOR"
    assert finding["location"]["paragraph_index"] == 2 + 1  # body + heading precede entries
    snapshot = payload["profile_snapshot"]
    assert snapshot["schema_version"] == 6
    assert snapshot["references"]["check_reference_order"] is True

    audit_id = payload["audit_id"]
    reopened = client.get(f"/api/audit/{audit_id}")
    assert reopened.status_code == 200
    assert reopened.json()["violations"] == payload["physical_layout_errors"]

    listing = client.get("/api/audits")
    assert listing.status_code == 200
    assert any(item["id"] == audit_id for item in listing.json())

    pdf = client.get(f"/api/audit/{audit_id}/export-pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    from pypdf import PdfReader

    text = re.sub(
        r"\s+", " ",
        "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(pdf.content)).pages
        ),
    )
    assert text.count("Reference List is not in alphabetical order") >= 1
    action = ("Move this Reference List entry to the correct alphabetical position "
              "according to the selected profile.")
    assert text.count(action) == 1
    assert f"Action {action}" in text
    assert finding["location"]["paragraph_index"] == 3
    assert "Paragraph 4" in text
    assert finding["actual_value"] in text
    assert "automatically reordered" not in text.lower()


def test_api_builtin_profile_has_no_order_finding(client):
    post = client.post(
        "/api/audit",
        files={"file": ("order.docx", _doc_bytes([SMITH, ADAMS]), "application/octet-stream")},
        params={"profile_id": SUC_PROFILE_ID},
    )
    assert post.status_code == 200, post.text
    codes = [v["rule_code"] for v in post.json()["physical_layout_errors"]]
    assert RULE_REFERENCE_ORDER not in codes


def test_api_score_reflects_one_citation_apa_minor(client):
    post = _post_audit(client, [SMITH, ADAMS], True)
    assert post.status_code == 200, post.text
    payload = post.json()
    breakdown = {b["category"]: b for b in payload["score_breakdown"]}
    assert breakdown["citation_apa"]["minor"] == 1
    assert payload["weighted_compliance_score"] == 100 - sum(
        b["deduction"] for b in payload["score_breakdown"]
    )


def test_validate_endpoint_accepts_the_new_toggle(client):
    resp = client.post(
        "/api/formatting-profiles/validate", json=_api_profile(True),
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["valid"] is True
