"""Reference-section structural boundaries — one shared termination contract.

Proves reference_section_bounds, role classification, Phase 2 matching
bounds, and Phase 3 entry selection agree on a single structural rule:

  start: first supported whole-paragraph References heading alias
  end:   first Appendix heading (any style), rubric heading (heading
         style + rubric text), or explicit Heading 1/2/3 paragraph
  entries: non-empty paragraphs strictly inside

Bold/centred/short Normal paragraphs, list paragraphs, punctuation,
author/year parsing, and DOI/URL presence never terminate the span.
Phase 3 formatting eligibility stays broader than citation-key parsing:
dateless, institution-authored, malformed, wrapped, DOI-less, URL-less,
and numbered entries inside a valid span remain eligible.
"""
import io

from docx import Document
from docx.shared import Inches

from app.services.citation_sensor import run_citation_sensor
from app.services.document_parser import extract_paragraphs
from app.services.reference_section import (
    is_references_span_terminator,
    reference_entry_indexes,
    reference_section_bounds,
)
from app.services.reference_rules import run_reference_rules
from app.services.role_classifier import classify_paragraphs
from app.services.profile_preset_adapter import EffectiveProfileConfig
from app.services.profile_registry import SUC_PROFILE_ID, get_builtin_profile
from app.services.profile_schema import new_custom_profile
from app.services.profile_snapshot import resolve_snapshot

ENTRY = "Smith, J. (2020). A synthetic source. Test Journal, 1(1), 1-10."
ENTRY2 = "Nguyen, T. (2021). Second entry. Other Press."


def _doc(parts):
    """Build a docx from (text, style|None) tuples; style None means default."""
    doc = Document()
    for text, style in parts:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plain(text, style=None):
    return (text, style)


def _parsed(data):
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    return doc, paragraphs


def _profile(**refs):
    profile = new_custom_profile(
        "Boundary Tests", base=get_builtin_profile(SUC_PROFILE_ID),
    )
    for name, value in refs.items():
        setattr(profile.references, name, value)
    return EffectiveProfileConfig(resolve_snapshot(profile))


def _no_indent(doc):
    for p in doc.paragraphs:
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _hanging_config():
    return _profile(check_hanging_indent=True, hanging_indent_in=0.5)


# --- terminator predicate ----------------------------------------------------


def test_appendix_text_terminates_in_any_style():
    assert is_references_span_terminator("Appendix A", None) is True
    assert is_references_span_terminator("appendix", "Normal") is True
    assert is_references_span_terminator("Appendix B", "Heading 1") is True


def test_rubric_heading_terminates_but_plain_rubric_text_does_not():
    assert is_references_span_terminator("Marking Rubric", "Heading 1") is True
    assert is_references_span_terminator("Rubric", "Normal") is False
    assert is_references_span_terminator("Marking Rubric", "Normal") is False


def test_heading_styles_terminate_but_lookalikes_do_not():
    for style in ("Heading 1", "Heading 2", "Heading 3", "HEADING 2"):
        assert is_references_span_terminator("Expected Results", style) is True
    for style in (None, "Normal", "List Bullet", "Title", "Heading 4"):
        assert is_references_span_terminator("Expected Results", style) is False
    assert is_references_span_terminator("", "Heading 1") is False
    assert is_references_span_terminator(None, "Heading 1") is False


# --- A/B: document end + appendix --------------------------------------------


def test_A_references_at_document_end_unchanged():
    data = _doc([_plain("Body (Garcia, 2018)."), _plain("References"), _plain(ENTRY)])
    _, paragraphs = _parsed(data)
    bounds = reference_section_bounds(paragraphs)
    assert bounds is not None
    assert bounds.end_index == len(paragraphs)
    assert tuple(bounds.entry_indexes) == (2,)


def test_B_appendix_termination_compatible():
    data = _doc([
        _plain("References"), _plain(ENTRY), _plain(ENTRY2),
        _plain("Appendix A"), _plain("Smith, J. (2020). Not an entry. X."),
    ])
    _, paragraphs = _parsed(data)
    bounds = reference_section_bounds(paragraphs)
    assert bounds is not None
    assert bounds.end_index == 3
    assert tuple(bounds.entry_indexes) == (1, 2)


# --- C: rubric parity ----------------------------------------------------------


def test_C_rubric_terminates_bounds_roles_and_pairs_together():
    data = _doc([
        _plain("References"), _plain(ENTRY),
        _plain("Marking Rubric", "Heading 1"), _plain("Smith, J. (2020). Tail. X."),
    ])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    bounds = reference_section_bounds(paragraphs)
    assert bounds is not None
    assert bounds.end_index == 2
    assert tuple(bounds.entry_indexes) == (1,)
    roles = classify_paragraphs(doc, paragraphs)
    assert roles[2] != "REFERENCE_ENTRY"
    assert roles[3] != "REFERENCE_ENTRY"


# --- D/E/F: heading 1/2/3 termination ------------------------------------------


def test_DEF_heading_levels_terminate_the_span():
    for level in ("Heading 1", "Heading 2", "Heading 3"):
        data = _doc([
            _plain("References"), _plain(ENTRY),
            _plain("Expected Results", level), _plain("Smith, J. (2020). Tail. X."),
        ])
        _, paragraphs = _parsed(data)
        bounds = reference_section_bounds(paragraphs)
        assert bounds is not None, level
        assert bounds.end_index == 2, level
        assert tuple(bounds.entry_indexes) == (1,), level


# --- G/O: heading + bullets excluded -------------------------------------------


def test_G_heading_and_bullets_are_outside_the_span():
    data = _doc([
        _plain("References"), _plain(ENTRY),
        _plain("Expected Results", "Heading 1"),
        _plain("First explanatory bullet.", "List Bullet"),
        _plain("Second explanatory bullet.", "List Bullet"),
    ])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    bounds = reference_section_bounds(paragraphs)
    assert bounds is not None
    assert tuple(bounds.entry_indexes) == (1,)
    roles = classify_paragraphs(doc, paragraphs)
    assert roles[2] == "HEADING_1"
    assert roles[3] != "REFERENCE_ENTRY"
    assert roles[4] != "REFERENCE_ENTRY"
    violations = run_reference_rules(
        doc, paragraphs, _hanging_config(),
    )
    assert all(
        (v.location or {}).get("paragraph_index", 1) < 2 for v in violations
    ), [v.rule_code for v in violations]


# --- H/I/J/K/L/M: broad eligibility inside a valid span -------------------------


def _hanging_codes(data):
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    violations = run_reference_rules(doc, paragraphs, _hanging_config())
    return violations


def test_H_dateless_entry_remains_eligible():
    data = _doc([
        _plain("References"),
        _plain("Smith. A study without any year whatsoever, still an entry."),
    ])
    violations = _hanging_codes(_no_indent(Document(io.BytesIO(data))))
    assert any(v.rule_code == "REFERENCE_HANGING_INDENT" for v in violations)


def test_I_institution_authored_entry_remains_eligible():
    data = _doc([
        _plain("References"),
        _plain("World Health Organization. Global health report. Geneva: WHO."),
    ])
    violations = _hanging_codes(_no_indent(Document(io.BytesIO(data))))
    assert any(v.rule_code == "REFERENCE_HANGING_INDENT" for v in violations)


def test_J_malformed_but_located_entry_remains_eligible():
    data = _doc([
        _plain("References"),
        _plain("just some prose words without author year structure at all"),
    ])
    violations = _hanging_codes(_no_indent(Document(io.BytesIO(data))))
    assert any(v.rule_code == "REFERENCE_HANGING_INDENT" for v in violations)


def test_K_doiless_urlless_entries_remain_eligible():
    data = _doc([_plain("References"), _plain(ENTRY)])
    violations = _hanging_codes(_no_indent(Document(io.BytesIO(data))))
    assert any(v.rule_code == "REFERENCE_HANGING_INDENT" for v in violations)


def test_L_numbered_list_entry_remains_supported():
    data = _doc([
        _plain("References"),
        _plain("Smith, J. (2020). Numbered entry. Test Journal.", "List Number"),
    ])
    _, paragraphs = _parsed(data)
    assert 1 in reference_entry_indexes(paragraphs)


def test_M_wrapped_entry_remains_eligible():
    long_entry = (
        "Smith, J. (2020). A very long reference entry that wraps across "
        "several visual lines but remains one paragraph. Test Journal, 1(1), 1-10."
    )
    data = _doc([_plain("References"), _plain(long_entry)])
    violations = _hanging_codes(_no_indent(Document(io.BytesIO(data))))
    assert any(v.rule_code == "REFERENCE_HANGING_INDENT" for v in violations)


# --- N: structurally ambiguous Normal prose stays inside -------------------------


def test_N_normal_prose_inside_span_is_conservatively_retained():
    # Structurally ambiguous by design (no heading/appendix/rubric
    # terminator): the prose stays eligible. Documented limitation —
    # no text-specific exclusion is applied.
    data = _doc([
        _plain("References"), _plain(ENTRY),
        _plain("Summary of findings and next steps."),
    ])
    _, paragraphs = _parsed(data)
    assert tuple(reference_section_bounds(paragraphs).entry_indexes) == (1, 2)
    doc = Document(io.BytesIO(data))
    roles = classify_paragraphs(doc, extract_paragraphs(doc))
    assert roles[2] == "REFERENCE_ENTRY"


# --- P/Q/R: no false termination -----------------------------------------------


def test_P_heading_like_words_inside_a_normal_entry_do_not_terminate():
    entry = "References and further reading discussed by Smith, J. (2020). X."
    data = _doc([
        _plain("References"), _plain(entry), _plain(ENTRY2),
    ])
    _, paragraphs = _parsed(data)
    bounds = reference_section_bounds(paragraphs)
    assert tuple(bounds.entry_indexes) == (1, 2)


def test_Q_first_terminator_wins_over_later_headings():
    data = _doc([
        _plain("References"), _plain(ENTRY),
        _plain("Notes", "Heading 1"), _plain("Appendix", "Heading 2"),
    ])
    _, paragraphs = _parsed(data)
    bounds = reference_section_bounds(paragraphs)
    assert bounds.end_index == 2
    assert tuple(bounds.entry_indexes) == (1,)


def test_R_alias_vocabulary_still_starts_spans_with_terminators():
    data = _doc([
        _plain("6 References"), _plain(ENTRY),
        _plain("Notes", "Heading 2"),
    ])
    _, paragraphs = _parsed(data)
    bounds = reference_section_bounds(paragraphs)
    assert bounds is not None
    assert bounds.heading_alias == "references"
    assert tuple(bounds.entry_indexes) == (1,)


# --- S: Phase 2 directionality ---------------------------------------------------


def test_S_excluded_end_matter_leaves_the_matching_index():
    cited = "Body cites Quinn (2019) here."
    good = "Quinn, A. (2019). Cited work. Test Journal."
    data = _doc([
        _plain(cited), _plain("References"), _plain(good),
        _plain("Notes", "Heading 1"), _plain("Quinn, A. (2019). End-matter copy. X."),
    ])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    roles = classify_paragraphs(doc, paragraphs)
    violations = run_citation_sensor(doc, paragraphs, config=None, roles=roles)
    codes = [v.rule_code for v in violations]
    assert "CITATION_MISMATCH" not in codes
    assert "REFERENCE_DUPLICATE" not in codes


def test_S_citation_to_only_excluded_entry_still_mismatches():
    cited = "Body cites Quinn (2019) here."
    data = _doc([
        _plain(cited), _plain("References"),
        _plain("Notes", "Heading 1"), _plain("Quinn, A. (2019). End-matter copy. X."),
    ])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    roles = classify_paragraphs(doc, paragraphs)
    violations = run_citation_sensor(doc, paragraphs, config=None, roles=roles)
    assert [v.rule_code for v in violations] == ["CITATION_MISMATCH"]


# --- T/U: no findings outside the span, locations stay truthful ------------------


def test_TU_excluded_content_gets_no_reference_findings_at_true_locations():
    data = _doc([
        _plain("References"), _plain(ENTRY),
        _plain("Expected Results", "Heading 1"),
        _plain("First explanatory bullet.", "List Bullet"),
    ])
    doc = Document(io.BytesIO(data))
    paragraphs = extract_paragraphs(doc)
    violations = run_reference_rules(doc, paragraphs, _hanging_config())
    for v in violations:
        idx = (v.location or {}).get("paragraph_index")
        assert idx is not None and idx < 2, (v.rule_code, idx)
