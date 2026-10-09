"""Synthetic DOCX builders for Reference List structure/evidence tests.

Every builder returns raw `.docx` bytes containing synthetic content only —
no real student document, no personal data, and nothing downloaded.

These fixtures deliberately cover the shapes a later rules branch must not
get wrong: style-inherited indents, numbered entries, mixed run formatting,
explicit page breaks, wrapped entries and excluded structures (tables,
text boxes).
"""

import io
from typing import Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

REFERENCE_TEXT = "Garcia, R. (2018). Adaptive compliance in practice. Journal of Tests, 4(2), 11-20."


def _doc_to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def body_with_heading(heading: str = "References", entries: Optional[List[str]] = None,
                      *, trailing: Optional[List[str]] = None) -> bytes:
    """Body paragraph, then a heading, then entries, then trailing paragraphs."""
    doc = Document()
    doc.add_paragraph("This paragraph discusses prior work (Garcia, 2018).")
    doc.add_paragraph(heading)
    for entry in entries or [REFERENCE_TEXT]:
        doc.add_paragraph(entry)
    for extra in trailing or []:
        doc.add_paragraph(extra)
    return _doc_to_bytes(doc)


def heading_variants() -> Dict[str, bytes]:
    """Each accepted heading alias, bare and with a numeric section prefix."""
    variants: Dict[str, bytes] = {}
    for alias in ("References", "Bibliography", "Works Cited", "Reference List"):
        variants[alias] = body_with_heading(alias)
        variants[f"6 {alias}"] = body_with_heading(f"6 {alias}")
        variants[f"6.1 {alias}"] = body_with_heading(f"6.1 {alias}")
    # Singular "Reference" is NOT an accepted alias — pinned so a future
    # change is deliberate rather than accidental.
    variants["Reference (singular, not an alias)"] = body_with_heading("Reference")
    variants["Prose starting with References"] = body_with_heading(
        "References are listed below"
    )
    return variants


def appendix_terminates_section() -> bytes:
    return body_with_heading(
        "References",
        [REFERENCE_TEXT, "Nguyen, T. (2020). Second entry. Other Press."],
        trailing=["Appendix A", "Rubric table follows."],
    )


def empty_paragraphs_inside_section() -> bytes:
    doc = Document()
    doc.add_paragraph("Body cites (Garcia, 2018).")
    doc.add_paragraph("References")
    doc.add_paragraph(REFERENCE_TEXT)
    doc.add_paragraph("")          # separator
    doc.add_paragraph("   ")       # whitespace-only separator
    doc.add_paragraph("Nguyen, T. (2020). Second entry. Other Press.")
    return _doc_to_bytes(doc)


def long_wrapped_entry() -> bytes:
    """One long entry. Visual wrapping must stay ONE paragraph."""
    return body_with_heading(
        "References",
        [REFERENCE_TEXT * 6],
    )


def manual_line_break_entry() -> bytes:
    """A `w:br` inside one paragraph still yields one entry."""
    doc = Document()
    doc.add_paragraph("Body cites (Garcia, 2018).")
    doc.add_paragraph("References")
    p = doc.add_paragraph("Garcia, R. (2018). Adaptive compliance in practice.")
    p.add_run().add_break()
    p.add_run("Journal of Tests, 4(2), 11-20.")
    return _doc_to_bytes(doc)


# --- Indentation ------------------------------------------------------------


def direct_hanging_indent(inches: float = 0.5) -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph(REFERENCE_TEXT)
    p.paragraph_format.left_indent = Inches(inches)
    p.paragraph_format.first_line_indent = Inches(-inches)
    return _doc_to_bytes(doc)


def direct_first_line_indent(inches: float = 0.5) -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph(REFERENCE_TEXT)
    p.paragraph_format.first_line_indent = Inches(inches)
    return _doc_to_bytes(doc)


def no_explicit_indent() -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    doc.add_paragraph(REFERENCE_TEXT)
    return _doc_to_bytes(doc)


def style_inherited_hanging_indent(inches: float = 0.5) -> bytes:
    """Indent lives on the style; the paragraph carries no <w:ind>."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.paragraph_format.left_indent = Inches(inches)
    normal.paragraph_format.first_line_indent = Inches(-inches)
    doc.add_paragraph("References")
    doc.add_paragraph(REFERENCE_TEXT)
    return _doc_to_bytes(doc)


def base_style_inherited_hanging_indent(inches: float = 0.5) -> bytes:
    """Two levels deep: a custom style based on Normal, which carries it."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.paragraph_format.left_indent = Inches(inches)
    normal.paragraph_format.first_line_indent = Inches(-inches)
    ref_style = doc.styles.add_style("Reference Entry", 1)
    ref_style.base_style = normal
    doc.add_paragraph("References")
    doc.add_paragraph(REFERENCE_TEXT, style="Reference Entry")
    return _doc_to_bytes(doc)


def numbered_reference() -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph(REFERENCE_TEXT)
    # Attach explicit numbering so the paragraph's own <w:pPr> carries
    # <w:numPr>, which is what the existing is_list_item check detects.
    pPr = p._p.get_or_add_pPr()
    from docx.oxml import OxmlElement

    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn_val("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn_val("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    return _doc_to_bytes(doc)


def qn_val(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


# --- Paragraph / run formatting ---------------------------------------------


def direct_alignment_and_spacing() -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph(REFERENCE_TEXT)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    return _doc_to_bytes(doc)


def inherited_alignment_and_spacing() -> bytes:
    """Values live on the style; the paragraph carries none."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(6)
    doc.add_paragraph("References")
    doc.add_paragraph(REFERENCE_TEXT)
    return _doc_to_bytes(doc)


def effective_font_family_and_size(name: str = "Times New Roman", size: float = 12.0) -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph()
    run = p.add_run(REFERENCE_TEXT)
    run.font.name = name
    run.font.size = Pt(size)
    return _doc_to_bytes(doc)


def mixed_font_family() -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph()
    a = p.add_run("Garcia, R. (2018). ")
    a.font.name = "Arial"
    a.font.size = Pt(12)
    b = p.add_run("Adaptive compliance in practice.")
    b.font.name = "Times New Roman"
    b.font.size = Pt(12)
    return _doc_to_bytes(doc)


def mixed_font_size() -> bytes:
    doc = Document()
    doc.add_paragraph("References")
    p = doc.add_paragraph()
    a = p.add_run("Garcia, R. (2018). ")
    a.font.name = "Times New Roman"
    a.font.size = Pt(12)
    b = p.add_run("Adaptive compliance in practice.")
    b.font.name = "Times New Roman"
    b.font.size = Pt(14)
    return _doc_to_bytes(doc)


# --- Page-start evidence ----------------------------------------------------


def explicit_page_break_run_before_heading() -> bytes:
    """A `w:br w:type="page"` run inside the heading paragraph."""
    doc = Document()
    doc.add_paragraph("Body cites (Garcia, 2018).")
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("References")
    doc.add_paragraph(REFERENCE_TEXT)
    return _doc_to_bytes(doc)


def page_break_before_property() -> bytes:
    """`w:pageBreakBefore` on the heading paragraph."""
    doc = Document()
    doc.add_paragraph("Body cites (Garcia, 2018).")
    h = doc.add_paragraph("References")
    h.paragraph_format.page_break_before = True
    doc.add_paragraph(REFERENCE_TEXT)
    return _doc_to_bytes(doc)


def no_page_break() -> bytes:
    """Natural pagination only — no explicit break anywhere."""
    return body_with_heading("References")


# --- Excluded structures ----------------------------------------------------


def table_based_reference() -> bytes:
    """A reference living inside a table cell — not a parser paragraph."""
    doc = Document()
    doc.add_paragraph("References")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = REFERENCE_TEXT
    return _doc_to_bytes(doc)


def textbox_reference() -> bytes:
    """A reference inside a real text box.

    A genuine text box nests its paragraph inside a shape
    (`v:textbox/w:txbxContent`), which is NOT a direct child of `w:body`.
    That is the point of this fixture: `doc.paragraphs` must not surface it
    as a Reference entry.
    """
    doc = Document()
    doc.add_paragraph("References")
    from docx.oxml import parse_xml

    xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:r><w:pict>"
        '<v:shape id="tb1" type="#_x0000_t202" style="width:200pt;height:40pt">'
        "<v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>Garcia, R. (2018). Adaptive compliance in practice."
        "</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox>"
        "</v:shape></w:pict></w:r></w:p>"
    )
    doc.element.body.append(parse_xml(xml))
    return _doc_to_bytes(doc)
